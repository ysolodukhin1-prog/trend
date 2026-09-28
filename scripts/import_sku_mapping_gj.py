from __future__ import annotations

import hashlib
import sys
import time
from pathlib import Path

import psycopg2.extras
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "ozon_category_dashboard"))

import app  # noqa: E402


SOURCE_FILE = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\SKU_mapping_Gj.xlsx"
)
BATCH_SIZE = 10_000


def clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def row_hash(*values):
    payload = "|".join("" if value is None else str(value) for value in values)
    return hashlib.md5(payload.encode("utf-8")).hexdigest()


def format_duration(seconds):
    seconds = int(seconds)
    minutes, sec = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours}h {minutes:02d}m {sec:02d}s"
    if minutes:
        return f"{minutes}m {sec:02d}s"
    return f"{sec}s"


def ensure_schema(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS public.sku_mapping_gj (
            mapping_hash text PRIMARY KEY,
            barcode text NOT NULL,
            wb_article text,
            wb_nmid text,
            ozon_sku text,
            gj_model text,
            assortment_bia text,
            tg text,
            tg_plus text,
            cg text,
            season text,
            source_file text,
            source_wb_row integer,
            source_ozon_row integer,
            imported_at timestamp without time zone DEFAULT now()
        )
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sku_mapping_gj_barcode ON public.sku_mapping_gj (barcode)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sku_mapping_gj_wb_article ON public.sku_mapping_gj (wb_article)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sku_mapping_gj_ozon_sku ON public.sku_mapping_gj (ozon_sku)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sku_mapping_gj_gj_model ON public.sku_mapping_gj (gj_model)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sku_mapping_gj_assortment ON public.sku_mapping_gj (assortment_bia)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sku_mapping_gj_tg ON public.sku_mapping_gj (tg)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sku_mapping_gj_tg_plus ON public.sku_mapping_gj (tg_plus)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sku_mapping_gj_cg ON public.sku_mapping_gj (cg)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_sku_mapping_gj_season ON public.sku_mapping_gj (season)")
    cur.execute(
        """
        CREATE MATERIALIZED VIEW IF NOT EXISTS public.mv_sku_mapping_gj_wb_article AS
        SELECT
            wb_article,
            max(wb_nmid) AS wb_nmid,
            min(barcode) AS barcode,
            min(ozon_sku) AS ozon_sku,
            max(gj_model) AS gj_model,
            max(assortment_bia) AS assortment_bia,
            max(tg) AS tg,
            max(tg_plus) AS tg_plus,
            max(cg) AS cg,
            max(season) AS season
        FROM public.sku_mapping_gj
        WHERE wb_article IS NOT NULL
        GROUP BY wb_article
        """
    )
    cur.execute(
        """
        CREATE MATERIALIZED VIEW IF NOT EXISTS public.mv_sku_mapping_gj_wb_nmid AS
        SELECT
            wb_nmid,
            min(wb_article) AS wb_article,
            min(barcode) AS barcode,
            min(ozon_sku) AS ozon_sku,
            max(gj_model) AS gj_model,
            max(assortment_bia) AS assortment_bia,
            max(tg) AS tg,
            max(tg_plus) AS tg_plus,
            max(cg) AS cg,
            max(season) AS season
        FROM public.sku_mapping_gj
        WHERE wb_nmid IS NOT NULL
        GROUP BY wb_nmid
        """
    )
    cur.execute(
        """
        CREATE MATERIALIZED VIEW IF NOT EXISTS public.mv_sku_mapping_gj_ozon_sku AS
        SELECT
            ozon_sku,
            min(barcode) AS barcode,
            min(wb_article) AS wb_article,
            max(wb_nmid) AS wb_nmid,
            max(gj_model) AS gj_model,
            max(assortment_bia) AS assortment_bia,
            max(tg) AS tg,
            max(tg_plus) AS tg_plus,
            max(cg) AS cg,
            max(season) AS season
        FROM public.sku_mapping_gj
        WHERE ozon_sku IS NOT NULL
        GROUP BY ozon_sku
        """
    )
    cur.execute(
        """
        CREATE MATERIALIZED VIEW IF NOT EXISTS public.mv_sku_mapping_gj_barcode AS
        SELECT
            barcode,
            min(wb_article) AS wb_article,
            max(wb_nmid) AS wb_nmid,
            min(ozon_sku) AS ozon_sku,
            max(gj_model) AS gj_model,
            max(assortment_bia) AS assortment_bia,
            max(tg) AS tg,
            max(tg_plus) AS tg_plus,
            max(cg) AS cg,
            max(season) AS season
        FROM public.sku_mapping_gj
        GROUP BY barcode
        """
    )
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_mv_sku_mapping_gj_wb_article_wb_article ON public.mv_sku_mapping_gj_wb_article (wb_article)")
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_mv_sku_mapping_gj_wb_nmid ON public.mv_sku_mapping_gj_wb_nmid (wb_nmid)")
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_mv_sku_mapping_gj_ozon_sku_ozon_sku ON public.mv_sku_mapping_gj_ozon_sku (ozon_sku)")
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_mv_sku_mapping_gj_barcode_barcode ON public.mv_sku_mapping_gj_barcode (barcode)")


def refresh_mapping_views(cur):
    for view_name in (
        "mv_sku_mapping_gj_wb_article",
        "mv_sku_mapping_gj_wb_nmid",
        "mv_sku_mapping_gj_ozon_sku",
        "mv_sku_mapping_gj_barcode",
    ):
        cur.execute(f"REFRESH MATERIALIZED VIEW public.{view_name}")
        cur.execute(f"ANALYZE public.{view_name}")


def sheet_dict(ws):
    headers = [clean(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    return {name: index for index, name in enumerate(headers) if name}


def read_sheet(ws, kind):
    idx = sheet_dict(ws)
    required = ["ID товара", "Модель GJ", "Осн. Ассорт / БиА", "ТГ", "ТГ+", "ЦГ", "Сезон"]
    missing = [name for name in required if name not in idx]
    if kind == "wb":
        missing.extend(name for name in ("Артикул", "nmid") if name not in idx)
    else:
        missing.extend(name for name in ("sku_OZON",) if name not in idx)
    if missing:
        raise RuntimeError(f"Sheet {ws.title}: missing headers {missing}")

    rows = {}
    total = max(ws.max_row - 1, 0)
    started = time.time()
    for position, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        barcode = clean(row[idx["ID товара"]])
        if not barcode:
            continue
        rows.setdefault(barcode, []).append(
            {
                "barcode": barcode,
                "wb_article": clean(row[idx["Артикул"]]) if kind == "wb" else None,
                "wb_nmid": clean(row[idx["nmid"]]) if kind == "wb" else None,
                "ozon_sku": clean(row[idx["sku_OZON"]]) if kind == "ozon" else None,
                "gj_model": clean(row[idx["Модель GJ"]]),
                "assortment_bia": clean(row[idx["Осн. Ассорт / БиА"]]),
                "tg": clean(row[idx["ТГ"]]),
                "tg_plus": clean(row[idx["ТГ+"]]),
                "cg": clean(row[idx["ЦГ"]]),
                "season": clean(row[idx["Сезон"]]),
                "row_num": position,
            }
        )
        if position % 25_000 == 0:
            done = position - 1
            elapsed = time.time() - started
            speed = done / elapsed if elapsed else 0
            print(
                f"{ws.title}: {done:,}/{total:,} rows | elapsed={format_duration(elapsed)} | "
                f"speed={speed:,.0f} rows/s",
                flush=True,
            )
    return rows


def first_non_empty(*values):
    for value in values:
        if value:
            return value
    return None


def build_records(wb_rows, ozon_rows):
    records = []
    for barcode in sorted(set(wb_rows) | set(ozon_rows)):
        wb_items = wb_rows.get(barcode) or [None]
        ozon_items = ozon_rows.get(barcode) or [None]
        for wb_item in wb_items:
            for ozon_item in ozon_items:
                category_source = wb_item or ozon_item or {}
                if wb_item and ozon_item:
                    category_source = {
                        key: first_non_empty(wb_item.get(key), ozon_item.get(key))
                        for key in ("gj_model", "assortment_bia", "tg", "tg_plus", "cg", "season")
                    }
                values = {
                    "barcode": barcode,
                    "wb_article": wb_item.get("wb_article") if wb_item else None,
                    "wb_nmid": wb_item.get("wb_nmid") if wb_item else None,
                    "ozon_sku": ozon_item.get("ozon_sku") if ozon_item else None,
                    "gj_model": category_source.get("gj_model"),
                    "assortment_bia": category_source.get("assortment_bia"),
                    "tg": category_source.get("tg"),
                    "tg_plus": category_source.get("tg_plus"),
                    "cg": category_source.get("cg"),
                    "season": category_source.get("season"),
                    "source_wb_row": wb_item.get("row_num") if wb_item else None,
                    "source_ozon_row": ozon_item.get("row_num") if ozon_item else None,
                }
                values["mapping_hash"] = row_hash(
                    values["barcode"],
                    values["wb_article"],
                    values["wb_nmid"],
                    values["ozon_sku"],
                    values["gj_model"],
                    values["assortment_bia"],
                    values["tg"],
                    values["tg_plus"],
                    values["cg"],
                    values["season"],
                )
                records.append(values)
    return records


def main():
    started = time.time()
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(SOURCE_FILE)

    print(f"Source: {SOURCE_FILE}", flush=True)
    workbook = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    wb_rows = read_sheet(workbook["wb"], "wb")
    ozon_rows = read_sheet(workbook["ozon"], "ozon")
    records = build_records(wb_rows, ozon_rows)
    print(
        f"Prepared records: {len(records):,} | WB barcodes={len(wb_rows):,} | Ozon barcodes={len(ozon_rows):,}",
        flush=True,
    )

    insert_sql = """
        INSERT INTO public.sku_mapping_gj (
            mapping_hash, barcode, wb_article, wb_nmid, ozon_sku, gj_model,
            assortment_bia, tg, tg_plus, cg, season, source_file,
            source_wb_row, source_ozon_row, imported_at
        )
        VALUES %s
        ON CONFLICT (mapping_hash) DO UPDATE SET
            barcode = EXCLUDED.barcode,
            wb_article = EXCLUDED.wb_article,
            wb_nmid = EXCLUDED.wb_nmid,
            ozon_sku = EXCLUDED.ozon_sku,
            gj_model = EXCLUDED.gj_model,
            assortment_bia = EXCLUDED.assortment_bia,
            tg = EXCLUDED.tg,
            tg_plus = EXCLUDED.tg_plus,
            cg = EXCLUDED.cg,
            season = EXCLUDED.season,
            source_file = EXCLUDED.source_file,
            source_wb_row = EXCLUDED.source_wb_row,
            source_ozon_row = EXCLUDED.source_ozon_row,
            imported_at = now()
    """
    template = (
        "(%(mapping_hash)s, %(barcode)s, %(wb_article)s, %(wb_nmid)s, %(ozon_sku)s, %(gj_model)s, "
        "%(assortment_bia)s, %(tg)s, %(tg_plus)s, %(cg)s, %(season)s, %(source_file)s, "
        "%(source_wb_row)s, %(source_ozon_row)s, now())"
    )
    for record in records:
        record["source_file"] = str(SOURCE_FILE)

    with app.get_conn() as conn, conn.cursor() as cur:
        ensure_schema(cur)
        cur.execute("TRUNCATE public.sku_mapping_gj")
        for start in range(0, len(records), BATCH_SIZE):
            batch = records[start : start + BATCH_SIZE]
            psycopg2.extras.execute_values(cur, insert_sql, batch, template=template, page_size=BATCH_SIZE)
            conn.commit()
            done = min(start + BATCH_SIZE, len(records))
            elapsed = time.time() - started
            speed = done / elapsed if elapsed else 0
            print(
                f"DB import: {done:,}/{len(records):,} | elapsed={format_duration(elapsed)} | "
                f"speed={speed:,.0f} rows/s",
                flush=True,
            )

        cur.execute("ANALYZE public.sku_mapping_gj")
        refresh_mapping_views(cur)
        conn.commit()
        cur.execute(
            """
            SELECT
                count(*) AS rows,
                count(DISTINCT barcode) AS barcodes,
                count(DISTINCT wb_article) AS wb_articles,
                count(DISTINCT ozon_sku) AS ozon_skus
            FROM public.sku_mapping_gj
            """
        )
        print(dict(cur.fetchone()), flush=True)

    print(f"Done in {format_duration(time.time() - started)}", flush=True)


if __name__ == "__main__":
    main()
