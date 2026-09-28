#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Import Wildberries media advertising daily XLSX snapshots into PostgreSQL."""

from __future__ import annotations

import argparse
import csv
import io
import re
import sys
import time
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DASHBOARD_ROOT = PROJECT_ROOT / "ozon_category_dashboard"
sys.path.insert(0, str(DASHBOARD_ROOT))

import app  # noqa: E402


DEFAULT_SOURCE_DIR = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Data\Wb\Media_adv"
)

LEVEL_BY_DIR = {
    "camp": "campaign",
    "creative": "creative",
    "group": "group",
}
SOURCE_SHEET = "Sheet1"

RAW_COLUMNS = [
    "report_date",
    "media_level",
    "campaign_id",
    "group_id",
    "creative_id",
    "campaign_name",
    "entity_name",
    "creative_type",
    "campaign_status",
    "preview_url",
    "reach",
    "budget_rub",
    "balance_rub",
    "expense_rub",
    "views",
    "clicks",
    "ctr_pct",
    "cpm_rub",
    "cpc_rub",
    "vr_pct",
    "vp25",
    "vp50",
    "vp75",
    "video_started",
    "video_completed",
    "video_paused",
    "video_resumed",
    "drr_pct",
    "baskets",
    "orders",
    "orders_post_click",
    "orders_qty_post_click",
    "orders_sum_post_click",
    "baskets_post_click",
    "baskets_qty_post_click",
    "baskets_sum_post_click",
    "orders_post_view",
    "orders_qty_post_view",
    "orders_sum_post_view",
    "baskets_post_view",
    "baskets_qty_post_view",
    "baskets_sum_post_view",
    "source_file",
    "source_sheet",
    "source_row_num",
    "file_size_bytes",
    "file_mtime",
]

COLUMN_MAP = {
    "campaignid": "campaign_id",
    "groupid": "group_id",
    "creativeid": "creative_id",
    "name": "entity_name",
    "creativetype": "creative_type",
    "campaignstatus": "campaign_status",
    "preview": "preview_url",
    "previews": "preview_url",
    "reach": "reach",
    "budget": "budget_rub",
    "balance": "balance_rub",
    "cost": "expense_rub",
    "views": "views",
    "clicks": "clicks",
    "ctr": "ctr_pct",
    "cpm": "cpm_rub",
    "cpc": "cpc_rub",
    "vr": "vr_pct",
    "vp25": "vp25",
    "vp50": "vp50",
    "vp75": "vp75",
    "videostarted": "video_started",
    "videocompleted": "video_completed",
    "videopaused": "video_paused",
    "videoresumed": "video_resumed",
    "drr": "drr_pct",
    "baskets": "baskets",
    "orders": "orders",
    "orderspostclick": "orders_post_click",
    "ordersqtypostclick": "orders_qty_post_click",
    "orderssumpostclick": "orders_sum_post_click",
    "basketspostclick": "baskets_post_click",
    "basketsqtypostclick": "baskets_qty_post_click",
    "basketssumpostclick": "baskets_sum_post_click",
    "orderspostview": "orders_post_view",
    "ordersqtypostview": "orders_qty_post_view",
    "orderssumpostview": "orders_sum_post_view",
    "basketspostview": "baskets_post_view",
    "basketsqtypostview": "baskets_qty_post_view",
    "basketssumpostview": "baskets_sum_post_view",
}

TEXT_COLUMNS = {
    "campaign_id",
    "group_id",
    "creative_id",
    "campaign_name",
    "entity_name",
    "creative_type",
    "campaign_status",
    "preview_url",
}
NUMERIC_COLUMNS = set(RAW_COLUMNS[10:42])


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if text.lower() in {"", "nan", "none", "null", "-", "–", "—"}:
        return None
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def normalize_header(value: Any) -> str:
    text = clean_text(value) or ""
    text = text.replace("\xa0", " ").replace("\n", " ").replace("\r", " ")
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def to_decimal(value: Any) -> Decimal:
    text = clean_text(value)
    if text is None:
        return Decimal("0")
    text = text.replace("\xa0", "").replace(" ", "").replace(",", ".").rstrip("%")
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text:
        return Decimal("0")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0")


def source_file_value(source_dir: Path, file_path: Path) -> str:
    try:
        return str(file_path.relative_to(source_dir))
    except ValueError:
        return str(file_path)


def report_date_from_filename(file_path: Path) -> date:
    match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", file_path.name)
    if not match:
        raise ValueError(f"Cannot infer report date from filename: {file_path.name}")
    day, month, year = match.groups()
    return date(int(year), int(month), int(day))


def media_level_from_path(source_dir: Path, file_path: Path) -> str:
    try:
        relative = file_path.relative_to(source_dir)
        first = relative.parts[0].lower()
    except (ValueError, IndexError):
        first = file_path.parent.parent.name.lower()
    level = LEVEL_BY_DIR.get(first)
    if not level:
        raise ValueError(f"Cannot infer WB media level from path: {file_path}")
    return level


def header_positions(header_row: tuple[Any, ...], file_path: Path) -> dict[str, int]:
    positions: dict[str, int] = {}
    for index, value in enumerate(header_row):
        db_name = COLUMN_MAP.get(normalize_header(value))
        if db_name and db_name not in positions:
            positions[db_name] = index
    required = ["campaign_id", "entity_name", "creative_type", "campaign_status", "views", "clicks"]
    missing = [name for name in required if name not in positions]
    if missing:
        raise ValueError(f"Missing WB media columns in {file_path}: {missing}")
    return positions


def row_cell(row: tuple[Any, ...], positions: dict[str, int], column: str) -> Any:
    index = positions.get(column)
    if index is None or index >= len(row):
        return None
    return row[index]


def read_wb_media_adv_rows(source_dir: Path, file_path: Path):
    stat = file_path.stat()
    source_file = source_file_value(source_dir, file_path)
    report_date = report_date_from_filename(file_path)
    media_level = media_level_from_path(source_dir, file_path)
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return
        positions = header_positions(header, file_path)
        for row_num, row in enumerate(rows, start=2):
            campaign_id = clean_text(row_cell(row, positions, "campaign_id"))
            entity_name = clean_text(row_cell(row, positions, "entity_name"))
            if not campaign_id and not entity_name:
                continue
            parsed: dict[str, Any] = {
                "report_date": report_date,
                "media_level": media_level,
                "campaign_name": entity_name if media_level == "campaign" else None,
                "entity_name": entity_name,
                "source_file": source_file,
                "source_sheet": SOURCE_SHEET,
                "source_row_num": row_num,
                "file_size_bytes": stat.st_size,
                "file_mtime": datetime.fromtimestamp(stat.st_mtime),
            }
            for column in RAW_COLUMNS[2:42]:
                if column in {"campaign_name", "entity_name"}:
                    continue
                value = row_cell(row, positions, column)
                if column in TEXT_COLUMNS:
                    parsed[column] = clean_text(value)
                elif column in NUMERIC_COLUMNS:
                    parsed[column] = to_decimal(value)
            yield tuple(parsed.get(column) for column in RAW_COLUMNS)
    finally:
        workbook.close()


class WbMediaAdvImporter:
    def __init__(
        self,
        source_dir: Path,
        *,
        dry_run: bool = False,
        force_reimport: bool = False,
        max_files: int | None = None,
        skip_import: bool = False,
        skip_views: bool = False,
    ) -> None:
        self.source_dir = source_dir
        self.dry_run = dry_run
        self.force_reimport = force_reimport
        self.max_files = max_files
        self.skip_import = skip_import
        self.skip_views = skip_views
        self.conn = None
        self.cur = None

    def connect(self) -> None:
        self.conn = app.get_conn()
        self.cur = self.conn.cursor()

    def close(self) -> None:
        if self.cur is not None:
            self.cur.close()
        if self.conn is not None:
            self.conn.close()

    def execute(self, query: str, params: tuple[Any, ...] | None = None) -> None:
        assert self.cur is not None
        self.cur.execute(query, params)

    def create_schema(self, *, drop_views: bool = True) -> None:
        if drop_views:
            for view in [
                "mv_wb_media_adv_campaign_daily",
                "mv_wb_media_adv_group_daily",
                "mv_wb_media_adv_creative_daily",
            ]:
                self.execute(f"DROP MATERIALIZED VIEW IF EXISTS public.{view}")
        self.execute(
            """
            CREATE TABLE IF NOT EXISTS public.wb_media_adv_daily_raw (
                id bigserial PRIMARY KEY,
                report_date date NOT NULL,
                media_level text NOT NULL,
                campaign_id text,
                group_id text,
                creative_id text,
                campaign_name text,
                entity_name text,
                creative_type text,
                campaign_status text,
                preview_url text,
                reach numeric,
                budget_rub numeric,
                balance_rub numeric,
                expense_rub numeric,
                views numeric,
                clicks numeric,
                ctr_pct numeric,
                cpm_rub numeric,
                cpc_rub numeric,
                vr_pct numeric,
                vp25 numeric,
                vp50 numeric,
                vp75 numeric,
                video_started numeric,
                video_completed numeric,
                video_paused numeric,
                video_resumed numeric,
                drr_pct numeric,
                baskets numeric,
                orders numeric,
                orders_post_click numeric,
                orders_qty_post_click numeric,
                orders_sum_post_click numeric,
                baskets_post_click numeric,
                baskets_qty_post_click numeric,
                baskets_sum_post_click numeric,
                orders_post_view numeric,
                orders_qty_post_view numeric,
                orders_sum_post_view numeric,
                baskets_post_view numeric,
                baskets_qty_post_view numeric,
                baskets_sum_post_view numeric,
                source_file text NOT NULL,
                source_sheet text NOT NULL,
                source_row_num integer NOT NULL,
                file_size_bytes bigint,
                file_mtime timestamp,
                imported_at timestamp NOT NULL DEFAULT now(),
                UNIQUE (source_file, source_sheet, source_row_num)
            )
            """
        )
        self.execute(
            """
            CREATE TABLE IF NOT EXISTS public.wb_media_adv_import_files (
                source_file text PRIMARY KEY,
                media_level text,
                rows_imported integer NOT NULL DEFAULT 0,
                report_date date,
                file_size_bytes bigint,
                file_mtime timestamp,
                imported_at timestamp NOT NULL DEFAULT now(),
                status text NOT NULL DEFAULT 'ok',
                error text
            )
            """
        )
        if self.force_reimport:
            self.execute("TRUNCATE TABLE public.wb_media_adv_daily_raw RESTART IDENTITY")
            self.execute("TRUNCATE TABLE public.wb_media_adv_import_files")
        self.conn.commit()

    def imported_file_state(self) -> dict[str, tuple[int | None, datetime | None, str | None]]:
        self.execute(
            """
            SELECT source_file, file_size_bytes, file_mtime, status
            FROM public.wb_media_adv_import_files
            """
        )
        return {
            row["source_file"]: (row["file_size_bytes"], row["file_mtime"], row["status"])
            for row in self.cur.fetchall()
        }

    def should_skip_file(self, file_path: Path, imported: dict[str, tuple[int | None, datetime | None, str | None]]) -> bool:
        if self.force_reimport:
            return False
        source_file = source_file_value(self.source_dir, file_path)
        state = imported.get(source_file)
        if not state:
            return False
        size, mtime, status = state
        stat = file_path.stat()
        return status == "ok" and size == stat.st_size and mtime and abs(mtime.timestamp() - stat.st_mtime) < 1

    def import_file(self, file_path: Path) -> int:
        assert self.cur is not None and self.conn is not None
        source_file = source_file_value(self.source_dir, file_path)
        stat = file_path.stat()
        try:
            rows = list(read_wb_media_adv_rows(self.source_dir, file_path))
            media_level = rows[0][1] if rows else media_level_from_path(self.source_dir, file_path)
            report_date = rows[0][0] if rows else report_date_from_filename(file_path)
            if self.dry_run:
                return len(rows)
            self.execute("DELETE FROM public.wb_media_adv_daily_raw WHERE source_file = %s", (source_file,))
            if rows:
                buffer = io.StringIO()
                writer = csv.writer(buffer, lineterminator="\n")
                writer.writerows(rows)
                buffer.seek(0)
                self.cur.copy_expert(
                    f"COPY public.wb_media_adv_daily_raw ({', '.join(RAW_COLUMNS)}) FROM STDIN WITH CSV",
                    buffer,
                )
            self.execute(
                """
                INSERT INTO public.wb_media_adv_import_files (
                    source_file, media_level, rows_imported, report_date, file_size_bytes,
                    file_mtime, imported_at, status, error
                )
                VALUES (%s, %s, %s, %s, %s, %s, now(), 'ok', NULL)
                ON CONFLICT (source_file) DO UPDATE SET
                    media_level = EXCLUDED.media_level,
                    rows_imported = EXCLUDED.rows_imported,
                    report_date = EXCLUDED.report_date,
                    file_size_bytes = EXCLUDED.file_size_bytes,
                    file_mtime = EXCLUDED.file_mtime,
                    imported_at = now(),
                    status = 'ok',
                    error = NULL
                """,
                (source_file, media_level, len(rows), report_date, stat.st_size, datetime.fromtimestamp(stat.st_mtime)),
            )
            self.conn.commit()
            return len(rows)
        except Exception as exc:
            if not self.dry_run:
                self.conn.rollback()
                self.execute(
                    """
                    INSERT INTO public.wb_media_adv_import_files (
                        source_file, rows_imported, file_size_bytes, file_mtime,
                        imported_at, status, error
                    )
                    VALUES (%s, 0, %s, %s, now(), 'error', %s)
                    ON CONFLICT (source_file) DO UPDATE SET
                        rows_imported = 0,
                        file_size_bytes = EXCLUDED.file_size_bytes,
                        file_mtime = EXCLUDED.file_mtime,
                        imported_at = now(),
                        status = 'error',
                        error = EXCLUDED.error
                    """,
                    (source_file, stat.st_size, datetime.fromtimestamp(stat.st_mtime), str(exc)[:2000]),
                )
                self.conn.commit()
            print(f"ERROR {file_path}: {exc}")
            return 0

    def create_level_view(self, view_name: str, media_level: str) -> None:
        self.execute(f"DROP MATERIALIZED VIEW IF EXISTS public.{view_name}")
        self.execute(
            f"""
            CREATE MATERIALIZED VIEW public.{view_name} AS
            SELECT
                report_date,
                media_level,
                campaign_id,
                group_id,
                creative_id,
                campaign_name,
                entity_name,
                coalesce(entity_name, campaign_name, campaign_id) AS product_name,
                coalesce(creative_id, group_id, campaign_id) AS product_artikul,
                creative_type AS category_name,
                creative_type AS campaign_format,
                campaign_status,
                CASE
                    WHEN coalesce(campaign_name, entity_name, '') ILIKE '%контекст%' THEN 'Контекстный баннер'
                    WHEN coalesce(campaign_name, entity_name, '') ILIKE '%брендовая полка%' THEN 'Брендовая полка'
                    WHEN coalesce(campaign_name, entity_name, '') ILIKE '%распродаж%' THEN 'Распродажа'
                    ELSE 'Прочее'
                END AS campaign_segment,
                NULLIF(trim(both ' _-' from regexp_replace(coalesce(entity_name, campaign_name, ''), '.*[_-]', '')), '') AS promoted_category,
                preview_url,
                reach,
                budget_rub,
                balance_rub,
                coalesce(expense_rub, 0) AS expense_rub,
                views AS impressions,
                clicks,
                ctr_pct,
                CASE WHEN coalesce(views, 0) <> 0
                    THEN round(coalesce(clicks, 0)::numeric / views::numeric * 100, 4)
                    ELSE 0 END AS ctr_calc_pct,
                cpm_rub,
                CASE WHEN coalesce(views, 0) <> 0
                    THEN round(coalesce(expense_rub, 0)::numeric / views::numeric * 1000, 4)
                    ELSE 0 END AS cpm_calc_rub,
                cpc_rub,
                CASE WHEN coalesce(clicks, 0) <> 0
                    THEN round(coalesce(expense_rub, 0)::numeric / clicks::numeric, 4)
                    ELSE 0 END AS cpc_calc_rub,
                vr_pct,
                vp25,
                vp50,
                vp75,
                video_started,
                video_completed,
                video_paused,
                video_resumed,
                baskets,
                orders,
                orders_qty_post_click AS orders_qty,
                orders_sum_post_click AS orders_amount_rub,
                orders_qty_post_view AS post_view_orders_qty,
                orders_sum_post_view AS post_view_revenue_rub,
                coalesce(orders_qty_post_click, 0) + coalesce(orders_qty_post_view, 0) AS attributed_orders_qty,
                coalesce(orders_sum_post_click, 0) + coalesce(orders_sum_post_view, 0) AS attributed_revenue_rub,
                baskets_qty_post_click,
                baskets_sum_post_click,
                baskets_qty_post_view,
                baskets_sum_post_view,
                drr_pct,
                CASE WHEN coalesce(orders_sum_post_click, 0) <> 0
                    THEN round(coalesce(expense_rub, 0)::numeric / orders_sum_post_click::numeric * 100, 4)
                    ELSE 0 END AS drr_direct_pct,
                CASE WHEN coalesce(orders_sum_post_click, 0) + coalesce(orders_sum_post_view, 0) <> 0
                    THEN round(coalesce(expense_rub, 0)::numeric / (coalesce(orders_sum_post_click, 0) + coalesce(orders_sum_post_view, 0))::numeric * 100, 4)
                    ELSE 0 END AS drr_attributed_pct,
                CASE WHEN coalesce(expense_rub, 0) <> 0
                    THEN round(coalesce(orders_sum_post_click, 0)::numeric / expense_rub::numeric, 4)
                    ELSE 0 END AS direct_roas,
                CASE WHEN coalesce(expense_rub, 0) <> 0
                    THEN round((coalesce(orders_sum_post_click, 0) + coalesce(orders_sum_post_view, 0))::numeric / expense_rub::numeric, 4)
                    ELSE 0 END AS attributed_roas,
                CASE WHEN coalesce(clicks, 0) <> 0
                    THEN round(coalesce(orders_qty_post_click, 0)::numeric / clicks::numeric * 100, 4)
                    ELSE 0 END AS click_to_order_pct,
                CASE WHEN coalesce(views, 0) <> 0
                    THEN round(coalesce(orders_qty_post_view, 0)::numeric / views::numeric * 1000, 4)
                    ELSE 0 END AS post_view_orders_per_1000_impressions,
                CASE WHEN coalesce(views, 0) <> 0
                    THEN round((coalesce(orders_sum_post_click, 0) + coalesce(orders_sum_post_view, 0))::numeric / views::numeric * 1000, 4)
                    ELSE 0 END AS revenue_per_1000_impressions,
                CASE WHEN coalesce(orders_sum_post_click, 0) + coalesce(orders_sum_post_view, 0) <> 0
                    THEN round(coalesce(orders_sum_post_view, 0)::numeric / (coalesce(orders_sum_post_click, 0) + coalesce(orders_sum_post_view, 0))::numeric * 100, 4)
                    ELSE 0 END AS post_view_revenue_share_pct,
                source_file,
                source_sheet,
                source_row_num,
                file_size_bytes,
                file_mtime,
                imported_at
            FROM public.wb_media_adv_daily_raw
            WHERE media_level = '{media_level}'
            """
        )
        self.execute(f"CREATE INDEX IF NOT EXISTS idx_{view_name}_date ON public.{view_name}(report_date)")
        self.execute(f"CREATE INDEX IF NOT EXISTS idx_{view_name}_campaign ON public.{view_name}(campaign_id)")
        self.execute(f"CREATE INDEX IF NOT EXISTS idx_{view_name}_format ON public.{view_name}(campaign_format)")

    def rebuild_views(self) -> None:
        self.create_level_view("mv_wb_media_adv_campaign_daily", "campaign")
        self.create_level_view("mv_wb_media_adv_group_daily", "group")
        self.create_level_view("mv_wb_media_adv_creative_daily", "creative")
        self.execute("CREATE INDEX IF NOT EXISTS idx_wb_media_adv_raw_source ON public.wb_media_adv_daily_raw(source_file, source_sheet, source_row_num)")
        self.execute("CREATE INDEX IF NOT EXISTS idx_wb_media_adv_raw_date ON public.wb_media_adv_daily_raw(report_date)")
        self.execute("CREATE INDEX IF NOT EXISTS idx_wb_media_adv_raw_level ON public.wb_media_adv_daily_raw(media_level)")
        self.execute("ANALYZE public.wb_media_adv_daily_raw")
        for view in [
            "mv_wb_media_adv_campaign_daily",
            "mv_wb_media_adv_group_daily",
            "mv_wb_media_adv_creative_daily",
        ]:
            self.execute(f"ANALYZE public.{view}")
        self.conn.commit()

    def check(self) -> None:
        for view in [
            "mv_wb_media_adv_campaign_daily",
            "mv_wb_media_adv_group_daily",
            "mv_wb_media_adv_creative_daily",
        ]:
            self.execute(
                f"""
                SELECT
                    '{view}' AS view_name,
                    count(*) AS rows_count,
                    count(DISTINCT report_date) AS days_count,
                    min(report_date) AS date_from,
                    max(report_date) AS date_to,
                    count(DISTINCT campaign_id) AS campaigns,
                    round(coalesce(sum(expense_rub), 0), 2) AS expense_rub,
                    round(coalesce(sum(attributed_revenue_rub), 0), 2) AS attributed_revenue_rub
                FROM public.{view}
                """
            )
            print(dict(self.cur.fetchone()))

    def run(self) -> None:
        files: list[Path] = []
        if not self.skip_import:
            files = sorted(self.source_dir.rglob("*.xlsx"))
            if self.max_files:
                files = files[: self.max_files]
            if not files:
                raise RuntimeError(f"No XLSX files found in {self.source_dir}")
            print(f"Files found: {len(files)}")
        else:
            print("Skipping file import; rebuilding WB media adv materialized views from existing raw data.")
        self.connect()
        try:
            self.create_schema(drop_views=not self.skip_views)
            if not self.skip_import and not self.force_reimport:
                imported = self.imported_file_state()
                original_count = len(files)
                files = [file for file in files if not self.should_skip_file(file, imported)]
                print(f"Already imported skipped: {original_count - len(files)}")
            total = 0
            started = time.perf_counter()
            if not self.skip_import:
                for file_path in files:
                    rows_count = self.import_file(file_path)
                    total += rows_count
                    print(f"Imported {file_path.name}: {rows_count:,} rows")
            if not self.dry_run:
                if self.skip_views:
                    self.conn.commit()
                    print("Materialized views deferred; run view refresh after imports.")
                else:
                    print("Rebuilding WB media adv materialized views...")
                    self.rebuild_views()
                    self.check()
            print(f"Imported rows: {total:,} | seconds={time.perf_counter() - started:.1f}")
        finally:
            self.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--force-reimport", action="store_true")
    parser.add_argument("--max-files", type=int, default=None)
    parser.add_argument("--skip-import", action="store_true")
    parser.add_argument("--skip-views", action="store_true")
    args = parser.parse_args()
    WbMediaAdvImporter(
        args.source_dir,
        dry_run=args.dry_run,
        force_reimport=args.force_reimport,
        max_files=args.max_files,
        skip_import=args.skip_import,
        skip_views=args.skip_views,
    ).run()


if __name__ == "__main__":
    main()
