#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Import current Ozon stock workbook into PostgreSQL.

The workbook has four sheets:
- Товары: product-level current stock by SKU
- Товар-кластер: product stock by cluster
- Товар-склад: product stock by warehouse
- Кластеры: cluster-level summary
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import psycopg2
from openpyxl import load_workbook
from psycopg2.extras import execute_values


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DASHBOARD_ROOT = PROJECT_ROOT / "ozon_category_dashboard"
sys.path.insert(0, str(DASHBOARD_ROOT))

import app  # noqa: E402
import inventory_history  # noqa: E402
from build_ozon_abc_materialized_views import create_or_refresh_ozon_abc_materialized_views  # noqa: E402


DEFAULT_SOURCE = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Data\Ozon\Stock\Stock.xlsx"
)

DATA_START_ROW = 5
BATCH_SIZE = 5000


@dataclass(frozen=True)
class ImportResult:
    sheet_name: str
    table_name: str
    rows: int


PRODUCT_COLUMNS = [
    "article",
    "product_name",
    "sku",
    "product_sign",
    "placement_zone",
    "liquidity_status",
    "days_to_stockout",
    "avg_daily_sales_28d",
    "days_without_sales",
    "available_to_sell",
    "preparing_to_sell",
    "marked_awaiting_removal",
    "marked_awaiting_upd",
    "expiring",
    "defective_from_supply",
    "defective_from_stock",
    "surplus_from_supply",
    "checking",
    "in_supply_orders",
    "in_transit_supply",
    "returning_from_customers",
    "preparing_to_remove",
]

CLUSTER_COLUMNS = [
    "article",
    "product_name",
    "sku",
    "product_sign",
    "placement_zone",
    "cluster_name",
    "liquidity_status",
    "days_to_stockout",
    "avg_daily_sales_28d",
    "days_without_sales",
    "available_to_sell",
    "preparing_to_sell",
    "marked_awaiting_removal",
    "marked_awaiting_upd",
    "expiring",
    "defective_from_stock",
    "defective_from_supply",
    "surplus_from_supply",
    "checking",
    "in_supply_orders",
    "in_transit_supply",
    "returning_from_customers",
    "preparing_to_remove",
]

WAREHOUSE_COLUMNS = [
    "article",
    "product_name",
    "sku",
    "product_sign",
    "placement_zone",
    "cluster_name",
    "warehouse_name",
    "available_to_sell",
    "preparing_to_sell",
    "marked_awaiting_removal",
    "marked_awaiting_upd",
    "expiring",
    "defective_from_supply",
    "defective_from_stock",
    "surplus_from_supply",
    "checking",
    "in_supply_orders",
    "in_transit_supply",
    "returning_from_customers",
    "preparing_to_remove",
]

CLUSTER_SUMMARY_COLUMNS = [
    "cluster_name",
    "liquidity_status",
    "days_to_stockout",
    "avg_daily_sales_28d",
    "days_without_sales",
    "available_to_sell",
    "preparing_to_sell",
    "marked_awaiting_removal",
    "marked_awaiting_upd",
    "expiring",
    "defective_from_supply",
    "defective_from_stock",
    "surplus_from_supply",
    "checking",
    "in_supply_orders",
    "in_transit_supply",
    "returning_from_customers",
    "preparing_to_remove",
]

TEXT_COLUMNS = {
    "article",
    "product_name",
    "sku",
    "product_sign",
    "placement_zone",
    "cluster_name",
    "warehouse_name",
    "liquidity_status",
}

INT_COLUMNS = {
    "days_without_sales",
    "available_to_sell",
    "preparing_to_sell",
    "marked_awaiting_removal",
    "marked_awaiting_upd",
    "expiring",
    "defective_from_supply",
    "defective_from_stock",
    "surplus_from_supply",
    "checking",
    "in_supply_orders",
    "in_transit_supply",
    "returning_from_customers",
    "preparing_to_remove",
}

NUMERIC_COLUMNS = {
    "days_to_stockout",
    "avg_daily_sales_28d",
}

SHEETS = {
    "Товары": ("ozon_stock_products", PRODUCT_COLUMNS),
    "Товар-кластер": ("ozon_stock_product_clusters", CLUSTER_COLUMNS),
    "Товар-склад": ("ozon_stock_product_warehouses", WAREHOUSE_COLUMNS),
    "Кластеры": ("ozon_stock_clusters", CLUSTER_SUMMARY_COLUMNS),
}

def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).replace("\xa0", " ").strip()
    return text or None


def parse_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_text(value)
    if text is None:
        return None
    text = text.replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def parse_int(value: Any) -> int:
    number = parse_number(value)
    if number is None:
        return 0
    return int(round(number))


def convert_cell(column_name: str, value: Any) -> Any:
    if column_name in TEXT_COLUMNS:
        return normalize_text(value)
    if column_name in INT_COLUMNS:
        return parse_int(value)
    if column_name in NUMERIC_COLUMNS:
        return parse_number(value)
    return value


def iter_sheet_rows(ws: Any, columns: list[str]) -> Iterable[tuple[Any, ...]]:
    width = len(columns)
    for row_num, row in enumerate(ws.iter_rows(min_row=DATA_START_ROW, values_only=True), DATA_START_ROW):
        values = list(row[:width])
        if not any(value not in (None, "") for value in values):
            continue
        converted = [convert_cell(column_name, value) for column_name, value in zip(columns, values)]
        yield (row_num, *converted)


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def format_duration(seconds: float) -> str:
    seconds = int(seconds)
    minutes, sec = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours}h {minutes:02d}m {sec:02d}s"
    if minutes:
        return f"{minutes}m {sec:02d}s"
    return f"{sec}s"


def format_clock(seconds: float) -> str:
    seconds = max(0, int(seconds))
    minutes, sec = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours:02d}:{minutes:02d}:{sec:02d}"


def print_progress(
    current: int,
    total: int,
    item: str,
    started: float,
    rows_done: int | None = None,
    rows_total: int | None = None,
    extra: str | None = None,
) -> None:
    total = max(total, 1)
    current = max(0, min(current, total))
    pct = (current / total) * 100
    elapsed = time.monotonic() - started
    parts = [
        f"ПРОГРЕСС: {current}/{total} ({pct:.1f}%)",
        item,
    ]
    if rows_done is not None:
        rows_text = f"строки {rows_done:,}"
        if rows_total and rows_total > 0:
            rows_text += f" / {rows_total:,}"
        parts.append(rows_text)
    parts.append(f"прошло {format_clock(elapsed)}")
    if current > 0 and current < total:
        remaining = elapsed * (total - current) / current
        parts.append(f"ETA {format_clock(remaining)}")
    if extra:
        parts.append(extra)
    print(" | ".join(parts), flush=True)


def estimated_sheet_rows(ws: Any) -> int | None:
    max_row = getattr(ws, "max_row", None)
    if not max_row:
        return None
    return max(0, int(max_row) - DATA_START_ROW + 1)


def create_schema(cur: Any) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS public.ozon_stock_import_files (
            source_file text PRIMARY KEY,
            file_hash text NOT NULL,
            file_size_bytes bigint NOT NULL,
            modified_at timestamp without time zone NOT NULL,
            imported_at timestamp without time zone NOT NULL DEFAULT now(),
            product_rows integer NOT NULL DEFAULT 0,
            product_cluster_rows integer NOT NULL DEFAULT 0,
            product_warehouse_rows integer NOT NULL DEFAULT 0,
            cluster_rows integer NOT NULL DEFAULT 0
        )
        """
    )
    create_stock_table(cur, "ozon_stock_products", PRODUCT_COLUMNS)
    create_stock_table(cur, "ozon_stock_product_clusters", CLUSTER_COLUMNS)
    create_stock_table(cur, "ozon_stock_product_warehouses", WAREHOUSE_COLUMNS)
    create_stock_table(cur, "ozon_stock_clusters", CLUSTER_SUMMARY_COLUMNS)
    cur.execute(
        """
        CREATE OR REPLACE VIEW public.vw_ozon_current_stock_by_sku AS
        SELECT
            sku,
            max(article) AS article,
            max(product_name) AS product_name,
            max(placement_zone) AS placement_zone,
            max(liquidity_status) AS liquidity_status,
            sum(available_to_sell)::numeric AS available_to_sell,
            sum(preparing_to_sell)::numeric AS preparing_to_sell,
            sum(marked_awaiting_removal)::numeric AS marked_awaiting_removal,
            sum(marked_awaiting_upd)::numeric AS marked_awaiting_upd,
            sum(expiring)::numeric AS expiring,
            sum(defective_from_supply)::numeric AS defective_from_supply,
            sum(defective_from_stock)::numeric AS defective_from_stock,
            sum(surplus_from_supply)::numeric AS surplus_from_supply,
            sum(checking)::numeric AS checking,
            sum(in_supply_orders)::numeric AS in_supply_orders,
            sum(in_transit_supply)::numeric AS in_transit_supply,
            sum(returning_from_customers)::numeric AS returning_from_customers,
            sum(preparing_to_remove)::numeric AS preparing_to_remove,
            max(avg_daily_sales_28d) AS avg_daily_sales_28d,
            max(days_without_sales) AS days_without_sales,
            max(days_to_stockout) AS days_to_stockout,
            max(imported_at) AS imported_at
        FROM public.ozon_stock_products
        WHERE sku IS NOT NULL
        GROUP BY sku
        """
    )
    rebuild_ozon_views(cur)


def create_stock_table(cur: Any, table_name: str, columns: list[str]) -> None:
    field_defs = [
        "id bigserial PRIMARY KEY",
        "source_file text NOT NULL",
        "source_sheet text NOT NULL",
        "source_row_num integer NOT NULL",
        "imported_at timestamp without time zone NOT NULL DEFAULT now()",
    ]
    for column_name in columns:
        if column_name in TEXT_COLUMNS:
            field_type = "text"
        elif column_name in INT_COLUMNS:
            field_type = "integer NOT NULL DEFAULT 0"
        elif column_name in NUMERIC_COLUMNS:
            field_type = "numeric"
        else:
            field_type = "text"
        field_defs.append(f"{column_name} {field_type}")
    cur.execute(f"CREATE TABLE IF NOT EXISTS public.{table_name} ({', '.join(field_defs)})")
    if "sku" in columns:
        cur.execute(f"CREATE INDEX IF NOT EXISTS idx_{table_name}_sku ON public.{table_name} (sku)")
    if "article" in columns:
        cur.execute(f"CREATE INDEX IF NOT EXISTS idx_{table_name}_article ON public.{table_name} (article)")
    if "cluster_name" in columns:
        cur.execute(f"CREATE INDEX IF NOT EXISTS idx_{table_name}_cluster ON public.{table_name} (cluster_name)")
    if "warehouse_name" in columns:
        cur.execute(f"CREATE INDEX IF NOT EXISTS idx_{table_name}_warehouse ON public.{table_name} (warehouse_name)")


def rebuild_ozon_views(cur: Any) -> None:
    cur.execute(
        """
        CREATE OR REPLACE VIEW public.vw_ozon_sku_sales_90d AS
        WITH sales AS (
            SELECT
                sku AS ozon_sku,
                max(kategoriya_3_urovnya) FILTER (WHERE kategoriya_3_urovnya IS NOT NULL AND kategoriya_3_urovnya <> '') AS ozon_category,
                max(kategoriya_1_urovnya) FILTER (WHERE kategoriya_1_urovnya IS NOT NULL AND kategoriya_1_urovnya <> '') AS ozon_category_1,
                max(kategoriya_2_urovnya) FILTER (WHERE kategoriya_2_urovnya IS NOT NULL AND kategoriya_2_urovnya <> '') AS ozon_category_2,
                max(artikul) FILTER (WHERE artikul IS NOT NULL AND artikul <> '') AS artikul_prodavtsa,
                max(tovary) FILTER (WHERE tovary IS NOT NULL AND tovary <> '') AS naimenovanie,
                max(brend) FILTER (WHERE brend IS NOT NULL AND brend <> '') AS brend,
                max(model) FILTER (WHERE model IS NOT NULL AND model <> '') AS model,
                max(shema_raboty) FILTER (WHERE shema_raboty IS NOT NULL AND shema_raboty <> '') AS shema_raboty,
                sum((voronka_prodazh_zakazano_tovarov)::numeric) AS zakazano_sht,
                sum((zakazano_na_summu)::numeric) AS zakazano_rub,
                sum((vykupleno_tovarov)::numeric) AS vykupleno_sht
            FROM public.ozon_stock_sales_90d
            WHERE voronka_prodazh_zakazano_tovarov ~ '^[0-9]+(\\.[0-9]+)?$'
              AND zakazano_na_summu ~ '^[0-9]+(\\.[0-9]+)?$'
              AND sku IS NOT NULL
            GROUP BY sku
        ),
        cat_products AS (
            SELECT
                sku,
                max(category_name) FILTER (WHERE category_name IS NOT NULL AND category_name <> '') AS category_name,
                max(artikul) FILTER (WHERE artikul IS NOT NULL AND artikul <> '') AS artikul,
                max(nazvanie_tovara) FILTER (WHERE nazvanie_tovara IS NOT NULL AND nazvanie_tovara <> '') AS nazvanie_tovara
            FROM public.ozon_cat_products
            WHERE sku IS NOT NULL
            GROUP BY sku
        ),
        products AS (
            SELECT
                sku,
                max(tip) FILTER (WHERE tip IS NOT NULL AND tip <> '') AS tip,
                max(artikul) FILTER (WHERE artikul IS NOT NULL AND artikul <> '') AS artikul,
                max(nazvanie_tovara) FILTER (WHERE nazvanie_tovara IS NOT NULL AND nazvanie_tovara <> '') AS nazvanie_tovara,
                max(brend) FILTER (WHERE brend IS NOT NULL AND brend <> '') AS brend
            FROM public.ozon_products
            WHERE sku IS NOT NULL
            GROUP BY sku
        ),
        all_skus AS (
            SELECT ozon_sku AS sku FROM sales
            UNION
            SELECT sku FROM public.vw_ozon_current_stock_by_sku
        )
        SELECT
            COALESCE(NULLIF(s.ozon_category, ''), cp.category_name, op.tip, 'Без категории') AS ozon_category,
            s.ozon_category_1,
            s.ozon_category_2,
            k.sku AS ozon_sku,
            COALESCE(s.artikul_prodavtsa, st.article, cp.artikul, op.artikul) AS artikul_prodavtsa,
            COALESCE(s.naimenovanie, st.product_name, cp.nazvanie_tovara, op.nazvanie_tovara) AS naimenovanie,
            COALESCE(s.brend, op.brend) AS brend,
            s.model,
            s.shema_raboty,
            COALESCE(s.zakazano_sht, 0::numeric) AS zakazano_sht,
            COALESCE(s.zakazano_rub, 0::numeric) AS zakazano_rub,
            COALESCE(s.vykupleno_sht, 0::numeric) AS vykupleno_sht,
            COALESCE(st.available_to_sell, 0::numeric) AS ostatok_na_konets,
            CASE
                WHEN COALESCE(s.zakazano_sht, 0::numeric) > 0
                THEN round((s.vykupleno_sht / s.zakazano_sht) * 100, 2)
                ELSE 0::numeric
            END AS vykup_pct
        FROM all_skus k
        LEFT JOIN sales s
            ON s.ozon_sku = k.sku
        LEFT JOIN public.vw_ozon_current_stock_by_sku st
            ON st.sku = k.sku
        LEFT JOIN cat_products cp
            ON cp.sku = k.sku
        LEFT JOIN products op
            ON op.sku = k.sku
        ORDER BY COALESCE(s.zakazano_rub, 0::numeric) DESC
        """
    )
    cur.execute(
        """
        CREATE OR REPLACE VIEW public.vw_ozon_category_stock_sku_attribute_stats AS
        WITH sales AS (
            SELECT
                TRIM(BOTH FROM ozon_category) AS category_name,
                count(DISTINCT ozon_sku) AS sku_count,
                sum(ostatok_na_konets) AS total_stock_qty,
                sum(zakazano_sht) AS zakazano_sht,
                sum(zakazano_rub) AS zakazano_rub,
                sum(vykupleno_sht) AS vykupleno_sht
            FROM public.vw_ozon_sku_sales_90d
            WHERE ozon_category IS NOT NULL
            GROUP BY TRIM(BOTH FROM ozon_category)
        ),
        attrs AS (
            SELECT
                TRIM(BOTH FROM category) AS category_name,
                count(DISTINCT attribute_name) AS category_attribute_count
            FROM public.ozon_attribute_definitions
            GROUP BY TRIM(BOTH FROM category)
        ),
        base AS (
            SELECT
                s.category_name,
                COALESCE(s.total_stock_qty, 0::numeric) AS total_stock_qty,
                COALESCE(s.sku_count, 0::bigint) AS sku_count,
                COALESCE(a.category_attribute_count, 0::bigint) AS category_attribute_count,
                COALESCE(s.zakazano_sht, 0::numeric) AS zakazano_sht,
                COALESCE(s.zakazano_rub, 0::numeric) AS zakazano_rub,
                COALESCE(s.vykupleno_sht, 0::numeric) AS vykupleno_sht,
                CASE
                    WHEN COALESCE(s.zakazano_sht, 0::numeric) > 0
                    THEN round((s.vykupleno_sht / s.zakazano_sht) * 100, 2)
                    ELSE 0::numeric
                END AS vykup_pct_sht,
                CASE
                    WHEN COALESCE(s.zakazano_rub, 0::numeric) > 0
                     AND COALESCE(s.zakazano_sht, 0::numeric) > 0
                    THEN round(((s.vykupleno_sht * (s.zakazano_rub / NULLIF(s.zakazano_sht, 0::numeric))) / s.zakazano_rub) * 100, 2)
                    ELSE 0::numeric
                END AS vykup_pct_rub
            FROM sales s
            LEFT JOIN attrs a ON a.category_name = s.category_name
        ),
        totals AS (
            SELECT
                sum(zakazano_rub) AS total_rub,
                sum(total_stock_qty) AS total_stock
            FROM base
        ),
        ranked AS (
            SELECT
                b.category_name,
                b.total_stock_qty,
                b.sku_count,
                b.category_attribute_count,
                b.zakazano_sht,
                b.zakazano_rub,
                b.vykupleno_sht,
                b.vykup_pct_sht,
                b.vykup_pct_rub,
                round((b.zakazano_rub / NULLIF(t.total_rub, 0::numeric)) * 100, 2) AS orders_share_pct,
                sum(round((b.zakazano_rub / NULLIF(t.total_rub, 0::numeric)) * 100, 2)) OVER (ORDER BY b.zakazano_rub DESC) AS orders_cumulative_pct,
                round((b.zakazano_rub / NULLIF(t.total_rub, 0::numeric)) * 100, 2) AS sales_share_pct,
                sum(round((b.zakazano_rub / NULLIF(t.total_rub, 0::numeric)) * 100, 2)) OVER (ORDER BY b.zakazano_rub DESC) AS sales_cumulative_pct,
                round((b.total_stock_qty / NULLIF(t.total_stock, 0::numeric)) * 100, 2) AS stock_share_pct,
                sum(round((b.total_stock_qty / NULLIF(t.total_stock, 0::numeric)) * 100, 2)) OVER (ORDER BY b.total_stock_qty DESC) AS stock_cumulative_pct
            FROM base b
            CROSS JOIN totals t
        )
        SELECT
            category_name,
            total_stock_qty,
            sku_count,
            category_attribute_count,
            zakazano_sht,
            zakazano_rub,
            vykupleno_sht,
            vykup_pct_sht,
            vykup_pct_rub,
            COALESCE(orders_share_pct, 0::numeric) AS orders_share_pct,
            COALESCE(orders_cumulative_pct, 0::numeric) AS orders_cumulative_pct,
            CASE
                WHEN COALESCE(zakazano_rub, 0::numeric) = 0 THEN 'Без ABC'
                WHEN orders_cumulative_pct <= 80 THEN 'A'
                WHEN orders_cumulative_pct <= 95 THEN 'B'
                ELSE 'C'
            END AS abc_orders,
            COALESCE(sales_share_pct, 0::numeric) AS sales_share_pct,
            COALESCE(sales_cumulative_pct, 0::numeric) AS sales_cumulative_pct,
            CASE
                WHEN COALESCE(zakazano_rub, 0::numeric) = 0 THEN 'Без ABC'
                WHEN sales_cumulative_pct <= 80 THEN 'A'
                WHEN sales_cumulative_pct <= 95 THEN 'B'
                ELSE 'C'
            END AS abc_sales,
            COALESCE(stock_share_pct, 0::numeric) AS stock_share_pct,
            COALESCE(stock_cumulative_pct, 0::numeric) AS stock_cumulative_pct,
            CASE
                WHEN COALESCE(total_stock_qty, 0::numeric) = 0 THEN 'Без ABC'
                WHEN stock_cumulative_pct <= 80 THEN 'A'
                WHEN stock_cumulative_pct <= 95 THEN 'B'
                ELSE 'C'
            END AS abc_stock,
            concat(
                CASE
                    WHEN COALESCE(zakazano_rub, 0::numeric) = 0 THEN 'Без ABC'
                    WHEN orders_cumulative_pct <= 80 THEN 'A'
                    WHEN orders_cumulative_pct <= 95 THEN 'B'
                    ELSE 'C'
                END,
                CASE
                    WHEN COALESCE(zakazano_rub, 0::numeric) = 0 THEN 'Без ABC'
                    WHEN sales_cumulative_pct <= 80 THEN 'A'
                    WHEN sales_cumulative_pct <= 95 THEN 'B'
                    ELSE 'C'
                END,
                CASE
                    WHEN COALESCE(total_stock_qty, 0::numeric) = 0 THEN 'Без ABC'
                    WHEN stock_cumulative_pct <= 80 THEN 'A'
                    WHEN stock_cumulative_pct <= 95 THEN 'B'
                    ELSE 'C'
                END
            ) AS abc_combined
        FROM ranked
        """
    )


def truncate_stock_tables(cur: Any) -> None:
    cur.execute(
        """
        TRUNCATE TABLE
            public.ozon_stock_products,
            public.ozon_stock_product_clusters,
            public.ozon_stock_product_warehouses,
            public.ozon_stock_clusters
        RESTART IDENTITY
        """
    )


def already_imported(cur: Any, source: Path, digest: str) -> bool:
    cur.execute(
        """
        SELECT 1
        FROM public.ozon_stock_import_files
        WHERE source_file = %s
          AND file_hash = %s
          AND file_size_bytes = %s
        """,
        (str(source), digest, source.stat().st_size),
    )
    return cur.fetchone() is not None


def import_sheet(
    cur: Any,
    ws: Any,
    table_name: str,
    columns: list[str],
    source_file: str,
    imported_at: datetime,
    progress_started: float,
    progress_step: int,
    progress_total: int,
    estimated_rows: int | None = None,
) -> int:
    started = time.monotonic()
    rows_inserted = 0
    insert_columns = ["source_file", "source_sheet", "source_row_num", "imported_at", *columns]
    template = "(" + ", ".join(["%s"] * len(insert_columns)) + ")"
    sql = f"INSERT INTO public.{table_name} ({', '.join(insert_columns)}) VALUES %s"
    batch = []
    for row in iter_sheet_rows(ws, columns):
        batch.append((source_file, ws.title, row[0], imported_at, *row[1:]))
        if len(batch) >= BATCH_SIZE:
            execute_values(cur, sql, batch, template=template)
            rows_inserted += len(batch)
            elapsed = time.monotonic() - started
            speed = rows_inserted / elapsed if elapsed > 0 else 0
            print_progress(
                progress_step,
                progress_total,
                f"лист {ws.title}",
                progress_started,
                rows_done=rows_inserted,
                rows_total=estimated_rows,
                extra=f"speed={speed:,.0f} rows/s",
            )
            batch.clear()
    if batch:
        execute_values(cur, sql, batch, template=template)
        rows_inserted += len(batch)
        elapsed = time.monotonic() - started
        speed = rows_inserted / elapsed if elapsed > 0 else 0
        print_progress(
            progress_step,
            progress_total,
            f"лист {ws.title}",
            progress_started,
            rows_done=rows_inserted,
            rows_total=estimated_rows,
            extra=f"speed={speed:,.0f} rows/s",
        )
    return rows_inserted


def refresh_materialized_views(cur: Any, started: float, progress_step: int, progress_total: int) -> int:
    print_progress(
        progress_step,
        progress_total,
        "обновление Ozon ABC materialized views",
        started,
        extra="может занять несколько минут",
    )
    create_or_refresh_ozon_abc_materialized_views(cur)
    print_progress(progress_step, progress_total, "Ozon ABC materialized views готовы", started)
    progress_step += 1
    print_progress(
        progress_step,
        progress_total,
        "обновление витрины остатков по категориям",
        started,
        extra="может занять несколько минут",
    )
    ensure_ozon_category_stock_materialized_view(cur)
    print_progress(progress_step, progress_total, "витрина остатков по категориям готова", started)
    return progress_step + 1


def ensure_ozon_category_stock_materialized_view(cur: Any) -> None:
    cur.execute("SELECT to_regclass('public.vw_ozon_category_stock_sku_attribute_stats') AS view_name")
    if not cur.fetchone()["view_name"]:
        return
    cur.execute("DROP MATERIALIZED VIEW IF EXISTS public.mv_ozon_category_stock_sku_attribute_stats")
    cur.execute(
        """
        CREATE MATERIALIZED VIEW public.mv_ozon_category_stock_sku_attribute_stats AS
        SELECT * FROM public.vw_ozon_category_stock_sku_attribute_stats
        """
    )
    for column in ("category_name", "total_stock_qty", "zakazano_rub", "abc_combined"):
        cur.execute(
            f"CREATE INDEX IF NOT EXISTS idx_mv_ozon_category_stock_{column} "
            f"ON public.mv_ozon_category_stock_sku_attribute_stats ({column})"
        )
    cur.execute("ANALYZE public.mv_ozon_category_stock_sku_attribute_stats")


def upsert_import_file(cur: Any, source: Path, digest: str, imported_at: datetime, results: list[ImportResult]) -> None:
    counts = {result.table_name: result.rows for result in results}
    stat = source.stat()
    cur.execute(
        """
        INSERT INTO public.ozon_stock_import_files (
            source_file,
            file_hash,
            file_size_bytes,
            modified_at,
            imported_at,
            product_rows,
            product_cluster_rows,
            product_warehouse_rows,
            cluster_rows
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (source_file) DO UPDATE SET
            file_hash = EXCLUDED.file_hash,
            file_size_bytes = EXCLUDED.file_size_bytes,
            modified_at = EXCLUDED.modified_at,
            imported_at = EXCLUDED.imported_at,
            product_rows = EXCLUDED.product_rows,
            product_cluster_rows = EXCLUDED.product_cluster_rows,
            product_warehouse_rows = EXCLUDED.product_warehouse_rows,
            cluster_rows = EXCLUDED.cluster_rows
        """,
        (
            str(source),
            digest,
            stat.st_size,
            datetime.fromtimestamp(stat.st_mtime),
            imported_at,
            counts.get("ozon_stock_products", 0),
            counts.get("ozon_stock_product_clusters", 0),
            counts.get("ozon_stock_product_warehouses", 0),
            counts.get("ozon_stock_clusters", 0),
        ),
    )



def refresh_views_only(refresh: bool) -> None:
    started = time.monotonic()
    progress_total = 2 + (2 if refresh else 0)
    print(
        "ПЛАН: импорт Ozon stock пропущен | "
        f"обновление materialized views: {'да' if refresh else 'нет'} | без API-пауз",
        flush=True,
    )
    progress_step = 1
    with app.get_conn() as conn:
        with conn.cursor() as cur:
            print_progress(
                progress_step,
                progress_total,
                "подготовка схемы и Ozon stock/category regular views",
                started,
            )
            create_schema(cur)
            if refresh:
                progress_step += 1
                progress_step = refresh_materialized_views(cur, started, progress_step, progress_total)
            else:
                progress_step += 1
            print_progress(progress_step, progress_total, "commit", started, rows_done=0)
        conn.commit()
    elapsed = time.monotonic() - started
    print(f"Imported rows: 0 | errors: 0 | total time: {format_clock(elapsed)}")
    print(
        "ИТОГ: импорт Ozon stock пропущен | строк 0 | ошибок 0 | "
        f"витрины {'обновлены' if refresh else 'regular views обновлены, materialized views пропущены'} | "
        f"stopped no | partial no | elapsed {format_clock(elapsed)}"
    )
def import_workbook(source: Path, refresh: bool, force: bool = False) -> list[ImportResult]:
    if not source.exists():
        raise FileNotFoundError(source)
    started = time.monotonic()
    progress_total = 6 + len(SHEETS) + (2 if refresh else 0)
    print(
        "ПЛАН: 1 файл | "
        f"{len(SHEETS)} листа | батч {BATCH_SIZE:,} строк | "
        f"обновление витрин: {'да' if refresh else 'нет'} | без API-пауз",
        flush=True,
    )
    print(f"ПЛАН: источник {source} | размер {source.stat().st_size:,} bytes", flush=True)
    progress_step = 1
    print_progress(progress_step, progress_total, "расчет хеша файла", started)
    digest = file_hash(source)
    imported_at = datetime.now()
    results: list[ImportResult] = []
    with app.get_conn() as conn:
        with conn.cursor() as cur:
            progress_step += 1
            print_progress(progress_step, progress_total, "подготовка схемы БД", started)
            create_schema(cur)
            if not force and already_imported(cur, source, digest):
                elapsed = time.monotonic() - started
                print("File is unchanged and already imported. Use --force to reload it anyway.")
                print(f"Imported rows: 0 | errors: 0 | total time: {format_clock(elapsed)}")
                print("ИТОГ: файл без изменений | строк 0 | ошибок 0 | stopped no | partial no")
                return []
            progress_step += 1
            print_progress(progress_step, progress_total, "чтение Excel workbook", started)
            wb = load_workbook(source, read_only=True, data_only=True)
            sheet_estimates = {sheet_name: estimated_sheet_rows(wb[sheet_name]) for sheet_name in wb.sheetnames}
            progress_step += 1
            print_progress(progress_step, progress_total, "очистка таблиц остатков", started)
            truncate_stock_tables(cur)
            for sheet_name, (table_name, columns) in SHEETS.items():
                if sheet_name not in wb.sheetnames:
                    raise ValueError(f"Sheet not found: {sheet_name}")
                progress_step += 1
                estimated_rows = sheet_estimates.get(sheet_name)
                print_progress(
                    progress_step,
                    progress_total,
                    f"импорт листа {sheet_name} -> {table_name}",
                    started,
                    rows_done=0,
                    rows_total=estimated_rows,
                )
                rows = import_sheet(
                    cur,
                    wb[sheet_name],
                    table_name,
                    columns,
                    str(source),
                    imported_at,
                    started,
                    progress_step,
                    progress_total,
                    estimated_rows,
                )
                results.append(ImportResult(sheet_name, table_name, rows))
                print_progress(
                    progress_step,
                    progress_total,
                    f"лист {sheet_name} импортирован",
                    started,
                    rows_done=rows,
                    rows_total=estimated_rows,
                )
                print(f"[{len(results)}/{len(SHEETS)}] {sheet_name}: imported {rows:,} rows, errors 0", flush=True)
            progress_step += 1
            total_rows = sum(result.rows for result in results)
            print_progress(
                progress_step,
                progress_total,
                "запись метаданных и ANALYZE",
                started,
                rows_done=total_rows,
            )
            upsert_import_file(cur, source, digest, imported_at, results)
            for table_name, _columns in SHEETS.values():
                cur.execute(f"ANALYZE public.{table_name}")
            cur.execute("ANALYZE public.ozon_stock_import_files")
            inventory_history.store_ozon_current_tables_snapshot(
                conn,
                datetime.fromtimestamp(source.stat().st_mtime).date(),
                str(source),
            )
            if refresh:
                progress_step += 1
                progress_step = refresh_materialized_views(cur, started, progress_step, progress_total)
            else:
                progress_step += 1
            print_progress(progress_step, progress_total, "commit", started, rows_done=sum(result.rows for result in results))
        conn.commit()
    elapsed = time.monotonic() - started
    total_rows = sum(result.rows for result in results)
    print(f"Imported rows: {total_rows:,} | errors: 0 | total time: {format_clock(elapsed)}")
    print(
        "ИТОГ: "
        f"файл {source.name} | листов {len(results)}/{len(SHEETS)} | "
        f"строк {total_rows:,} | ошибок 0 | "
        f"витрины {'обновлены' if refresh else 'пропущены'} | stopped no | partial no | "
        f"elapsed {format_clock(elapsed)}"
    )
    return results


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import Ozon Stock.xlsx into PostgreSQL")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--skip-import", action="store_true", help="Do not read Stock.xlsx; rebuild Ozon stock/category views only")
    parser.add_argument("--no-refresh", action="store_true", help="Do not refresh dependent Ozon materialized views")
    parser.add_argument("--force", action="store_true", help="Reload the workbook even if the file hash is unchanged")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.skip_import:
        refresh_views_only(refresh=not args.no_refresh)
        print("Summary:")
        print("- import skipped; Ozon stock/category views refreshed")
        return
    results = import_workbook(args.source, refresh=not args.no_refresh, force=args.force)
    print("Summary:")
    if results:
        for result in results:
            print(f"- {result.sheet_name}: {result.rows:,} rows -> public.{result.table_name}")
    else:
        print("- no rows imported")

if __name__ == "__main__":
    main()
