from __future__ import annotations

import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl
import psycopg2
from openpyxl.utils import get_column_letter
from psycopg2.extras import execute_values


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "ozon_category_dashboard"))

import app  # noqa: E402


SEO_ROOT = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\SEO\Готовые SEO"
)
TAG = "May_2026"
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "seo_may_2026"
OUTPUT_XLSX = OUTPUT_DIR / "gloria_jeans_seo_sku_may_2026.xlsx"


def normalized_header(value: object) -> str:
    return str(value or "").strip().lower()


def clean_sku(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def iter_source_files(marketplace: str) -> list[Path]:
    source_dir = SEO_ROOT / marketplace.upper()
    files = []
    for path in sorted(source_dir.rglob("*.xlsx")):
        if any(part.upper() == "BACKUP" for part in path.parts):
            continue
        if path.name.startswith("~$"):
            continue
        files.append(path)
    return files


def extract_marketplace_rows(marketplace: str) -> list[dict[str, str]]:
    rows_by_sku: OrderedDict[str, dict[str, str]] = OrderedDict()
    preferred_title_headers = (
        ["title"]
        if marketplace == "wb"
        else ["original feed title", "title", "seo name"]
    )
    for path in iter_source_files(marketplace):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {normalized_header(value): index for index, value in enumerate(header_values)}
        sku_index = headers.get("sku")
        if sku_index is None:
            continue
        title_index = next(
            (headers[name] for name in preferred_title_headers if name in headers),
            None,
        )
        for raw_row in sheet.iter_rows(min_row=2, values_only=True):
            sku = clean_sku(raw_row[sku_index] if sku_index < len(raw_row) else "")
            if not sku:
                continue
            product_name = ""
            if title_index is not None and title_index < len(raw_row):
                product_name = str(raw_row[title_index] or "").strip()
            item = rows_by_sku.setdefault(
                sku,
                {
                    "marketplace": marketplace,
                    "sku": sku,
                    "product_name": product_name,
                    "tag": TAG,
                    "source_files": path.name,
                },
            )
            if not item["product_name"] and product_name:
                item["product_name"] = product_name
            if path.name not in item["source_files"].split("; "):
                item["source_files"] = f"{item['source_files']}; {path.name}"
    return list(rows_by_sku.values())


def ensure_tag_table(cur) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS public.marketplace_sku_tags (
            marketplace text NOT NULL,
            sku text NOT NULL,
            tag text NOT NULL,
            product_name text,
            source_file text,
            tagged_at timestamp without time zone NOT NULL DEFAULT now(),
            PRIMARY KEY (marketplace, sku, tag)
        )
        """
    )
    cur.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_marketplace_sku_tags_tag
        ON public.marketplace_sku_tags (marketplace, tag)
        """
    )
    cur.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_marketplace_sku_tags_sku
        ON public.marketplace_sku_tags (marketplace, sku)
        """
    )


def upsert_tags(rows: list[dict[str, str]]) -> None:
    values = [
        (
            row["marketplace"],
            row["sku"],
            row["tag"],
            row["product_name"],
            row["source_files"],
        )
        for row in rows
    ]
    with psycopg2.connect(**app.read_db_config()) as conn, conn.cursor() as cur:
        ensure_tag_table(cur)
        execute_values(
            cur,
            """
            INSERT INTO public.marketplace_sku_tags
                (marketplace, sku, tag, product_name, source_file)
            VALUES %s
            ON CONFLICT (marketplace, sku, tag) DO UPDATE SET
                product_name = EXCLUDED.product_name,
                source_file = EXCLUDED.source_file,
                tagged_at = now()
            """,
            values,
        )


def write_report(rows_by_marketplace: dict[str, list[dict[str, str]]]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for marketplace, rows in rows_by_marketplace.items():
        sheet = workbook.create_sheet(marketplace.upper())
        sheet.append(["Артикул", "Наименование товара", "Тег", "Источник"])
        for row in rows:
            sheet.append([row["sku"], row["product_name"], row["tag"], row["source_files"]])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = [18, 72, 14, 56]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    summary = workbook.create_sheet("Summary", 0)
    summary.append(["Маркетплейс", "Уникальных SKU", "Тег"])
    for marketplace, rows in rows_by_marketplace.items():
        summary.append([marketplace.upper(), len(rows), TAG])
    summary.freeze_panes = "A2"
    for index, width in enumerate([18, 18, 14], start=1):
        summary.column_dimensions[get_column_letter(index)].width = width
    workbook.save(OUTPUT_XLSX)


def main() -> int:
    rows_by_marketplace = {
        "wb": extract_marketplace_rows("wb"),
        "ozon": extract_marketplace_rows("ozon"),
    }
    all_rows = [row for rows in rows_by_marketplace.values() for row in rows]
    if not all_rows:
        raise RuntimeError(f"No SEO rows found in {SEO_ROOT}")
    upsert_tags(all_rows)
    write_report(rows_by_marketplace)
    for marketplace, rows in rows_by_marketplace.items():
        print(f"{marketplace.upper()}: {len(rows)} unique SKU tagged as {TAG}")
    print(f"Report: {OUTPUT_XLSX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
