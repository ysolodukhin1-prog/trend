from __future__ import annotations

import shutil
import sys
from copy import copy
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
import psycopg2


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "ozon_category_dashboard"))

import app  # noqa: E402


WORKBOOK_PATH = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\SEO\GJ лето 2026 МП тотал.xlsx"
)
BACKUP_DIR = PROJECT_ROOT / "outputs" / "seo_may_2026" / "backups"
SEO_STATUS_HEADER = "Статус Seo"


SHEET_CONFIG = {
    "озон": {"marketplace": "ozon", "sku_header": "Ozon ID"},
    "вб": {"marketplace": "wb", "sku_header": "Артикул WB"},
}


def clean_sku(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_tags() -> dict[tuple[str, str], list[str]]:
    rows: dict[tuple[str, str], list[str]] = defaultdict(list)
    with psycopg2.connect(**app.read_db_config()) as conn, conn.cursor() as cur:
        cur.execute(
            """
            SELECT marketplace, sku, tag
            FROM public.marketplace_sku_tags
            ORDER BY marketplace, sku, tag
            """
        )
        for marketplace, sku, tag in cur.fetchall():
            rows[(marketplace, sku)].append(tag)
    return rows


def header_map(sheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in sheet[1]
        if cell.value is not None and str(cell.value).strip()
    }


def get_or_create_status_column(sheet) -> int:
    headers = header_map(sheet)
    existing = headers.get(SEO_STATUS_HEADER)
    if existing:
        return existing
    column = sheet.max_column + 1
    sheet.cell(row=1, column=column, value=SEO_STATUS_HEADER)
    source = sheet.cell(row=1, column=sheet.max_column - 1)
    target = sheet.cell(row=1, column=column)
    if source.has_style:
        target._style = copy(source._style)
    return column


def make_backup() -> Path:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = BACKUP_DIR / f"{WORKBOOK_PATH.stem}_{timestamp}{WORKBOOK_PATH.suffix}"
    shutil.copy2(WORKBOOK_PATH, backup_path)
    return backup_path


def mark_workbook() -> dict[str, int | str]:
    tags = load_tags()
    workbook = openpyxl.load_workbook(WORKBOOK_PATH)
    stats: dict[str, int | str] = {}
    for sheet_name, config in SHEET_CONFIG.items():
        sheet = workbook[sheet_name]
        headers = header_map(sheet)
        sku_column = headers[config["sku_header"]]
        status_column = get_or_create_status_column(sheet)
        matched = 0
        for row_index in range(2, sheet.max_row + 1):
            sku = clean_sku(sheet.cell(row=row_index, column=sku_column).value)
            row_tags = tags.get((config["marketplace"], sku), [])
            status = ", ".join(row_tags)
            sheet.cell(row=row_index, column=status_column, value=status)
            if status:
                matched += 1
        sheet.column_dimensions[openpyxl.utils.get_column_letter(status_column)].width = 16
        stats[f"{sheet_name}_matched"] = matched
        stats[f"{sheet_name}_rows"] = sheet.max_row - 1
    backup_path = make_backup()
    workbook.save(WORKBOOK_PATH)
    stats["backup"] = str(backup_path)
    stats["workbook"] = str(WORKBOOK_PATH)
    return stats


def main() -> int:
    stats = mark_workbook()
    for key, value in stats.items():
        print(f"{key}: {value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
