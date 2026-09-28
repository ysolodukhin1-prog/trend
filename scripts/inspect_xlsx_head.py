#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("--rows", type=int, default=10)
    args = parser.parse_args()

    workbook = load_workbook(args.path, read_only=True, data_only=True)
    try:
        print("sheets:", workbook.sheetnames)
        for sheet_name in workbook.sheetnames:
            print(f"=== {sheet_name} ===")
            ws = workbook[sheet_name]
            for row in ws.iter_rows(min_row=1, max_row=args.rows, values_only=True):
                print(row)
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
