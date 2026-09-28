"""Primary-source adapters for Gloria Jeans marketplace plan/fact facts."""

from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook


DEFAULT_WB_REPORT = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Data\Wb\repport\report.xlsx"
)
DEFAULT_WB_FIX_DIR = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Data\Wb\Fix"
)
DEFAULT_OZON_FUNNEL_DIR = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Data\Ozon\Funel"
)
DEFAULT_OZON_SPEND_DIR = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Data\Ozon\Spend"
)


DailyRow = tuple[Any, ...]


def _number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".").strip()
    return float(text) if text else 0.0


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError(f"Cannot parse source date: {value!r}")


def _excel_date(value: Any) -> date:
    """Parse a date emitted either by openpyxl or directly from XLSX XML."""
    if isinstance(value, (int, float)) or re.fullmatch(r"\d+(?:\.\d+)?", str(value or "").strip()):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    return _date(value)


def read_wb_report(path: Path) -> dict[date, tuple[float, float]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    # The WB export has an incorrect A1 dimension despite containing thousands of rows.
    worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    headers = list(next(rows))
    order_index = headers.index(
        "Сумма заказов по розничным ценам с учётом согласованной скидки, руб."
    )
    sales_index = headers.index(
        "Сумма продаж по розничным ценам с учётом согласованной скидки, руб."
    )
    result: dict[date, tuple[float, float]] = {}
    for row in rows:
        if len(row) <= max(order_index, sales_index, 2) or row[2] in (None, ""):
            continue
        report_date = _date(row[2])
        result[report_date] = (_number(row[order_index]), _number(row[sales_index]))
    workbook.close()
    return result


def read_wb_fix(directory: Path, sales: dict[date, float] | None = None) -> dict[date, float]:
    """Actual ad expense per day; the actual turnover is collected into `sales` when given."""
    result: dict[date, float] = {}
    sales = sales if sales is not None else {}
    files = sorted(directory.glob("*.zip"), key=lambda item: (item.stat().st_mtime_ns, item.name))
    for archive_path in files:
        with zipfile.ZipFile(archive_path) as archive:
            members = [name for name in archive.namelist() if name.lower().endswith(".xlsx")]
            if not members:
                continue
            with archive.open(members[0]) as source:
                workbook = load_workbook(io.BytesIO(source.read()), read_only=True, data_only=True)
        worksheet = workbook["Детальная информация"]
        # Some WB Fix exports declare an empty worksheet dimension even though rows exist.
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        next(rows, None)
        headers = list(next(rows))
        year_index = headers.index("Год")
        period_index = headers.index("Период")
        expense_index = headers.index("Реклама. Фактические затраты")
        sales_index = headers.index("Продажи. Фактический оборот") if "Продажи. Фактический оборот" in headers else None
        for row in rows:
            if len(row) <= expense_index or row[year_index] in (None, "") or row[period_index] in (None, ""):
                continue
            report_date = datetime.strptime(
                f"{str(row[period_index]).strip()}.{int(row[year_index])}", "%d.%m.%Y"
            ).date()
            result[report_date] = _number(row[expense_index])
            if sales_index is not None and len(row) > sales_index and row[sales_index] not in (None, ""):
                sales[report_date] = _number(row[sales_index])
        workbook.close()
    return result


def build_wb_month_rows(
    wb_report: dict[date, tuple[float, float]],
    wb_fix: dict[date, float],
    month_start: date,
    *,
    wb_source: str = "primary://wb/report+fix",
) -> tuple[list[DailyRow], dict[str, Any]]:
    """Build one complete WB month from report orders/sales and Fix expense.

    The refresh is intentionally all-or-nothing. A missing Fix day must not turn
    into a zero expense or partially replace the month already stored in BI.
    """
    month_start = month_start.replace(day=1)
    next_month = date(
        month_start.year + (1 if month_start.month == 12 else 0),
        1 if month_start.month == 12 else month_start.month + 1,
        1,
    )
    report_dates = sorted(
        report_date for report_date in wb_report if month_start <= report_date < next_month
    )
    if not report_dates:
        raise RuntimeError(f"WB report has no rows for {month_start:%Y-%m}")
    missing_fix_dates = [report_date for report_date in report_dates if report_date not in wb_fix]
    if missing_fix_dates:
        preview = ", ".join(value.isoformat() for value in missing_fix_dates[:10])
        suffix = " ..." if len(missing_fix_dates) > 10 else ""
        raise RuntimeError(
            f"WB Fix does not cover {len(missing_fix_dates)} report days for "
            f"{month_start:%Y-%m}: {preview}{suffix}. Existing BI month was preserved."
        )

    rows = [
        (
            "wb",
            report_date,
            wb_report[report_date][0],
            wb_report[report_date][1],
            wb_fix[report_date],
            wb_source,
            "Primary",
            None,
        )
        for report_date in report_dates
    ]
    return rows, {
        "month_start": month_start,
        "date_from": report_dates[0],
        "date_to": report_dates[-1],
        "rows": len(rows),
        "orders_rub": round(sum(float(row[2]) for row in rows), 2),
        "sales_rub": round(sum(float(row[3]) for row in rows), 2),
        "ad_spend_rub": round(sum(float(row[4]) for row in rows), 2),
    }


def _xlsx_rows(path: Path):
    with zipfile.ZipFile(path) as archive, archive.open("xl/worksheets/sheet1.xml") as stream:
        for _, element in ET.iterparse(stream, events=("end",)):
            if not element.tag.endswith("}row"):
                continue
            row: list[str | None] = []
            for cell in list(element):
                values = [node.text for node in cell.iter() if node.text is not None]
                row.append(values[-1] if values else None)
            element.clear()
            yield row


# Ozon seller cabinet shows «Продажи и возвраты» as the sum of these accrual types.
# «Возврат выручки» arrives as a negative amount and nets the returns out, exactly as
# the cabinet does; «Возврат вознаграждения» is a commission refund and stays out.
OZON_REVENUE_TYPES = ("Выручка", "Возврат выручки", "Баллы за скидки", "Программы партнёров")


def read_ozon_spend(directory: Path) -> dict[date, float]:
    result: dict[date, float] = {}
    # Ozon exports may be cumulative (01.MM-N.MM). Process older files first so
    # the newest export becomes the authoritative snapshot for every included day.
    files = sorted(directory.glob("*/*.xlsx"), key=lambda item: (item.stat().st_mtime_ns, item.name))
    for path in files:
        rows = _xlsx_rows(path)
        period_row = next(rows, [])
        headers = list(next(rows, []))
        required_headers = {"Дата начисления", "Тип начисления", "Сумма итого, руб."}
        if not period_row or not required_headers.issubset(headers):
            continue
        date_index = headers.index("Дата начисления")
        type_index = headers.index("Тип начисления")
        amount_index = headers.index("Сумма итого, руб.")
        daily_revenue: dict[date, float] = {}
        for row in rows:
            if len(row) <= max(date_index, type_index, amount_index):
                continue
            if row[type_index] not in OZON_REVENUE_TYPES or row[date_index] in (None, ""):
                continue
            report_date = _excel_date(row[date_index])
            daily_revenue[report_date] = daily_revenue.get(report_date, 0.0) + _number(row[amount_index])
        for report_date, revenue in daily_revenue.items():
            result[report_date] = round(revenue, 2)
    return result


def _row_value(row: Any, key: str, position: int) -> Any:
    return row.get(key) if isinstance(row, dict) else row[position]


def read_ozon_db(connection_factory: Callable[[], Any]) -> tuple[dict[date, float], dict[date, float]]:
    orders: dict[date, float] = {}
    product_ads: dict[date, float] = {}
    media_ads: dict[date, float] = {}
    with connection_factory() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT report_date, sum(coalesce(ordered_amount_rub, 0)) AS amount
            FROM public.ozon_funnel_daily GROUP BY report_date ORDER BY report_date
            """
        )
        for row in cursor.fetchall():
            orders[_row_value(row, "report_date", 0)] = _number(_row_value(row, "amount", 1))

        cursor.execute(
            """
            SELECT report_date, sum(coalesce(expense_rub, 0)) AS amount
            FROM public.ozon_adv_daily_raw GROUP BY report_date ORDER BY report_date
            """
        )
        for row in cursor.fetchall():
            product_ads[_row_value(row, "report_date", 0)] = _number(_row_value(row, "amount", 1))

        cursor.execute("SELECT to_regclass('public.mv_ozon_media_adv_daily') AS relation_name")
        relation_row = cursor.fetchone()
        relation_name = _row_value(relation_row, "relation_name", 0)
        if relation_name:
            cursor.execute(
                """
                SELECT report_date, sum(coalesce(expense_rub, 0)) AS amount
                FROM public.mv_ozon_media_adv_daily GROUP BY report_date ORDER BY report_date
                """
            )
            for row in cursor.fetchall():
                media_ads[_row_value(row, "report_date", 0)] = _number(_row_value(row, "amount", 1))

    ad_spend = {
        report_date: round(product_ads.get(report_date, 0) + media_ads.get(report_date, 0), 2)
        for report_date in set(product_ads) | set(media_ads)
    }
    return orders, ad_spend


def build_ozon_month_rows(
    ozon_sales: dict[date, float],
    ozon_orders: dict[date, float],
    ozon_ad_spend: dict[date, float],
    month_start: date,
    *,
    source: str = "primary://ozon/funnel+spend+advertising-db",
) -> tuple[list[DailyRow], dict[str, Any]]:
    """Build Ozon rows only where all three daily primary sources agree.

    A missing seller settlement is source lag, not zero revenue.  Therefore the
    caller may upsert only the returned dates and retain earlier proven values.
    """
    month_start = month_start.replace(day=1)
    month_end = date(
        month_start.year + (1 if month_start.month == 12 else 0),
        1 if month_start.month == 12 else month_start.month + 1,
        1,
    )
    covered_dates = sorted(
        report_date
        for report_date in (set(ozon_sales) & set(ozon_orders) & set(ozon_ad_spend))
        if month_start <= report_date < month_end
    )
    rows: list[DailyRow] = [
        (
            "ozon",
            report_date,
            ozon_orders[report_date],
            ozon_sales[report_date],
            ozon_ad_spend[report_date],
            source,
            "Primary",
            None,
        )
        for report_date in covered_dates
    ]
    return rows, {
        "month_start": month_start,
        "rows": len(rows),
        "date_from": min(covered_dates) if covered_dates else None,
        "date_to": max(covered_dates) if covered_dates else None,
        "orders_rub": round(sum(row[2] for row in rows), 2),
        "sales_rub": round(sum(row[3] for row in rows), 2),
        "ad_spend_rub": round(sum(row[4] for row in rows), 2),
    }


def read_wb_db(connection_factory: Callable[[], Any]) -> dict[date, float]:
    """Daily WB ad spend from the database: product plus media advertising.

    The `Fix` archive holds the corrected fixed-rate expense but is exported per
    month or even per single day, so it cannot drive the daily report on its own.
    """
    product_ads: dict[date, float] = {}
    media_ads: dict[date, float] = {}
    with connection_factory() as connection, connection.cursor() as cursor:
        for table, target in (("wb_adv_daily_raw", product_ads), ("wb_media_adv_daily_raw", media_ads)):
            cursor.execute("SELECT to_regclass(%s) AS relation_name", (f"public.{table}",))
            if not _row_value(cursor.fetchone(), "relation_name", 0):
                continue
            cursor.execute(
                f"""
                SELECT report_date, sum(coalesce(expense_rub, 0)) AS amount
                FROM public.{table} GROUP BY report_date ORDER BY report_date
                """
            )
            for row in cursor.fetchall():
                target[_row_value(row, "report_date", 0)] = _number(_row_value(row, "amount", 1))
    return {
        report_date: round(product_ads.get(report_date, 0) + media_ads.get(report_date, 0), 2)
        for report_date in set(product_ads) | set(media_ads)
    }


def merge_primary_rows(
    legacy_rows: list[DailyRow],
    wb_report: dict[date, tuple[float, float]],
    wb_fix: dict[date, float],
    ozon_orders: dict[date, float],
    ozon_sales: dict[date, float],
    ozon_ad_spend: dict[date, float],
    wb_db_ad_spend: dict[date, float] | None = None,
    wb_fix_sales: dict[date, float] | None = None,
    *,
    wb_source: str = "primary://wb/report+fix",
    ozon_source: str = "primary://ozon/funnel+spend+advertising-db",
) -> tuple[list[DailyRow], dict[str, Any]]:
    merged = {(row[0], row[1]): list(row) for row in legacy_rows}

    # Every day present in the WB report becomes a row. Fixed-rate expense wins when the
    # `Fix` export covers that day, otherwise the daily spend from the database is used.
    wb_db_ad_spend = wb_db_ad_spend or {}
    wb_fix_sales = wb_fix_sales or {}
    wb_dates = sorted(wb_report)
    wb_fix_days = 0
    wb_db_days = 0
    for report_date in wb_dates:
        orders, sales = wb_report[report_date]
        if report_date in wb_fix:
            # Fix export supplies the actual ad expense; sales always come from the WB report.
            ad_spend = wb_fix[report_date]
            source_sheet = "Primary"
            wb_fix_days += 1
        else:
            ad_spend = wb_db_ad_spend.get(report_date, 0.0)
            source_sheet = "Primary+DbSpend"
            wb_db_days += 1
        merged[("wb", report_date)] = [
            "wb", report_date, orders, sales, ad_spend, wb_source, source_sheet, None
        ]

    ozon_cutoffs = [max(values) for values in (ozon_orders, ozon_sales, ozon_ad_spend) if values]
    ozon_cutoff = min(ozon_cutoffs) if len(ozon_cutoffs) == 3 else None
    ozon_dates = sorted(set(ozon_orders) & set(ozon_ad_spend))
    if ozon_cutoff:
        ozon_dates = [report_date for report_date in ozon_dates if report_date <= ozon_cutoff]
    for report_date in ozon_dates:
        previous = merged.get(("ozon", report_date))
        sales = ozon_sales.get(report_date)
        uses_legacy_sales = sales is None and previous is not None
        if uses_legacy_sales:
            sales = previous[3]
        if sales is None:
            continue
        source = ozon_source
        source_sheet = "Primary"
        if uses_legacy_sales:
            source = f"{ozon_source} | legacy-sales:{previous[5]}"
            source_sheet = "Primary+LegacySalesFallback"
        merged[("ozon", report_date)] = [
            "ozon", report_date, ozon_orders[report_date], sales, ozon_ad_spend[report_date],
            source, source_sheet, None,
        ]

    stats = {
        "wb_rows": len(wb_dates),
        "wb_fix_days": wb_fix_days,
        "wb_db_spend_days": wb_db_days,
        "wb_max_date": max(wb_dates) if wb_dates else None,
        "ozon_rows": len(ozon_dates),
        "ozon_max_date": max(ozon_dates) if ozon_dates else None,
        "ozon_spend_rows": len(ozon_sales),
    }
    return [tuple(row) for _, row in sorted(merged.items())], stats


def load_primary_rows(
    legacy_rows: list[DailyRow],
    connection_factory: Callable[[], Any],
    *,
    wb_report_path: Path = DEFAULT_WB_REPORT,
    wb_fix_dir: Path = DEFAULT_WB_FIX_DIR,
    ozon_spend_dir: Path = DEFAULT_OZON_SPEND_DIR,
) -> tuple[list[DailyRow], dict[str, Any]]:
    print("ПРОГРЕСС: 1/5 | WB report | читаю заказы и продажи", flush=True)
    wb_report = read_wb_report(wb_report_path)
    print("ПРОГРЕСС: 2/5 | WB Fix | читаю фактический оборот и расходы", flush=True)
    wb_fix_sales: dict[date, float] = {}
    wb_fix = read_wb_fix(wb_fix_dir, wb_fix_sales)
    print("ПРОГРЕСС: 3/5 | Ozon Spend | читаю выручку", flush=True)
    ozon_sales = read_ozon_spend(ozon_spend_dir)
    print("ПРОГРЕСС: 4/5 | Ozon DB | заказы + товарная и медийная реклама", flush=True)
    ozon_orders, ozon_ad_spend = read_ozon_db(connection_factory)
    print("ПРОГРЕСС: 5/5 | WB DB | товарная и медийная реклама по дням", flush=True)
    wb_db_ad_spend = read_wb_db(connection_factory)
    return merge_primary_rows(
        legacy_rows,
        wb_report,
        wb_fix,
        ozon_orders,
        ozon_sales,
        ozon_ad_spend,
        wb_db_ad_spend,
        wb_fix_sales,
        wb_source=f"{wb_report_path} | {wb_fix_dir}",
        ozon_source=f"{DEFAULT_OZON_FUNNEL_DIR} | {ozon_spend_dir} | db:ozon_adv+ozon_media_adv",
    )
