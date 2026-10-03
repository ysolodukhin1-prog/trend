#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Load KM Trade Ozon finance history from XLSX and Seller API."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import time
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import psycopg2
from openpyxl import load_workbook
from psycopg2.extras import Json, execute_batch, execute_values


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DASHBOARD_ROOT = PROJECT_ROOT / "ozon_category_dashboard"
TARGET_DB = os.environ.get("DASHBOARD_DB_NAME", "km_trade_products")
CLIENT_KEY = os.environ.get("DASHBOARD_CLIENT", "km_trade")
CLIENT_LABEL = os.environ.get("DASHBOARD_CLIENT_LABEL", CLIENT_KEY)
CLIENT_SUFFIX = re.sub(r"[^A-Z0-9]+", "_", CLIENT_KEY.upper()).strip("_")
API_BASE_URL = os.environ.get(
    "OZON_SELLER_API_BASE_URL", "https://api-seller.ozon.ru"
).rstrip("/")
FIN_DIR = Path(
    os.environ.get(f"OZON_FIN_DIR_{CLIENT_SUFFIX}")
    or os.environ.get("KM_OZON_FIN_DIR")
    or r"G:\Общие диски\Kokoc Marketplaces\Clients\KM Trade\Аналитика\Дашборды\Data\Ozon\Fin"
)
FINANCE_INTERVAL_SECONDS = float(os.environ.get("OZON_FINANCE_INTERVAL_SECONDS", "7"))
MAX_RETRIES = int(os.environ.get("OZON_API_MAX_RETRIES", "6"))
HTTP_TIMEOUT_SECONDS = int(os.environ.get("OZON_API_HTTP_TIMEOUT_SECONDS", "90"))
PRICE_PAGE_SIZE = 1000
ATTRIBUTES_PAGE_SIZE = 100
SOURCE_API = "ozon_api"
SOURCE_XLSX = "ozon_xlsx"


os.environ["PGDATABASE"] = TARGET_DB
os.environ["DASHBOARD_CLIENT"] = CLIENT_KEY
sys.path.insert(0, str(DASHBOARD_ROOT))

import app  # noqa: E402


class FinancePipelineError(RuntimeError):
    pass


def money(value: Any) -> Decimal:
    if isinstance(value, dict):
        value = value.get("amount")
    if value in (None, ""):
        return Decimal("0")
    text = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0")


def text(value: Any) -> str:
    if value is None:
        return ""
    result = str(value).strip()
    return result[1:].strip() if result.startswith("'") else result


def normalized_header(value: Any) -> str:
    return re.sub(r"\s+", " ", text(value)).lower()


def parse_day(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = text(value)
    if not raw:
        return None
    for candidate in (raw[:10], raw.replace(".", "-")[:10]):
        try:
            if re.fullmatch(r"\d{2}-\d{2}-\d{4}", candidate):
                return datetime.strptime(candidate, "%d-%m-%Y").date()
            return date.fromisoformat(candidate)
        except ValueError:
            continue
    return None


def stable_key(*parts: Any) -> str:
    raw = "\x1f".join(text(part) for part in parts)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def money_amount(value: Any) -> Decimal:
    return money(value.get("amount")) if isinstance(value, dict) else money(value)


def money_currency(value: Any) -> str:
    return text(value.get("currency")) if isinstance(value, dict) else "RUB"


def seller_credentials() -> tuple[str, str]:
    env_values = app.read_app_env_file()
    client_id = os.environ.get(
        f"OZON_SELLER_CLIENT_ID_{CLIENT_SUFFIX}"
    ) or env_values.get(f"OZON_SELLER_CLIENT_ID_{CLIENT_SUFFIX}")
    api_key = os.environ.get(
        f"OZON_SELLER_API_KEY_{CLIENT_SUFFIX}"
    ) or env_values.get(f"OZON_SELLER_API_KEY_{CLIENT_SUFFIX}")
    if (not client_id or not api_key) and CLIENT_KEY in app.ADMIN_CLIENTS:
        client_id, api_key = app.ozon_seo_credentials_value(CLIENT_KEY)
    if not client_id or not api_key:
        raise FinancePipelineError(
            f"Для {CLIENT_LABEL} не сохранены Seller API Client-Id/Api-Key"
        )
    return client_id, api_key


def connect_km():
    config = dict(app.read_db_config())
    config["database"] = TARGET_DB
    conn = psycopg2.connect(**config)
    conn.autocommit = False
    with conn.cursor() as cur:
        cur.execute("SELECT current_database()")
        actual = cur.fetchone()[0]
    if actual != TARGET_DB:
        conn.close()
        raise FinancePipelineError(
            f"Защита клиента: ожидалась БД {TARGET_DB}, открылась {actual}"
        )
    return conn


def ensure_schema(conn) -> None:
    with conn.cursor() as cur:
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS public.ozon_finance_accrual_types (
                type_id integer PRIMARY KEY,
                type_name text NOT NULL,
                description text,
                updated_at timestamp without time zone NOT NULL DEFAULT now()
            );

            CREATE TABLE IF NOT EXISTS public.ozon_finance_events (
                event_key text PRIMARY KEY,
                source_kind text NOT NULL,
                source_ref text NOT NULL,
                operation_date date NOT NULL,
                accrual_id bigint,
                unit_number text,
                posting_number text,
                accrued_category text,
                operation_type_code text,
                operation_type_name text,
                total_amount numeric NOT NULL DEFAULT 0,
                currency text NOT NULL DEFAULT 'RUB',
                raw_json jsonb,
                imported_at timestamp without time zone NOT NULL DEFAULT now()
            );

            CREATE TABLE IF NOT EXISTS public.ozon_finance_lines (
                line_key text PRIMARY KEY,
                event_key text NOT NULL,
                source_kind text NOT NULL,
                operation_date date NOT NULL,
                posting_number text,
                sku text,
                article text,
                product_name text,
                delivery_schema text,
                line_kind text NOT NULL,
                type_id integer,
                type_name text,
                amount numeric NOT NULL DEFAULT 0,
                currency text NOT NULL DEFAULT 'RUB',
                quantity numeric,
                commission_rate_pct numeric,
                localization_index numeric,
                delivery_time_hours numeric,
                allocation_scope text NOT NULL DEFAULT 'exact',
                raw_json jsonb,
                imported_at timestamp without time zone NOT NULL DEFAULT now()
            );

            CREATE TABLE IF NOT EXISTS public.ozon_product_price_snapshots (
                snapshot_date date NOT NULL,
                sku text NOT NULL,
                offer_id text,
                product_id text,
                currency text,
                price numeric,
                marketing_seller_price numeric,
                min_price numeric,
                old_price numeric,
                retail_price numeric,
                net_price numeric,
                vat_pct numeric,
                acquiring numeric,
                sales_percent_fbo numeric,
                sales_percent_fbs numeric,
                sales_percent_rfbs numeric,
                fbo_delivery_amount numeric,
                fbo_logistics_min numeric,
                fbo_logistics_max numeric,
                fbo_return_amount numeric,
                fbs_delivery_amount numeric,
                fbs_logistics_min numeric,
                fbs_logistics_max numeric,
                fbs_first_mile_min numeric,
                fbs_first_mile_max numeric,
                fbs_return_amount numeric,
                price_index_color text,
                ozon_price_index numeric,
                market_price_index numeric,
                self_price_index numeric,
                volume_weight numeric,
                depth numeric,
                width numeric,
                height numeric,
                dimension_unit text,
                weight numeric,
                weight_unit text,
                raw_json jsonb,
                imported_at timestamp without time zone NOT NULL DEFAULT now(),
                PRIMARY KEY (snapshot_date, sku)
            );

            CREATE INDEX IF NOT EXISTS idx_ozon_price_snapshots_sku_latest
                ON public.ozon_product_price_snapshots (sku, snapshot_date DESC);
            CREATE INDEX IF NOT EXISTS idx_ozon_price_snapshots_offer_latest
                ON public.ozon_product_price_snapshots (offer_id, snapshot_date DESC)
                WHERE offer_id IS NOT NULL;

            CREATE TABLE IF NOT EXISTS public.ozon_unit_product_settings (
                sku text PRIMARY KEY,
                cogs_per_unit numeric,
                fulfillment_per_unit numeric NOT NULL DEFAULT 0,
                inbound_per_unit numeric NOT NULL DEFAULT 0,
                crossdock_per_unit numeric NOT NULL DEFAULT 0,
                acceptance_per_unit numeric NOT NULL DEFAULT 0,
                storage_per_unit numeric NOT NULL DEFAULT 0,
                other_per_unit numeric NOT NULL DEFAULT 0,
                planned_price numeric,
                return_rate_pct numeric,
                target_margin_pct numeric,
                target_advertising_pct numeric,
                depth numeric,
                width numeric,
                height numeric,
                dimension_unit text,
                weight numeric,
                weight_unit text,
                notes text,
                updated_at timestamp without time zone NOT NULL DEFAULT now()
            );

            CREATE TABLE IF NOT EXISTS public.ozon_unit_global_settings (
                singleton boolean PRIMARY KEY DEFAULT true CHECK (singleton),
                tax_pct numeric,
                vat_pct numeric,
                acquiring_pct numeric,
                capital_days numeric,
                capital_rate_pct numeric,
                advertising_pct numeric,
                return_rate_pct numeric,
                target_margin_pct numeric,
                mrc_margin_pct numeric,
                rrc_margin_pct numeric,
                updated_at timestamp without time zone NOT NULL DEFAULT now()
            );

            ALTER TABLE public.ozon_product_price_snapshots
                ADD COLUMN IF NOT EXISTS depth numeric,
                ADD COLUMN IF NOT EXISTS width numeric,
                ADD COLUMN IF NOT EXISTS height numeric,
                ADD COLUMN IF NOT EXISTS dimension_unit text,
                ADD COLUMN IF NOT EXISTS weight numeric,
                ADD COLUMN IF NOT EXISTS weight_unit text;

            ALTER TABLE public.ozon_unit_product_settings
                ADD COLUMN IF NOT EXISTS planned_price numeric,
                ADD COLUMN IF NOT EXISTS return_rate_pct numeric,
                ADD COLUMN IF NOT EXISTS target_advertising_pct numeric,
                ADD COLUMN IF NOT EXISTS crossdock_per_unit numeric NOT NULL DEFAULT 0,
                ADD COLUMN IF NOT EXISTS acceptance_per_unit numeric NOT NULL DEFAULT 0,
                ADD COLUMN IF NOT EXISTS storage_per_unit numeric NOT NULL DEFAULT 0,
                ADD COLUMN IF NOT EXISTS depth numeric,
                ADD COLUMN IF NOT EXISTS width numeric,
                ADD COLUMN IF NOT EXISTS height numeric,
                ADD COLUMN IF NOT EXISTS dimension_unit text,
                ADD COLUMN IF NOT EXISTS weight numeric,
                ADD COLUMN IF NOT EXISTS weight_unit text;

            ALTER TABLE public.ozon_unit_global_settings
                ADD COLUMN IF NOT EXISTS return_rate_pct numeric,
                ADD COLUMN IF NOT EXISTS vat_pct numeric,
                ADD COLUMN IF NOT EXISTS acquiring_pct numeric,
                ADD COLUMN IF NOT EXISTS capital_days numeric,
                ADD COLUMN IF NOT EXISTS capital_rate_pct numeric,
                ADD COLUMN IF NOT EXISTS mrc_margin_pct numeric,
                ADD COLUMN IF NOT EXISTS rrc_margin_pct numeric;

            INSERT INTO public.ozon_unit_global_settings (
                singleton, tax_pct, vat_pct, acquiring_pct, capital_days,
                capital_rate_pct, advertising_pct, return_rate_pct, target_margin_pct,
                mrc_margin_pct, rrc_margin_pct
            ) VALUES (true, NULL, 0, NULL, 30, 20, NULL, NULL, 20, 0, 20)
            ON CONFLICT (singleton) DO NOTHING;

            UPDATE public.ozon_unit_global_settings
            SET mrc_margin_pct = coalesce(mrc_margin_pct, 0),
                rrc_margin_pct = coalesce(rrc_margin_pct, target_margin_pct, 20)
            WHERE mrc_margin_pct IS NULL OR rrc_margin_pct IS NULL;

            CREATE TABLE IF NOT EXISTS public.ozon_pl_expenses (
                expense_id bigserial PRIMARY KEY,
                date_from date NOT NULL,
                date_to date NOT NULL,
                expense_name text NOT NULL,
                amount numeric NOT NULL,
                allocation_method text NOT NULL DEFAULT 'period',
                notes text,
                updated_at timestamp without time zone NOT NULL DEFAULT now()
            );

            CREATE INDEX IF NOT EXISTS idx_ozon_finance_events_date
                ON public.ozon_finance_events(operation_date);
            CREATE INDEX IF NOT EXISTS idx_ozon_finance_events_source
                ON public.ozon_finance_events(source_kind, operation_date);
            CREATE INDEX IF NOT EXISTS idx_ozon_finance_lines_date
                ON public.ozon_finance_lines(operation_date);
            CREATE INDEX IF NOT EXISTS idx_ozon_finance_lines_sku
                ON public.ozon_finance_lines(sku, operation_date);
            CREATE INDEX IF NOT EXISTS idx_ozon_finance_lines_type
                ON public.ozon_finance_lines(line_kind, type_name);
            """
        )
    conn.commit()


class FinanceApiClient:
    def __init__(self) -> None:
        client_id, api_key = seller_credentials()
        self.headers = {
            "Client-Id": client_id,
            "Api-Key": api_key,
            "Accept": "application/json",
            "Content-Type": "application/json",
        }
        self.last_request_at = 0.0
        self.request_count = 0

    def _wait(self) -> None:
        remaining = FINANCE_INTERVAL_SECONDS - (time.monotonic() - self.last_request_at)
        while remaining > 0:
            print(
                f"ПРОГРЕСС: лимит Ozon Finance API | пауза {remaining:.1f} сек. | "
                f"запросов {self.request_count}",
                flush=True,
            )
            time.sleep(min(remaining, 10))
            remaining = FINANCE_INTERVAL_SECONDS - (time.monotonic() - self.last_request_at)

    def post(self, path: str, payload: dict[str, Any]) -> dict[str, Any]:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        for attempt in range(1, MAX_RETRIES + 1):
            self._wait()
            self.request_count += 1
            try:
                with urlopen(
                    Request(
                        API_BASE_URL + path,
                        data=body,
                        headers=self.headers,
                        method="POST",
                    ),
                    timeout=HTTP_TIMEOUT_SECONDS,
                ) as response:
                    raw = response.read()
                self.last_request_at = time.monotonic()
                parsed = json.loads(raw.decode("utf-8-sig"))
                if not isinstance(parsed, dict):
                    raise FinancePipelineError(f"{path} вернул не объект JSON")
                return parsed
            except HTTPError as exc:
                self.last_request_at = time.monotonic()
                retryable = exc.code == 429 or 500 <= exc.code < 600
                detail = exc.read().decode("utf-8", errors="replace")[:800]
                if not retryable or attempt >= MAX_RETRIES:
                    raise FinancePipelineError(
                        f"Ozon Finance API HTTP {exc.code} для {path}: {detail}"
                    ) from exc
                retry_after = exc.headers.get("Retry-After")
                try:
                    wait_seconds = max(float(retry_after or 0), 60 if exc.code == 429 else 2**attempt)
                except ValueError:
                    wait_seconds = 60 if exc.code == 429 else 2**attempt
                print(
                    f"ПРОГРЕСС: повтор {attempt}/{MAX_RETRIES} | {path} | "
                    f"HTTP {exc.code} | ожидание {wait_seconds:.0f} сек.",
                    flush=True,
                )
                time.sleep(min(wait_seconds, 120))
            except (URLError, TimeoutError, json.JSONDecodeError) as exc:
                self.last_request_at = time.monotonic()
                if attempt >= MAX_RETRIES:
                    raise FinancePipelineError(f"Ошибка запроса {path}: {exc}") from exc
                wait_seconds = min(2**attempt, 60)
                print(
                    f"ПРОГРЕСС: повтор {attempt}/{MAX_RETRIES} | {path} | "
                    f"сеть/JSON | ожидание {wait_seconds} сек.",
                    flush=True,
                )
                time.sleep(wait_seconds)
        raise FinancePipelineError(f"Исчерпаны повторы {path}")


MANUAL_COMPONENTS = [
    ("Вознаграждение Ozon", "commission"),
    ("Сборка заказа", "fulfillment_ozon"),
    ("Обработка отправления (Drop-off/Pick-up) (разбивается по товарам пропорционально количеству в отправлении)", "logistics"),
    ("Магистраль", "logistics"),
    ("Последняя миля (разбивается по товарам пропорционально доле цены товара в сумме отправления)", "last_mile"),
    ("Обратная магистраль", "reverse_logistics"),
    ("Обработка возврата", "reverse_logistics"),
    ("Обработка отмененного или невостребованного товара (разбивается по товарам в отправлении в одинаковой пропорции)", "reverse_logistics"),
    ("Обработка невыкупленного товара", "reverse_logistics"),
    ("Логистика", "logistics"),
    ("Обратная логистика", "reverse_logistics"),
]


def xlsx_value(row: tuple[Any, ...], headers: dict[str, int], name: str) -> Any:
    index = headers.get(normalized_header(name))
    return row[index] if index is not None and index < len(row) else None


def manual_records() -> tuple[list[tuple[Any, ...]], list[tuple[Any, ...]], date | None, date | None]:
    if not FIN_DIR.exists():
        raise FileNotFoundError(FIN_DIR)
    files = sorted(FIN_DIR.glob("*.xlsx"))
    events: list[tuple[Any, ...]] = []
    lines: list[tuple[Any, ...]] = []
    min_day: date | None = None
    max_day: date | None = None
    print(
        f"ПЛАН XLSX: файлов {len(files)} | все строки начислений | "
        "комиссия и 10 логистических компонентов сохраняются отдельно.",
        flush=True,
    )
    started = time.monotonic()
    for file_index, path in enumerate(files, start=1):
        workbook = load_workbook(path, read_only=True, data_only=True)
        file_events = 0
        file_lines = 0
        try:
            sheet = workbook[workbook.sheetnames[0]]
            header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            headers = {
                normalized_header(value): index
                for index, value in enumerate(header_row)
                if text(value)
            }
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                operation_day = parse_day(xlsx_value(row, headers, "Дата начисления"))
                if operation_day is None:
                    continue
                accrual_type = text(xlsx_value(row, headers, "Тип начисления"))
                posting = text(
                    xlsx_value(
                        row,
                        headers,
                        "Номер отправления или идентификатор услуги",
                    )
                )
                sku = text(xlsx_value(row, headers, "SKU"))
                article = text(xlsx_value(row, headers, "Артикул"))
                product_name = text(xlsx_value(row, headers, "Название товара или услуги"))
                quantity = money(xlsx_value(row, headers, "Количество"))
                gross = money(
                    xlsx_value(
                        row,
                        headers,
                        "За продажу или возврат до вычета комиссий и услуг",
                    )
                )
                total = money(xlsx_value(row, headers, "Итого, руб."))
                commission_rate = money(
                    xlsx_value(row, headers, "Вознаграждение Ozon, %")
                )
                localization = money(xlsx_value(row, headers, "Индекс локализации"))
                delivery_hours = money(
                    xlsx_value(row, headers, "Среднее время доставки, часы")
                )
                event_key = "xlsx:" + stable_key(path.name, sheet.title, row_number)
                raw_meta = {
                    "source_file": path.name,
                    "source_sheet": sheet.title,
                    "source_row": row_number,
                }
                events.append(
                    (
                        event_key,
                        SOURCE_XLSX,
                        path.name,
                        operation_day,
                        None,
                        posting or None,
                        posting or None,
                        "ITEM" if sku else "NON_ITEM",
                        None,
                        accrual_type or None,
                        total,
                        "RUB",
                        Json(raw_meta),
                    )
                )
                file_events += 1
                min_day = operation_day if min_day is None else min(min_day, operation_day)
                max_day = operation_day if max_day is None else max(max_day, operation_day)

                def append_line(
                    suffix: str,
                    line_kind: str,
                    amount: Decimal,
                    type_name: str,
                    *,
                    line_quantity: Decimal | None = None,
                ) -> None:
                    nonlocal file_lines
                    if amount == 0 and not line_quantity:
                        return
                    lines.append(
                        (
                            f"{event_key}:{suffix}",
                            event_key,
                            SOURCE_XLSX,
                            operation_day,
                            posting or None,
                            sku or None,
                            article or None,
                            product_name or None,
                            None,
                            line_kind,
                            None,
                            type_name,
                            amount,
                            "RUB",
                            line_quantity,
                            commission_rate if line_kind == "commission" else None,
                            localization or None,
                            delivery_hours or None,
                            "exact",
                            Json(raw_meta),
                        )
                    )
                    file_lines += 1

                append_line("revenue", "revenue", gross, accrual_type, line_quantity=quantity)
                for component_index, (header, line_kind) in enumerate(MANUAL_COMPONENTS):
                    append_line(
                        f"component:{component_index}",
                        line_kind,
                        money(xlsx_value(row, headers, header)),
                        header,
                    )
                component_sum = gross + sum(
                    money(xlsx_value(row, headers, header))
                    for header, _line_kind in MANUAL_COMPONENTS
                )
                residual = total - component_sum
                append_line("residual", "other_service", residual, f"{accrual_type}: прочее")
        finally:
            workbook.close()
        elapsed = time.monotonic() - started
        eta = elapsed / file_index * (len(files) - file_index) if file_index else 0
        print(
            f"[{file_index}/{len(files)}] {path.name}: событий {file_events:,}, "
            f"строк {file_lines:,}, ошибок 0",
            flush=True,
        )
        print(
            f"ПРОГРЕСС: {file_index}/{len(files)} "
            f"({file_index / max(len(files), 1) * 100:.1f}%) | {path.name} | "
            f"накоплено событий {len(events):,}, строк {len(lines):,} | "
            f"elapsed={elapsed:.1f}s | ETA={eta:.1f}s",
            flush=True,
        )
    return events, lines, min_day, max_day


def replace_range(
    conn,
    source_kind: str,
    date_from: date,
    date_to: date,
    events: list[tuple[Any, ...]],
    lines: list[tuple[Any, ...]],
    *,
    dry_run: bool,
) -> tuple[int, int]:
    if dry_run:
        return len(events), len(lines)
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                DELETE FROM public.ozon_finance_lines
                WHERE source_kind = %s AND operation_date BETWEEN %s AND %s
                """,
                (source_kind, date_from, date_to),
            )
            cur.execute(
                """
                DELETE FROM public.ozon_finance_events
                WHERE source_kind = %s AND operation_date BETWEEN %s AND %s
                """,
                (source_kind, date_from, date_to),
            )
            if events:
                execute_values(
                    cur,
                    """
                    INSERT INTO public.ozon_finance_events (
                        event_key, source_kind, source_ref, operation_date,
                        accrual_id, unit_number, posting_number, accrued_category,
                        operation_type_code, operation_type_name, total_amount,
                        currency, raw_json
                    ) VALUES %s
                    """,
                    events,
                    page_size=1000,
                )
            if lines:
                execute_values(
                    cur,
                    """
                    INSERT INTO public.ozon_finance_lines (
                        line_key, event_key, source_kind, operation_date,
                        posting_number, sku, article, product_name, delivery_schema,
                        line_kind, type_id, type_name, amount, currency, quantity,
                        commission_rate_pct, localization_index, delivery_time_hours,
                        allocation_scope, raw_json
                    ) VALUES %s
                    """,
                    lines,
                    page_size=1000,
                )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return len(events), len(lines)


def fetch_accrual_types(api: FinanceApiClient) -> dict[int, dict[str, Any]]:
    payload = api.post("/v1/finance/accrual/types", {})
    rows = payload.get("accrual_types") or []
    result: dict[int, dict[str, Any]] = {}
    for item in rows:
        if not isinstance(item, dict):
            continue
        try:
            type_id = int(item.get("id"))
        except (TypeError, ValueError):
            continue
        result[type_id] = {
            "name": text(item.get("name")),
            "description": text(item.get("description")),
        }
    if not result:
        raise FinancePipelineError("/v1/finance/accrual/types вернул пустой справочник")
    return result


def store_accrual_types(
    conn, types: dict[int, dict[str, Any]], *, dry_run: bool
) -> None:
    if dry_run:
        return
    with conn.cursor() as cur:
        execute_values(
            cur,
            """
            INSERT INTO public.ozon_finance_accrual_types (
                type_id, type_name, description
            ) VALUES %s
            ON CONFLICT (type_id) DO UPDATE SET
                type_name = EXCLUDED.type_name,
                description = EXCLUDED.description,
                updated_at = now()
            """,
            [
                (type_id, item["name"], item["description"])
                for type_id, item in sorted(types.items())
            ],
        )
    conn.commit()


def fetch_day_accruals(api: FinanceApiClient, operation_day: date) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    last_id = ""
    page = 0
    while True:
        page += 1
        request_payload: dict[str, Any] = {"date": operation_day.isoformat()}
        if last_id:
            request_payload["last_id"] = last_id
        print(
            f"ПРОГРЕСС: Finance API | {operation_day} | страница {page} | "
            f"накоплено {len(rows):,} | запрос {api.request_count + 1}",
            flush=True,
        )
        payload = api.post("/v1/finance/accrual/by-day", request_payload)
        page_rows = payload.get("accruals") or []
        if not isinstance(page_rows, list):
            raise FinancePipelineError("accruals в /by-day имеет неверный тип")
        rows.extend(item for item in page_rows if isinstance(item, dict))
        new_last_id = text(payload.get("last_id"))
        if not new_last_id or new_last_id == last_id:
            break
        last_id = new_last_id
    return rows


def append_api_line(
    lines: list[tuple[Any, ...]],
    event_key: str,
    operation_day: date,
    posting_number: str,
    sku: str,
    delivery_schema: str,
    line_kind: str,
    type_id: int | None,
    type_name: str,
    amount: Decimal,
    currency: str,
    *,
    quantity: Decimal | None = None,
    commission_rate: Decimal | None = None,
    scope: str = "exact",
    raw: Any = None,
    suffix: str,
) -> None:
    if amount == 0 and not quantity:
        return
    lines.append(
        (
            f"{event_key}:{suffix}",
            event_key,
            SOURCE_API,
            operation_day,
            posting_number or None,
            sku or None,
            None,
            None,
            delivery_schema or None,
            line_kind,
            type_id,
            type_name or None,
            amount,
            currency or "RUB",
            quantity,
            commission_rate,
            None,
            None,
            scope,
            Json(raw) if raw is not None else None,
        )
    )


def infer_line_kind(type_name: str, amount: Decimal) -> str:
    value = re.sub(r"[^a-zа-я0-9]+", "", type_name.lower())
    if any(word in value for word in (
        "реклам", "продвиж", "трафарет", "выводвтоп", "оплатазаклик",
        "payperclick", "promotion", "mailing",
    )):
        return "advertising"
    if any(word in value for word in (
        "возврат", "обратн", "невыкуп", "отмен",
        "returnflow", "sellerreturns", "returnacceptance",
    )):
        return "reverse_logistics"
    if any(word in value for word in ("packingfee", "packagecost", "crossdock", "supplyinbound")):
        return "fulfillment_ozon"
    if any(word in value for word in ("placements", "размещен", "склад")):
        return "storage"
    if any(word in value for word in (
        "логист", "достав", "магистрал", "мил", "обработ",
        "logistic", "lastmile", "deliverytohandover", "replenishment",
    )):
        return "logistics"
    if "acquiring" in value:
        return "acquiring"
    if "subscription" in value:
        return "subscription"
    if any(word in value for word in ("компенсац", "возмещ", "itemcompensation")) and amount > 0:
        return "compensation"
    if any(word in value for word in ("штраф", "удержан", "пеня", "антифрод", "fine")):
        return "fines"
    return "other_service"


def api_records_for_day(
    operation_day: date,
    accruals: list[dict[str, Any]],
    types: dict[int, dict[str, Any]],
) -> tuple[list[tuple[Any, ...]], list[tuple[Any, ...]]]:
    events: list[tuple[Any, ...]] = []
    lines: list[tuple[Any, ...]] = []
    for event_index, item in enumerate(accruals, start=1):
        accrual_id_raw = item.get("accrual_id")
        try:
            accrual_id = int(accrual_id_raw)
        except (TypeError, ValueError):
            accrual_id = None
        unit_number = text(item.get("unit_number"))
        category = text(item.get("accrued_category")) or "UNSPECIFIED"
        total_money = item.get("total_amount") or {}
        total_amount = money_amount(total_money)
        currency = money_currency(total_money) or "RUB"
        source_identity = accrual_id if accrual_id is not None else stable_key(
            operation_day, unit_number, category, total_amount, event_index
        )
        event_key = f"api:{operation_day}:{source_identity}"
        posting_number = unit_number if category == "POSTING" else ""
        events.append(
            (
                event_key,
                SOURCE_API,
                "/v1/finance/accrual/by-day",
                operation_day,
                accrual_id,
                unit_number or None,
                posting_number or None,
                category,
                None,
                None,
                total_amount,
                currency,
                Json(item),
            )
        )

        posting = item.get("posting")
        if isinstance(posting, dict):
            delivery_schema = text(posting.get("delivery_schema"))
            for product_index, product in enumerate(posting.get("products") or [], start=1):
                if not isinstance(product, dict):
                    continue
                sku = text(product.get("sku"))
                commission = product.get("commission") or {}
                seller_price = money_amount(commission.get("seller_price"))
                sale_amount = money_amount(commission.get("sale_amount"))
                sale_price = money_amount(commission.get("sale_price"))
                quantity = (
                    abs(sale_amount / seller_price)
                    if seller_price
                    else (abs(sale_amount / sale_price) if sale_price else None)
                )
                ratio = money(commission.get("commission_ratio"))
                append_api_line(
                    lines,
                    event_key,
                    operation_day,
                    posting_number,
                    sku,
                    delivery_schema,
                    "revenue",
                    None,
                    "Реализовано на сумму",
                    sale_amount,
                    money_currency(commission.get("sale_amount")),
                    quantity=quantity,
                    commission_rate=ratio * 100 if abs(ratio) <= 1 else ratio,
                    raw=product,
                    suffix=f"product:{product_index}:revenue",
                )
                final_commission = money_amount(commission.get("commission"))
                if final_commission == 0:
                    final_commission = money_amount(commission.get("sale_commission"))
                append_api_line(
                    lines,
                    event_key,
                    operation_day,
                    posting_number,
                    sku,
                    delivery_schema,
                    "commission",
                    None,
                    "Комиссия Ozon",
                    final_commission,
                    money_currency(
                        commission.get("commission") or commission.get("sale_commission")
                    ),
                    commission_rate=ratio * 100 if abs(ratio) <= 1 else ratio,
                    raw=commission,
                    suffix=f"product:{product_index}:commission",
                )
                # bonus и coinvestment описывают состав цены и уже учтены
                # в sale_amount/commission. Это не отдельные денежные начисления.
                delivery = product.get("delivery") or {}
                services = delivery.get("services") or []
                if services:
                    for service_index, service in enumerate(services, start=1):
                        if not isinstance(service, dict):
                            continue
                        try:
                            type_id = int(service.get("type_id"))
                        except (TypeError, ValueError):
                            type_id = None
                        type_name = types.get(type_id or -1, {}).get("name", "")
                        accrued = service.get("accrued") or {}
                        amount = money_amount(accrued)
                        append_api_line(
                            lines,
                            event_key,
                            operation_day,
                            posting_number,
                            sku,
                            delivery_schema,
                            infer_line_kind(type_name, amount),
                            type_id,
                            type_name,
                            amount,
                            money_currency(accrued),
                            raw=service,
                            suffix=f"product:{product_index}:delivery:{service_index}",
                        )
                else:
                    accrued = delivery.get("total_accrued") or {}
                    amount = money_amount(accrued)
                    append_api_line(
                        lines,
                        event_key,
                        operation_day,
                        posting_number,
                        sku,
                        delivery_schema,
                        "logistics",
                        None,
                        "Доставка Ozon",
                        amount,
                        money_currency(accrued),
                        raw=delivery,
                        suffix=f"product:{product_index}:delivery-total",
                    )

        item_fees = item.get("item_fees")
        if isinstance(item_fees, dict):
            for sku_index, sku_fees in enumerate(item_fees.get("fees") or [], start=1):
                if not isinstance(sku_fees, dict):
                    continue
                sku = text(sku_fees.get("sku"))
                for fee_index, fee in enumerate(sku_fees.get("fees") or [], start=1):
                    if not isinstance(fee, dict):
                        continue
                    try:
                        type_id = int(fee.get("type_id"))
                    except (TypeError, ValueError):
                        type_id = None
                    type_name = types.get(type_id or -1, {}).get("name", "")
                    accrued = fee.get("accrued") or {}
                    amount = money_amount(accrued)
                    append_api_line(
                        lines,
                        event_key,
                        operation_day,
                        unit_number,
                        sku,
                        "",
                        infer_line_kind(type_name, amount),
                        type_id,
                        type_name,
                        amount,
                        money_currency(accrued),
                        raw=fee,
                        suffix=f"item:{sku_index}:{fee_index}",
                    )

        non_item = item.get("non_item_fee")
        if isinstance(non_item, dict):
            try:
                type_id = int(non_item.get("type_id"))
            except (TypeError, ValueError):
                type_id = None
            type_name = types.get(type_id or -1, {}).get("name", "")
            accrued = non_item.get("accrued") or {}
            amount = money_amount(accrued)
            append_api_line(
                lines,
                event_key,
                operation_day,
                unit_number,
                "",
                "",
                infer_line_kind(type_name, amount),
                type_id,
                type_name,
                amount,
                money_currency(accrued),
                scope="account",
                raw=non_item,
                suffix="non-item",
            )

        container_fees = item.get("container_fees")
        if isinstance(container_fees, dict):
            for fee_index, fee in enumerate(container_fees.get("fees") or [], start=1):
                if not isinstance(fee, dict):
                    continue
                try:
                    type_id = int(fee.get("type_id"))
                except (TypeError, ValueError):
                    type_id = None
                type_name = types.get(type_id or -1, {}).get("name", "")
                accrued = fee.get("accrued") or {}
                amount = money_amount(accrued)
                append_api_line(
                    lines,
                    event_key,
                    operation_day,
                    unit_number,
                    "",
                    "",
                    infer_line_kind(type_name, amount),
                    type_id,
                    type_name,
                    amount,
                    money_currency(accrued),
                    scope="account",
                    raw=fee,
                    suffix=f"container:{fee_index}",
                )
    return events, lines


def index_value(container: Any) -> Decimal | None:
    if not isinstance(container, dict):
        return None
    for key in ("price_index_value", "minimal_price", "price"):
        if container.get(key) not in (None, ""):
            return money(container.get(key))
    return None


def fetch_product_dimensions(api: FinanceApiClient) -> dict[str, dict[str, Any]]:
    last_id = ""
    page = 0
    dimensions: dict[str, dict[str, Any]] = {}
    while True:
        page += 1
        print(
            f"ПРОГРЕСС: габариты | страница {page} | товаров {len(dimensions):,} | "
            f"запрос {api.request_count + 1}",
            flush=True,
        )
        payload = api.post(
            "/v4/product/info/attributes",
            {
                "filter": {"visibility": "ALL"},
                "last_id": last_id,
                "limit": ATTRIBUTES_PAGE_SIZE,
                "sort_dir": "ASC",
            },
        )
        items = payload.get("result") or []
        for item in items:
            if not isinstance(item, dict):
                continue
            value = {
                "depth": money(item.get("depth")),
                "width": money(item.get("width")),
                "height": money(item.get("height")),
                "dimension_unit": text(item.get("dimension_unit")) or None,
                "weight": money(item.get("weight")),
                "weight_unit": text(item.get("weight_unit")) or None,
            }
            for key in (
                text(item.get("id")),
                text(item.get("offer_id")),
                text(item.get("sku")),
            ):
                if key:
                    dimensions[key] = value
        new_last_id = text(payload.get("last_id"))
        total = int(payload.get("total") or len(items))
        if not items or not new_last_id or new_last_id == last_id or len(items) < ATTRIBUTES_PAGE_SIZE:
            break
        last_id = new_last_id
        if len(dimensions) >= total * 2:
            break
    return dimensions


def fetch_price_snapshots(api: FinanceApiClient, dimensions: dict[str, dict[str, Any]]) -> list[tuple[Any, ...]]:
    cursor = ""
    rows: list[tuple[Any, ...]] = []
    snapshot_day = app.marketplace_today()
    page = 0
    while True:
        page += 1
        print(
            f"ПРОГРЕСС: цены/тарифы | страница {page} | накоплено {len(rows):,} | "
            f"запрос {api.request_count + 1}",
            flush=True,
        )
        payload = api.post(
            "/v5/product/info/prices",
            {"filter": {"visibility": "ALL"}, "cursor": cursor, "limit": PRICE_PAGE_SIZE},
        )
        items = payload.get("items") or []
        for item in items:
            if not isinstance(item, dict):
                continue
            sku = text(item.get("product_id"))
            if not sku:
                continue
            price = item.get("price") or {}
            commissions = item.get("commissions") or {}
            indexes = item.get("price_indexes") or {}
            product_dimensions = (
                dimensions.get(text(item.get("product_id")))
                or dimensions.get(text(item.get("offer_id")))
                or {}
            )
            rows.append(
                (
                    snapshot_day,
                    sku,
                    text(item.get("offer_id")) or None,
                    text(item.get("product_id")) or None,
                    text(price.get("currency_code")) or "RUB",
                    money(price.get("price")),
                    money(price.get("marketing_seller_price")),
                    money(price.get("min_price")),
                    money(price.get("old_price")),
                    money(price.get("retail_price")),
                    money(price.get("net_price")),
                    money(price.get("vat")),
                    money(item.get("acquiring")),
                    money(commissions.get("sales_percent_fbo")),
                    money(commissions.get("sales_percent_fbs")),
                    money(commissions.get("sales_percent_rfbs")),
                    money(commissions.get("fbo_deliv_to_customer_amount")),
                    money(commissions.get("fbo_direct_flow_trans_min_amount")),
                    money(commissions.get("fbo_direct_flow_trans_max_amount")),
                    money(commissions.get("fbo_return_flow_amount")),
                    money(commissions.get("fbs_deliv_to_customer_amount")),
                    money(commissions.get("fbs_direct_flow_trans_min_amount")),
                    money(commissions.get("fbs_direct_flow_trans_max_amount")),
                    money(commissions.get("fbs_first_mile_min_amount")),
                    money(commissions.get("fbs_first_mile_max_amount")),
                    money(commissions.get("fbs_return_flow_amount")),
                    text(indexes.get("color_index")) or None,
                    index_value(indexes.get("ozon_index_data")),
                    index_value(indexes.get("external_index_data")),
                    index_value(indexes.get("self_marketplaces_index_data")),
                    money(item.get("volume_weight")),
                    product_dimensions.get("depth"),
                    product_dimensions.get("width"),
                    product_dimensions.get("height"),
                    product_dimensions.get("dimension_unit"),
                    product_dimensions.get("weight"),
                    product_dimensions.get("weight_unit"),
                    Json(item),
                )
            )
        new_cursor = text(payload.get("cursor"))
        total = int(payload.get("total") or len(rows))
        if not items or not new_cursor or new_cursor == cursor or len(rows) >= total:
            break
        cursor = new_cursor
    return rows


def store_price_snapshots(
    conn, rows: list[tuple[Any, ...]], *, dry_run: bool
) -> int:
    if dry_run:
        return len(rows)
    if not rows:
        raise FinancePipelineError("Ozon не вернул цены; снимок не изменён")
    snapshot_day = rows[0][0]
    try:
        with conn.cursor() as cur:
            cur.execute(
                "DELETE FROM public.ozon_product_price_snapshots WHERE snapshot_date = %s",
                (snapshot_day,),
            )
            execute_values(
                cur,
                """
                INSERT INTO public.ozon_product_price_snapshots (
                    snapshot_date, sku, offer_id, product_id, currency, price,
                    marketing_seller_price, min_price, old_price, retail_price,
                    net_price, vat_pct, acquiring, sales_percent_fbo,
                    sales_percent_fbs, sales_percent_rfbs, fbo_delivery_amount,
                    fbo_logistics_min, fbo_logistics_max, fbo_return_amount,
                    fbs_delivery_amount, fbs_logistics_min, fbs_logistics_max,
                    fbs_first_mile_min, fbs_first_mile_max, fbs_return_amount,
                    price_index_color, ozon_price_index, market_price_index,
                    self_price_index, volume_weight, depth, width, height,
                    dimension_unit, weight, weight_unit, raw_json
                ) VALUES %s
                """,
                rows,
                page_size=1000,
            )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return len(rows)


def normalize_existing_api_lines(conn, *, dry_run: bool) -> tuple[int, int]:
    with conn.cursor() as cur:
        cur.execute(
            '''
            SELECT line_key, coalesce(type_name, ''), amount, line_kind
            FROM public.ozon_finance_lines
            WHERE source_kind = %s
              AND line_kind NOT IN ('revenue', 'commission', 'bonus', 'coinvestment')
            ''',
            (SOURCE_API,),
        )
        updates = []
        for line_key, type_name, amount, current_kind in cur.fetchall():
            new_kind = infer_line_kind(type_name, money(amount))
            if new_kind != current_kind:
                updates.append((new_kind, line_key))
        cur.execute(
            '''
            SELECT count(*)
            FROM public.ozon_finance_lines
            WHERE source_kind = %s
              AND line_kind IN ('bonus', 'coinvestment')
            ''',
            (SOURCE_API,),
        )
        informational_rows = int(cur.fetchone()[0])
        if dry_run:
            conn.rollback()
            return informational_rows, len(updates)
        cur.execute(
            '''
            DELETE FROM public.ozon_finance_lines
            WHERE source_kind = %s
              AND line_kind IN ('bonus', 'coinvestment')
            ''',
            (SOURCE_API,),
        )
        if updates:
            execute_batch(
                cur,
                '''
                UPDATE public.ozon_finance_lines
                SET line_kind = %s, imported_at = now()
                WHERE line_key = %s
                ''',
                updates,
                page_size=1000,
            )
    conn.commit()
    return informational_rows, len(updates)


def default_api_range(conn) -> tuple[date, date]:
    yesterday = app.marketplace_today() - timedelta(days=1)
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT
                max(operation_date) FILTER (WHERE source_kind = %s),
                max(operation_date) FILTER (WHERE source_kind = %s)
            FROM public.ozon_finance_events
            """,
            (SOURCE_API, SOURCE_XLSX),
        )
        api_max, xlsx_max = cur.fetchone()
    last = api_max or xlsx_max
    start = last + timedelta(days=1) if last else yesterday
    return start, yesterday


def validate_range(date_from: date, date_to: date) -> None:
    if date_from > date_to:
        raise FinancePipelineError("Дата начала позже даты окончания")
    yesterday = app.marketplace_today() - timedelta(days=1)
    if date_to > yesterday:
        raise FinancePipelineError(f"Дата окончания не может быть позже {yesterday}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", choices=("manual", "api", "all", "reclassify"), default="api")
    parser.add_argument("--date-from", type=date.fromisoformat)
    parser.add_argument("--date-to", type=date.fromisoformat)
    parser.add_argument(
        "--skip-catalog-refresh",
        action="store_true",
        help=(
            "Не запрашивать повторно каталог габаритов и текущие цены/тарифы. "
            "Используется для исторической дозагрузки финансов."
        ),
    )
    parser.add_argument(
        "--with-product-snapshot",
        action="store_true",
        help="Обновить габариты и цены/тарифы; использовать только в ежемесячном запуске",
    )
    parser.add_argument("--dry-run", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    started = time.monotonic()
    conn = connect_km()
    api: FinanceApiClient | None = None
    total_events = 0
    total_lines = 0
    total_prices = 0
    errors = 0
    try:
        ensure_schema(conn)
        steps = ["manual", "api"] if args.source == "all" else [args.source]
        print(
            f"ПЛАН: клиент {CLIENT_LABEL} | БД {TARGET_DB} | шагов {len(steps)} | "
            f"источник {args.source} | Finance API последовательно, пауза "
            f"{FINANCE_INTERVAL_SECONDS:.1f} сек. | 429: ожидание не менее 60 сек. | "
            f"запись только в {TARGET_DB}.",
            flush=True,
        )
        for step_index, step in enumerate(steps, start=1):
            print(
                f"ПРОГРЕСС: {step_index}/{len(steps)} "
                f"({step_index / len(steps) * 100:.1f}%) | шаг {step} | "
                f"событий {total_events:,} | строк {total_lines:,}",
                flush=True,
            )
            if step == "reclassify":
                removed, updated = normalize_existing_api_lines(
                    conn, dry_run=args.dry_run
                )
                print(
                    f"[{step_index}/{len(steps)}] reclassify: удалено справочных "
                    f"bonus/coinvestment {removed:,}, переклассифицировано {updated:,}, "
                    "API-запросов 0, ошибок 0",
                    flush=True,
                )
                continue
            if step == "manual":
                events, lines, min_day, max_day = manual_records()
                if min_day is None or max_day is None:
                    raise FinancePipelineError("В ручных Fin-файлах нет строк с датами")
                loaded_events, loaded_lines = replace_range(
                    conn,
                    SOURCE_XLSX,
                    min_day,
                    max_day,
                    events,
                    lines,
                    dry_run=args.dry_run,
                )
                total_events += loaded_events
                total_lines += loaded_lines
                print(
                    f"[{step_index}/{len(steps)}] manual: период {min_day}..{max_day}, "
                    f"событий {loaded_events:,}, строк {loaded_lines:,}, ошибок 0",
                    flush=True,
                )
                continue

            date_from, date_to = default_api_range(conn)
            date_from = args.date_from or date_from
            date_to = args.date_to or date_to
            if date_from > date_to:
                print(
                    f"[{step_index}/{len(steps)}] api: новых дат нет; обновляю только "
                    "справочник начислений и снимок тарифов/индексов",
                    flush=True,
                )
            else:
                validate_range(date_from, date_to)
            api = api or FinanceApiClient()
            types = fetch_accrual_types(api)
            store_accrual_types(conn, types, dry_run=args.dry_run)
            days = (
                [
                    date_from + timedelta(days=offset)
                    for offset in range((date_to - date_from).days + 1)
                ]
                if date_from <= date_to
                else []
            )
            api_started = time.monotonic()
            for day_index, operation_day in enumerate(days, start=1):
                accruals = fetch_day_accruals(api, operation_day)
                events, lines = api_records_for_day(operation_day, accruals, types)
                loaded_events, loaded_lines = replace_range(
                    conn,
                    SOURCE_API,
                    operation_day,
                    operation_day,
                    events,
                    lines,
                    dry_run=args.dry_run,
                )
                total_events += loaded_events
                total_lines += loaded_lines
                elapsed = time.monotonic() - api_started
                eta = elapsed / day_index * (len(days) - day_index) if day_index else 0
                print(
                    f"[{day_index}/{len(days)}] {operation_day}: начислений "
                    f"{loaded_events:,}, строк {loaded_lines:,}, ошибок 0",
                    flush=True,
                )
                print(
                    f"ПРОГРЕСС: {day_index}/{len(days)} "
                    f"({day_index / max(len(days), 1) * 100:.1f}%) | {operation_day} | "
                    f"API-запросов {api.request_count} | накоплено событий "
                    f"{total_events:,}, строк {total_lines:,} | "
                    f"elapsed={elapsed:.1f}s | ETA={eta:.1f}s",
                    flush=True,
                )
            dimensions: dict[str, dict[str, Any]] = {}
            price_rows: list[tuple[Any, ...]] = []
            if args.with_product_snapshot and not args.skip_catalog_refresh:
                dimensions = fetch_product_dimensions(api)
                price_rows = fetch_price_snapshots(api, dimensions)
                total_prices += store_price_snapshots(conn, price_rows, dry_run=args.dry_run)
            else:
                print(
                    f"[{step_index}/{len(steps)}] api: габариты и текущие "
                    "цены/тарифы пропущены; снимок выполняется только ежемесячно",
                    flush=True,
                )
            print(
                f"[{step_index}/{len(steps)}] api: период "
                f"{date_from if days else '-'}..{date_to if days else '-'}, "
                f"тарифов/цен {len(price_rows):,}, габаритов {len(dimensions):,}, "
                f"запросов {api.request_count}, ошибок 0",
                flush=True,
            )
        elapsed = time.monotonic() - started
        print(
            f"ИТОГ: {CLIENT_LABEL} | БД {TARGET_DB} | событий {total_events:,} | "
            f"строк {total_lines:,} | тарифов/цен {total_prices:,} | "
            f"API-запросов {api.request_count if api else 0} | ошибок {errors} | "
            f"elapsed={elapsed:.1f}s | stopped=false | partial=false | "
            f"режим={'dry-run' if args.dry_run else 'запись завершена'}",
            flush=True,
        )
        return 0
    except Exception as exc:
        conn.rollback()
        errors += 1
        elapsed = time.monotonic() - started
        print(
            f"ИТОГ: {CLIENT_LABEL} | остановлено с ошибкой | БД {TARGET_DB} | "
            f"событий {total_events:,} | строк {total_lines:,} | "
            f"API-запросов {api.request_count if api else 0} | "
            f"elapsed={elapsed:.1f}s | {exc}",
            file=sys.stderr,
            flush=True,
        )
        return 1
    finally:
        conn.close()


if __name__ == "__main__":
    raise SystemExit(main())
