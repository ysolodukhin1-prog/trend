#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import time
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

import openpyxl
import psycopg2
from psycopg2.extras import RealDictCursor, execute_values


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "ozon_category_dashboard"))

import app  # noqa: E402


WORKBOOK_PATH = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Back to School 2026 МП 01 07.xlsx"
)
TARGET_SEASON = "Back to school 2026"
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "back_to_school_2026_season_sync"
SHEET_CONFIG = {
    "ВБ": {
        "marketplace": "wb",
        "sku_header": "Артикул",
        "barcode_header": None,
        "vendor_header": "Арт. Поставщика",
        "model_header": "Модель ГД",
        "assortment_header": "Осн. Ассорт / БиА",
        "tg_header": "Department",
        "tg_plus_header": "ТГ+",
        "cg_header": "ЦГ",
        "field": "wb_nmid",
    },
    "ОЗОН": {
        "marketplace": "ozon",
        "sku_header": "SKU (Sku)",
        "barcode_header": "ШК Ozon (Barcode)",
        "vendor_header": None,
        "model_header": "Модель ГД",
        "assortment_header": "Осн. Ассорт / БиА",
        "tg_header": "Department",
        "tg_plus_header": "ТГ+",
        "cg_header": "ЦГ",
        "field": "ozon_sku",
    },
}


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def row_hash(*values: object) -> str:
    payload = "|".join("" if value is None else str(value) for value in values)
    return hashlib.md5(payload.encode("utf-8")).hexdigest()


def header_index(sheet) -> dict[str, int]:
    headers = [clean(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    return {value.lower(): index for index, value in enumerate(headers) if value}


def value(row, index):
    if index is None or index >= len(row):
        return ""
    return clean(row[index])


def extract_rows(source: Path) -> tuple[list[dict[str, object]], dict[str, object]]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    rows: OrderedDict[tuple[str, str, str], dict[str, object]] = OrderedDict()
    summary = {}
    for sheet_name, config in SHEET_CONFIG.items():
        sheet = workbook[sheet_name]
        idx = header_index(sheet)
        required = [config["sku_header"], config["model_header"], config["assortment_header"], config["tg_header"], config["tg_plus_header"], config["cg_header"]]
        if config["barcode_header"]:
            required.append(config["barcode_header"])
        if config["vendor_header"]:
            required.append(config["vendor_header"])
        missing = [name for name in required if name.lower() not in idx]
        if missing:
            raise RuntimeError(f"{sheet_name}: missing columns {missing}")
        sheet_rows = 0
        source_seasons = {}
        for source_row, raw_row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            sku = value(raw_row, idx[config["sku_header"].lower()])
            if not sku:
                continue
            source_season = value(raw_row, idx.get("сезон"))
            source_seasons[source_season] = source_seasons.get(source_season, 0) + 1
            barcode = value(raw_row, idx.get(config["barcode_header"].lower())) if config["barcode_header"] else ""
            vendor_article = value(raw_row, idx.get(config["vendor_header"].lower())) if config["vendor_header"] else ""
            if config["marketplace"] == "wb" and not barcode:
                barcode = f"BTS2026-WB-{sku}"
            if config["marketplace"] == "ozon" and not barcode:
                barcode = f"BTS2026-OZON-{sku}"
            key = (config["marketplace"], sku, barcode)
            rows.setdefault(
                key,
                {
                    "marketplace": config["marketplace"],
                    "sku": sku,
                    "barcode": barcode,
                    "vendor_article": vendor_article,
                    "gj_model": value(raw_row, idx[config["model_header"].lower()]),
                    "assortment_bia": value(raw_row, idx[config["assortment_header"].lower()]),
                    "tg": value(raw_row, idx[config["tg_header"].lower()]),
                    "tg_plus": value(raw_row, idx[config["tg_plus_header"].lower()]),
                    "cg": value(raw_row, idx[config["cg_header"].lower()]),
                    "season": TARGET_SEASON,
                    "source_file": source.name,
                    "source_wb_row": source_row if config["marketplace"] == "wb" else None,
                    "source_ozon_row": source_row if config["marketplace"] == "ozon" else None,
                },
            )
            sheet_rows += 1
        summary[config["marketplace"]] = {
            "sheet": sheet_name,
            "rows": sheet_rows,
            "unique_sku": len({row[1] for row in rows if row[0] == config["marketplace"]}),
            "source_seasons": source_seasons,
        }
    return list(rows.values()), summary


def ensure_mapping_schema(cur) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS public.sku_mapping_gj (
            mapping_hash text PRIMARY KEY,
            barcode text NOT NULL,
            wb_article text,
            wb_nmid text,
            ozon_sku text,
            gj_model text,
            assortment_bia text,
            tg text,
            tg_plus text,
            cg text,
            season text,
            source_file text,
            source_wb_row integer,
            source_ozon_row integer,
            imported_at timestamp without time zone DEFAULT now()
        )
        """
    )


def refresh_mapping_views(cur) -> None:
    for view_name in (
        "mv_sku_mapping_gj_wb_article",
        "mv_sku_mapping_gj_wb_nmid",
        "mv_sku_mapping_gj_ozon_sku",
        "mv_sku_mapping_gj_barcode",
    ):
        print(f"ПРОГРЕСС: refresh {view_name}", flush=True)
        cur.execute(f"REFRESH MATERIALIZED VIEW public.{view_name}")
        cur.execute(f"ANALYZE public.{view_name}")


def backup_existing_rows(cur, rows: list[dict[str, object]]) -> Path:
    wb_skus = sorted({str(row["sku"]) for row in rows if row["marketplace"] == "wb"})
    ozon_skus = sorted({str(row["sku"]) for row in rows if row["marketplace"] == "ozon"})
    with cur.connection.cursor(cursor_factory=RealDictCursor) as dict_cur:
        dict_cur.execute(
            """
            SELECT *
            FROM public.sku_mapping_gj
            WHERE wb_nmid = ANY(%s) OR ozon_sku = ANY(%s)
            ORDER BY mapping_hash
            """,
            (wb_skus, ozon_skus),
        )
        existing = [dict(row) for row in dict_cur.fetchall()]
    for row in existing:
        for key, value in list(row.items()):
            if isinstance(value, datetime):
                row[key] = value.isoformat(sep=" ", timespec="seconds")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    backup_path = OUTPUT_DIR / f"sku_mapping_gj_back_to_school_backup_{datetime.now():%Y%m%d_%H%M%S}.json"
    backup_path.write_text(json.dumps(existing, ensure_ascii=False, indent=2), encoding="utf-8")
    return backup_path


def sync_rows(rows: list[dict[str, object]], dry_run: bool) -> dict[str, object]:
    wb_rows = [row for row in rows if row["marketplace"] == "wb"]
    ozon_rows = [row for row in rows if row["marketplace"] == "ozon"]
    result: dict[str, object] = {"updated": {}, "inserted": {}, "backup_path": ""}
    started = time.monotonic()
    with psycopg2.connect(**app.read_db_config()) as conn, conn.cursor() as cur:
        ensure_mapping_schema(cur)
        backup_path = backup_existing_rows(cur, rows)
        result["backup_path"] = str(backup_path)
        print(f"ПРОГРЕСС: backup saved | {backup_path}", flush=True)
        for index, (marketplace, field, marketplace_rows) in enumerate(
            (("wb", "wb_nmid", wb_rows), ("ozon", "ozon_sku", ozon_rows)),
            start=1,
        ):
            print(
                "ПРОГРЕСС: "
                f"{index}/2 ({index / 2:.0%}) | {marketplace.upper()} | "
                f"source_rows={len(marketplace_rows)} | elapsed={time.monotonic() - started:.1f}s",
                flush=True,
            )
            cur.execute("DROP TABLE IF EXISTS tmp_bts_season")
            cur.execute(
                """
                CREATE TEMP TABLE tmp_bts_season (
                    marketplace text NOT NULL,
                    sku text NOT NULL,
                    barcode text NOT NULL,
                    vendor_article text,
                    gj_model text,
                    assortment_bia text,
                    tg text,
                    tg_plus text,
                    cg text,
                    season text NOT NULL,
                    source_file text,
                    source_wb_row integer,
                    source_ozon_row integer
                ) ON COMMIT DROP
                """
            )
            execute_values(
                cur,
                """
                INSERT INTO tmp_bts_season (
                    marketplace, sku, barcode, vendor_article, gj_model,
                    assortment_bia, tg, tg_plus, cg, season, source_file,
                    source_wb_row, source_ozon_row
                ) VALUES %s
                """,
                [
                    (
                        row["marketplace"],
                        row["sku"],
                        row["barcode"],
                        row["vendor_article"],
                        row["gj_model"],
                        row["assortment_bia"],
                        row["tg"],
                        row["tg_plus"],
                        row["cg"],
                        row["season"],
                        row["source_file"],
                        row["source_wb_row"],
                        row["source_ozon_row"],
                    )
                    for row in marketplace_rows
                ],
            )
            cur.execute(
                f"""
                UPDATE public.sku_mapping_gj m
                SET season = %s
                WHERE m.{field} IN (SELECT DISTINCT sku FROM tmp_bts_season WHERE marketplace = %s)
                  AND coalesce(m.season, '') <> %s
                """,
                (TARGET_SEASON, marketplace, TARGET_SEASON),
            )
            result["updated"][marketplace] = cur.rowcount
            cur.execute(
                f"""
                SELECT DISTINCT t.*
                FROM tmp_bts_season t
                WHERE t.marketplace = %s
                  AND NOT EXISTS (
                      SELECT 1 FROM public.sku_mapping_gj m WHERE m.{field} = t.sku
                  )
                ORDER BY t.sku, t.barcode
                """,
                (marketplace,),
            )
            missing_rows = cur.fetchall()
            insert_values = []
            seen_keys = set()
            for row in missing_rows:
                (
                    row_marketplace,
                    sku,
                    barcode,
                    vendor_article,
                    gj_model,
                    assortment_bia,
                    tg,
                    tg_plus,
                    cg,
                    season,
                    source_file,
                    source_wb_row,
                    source_ozon_row,
                ) = row
                dedupe_key = (row_marketplace, sku, barcode)
                if dedupe_key in seen_keys:
                    continue
                seen_keys.add(dedupe_key)
                wb_article = vendor_article if marketplace == "wb" else None
                wb_nmid = sku if marketplace == "wb" else None
                ozon_sku = sku if marketplace == "ozon" else None
                mapping_hash = row_hash(
                    "back_to_school_2026",
                    row_marketplace,
                    sku,
                    barcode,
                    wb_article,
                    ozon_sku,
                )
                insert_values.append(
                    (
                        mapping_hash,
                        barcode,
                        wb_article,
                        wb_nmid,
                        ozon_sku,
                        gj_model,
                        assortment_bia,
                        tg,
                        tg_plus,
                        cg,
                        season,
                        source_file,
                        source_wb_row,
                        source_ozon_row,
                    )
                )
            if insert_values:
                execute_values(
                    cur,
                    """
                    INSERT INTO public.sku_mapping_gj (
                        mapping_hash, barcode, wb_article, wb_nmid, ozon_sku,
                        gj_model, assortment_bia, tg, tg_plus, cg, season,
                        source_file, source_wb_row, source_ozon_row
                    ) VALUES %s
                    ON CONFLICT (mapping_hash) DO UPDATE SET
                        season = EXCLUDED.season,
                        source_file = EXCLUDED.source_file,
                        imported_at = now()
                    """,
                    insert_values,
                )
            result["inserted"][marketplace] = len(insert_values)
        if dry_run:
            conn.rollback()
        else:
            refresh_mapping_views(cur)
            conn.commit()
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Sync Back to School 2026 SKU mapping season.")
    parser.add_argument("--source", type=Path, default=WORKBOOK_PATH)
    parser.add_argument("--dry-run", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    print(
        "ПЛАН: прочитать ВБ/ОЗОН из Excel | выставить season=Back to school 2026 | "
        "обновить sku_mapping_gj | refresh mapping views",
        flush=True,
    )
    rows, summary = extract_rows(args.source)
    result = sync_rows(rows, args.dry_run)
    payload = {
        "source": str(args.source),
        "dry_run": args.dry_run,
        "target_season": TARGET_SEASON,
        "summary": summary,
        "result": result,
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2), flush=True)
    print(
        "ИТОГ: "
        f"mode={'dry-run' if args.dry_run else 'write'} | "
        f"updated={result['updated']} | inserted={result['inserted']} | backup={result['backup_path']}",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
