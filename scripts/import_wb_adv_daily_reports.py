#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Import Wildberries product advertising monthly XLSX reports into PostgreSQL."""

from __future__ import annotations

import argparse
import csv
import io
import re
import sys
import time
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DASHBOARD_ROOT = PROJECT_ROOT / "ozon_category_dashboard"
sys.path.insert(0, str(DASHBOARD_ROOT))

import app  # noqa: E402
from wb_advertising_view import rebuild_wb_advertising_view  # noqa: E402


DEFAULT_SOURCE_DIR = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Data\Wb\Adv"
)
SHEET_NAME = "Сводный отчет по рекламе"


RAW_COLUMNS = [
    "report_date",
    "wb_marketplace_article",
    "seller_article",
    "impressions",
    "clicks",
    "ctr_pct",
    "expense_rub",
    "fact_expense_rub",
    "cpc_rub",
    "cpm_rub",
    "added_to_cart",
    "orders_qty",
    "cr_pct",
    "orders_amount_rub",
    "cpa_rub",
    "drr_pct",
    "total_orders_qty",
    "total_orders_amount_rub",
    "total_drr_pct",
    "total_cpa_rub",
    "source_file",
    "source_sheet",
    "source_row_num",
    "file_size_bytes",
    "file_mtime",
]


COLUMN_MAP = {
    "день": "report_date",
    "артикул маркетплейса": "wb_marketplace_article",
    "артикул продавца": "seller_article",
    "показы": "impressions",
    "клики": "clicks",
    "ctr %": "ctr_pct",
    "расход ₽": "expense_rub",
    "расход факт ₽": "fact_expense_rub",
    "cpc ₽": "cpc_rub",
    "cpm ₽ средняя": "cpm_rub",
    "добавлено в корзину": "added_to_cart",
    "оформлено заказов": "orders_qty",
    "cr %": "cr_pct",
    "сумма заказов ₽": "orders_amount_rub",
    "cpa ₽": "cpa_rub",
    "дрр %": "drr_pct",
    "общее кол во заказов": "total_orders_qty",
    "общая сумма заказов ₽": "total_orders_amount_rub",
    "общая дрр %": "total_drr_pct",
    "общий cpa ₽": "total_cpa_rub",
}

NUMERIC_COLUMNS = {
    "impressions",
    "clicks",
    "ctr_pct",
    "expense_rub",
    "fact_expense_rub",
    "cpc_rub",
    "cpm_rub",
    "added_to_cart",
    "orders_qty",
    "cr_pct",
    "orders_amount_rub",
    "cpa_rub",
    "drr_pct",
    "total_orders_qty",
    "total_orders_amount_rub",
    "total_drr_pct",
    "total_cpa_rub",
}


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if text.lower() in {"", "nan", "none", "null", "-", "–", "—"}:
        return None
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def normalize_header(value: Any) -> str:
    text = clean_text(value) or ""
    text = text.replace("\xa0", " ").replace("\n", " ").replace("\r", " ")
    text = text.replace("ё", "е").lower()
    text = re.sub(r"[(),;:]", " ", text)
    text = re.sub(r"[-–—]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def to_decimal(value: Any) -> Decimal:
    text = clean_text(value)
    if text is None:
        return Decimal("0")
    text = text.replace("\xa0", "").replace(" ", "").replace(",", ".").rstrip("%")
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text:
        return Decimal("0")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0")


def to_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if text is None:
        raise ValueError("empty report date")
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Cannot parse report date: {value!r}")


def source_file_value(source_dir: Path, file_path: Path) -> str:
    try:
        return str(file_path.relative_to(source_dir))
    except ValueError:
        return str(file_path)


def header_positions(header_row: tuple[Any, ...], file_path: Path) -> dict[str, int]:
    positions: dict[str, int] = {}
    for index, value in enumerate(header_row):
        db_name = COLUMN_MAP.get(normalize_header(value))
        if db_name and db_name not in positions:
            positions[db_name] = index
    missing = [db_name for db_name in RAW_COLUMNS[:20] if db_name not in positions]
    if missing:
        raise ValueError(f"Missing WB adv columns in {file_path}: {missing}")
    return positions


def read_wb_adv_rows(source_dir: Path, file_path: Path):
    stat = file_path.stat()
    source_file = source_file_value(source_dir, file_path)
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Sheet {SHEET_NAME!r} not found in {file_path}")
        sheet = workbook[SHEET_NAME]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return
        positions = header_positions(header, file_path)
        for row_num, row in enumerate(rows, start=2):
            wb_article = clean_text(row[positions["wb_marketplace_article"]])
            seller_article = clean_text(row[positions["seller_article"]])
            if not wb_article and not seller_article:
                continue
            parsed: dict[str, Any] = {}
            for column in RAW_COLUMNS[:20]:
                value = row[positions[column]]
                if column == "report_date":
                    parsed[column] = to_date(value)
                elif column in NUMERIC_COLUMNS:
                    parsed[column] = to_decimal(value)
                else:
                    parsed[column] = clean_text(value)
            parsed.update(
                {
                    "source_file": source_file,
                    "source_sheet": SHEET_NAME,
                    "source_row_num": row_num,
                    "file_size_bytes": stat.st_size,
                    "file_mtime": datetime.fromtimestamp(stat.st_mtime),
                }
            )
            yield tuple(parsed[column] for column in RAW_COLUMNS)
    finally:
        workbook.close()


class WbAdvImporter:
    def __init__(
        self,
        source_dir: Path,
        *,
        dry_run: bool = False,
        force_reimport: bool = False,
        max_files: int | None = None,
        skip_import: bool = False,
        skip_views: bool = False,
    ) -> None:
        self.source_dir = source_dir
        self.dry_run = dry_run
        self.force_reimport = force_reimport
        self.max_files = max_files
        self.skip_import = skip_import
        self.skip_views = skip_views
        self.conn = None
        self.cur = None

    def connect(self) -> None:
        self.conn = app.get_conn()
        self.cur = self.conn.cursor()

    def close(self) -> None:
        if self.cur is not None:
            self.cur.close()
        if self.conn is not None:
            self.conn.close()

    def execute(self, query: str, params: tuple[Any, ...] | None = None) -> None:
        assert self.cur is not None
        self.cur.execute(query, params)

    def create_schema(self, *, drop_views: bool = True) -> None:
        if drop_views:
            self.execute("DROP MATERIALIZED VIEW IF EXISTS public.mv_wb_adv_daily_by_article_category")
        self.execute(
            """
            CREATE TABLE IF NOT EXISTS public.wb_adv_daily_raw (
                id bigserial PRIMARY KEY,
                report_date date NOT NULL,
                wb_marketplace_article text,
                seller_article text,
                impressions numeric,
                clicks numeric,
                ctr_pct numeric,
                expense_rub numeric,
                fact_expense_rub numeric,
                cpc_rub numeric,
                cpm_rub numeric,
                added_to_cart numeric,
                orders_qty numeric,
                cr_pct numeric,
                orders_amount_rub numeric,
                cpa_rub numeric,
                drr_pct numeric,
                total_orders_qty numeric,
                total_orders_amount_rub numeric,
                total_drr_pct numeric,
                total_cpa_rub numeric,
                source_file text NOT NULL,
                source_sheet text NOT NULL,
                source_row_num integer NOT NULL,
                file_size_bytes bigint,
                file_mtime timestamp,
                imported_at timestamp NOT NULL DEFAULT now(),
                UNIQUE (source_file, source_sheet, source_row_num)
            )
            """
        )
        self.execute(
            """
            CREATE TABLE IF NOT EXISTS public.wb_adv_daily_import_files (
                source_file text PRIMARY KEY,
                rows_imported integer NOT NULL DEFAULT 0,
                date_from date,
                date_to date,
                file_size_bytes bigint,
                file_mtime timestamp,
                imported_at timestamp NOT NULL DEFAULT now(),
                status text NOT NULL DEFAULT 'ok',
                error text
            )
            """
        )
        if self.force_reimport:
            self.execute("TRUNCATE TABLE public.wb_adv_daily_raw RESTART IDENTITY")
            self.execute("TRUNCATE TABLE public.wb_adv_daily_import_files")
        self.conn.commit()

    def imported_file_state(self) -> dict[str, tuple[int | None, datetime | None, str | None]]:
        self.execute(
            """
            SELECT source_file, file_size_bytes, file_mtime, status
            FROM public.wb_adv_daily_import_files
            """
        )
        return {
            row["source_file"]: (row["file_size_bytes"], row["file_mtime"], row["status"])
            for row in self.cur.fetchall()
        }

    def should_skip_file(self, file_path: Path, imported: dict[str, tuple[int | None, datetime | None, str | None]]) -> bool:
        if self.force_reimport:
            return False
        source_file = source_file_value(self.source_dir, file_path)
        state = imported.get(source_file)
        if not state:
            return False
        size, mtime, status = state
        stat = file_path.stat()
        return status == "ok" and size == stat.st_size and mtime and abs(mtime.timestamp() - stat.st_mtime) < 1

    def import_file(self, file_path: Path) -> int:
        assert self.cur is not None and self.conn is not None
        source_file = source_file_value(self.source_dir, file_path)
        stat = file_path.stat()
        try:
            rows = list(read_wb_adv_rows(self.source_dir, file_path))
            if self.dry_run:
                return len(rows)
            self.execute("DELETE FROM public.wb_adv_daily_raw WHERE source_file = %s", (source_file,))
            if rows:
                buffer = io.StringIO()
                writer = csv.writer(buffer, lineterminator="\n")
                writer.writerows(rows)
                buffer.seek(0)
                self.cur.copy_expert(
                    f"COPY public.wb_adv_daily_raw ({', '.join(RAW_COLUMNS)}) FROM STDIN WITH CSV",
                    buffer,
                )
            dates = [row[0] for row in rows]
            self.execute(
                """
                INSERT INTO public.wb_adv_daily_import_files (
                    source_file, rows_imported, date_from, date_to, file_size_bytes,
                    file_mtime, imported_at, status, error
                )
                VALUES (%s, %s, %s, %s, %s, %s, now(), 'ok', NULL)
                ON CONFLICT (source_file) DO UPDATE SET
                    rows_imported = EXCLUDED.rows_imported,
                    date_from = EXCLUDED.date_from,
                    date_to = EXCLUDED.date_to,
                    file_size_bytes = EXCLUDED.file_size_bytes,
                    file_mtime = EXCLUDED.file_mtime,
                    imported_at = now(),
                    status = 'ok',
                    error = NULL
                """,
                (
                    source_file,
                    len(rows),
                    min(dates) if dates else None,
                    max(dates) if dates else None,
                    stat.st_size,
                    datetime.fromtimestamp(stat.st_mtime),
                ),
            )
            self.conn.commit()
            return len(rows)
        except Exception as exc:
            if not self.dry_run:
                self.conn.rollback()
                self.execute(
                    """
                    INSERT INTO public.wb_adv_daily_import_files (
                        source_file, rows_imported, file_size_bytes, file_mtime,
                        imported_at, status, error
                    )
                    VALUES (%s, 0, %s, %s, now(), 'error', %s)
                    ON CONFLICT (source_file) DO UPDATE SET
                        rows_imported = 0,
                        file_size_bytes = EXCLUDED.file_size_bytes,
                        file_mtime = EXCLUDED.file_mtime,
                        imported_at = now(),
                        status = 'error',
                        error = EXCLUDED.error
                    """,
                    (source_file, stat.st_size, datetime.fromtimestamp(stat.st_mtime), str(exc)[:2000]),
                )
                self.conn.commit()
            print(f"ERROR {file_path}: {exc}")
            return 0

    def rebuild_view(self) -> None:
        self.execute("DROP MATERIALIZED VIEW IF EXISTS public.mv_wb_adv_daily_by_article_category")
        self.execute(
            """
            CREATE MATERIALIZED VIEW public.mv_wb_adv_daily_by_article_category AS
            WITH funnel_products AS (
                SELECT
                    sku::text AS wb_marketplace_article,
                    max(seller_article) AS seller_article,
                    max(product_name) AS product_name,
                    max(product_artikul) AS product_artikul,
                    max(category_name) AS category_name,
                    max(category_level_2) AS subcategory_name,
                    max(brand) AS brand,
                    max(barcode) AS barcode
                FROM public.mv_wb_funnel_daily_by_article_category
                GROUP BY sku::text
            )
            SELECT
                r.report_date,
                r.wb_marketplace_article,
                r.wb_marketplace_article AS ozon_marketplace_article,
                r.wb_marketplace_article AS sku,
                r.seller_article,
                coalesce(fp.product_artikul, r.seller_article, r.wb_marketplace_article) AS product_artikul,
                coalesce(fp.product_name, r.seller_article, r.wb_marketplace_article) AS product_name,
                coalesce(fp.category_name, 'Без категории') AS category_name,
                fp.subcategory_name,
                fp.brand,
                fp.barcode,
                r.impressions,
                r.clicks,
                r.ctr_pct,
                r.expense_rub,
                r.fact_expense_rub,
                r.cpc_rub,
                r.cpm_rub,
                r.added_to_cart,
                r.orders_qty,
                r.cr_pct,
                r.orders_amount_rub,
                r.cpa_rub,
                r.drr_pct,
                r.total_orders_qty,
                r.total_orders_amount_rub,
                r.total_drr_pct,
                r.total_cpa_rub,
                r.source_file,
                r.source_sheet,
                r.source_row_num,
                r.file_size_bytes,
                r.file_mtime,
                r.imported_at
            FROM public.wb_adv_daily_raw r
            LEFT JOIN funnel_products fp
                ON fp.wb_marketplace_article = r.wb_marketplace_article
            """
        )
        # После ручного импорта не откатываем уже подключённые API-даты.
        assert self.cur is not None
        rebuild_wb_advertising_view(self.cur)
        for index_sql in [
            "CREATE INDEX IF NOT EXISTS idx_wb_adv_raw_source ON public.wb_adv_daily_raw(source_file, source_sheet, source_row_num)",
            "CREATE INDEX IF NOT EXISTS idx_wb_adv_raw_date ON public.wb_adv_daily_raw(report_date)",
            "CREATE INDEX IF NOT EXISTS idx_wb_adv_raw_article ON public.wb_adv_daily_raw(wb_marketplace_article)",
            "CREATE INDEX IF NOT EXISTS idx_mv_wb_adv_date ON public.mv_wb_adv_daily_by_article_category(report_date)",
            "CREATE INDEX IF NOT EXISTS idx_mv_wb_adv_article ON public.mv_wb_adv_daily_by_article_category(wb_marketplace_article)",
            "CREATE INDEX IF NOT EXISTS idx_mv_wb_adv_category ON public.mv_wb_adv_daily_by_article_category(category_name)",
            "CREATE INDEX IF NOT EXISTS idx_mv_wb_adv_expense ON public.mv_wb_adv_daily_by_article_category(expense_rub)",
            "CREATE INDEX IF NOT EXISTS idx_mv_wb_adv_orders_amount ON public.mv_wb_adv_daily_by_article_category(orders_amount_rub)",
        ]:
            self.execute(index_sql)
        self.execute("ANALYZE public.wb_adv_daily_raw")
        self.execute("ANALYZE public.mv_wb_adv_daily_by_article_category")
        self.conn.commit()

    def check(self) -> None:
        self.execute(
            """
            SELECT
                count(*) AS rows_count,
                count(DISTINCT report_date) AS days_count,
                min(report_date) AS date_from,
                max(report_date) AS date_to,
                count(DISTINCT wb_marketplace_article) AS sku_count,
                round(coalesce(sum(expense_rub), 0), 2) AS expense_rub,
                round(coalesce(sum(orders_amount_rub), 0), 2) AS orders_amount_rub
            FROM public.mv_wb_adv_daily_by_article_category
            """
        )
        print(dict(self.cur.fetchone()))

    def run(self) -> None:
        files: list[Path] = []
        if not self.skip_import:
            files = sorted(self.source_dir.rglob("*.xlsx"))
            if self.max_files:
                files = files[: self.max_files]
            if not files:
                raise RuntimeError(f"No XLSX files found in {self.source_dir}")
            print(f"Files found: {len(files)}")
        else:
            print("Skipping file import; rebuilding WB adv materialized view from existing raw data.")
        self.connect()
        try:
            self.create_schema(drop_views=not self.skip_views)
            if not self.skip_import and not self.force_reimport:
                imported = self.imported_file_state()
                original_count = len(files)
                files = [file for file in files if not self.should_skip_file(file, imported)]
                print(f"Already imported skipped: {original_count - len(files)}")
            total = 0
            started = time.perf_counter()
            if not self.skip_import:
                for file_path in files:
                    rows_count = self.import_file(file_path)
                    total += rows_count
                    print(f"Imported {file_path.name}: {rows_count:,} rows")
            if not self.dry_run:
                if self.skip_views:
                    self.conn.commit()
                    print("Materialized views deferred; run view refresh after imports.")
                else:
                    print("Rebuilding WB adv materialized view...")
                    self.rebuild_view()
                    self.check()
            print(f"Imported rows: {total:,} | seconds={time.perf_counter() - started:.1f}")
        finally:
            self.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--force-reimport", action="store_true")
    parser.add_argument("--max-files", type=int, default=None)
    parser.add_argument("--skip-import", action="store_true")
    parser.add_argument("--skip-views", action="store_true")
    args = parser.parse_args()
    WbAdvImporter(
        args.source_dir,
        dry_run=args.dry_run,
        force_reimport=args.force_reimport,
        max_files=args.max_files,
        skip_import=args.skip_import,
        skip_views=args.skip_views,
    ).run()


if __name__ == "__main__":
    main()
