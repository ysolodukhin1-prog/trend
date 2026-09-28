#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import io
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook


def non_empty_count(row: tuple[object, ...]) -> int:
    return sum(value is not None and str(value).strip() != "" for value in row)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("--rows", type=int, default=8)
    args = parser.parse_args()

    with zipfile.ZipFile(args.path) as archive:
        names = [name for name in archive.namelist() if name.lower().endswith(".xlsx")]
        if not names:
            raise FileNotFoundError(f"No .xlsx files in {args.path}")
        name = names[0]
        print(f"file={name}")
        workbook = load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=True)

    try:
        print("sheets:", workbook.sheetnames)
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            ws.reset_dimensions()
            print(f"=== {sheet_name} ===")
            printed = 0
            for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if non_empty_count(row) < 2:
                    continue
                print(row_index, row)
                printed += 1
                if printed >= args.rows:
                    break
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
