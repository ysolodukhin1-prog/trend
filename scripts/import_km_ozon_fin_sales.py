#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import psycopg2
from openpyxl import load_workbook
from psycopg2.extras import execute_values

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "ozon_category_dashboard"))

import app  # noqa: E402

SOURCE_DIR = Path(
    os.environ.get(
        "KM_OZON_FIN_DIR",
        r"G:\Общие диски\Kokoc Marketplaces\Clients\KM Trade\Аналитика\Дашборды\Data\Ozon\Fin",
    )
)
PRODUCTS_DIR = Path(
    os.environ.get(
        "KM_OZON_PRODUCT_CATEGORIES_DIR",
        r"G:\Общие диски\Kokoc Marketplaces\Clients\KM Trade\Аналитика\Дашборды\Data\Ozon\Product_catigories",
    )
)
TARGET_DB = os.environ.get("KM_DB_NAME", "km_trade_products")
OZON_PRODUCTS_COLUMNS = [
    "artikul",
    "ozon_product_id",
    "sku",
    "barcode",
    "nazvanie_tovara",
    "kontent_reyting",
    "brend",
    "status_tovara",
    "metki",
    "otzyvy",
    "reyting",
    "prichiny_skrytiya",
    "data_sozdaniya",
    "kategoriya",
    "tip",
    "obem_tovara_l",
    "obemnyy_ves_kg",
    "dostupno_k_prodazhe_po_sheme_fbo_sht",
    "zarezervirovano_sht",
    "dostupno_k_prodazhe_po_sheme_fbs_sht",
    "dostupno_k_prodazhe_po_sheme_realfbs_sht",
    "zarezervirovano_na_moih_skladah_sht",
    "tekuschaya_tsena_s_uchetom_skidki_rub",
    "tsena_do_skidki_perecherknutaya_tsena_rub",
    "tsena_premium_rub",
    "razmer_nds_pct",
    "oshibki",
    "preduprezhdeniya",
    "import_file",
]
OZON_PRODUCTS_HEADER_MAP = {
    "артикул": "artikul",
    "ozon product id": "ozon_product_id",
    "sku": "sku",
    "barcode": "barcode",
    "название товара": "nazvanie_tovara",
    "контент-рейтинг": "kontent_reyting",
    "бренд": "brend",
    "статус товара": "status_tovara",
    "метки": "metki",
    "отзывы": "otzyvy",
    "рейтинг": "reyting",
    "причины скрытия": "prichiny_skrytiya",
    "дата создания": "data_sozdaniya",
    "категория": "kategoriya",
    "тип": "tip",
    "объем товара, л": "obem_tovara_l",
    "объемный вес, кг": "obemnyy_ves_kg",
    "доступно к продаже по схеме fbo, шт.": "dostupno_k_prodazhe_po_sheme_fbo_sht",
    "зарезервировано, шт": "zarezervirovano_sht",
    "доступно к продаже по схеме fbs, шт.": "dostupno_k_prodazhe_po_sheme_fbs_sht",
    "доступно к продаже по схеме realfbs, шт.": "dostupno_k_prodazhe_po_sheme_realfbs_sht",
    "зарезервировано на моих складах, шт.": "zarezervirovano_na_moih_skladah_sht",
    "текущая цена с учетом скидки, руб.": "tekuschaya_tsena_s_uchetom_skidki_rub",
    "цена до скидки (перечеркнутая цена), руб.": "tsena_do_skidki_perecherknutaya_tsena_rub",
    "цена premium, руб.": "tsena_premium_rub",
    "размер ндс, %": "razmer_nds_pct",
    "ошибки": "oshibki",
    "предупреждения": "preduprezhdeniya",
}


def db_config():
    return {**app.read_db_config(), "database": TARGET_DB}


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text[1:].strip() if text.startswith("'") else text


def as_float(value) -> float:
    text = clean(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def normalize_header(value) -> str:
    return re.sub(r"\s+", " ", clean(value)).lower()


def read_rows():
    if not SOURCE_DIR.exists():
        raise FileNotFoundError(SOURCE_DIR)
    rows = []
    for path in sorted(SOURCE_DIR.glob("*.xlsx")):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            headers = {normalize_header(value): idx for idx, value in enumerate(header_row)}

            def col(name):
                return headers.get(normalize_header(name))

            idx_type = col("Тип начисления")
            idx_sku = col("SKU")
            idx_article = col("Артикул")
            idx_name = col("Название товара или услуги")
            idx_qty = col("Количество")
            idx_amount = col("За продажу или возврат до вычета комиссий и услуг")
            for row in sheet.iter_rows(min_row=2, values_only=True):
                sku = clean(row[idx_sku]) if idx_sku is not None else ""
                article = clean(row[idx_article]) if idx_article is not None else ""
                if not sku and not article:
                    continue
                accrual_type = clean(row[idx_type]) if idx_type is not None else ""
                amount = as_float(row[idx_amount]) if idx_amount is not None else 0.0
                qty = as_float(row[idx_qty]) if idx_qty is not None else 0.0
                if amount <= 0 and "доставка покупателю" not in accrual_type.lower():
                    continue
                rows.append(
                    (
                        clean(row[idx_name]) if idx_name is not None else "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        sku,
                        article,
                        f"{amount:.2f}",
                        f"{qty:.0f}",
                        f"{qty:.0f}",
                        "0",
                    )
                )
        finally:
            workbook.close()
    return rows


def read_ozon_products_rows():
    rows = []
    for path in sorted(PRODUCTS_DIR.glob("Товары*.xlsx")):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            header_to_index = {normalize_header(value): idx for idx, value in enumerate(header_row)}
            for row in sheet.iter_rows(min_row=2, values_only=True):
                record = {column: "" for column in OZON_PRODUCTS_COLUMNS}
                record["import_file"] = path.name
                for header, column in OZON_PRODUCTS_HEADER_MAP.items():
                    idx = header_to_index.get(header)
                    if idx is not None and idx < len(row):
                        record[column] = clean(row[idx])
                if record["artikul"] or record["sku"] or record["ozon_product_id"]:
                    rows.append(tuple(record[column] for column in OZON_PRODUCTS_COLUMNS))
        finally:
            workbook.close()
    return rows


def ensure_attribute_definitions(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS public.ozon_attribute_definitions (
            id bigserial PRIMARY KEY,
            category text,
            category_id integer,
            attribute_id integer,
            attribute_name text,
            imported_at timestamp DEFAULT now()
        )
        """
    )
    cur.execute("TRUNCATE TABLE public.ozon_attribute_definitions RESTART IDENTITY")
    cur.execute(
        """
        INSERT INTO public.ozon_attribute_definitions (
            category, category_id, attribute_id, attribute_name
        )
        SELECT DISTINCT
            COALESCE(NULLIF(p.category_name, ''), ca.category_id::text) AS category,
            ca.category_id,
            ca.attribute_id,
            ca.attribute_name
        FROM public.ozon_cat_category_attributes ca
        LEFT JOIN (
            SELECT DISTINCT category_id, category_name
            FROM public.ozon_cat_products
        ) p ON p.category_id = ca.category_id
        WHERE NULLIF(ca.attribute_name, '') IS NOT NULL
        """
    )
    cur.execute(
        "CREATE INDEX IF NOT EXISTS idx_ozon_attribute_definitions_category "
        "ON public.ozon_attribute_definitions(category)"
    )
    cur.execute("ANALYZE public.ozon_attribute_definitions")


def main():
    rows = read_rows()
    with psycopg2.connect(**db_config()) as conn, conn.cursor() as cur:
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS public.ozon_products (
                id bigserial PRIMARY KEY,
                artikul text,
                ozon_product_id text,
                sku text,
                barcode text,
                nazvanie_tovara text,
                kontent_reyting text,
                brend text,
                status_tovara text,
                metki text,
                otzyvy text,
                reyting text,
                prichiny_skrytiya text,
                data_sozdaniya text,
                kategoriya text,
                tip text,
                obem_tovara_l text,
                obemnyy_ves_kg text,
                dostupno_k_prodazhe_po_sheme_fbo_sht text,
                zarezervirovano_sht text,
                dostupno_k_prodazhe_po_sheme_fbs_sht text,
                dostupno_k_prodazhe_po_sheme_realfbs_sht text,
                zarezervirovano_na_moih_skladah_sht text,
                tekuschaya_tsena_s_uchetom_skidki_rub text,
                tsena_do_skidki_perecherknutaya_tsena_rub text,
                tsena_premium_rub text,
                razmer_nds_pct text,
                oshibki text,
                preduprezhdeniya text,
                import_file varchar(500),
                imported_at timestamp DEFAULT now()
            )
            """
        )
        cur.execute("TRUNCATE TABLE public.ozon_products RESTART IDENTITY")
        cur.execute(
            """
            INSERT INTO public.ozon_products (
                artikul, ozon_product_id, sku, barcode, nazvanie_tovara, brend,
                kategoriya, tip, tekuschaya_tsena_s_uchetom_skidki_rub,
                tsena_do_skidki_perecherknutaya_tsena_rub, razmer_nds_pct, import_file
            )
            SELECT
                artikul,
                p.product_id::text,
                sku,
                shtrihkod_seriynyy_nomer_ean,
                nazvanie_tovara,
                NULLIF(MAX(value_text) FILTER (WHERE attribute_id = 5274), ''),
                category_name,
                tip,
                tsena_rub,
                tsena_do_skidki_rub,
                nds_pct,
                import_file
            FROM public.ozon_cat_products p
            LEFT JOIN public.ozon_cat_product_attributes pa ON pa.product_id = p.product_id
            GROUP BY
                p.product_id, p.artikul, p.sku, p.shtrihkod_seriynyy_nomer_ean,
                p.nazvanie_tovara, p.category_name, p.tip, p.tsena_rub,
                p.tsena_do_skidki_rub, p.nds_pct, p.import_file
            """
        )
        product_count = cur.rowcount
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ozon_products_sku ON public.ozon_products(sku)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ozon_products_artikul ON public.ozon_products(artikul)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ozon_products_barcode ON public.ozon_products(barcode)")
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS public.ozon_stock_sales_90d (
                id bigserial PRIMARY KEY,
                tovary text,
                kategoriya_1_urovnya text,
                kategoriya_2_urovnya text,
                kategoriya_3_urovnya text,
                brend text,
                model text,
                shema_raboty text,
                sku text,
                artikul text,
                zakazano_na_summu text,
                voronka_prodazh_zakazano_tovarov text,
                vykupleno_tovarov text,
                faktory_prodazh_ostatok_na_konets_perioda text
            )
            """
        )
        cur.execute("TRUNCATE TABLE public.ozon_stock_sales_90d RESTART IDENTITY")
        if rows:
            execute_values(
                cur,
                """
                INSERT INTO public.ozon_stock_sales_90d (
                    tovary, kategoriya_1_urovnya, kategoriya_2_urovnya, kategoriya_3_urovnya,
                    brend, model, shema_raboty, sku, artikul, zakazano_na_summu,
                    voronka_prodazh_zakazano_tovarov, vykupleno_tovarov,
                    faktory_prodazh_ostatok_na_konets_perioda
                )
                VALUES %s
                """,
                rows,
                page_size=5000,
            )
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ozon_stock_sales_90d_sku ON public.ozon_stock_sales_90d(sku)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ozon_stock_sales_90d_artikul ON public.ozon_stock_sales_90d(artikul)")
        ensure_attribute_definitions(cur)
        cur.execute("ANALYZE public.ozon_stock_sales_90d")
        cur.execute("ANALYZE public.ozon_products")
        conn.commit()
    print(f"Done km_ozon_fin_sales: rows={len(rows)}, ozon_products={product_count}, errors=0")


if __name__ == "__main__":
    main()
