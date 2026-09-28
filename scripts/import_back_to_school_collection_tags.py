#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import json
import sys
import time
from collections import OrderedDict
from pathlib import Path

import openpyxl
import psycopg2
from psycopg2.extras import execute_values


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "ozon_category_dashboard"))

import app  # noqa: E402


WORKBOOK_PATH = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Back to School 2026 МП 01 07.xlsx"
)
SOURCE_TAGS = ("школа", "околошкола")
SHEET_CONFIG = {
    "ВБ": {
        "marketplace": "wb",
        "sku_header": "Артикул",
        "name_header": "Модель ГД",
        "tag_header": "признак",
        "match_view": "mv_product_abc_wb",
    },
    "ОЗОН": {
        "marketplace": "ozon",
        "sku_header": "SKU (Sku)",
        "name_header": "Модель ГД",
        "tag_header": "признак",
        "match_view": "mv_product_abc_ozon",
    },
}


def clean_text(value: object) -> str:
    return str(value or "").strip()


def clean_sku(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def header_map(sheet) -> dict[str, int]:
    return {
        clean_text(value).lower(): index
        for index, value in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        if clean_text(value)
    }


def ensure_required_headers(sheet_name: str, headers: dict[str, int], required: list[str]) -> None:
    missing = [header for header in required if header.lower() not in headers]
    if missing:
        raise RuntimeError(f"{sheet_name}: missing required columns: {', '.join(missing)}")


def extract_rows(workbook_path: Path) -> tuple[list[tuple[str, str, str, str, str]], dict[str, dict[str, object]]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    rows_by_key: OrderedDict[tuple[str, str], tuple[str, str, str, str, str]] = OrderedDict()
    sheet_stats: dict[str, dict[str, object]] = {}

    for sheet_name, config in SHEET_CONFIG.items():
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(f"Workbook sheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
        headers = header_map(sheet)
        ensure_required_headers(
            sheet_name,
            headers,
            [config["sku_header"], config["tag_header"]],
        )
        sku_index = headers[config["sku_header"].lower()]
        tag_index = headers[config["tag_header"].lower()]
        name_index = headers.get(config["name_header"].lower())
        marketplace = config["marketplace"]
        sheet_rows = 0
        skipped_empty_sku = 0
        skipped_empty_tag = 0
        skipped_unknown_tag = 0
        conflicts: dict[str, set[str]] = {}

        for raw_row in sheet.iter_rows(min_row=2, values_only=True):
            sku = clean_sku(raw_row[sku_index] if sku_index < len(raw_row) else "")
            tag = clean_text(raw_row[tag_index] if tag_index < len(raw_row) else "").lower()
            if not sku:
                skipped_empty_sku += 1
                continue
            if not tag:
                skipped_empty_tag += 1
                continue
            if tag not in SOURCE_TAGS:
                skipped_unknown_tag += 1
                continue
            product_name = ""
            if name_index is not None and name_index < len(raw_row):
                product_name = clean_text(raw_row[name_index])
            key = (marketplace, sku)
            existing = rows_by_key.get(key)
            if existing and existing[2] != tag:
                conflicts.setdefault(sku, set()).update([existing[2], tag])
                continue
            rows_by_key[key] = (
                marketplace,
                sku,
                tag,
                product_name,
                workbook_path.name,
            )
            sheet_rows += 1

        if conflicts:
            sample = ", ".join(f"{sku}: {sorted(tags)}" for sku, tags in list(conflicts.items())[:10])
            raise RuntimeError(f"{sheet_name}: conflicting tags for the same SKU: {sample}")

        sheet_stats[marketplace] = {
            "sheet": sheet_name,
            "rows": sheet_rows,
            "skipped_empty_sku": skipped_empty_sku,
            "skipped_empty_tag": skipped_empty_tag,
            "skipped_unknown_tag": skipped_unknown_tag,
        }

    summary: dict[str, dict[str, object]] = {}
    for sheet_name, config in SHEET_CONFIG.items():
        marketplace = config["marketplace"]
        marketplace_rows = [row for row in rows_by_key.values() if row[0] == marketplace]
        tag_counts = {tag: sum(1 for row in marketplace_rows if row[2] == tag) for tag in SOURCE_TAGS}
        summary[marketplace] = {
            **sheet_stats[marketplace],
            "unique_sku": len(marketplace_rows),
            "unique_tag_counts": tag_counts,
        }

    return list(rows_by_key.values()), summary


def ensure_tag_table(cur) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS public.marketplace_sku_tags (
            marketplace text NOT NULL,
            sku text NOT NULL,
            tag text NOT NULL,
            product_name text,
            source_file text,
            tagged_at timestamp without time zone NOT NULL DEFAULT now(),
            PRIMARY KEY (marketplace, sku, tag)
        )
        """
    )
    cur.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_marketplace_sku_tags_tag
        ON public.marketplace_sku_tags (marketplace, tag)
        """
    )
    cur.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_marketplace_sku_tags_sku
        ON public.marketplace_sku_tags (marketplace, sku)
        """
    )


def update_tags(rows: list[tuple[str, str, str, str, str]], dry_run: bool) -> dict[str, object]:
    by_marketplace = {
        marketplace: [row for row in rows if row[0] == marketplace]
        for marketplace in ("wb", "ozon")
    }
    result: dict[str, object] = {"deleted": 0, "inserted": {}, "matched": {}, "missing": {}}
    started = time.monotonic()
    with psycopg2.connect(**app.read_db_config()) as conn, conn.cursor() as cur:
        ensure_tag_table(cur)
        for index, (marketplace, marketplace_rows) in enumerate(by_marketplace.items(), start=1):
            view_name = SHEET_CONFIG["ВБ" if marketplace == "wb" else "ОЗОН"]["match_view"]
            print(
                "ПРОГРЕСС: "
                f"{index}/2 ({index / 2:.0%}) | {marketplace.upper()} | "
                f"source_sku={len(marketplace_rows)} | elapsed={time.monotonic() - started:.1f}s",
                flush=True,
            )
            cur.execute("DROP TABLE IF EXISTS tmp_back_to_school_tags")
            cur.execute(
                """
                CREATE TEMP TABLE tmp_back_to_school_tags (
                    marketplace text NOT NULL,
                    sku text NOT NULL,
                    tag text NOT NULL,
                    product_name text,
                    source_file text
                ) ON COMMIT DROP
                """
            )
            execute_values(
                cur,
                """
                INSERT INTO tmp_back_to_school_tags
                    (marketplace, sku, tag, product_name, source_file)
                VALUES %s
                """,
                marketplace_rows,
            )
            cur.execute(
                f"""
                SELECT count(DISTINCT t.sku)
                FROM tmp_back_to_school_tags t
                JOIN public.{view_name} v ON v.artikul_wb::text = t.sku
                WHERE t.marketplace = %s
                """,
                (marketplace,),
            )
            matched = cur.fetchone()[0]
            result["matched"][marketplace] = matched
            result["missing"][marketplace] = len(marketplace_rows) - matched

            if dry_run:
                continue

            cur.execute(
                """
                DELETE FROM public.marketplace_sku_tags
                WHERE marketplace = %s AND tag = ANY(%s)
                """,
                (marketplace, list(SOURCE_TAGS)),
            )
            result["deleted"] = int(result["deleted"]) + cur.rowcount
            cur.execute(
                f"""
                INSERT INTO public.marketplace_sku_tags
                    (marketplace, sku, tag, product_name, source_file)
                SELECT DISTINCT t.marketplace, t.sku, t.tag, t.product_name, t.source_file
                FROM tmp_back_to_school_tags t
                JOIN public.{view_name} v ON v.artikul_wb::text = t.sku
                WHERE t.marketplace = %s
                ON CONFLICT (marketplace, sku, tag) DO UPDATE SET
                    product_name = EXCLUDED.product_name,
                    source_file = EXCLUDED.source_file,
                    tagged_at = now()
                """,
                (marketplace,),
            )
            result["inserted"][marketplace] = cur.rowcount
        if dry_run:
            conn.rollback()
        else:
            conn.commit()
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Import Gloria Jeans Back to School collection tags into marketplace_sku_tags."
    )
    parser.add_argument("--source", type=Path, default=WORKBOOK_PATH, help="Source XLSX workbook.")
    parser.add_argument("--dry-run", action="store_true", help="Read and match data without writing to DB.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.source.exists():
        raise FileNotFoundError(args.source)
    print(
        "ПЛАН: 2 листа Excel | теги: школа, околошкола | "
        "матчинг по mv_product_abc_wb/ozon | обновление только этих двух тегов",
        flush=True,
    )
    rows, summary = extract_rows(args.source)
    if not rows:
        raise RuntimeError(f"No tagged SKU rows found in {args.source}")
    result = update_tags(rows, args.dry_run)
    payload = {
        "source": str(args.source),
        "dry_run": args.dry_run,
        "summary": summary,
        "result": result,
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2), flush=True)
    print(
        "ИТОГ: "
        f"mode={'dry-run' if args.dry_run else 'write'} | "
        f"wb_matched={result['matched'].get('wb', 0)} | "
        f"ozon_matched={result['matched'].get('ozon', 0)} | "
        f"deleted={result['deleted']} | inserted={result['inserted']}",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
