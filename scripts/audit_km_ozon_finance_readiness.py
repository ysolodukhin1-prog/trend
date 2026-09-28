#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Read-only audit of KM Trade Ozon finance sources and database readiness."""

from __future__ import annotations

import os
import sys
import time
from pathlib import Path

import psycopg2
from openpyxl import load_workbook
from psycopg2.extras import RealDictCursor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "ozon_category_dashboard"))

import app  # noqa: E402


TARGET_DB = "km_trade_products"
FIN_DIR = Path(
    os.environ.get(
        "KM_OZON_FIN_DIR",
        r"G:\Общие диски\Kokoc Marketplaces\Clients\KM Trade\Аналитика\Дашборды\Data\Ozon\Fin",
    )
)


def connect_km():
    config = dict(app.read_db_config())
    config["database"] = TARGET_DB
    conn = psycopg2.connect(**config, cursor_factory=RealDictCursor)
    conn.autocommit = True
    with conn.cursor() as cur:
        cur.execute("SELECT current_database() AS name")
        actual = cur.fetchone()["name"]
    if actual != TARGET_DB:
        conn.close()
        raise RuntimeError(f"Защита клиента: ожидалась {TARGET_DB}, открылась {actual}")
    return conn


def normalized_header(value: object) -> str:
    return " ".join(str(value or "").replace("\n", " ").split())


def main() -> int:
    started = time.monotonic()
    files = sorted(FIN_DIR.glob("*.xlsx")) if FIN_DIR.exists() else []
    print(
        f"ПЛАН: KM Trade | БД {TARGET_DB} | файлов Fin {len(files)} | "
        "прочитать только имена файлов, листов и заголовков; "
        "проверить только схемы/диапазоны таблиц; запись запрещена.",
        flush=True,
    )
    header_sets: dict[tuple[str, ...], list[str]] = {}
    for index, path in enumerate(files, start=1):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            headers = tuple(
                header
                for header in (
                    normalized_header(value)
                    for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
                )
                if header
            )
            header_sets.setdefault(headers, []).append(path.name)
            print(
                f"[{index}/{len(files)}] {path.name}: лист={sheet.title!r}, "
                f"строк≈{sheet.max_row if sheet.max_row is not None else 'unknown'}, "
                f"колонок={len(headers)}",
                flush=True,
            )
            elapsed = time.monotonic() - started
            eta = elapsed / index * (len(files) - index) if index else 0
            print(
                f"ПРОГРЕСС: {index}/{len(files)} "
                f"({index / max(len(files), 1) * 100:.1f}%) | "
                f"{path.name} | наборов заголовков {len(header_sets)} | "
                f"elapsed={elapsed:.1f}s | ETA={eta:.1f}s",
                flush=True,
            )
        finally:
            workbook.close()

    for index, (headers, names) in enumerate(header_sets.items(), start=1):
        print(
            f"HEADER_SET {index}: files={len(names)} | " + " || ".join(headers),
            flush=True,
        )

    with connect_km() as conn, conn.cursor() as cur:
        cur.execute(
            """
            SELECT c.relname AS relation, c.relkind
            FROM pg_class c
            JOIN pg_namespace n ON n.oid = c.relnamespace
            WHERE n.nspname = 'public'
              AND (
                c.relname ILIKE '%fin%'
                OR c.relname ILIKE '%cost%'
                OR c.relname ILIKE '%expense%'
                OR c.relname ILIKE '%price%'
                OR c.relname ILIKE '%unit%'
                OR c.relname ILIKE '%profit%'
              )
            ORDER BY c.relname
            """
        )
        relations = cur.fetchall()
        for row in relations:
            relation = row["relation"]
            cur.execute(
                """
                SELECT column_name, data_type
                FROM information_schema.columns
                WHERE table_schema = 'public' AND table_name = %s
                ORDER BY ordinal_position
                """,
                (relation,),
            )
            columns = cur.fetchall()
            print(
                f"RELATION {relation} kind={row['relkind']} | "
                + ", ".join(f"{item['column_name']}:{item['data_type']}" for item in columns),
                flush=True,
            )

        cur.execute(
            """
            SELECT table_name, column_name
            FROM information_schema.columns
            WHERE table_schema = 'public'
              AND (
                column_name ILIKE '%cost%'
                OR column_name ILIKE '%себест%'
                OR column_name ILIKE '%expense%'
                OR column_name ILIKE '%commission%'
                OR column_name ILIKE '%profit%'
              )
            ORDER BY table_name, ordinal_position
            """
        )
        candidates = cur.fetchall()
        print(
            "COST_CANDIDATES: "
            + (
                ", ".join(f"{row['table_name']}.{row['column_name']}" for row in candidates)
                if candidates
                else "нет"
            ),
            flush=True,
        )

    elapsed = time.monotonic() - started
    print(
        f"ИТОГ: KM Trade | read_only=true | database={TARGET_DB} | "
        f"files={len(files)} | header_sets={len(header_sets)} | "
        f"finance_like_relations={len(relations)} | cost_candidates={len(candidates)} | "
        f"elapsed={elapsed:.1f}s | stopped=false | partial=false",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
