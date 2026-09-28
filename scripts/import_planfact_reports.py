#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Import marketplace plan/fact from primary exports into PostgreSQL views."""

from __future__ import annotations

import argparse
import sys
import time
from datetime import date, datetime
import re
from pathlib import Path
from typing import Any

import psycopg2
from psycopg2.extras import execute_values
from openpyxl import load_workbook
from tqdm import tqdm


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DASHBOARD_ROOT = PROJECT_ROOT / "ozon_category_dashboard"
sys.path.insert(0, str(DASHBOARD_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

import app  # noqa: E402
from planfact_primary_sources import (  # noqa: E402
    DEFAULT_OZON_SPEND_DIR,
    DEFAULT_WB_FIX_DIR,
    DEFAULT_WB_REPORT,
    build_ozon_month_rows,
    build_wb_month_rows,
    load_primary_rows,
    read_ozon_db,
    read_ozon_spend,
    read_wb_fix,
    read_wb_report,
)


DEFAULT_SOURCE_DIR = Path(
    r"G:\Общие диски\Kokoc Marketplaces\Clients\Gloria Jeans\Аналитика"
    r"\Дашборды\Data\Pl_F"
)

MONTHS_RU = {
    "январь": 1,
    "февраль": 2,
    "март": 3,
    "апрель": 4,
    "май": 5,
    "июнь": 6,
    "июль": 7,
    "август": 8,
    "сентябрь": 9,
    "октябрь": 10,
    "ноябрь": 11,
    "декабрь": 12,
}


def format_duration(seconds: float) -> str:
    seconds = max(0, int(seconds))
    minutes, secs = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours}h {minutes:02d}m {secs:02d}s"
    if minutes:
        return f"{minutes}m {secs:02d}s"
    return f"{secs}s"


# Служебные копии рядом с исходником: бэкапы правок и файлы блокировки Excel.
SOURCE_IGNORE_RE = re.compile(r"(\.backup[-.]|\.bak$|[-_ ]копия|\(\d+\)$)", re.I)


def is_service_copy(path: Path) -> bool:
    if path.name.startswith("~$"):
        return True
    return bool(SOURCE_IGNORE_RE.search(path.stem))


def find_source_file(source_dir: Path) -> Path:
    candidates = sorted(source_dir.glob("*.xlsx"))
    files = [path for path in candidates if not is_service_copy(path)]
    skipped = [path.name for path in candidates if is_service_copy(path)]
    if skipped:
        print(
            f"ПРЕДУПРЕЖДЕНИЕ: пропущены служебные копии в {source_dir}: {', '.join(skipped)}",
            flush=True,
        )
    if not files:
        raise FileNotFoundError(f"No .xlsx files found in {source_dir}")
    if len(files) > 1:
        names = ", ".join(path.name for path in files)
        raise RuntimeError(f"Expected one .xlsx file in {source_dir}, found: {names}")
    return files[0]


def normalize_marketplace(value: Any) -> str:
    text = str(value or "").strip().lower()
    if text in {"ozon", "озон"}:
        return "ozon"
    if text in {"wb", "вб", "wildberries", "wldberries"}:
        return "wb"
    raise ValueError(f"Unknown marketplace: {value!r}")


def to_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Cannot parse date: {value!r}")


def to_decimal(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".").strip()
    return float(text) if text else 0.0


def parse_month(value: Any, fallback_year: int) -> date:
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    text = str(value or "").strip().lower()
    month = MONTHS_RU.get(text)
    if not month:
        for name, number in MONTHS_RU.items():
            if name in text:
                month = number
                break
    if not month:
        raise ValueError(f"Cannot parse month: {value!r}")
    return date(fallback_year, month, 1)


def read_rows(file_path: Path) -> tuple[list[tuple[Any, ...]], list[tuple[Any, ...]]]:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    if "Data" not in wb.sheetnames or "Plan" not in wb.sheetnames:
        raise RuntimeError(f"Workbook must contain sheets Data and Plan: {file_path}")

    data_ws = wb["Data"]
    daily_rows: list[tuple[Any, ...]] = []
    data_dates: list[date] = []
    for row_num, row in enumerate(data_ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        marketplace = normalize_marketplace(row[0])
        report_date = to_date(row[1])
        data_dates.append(report_date)
        daily_rows.append(
            (
                marketplace,
                report_date,
                to_decimal(row[2]),
                to_decimal(row[3]),
                to_decimal(row[4]),
                str(file_path),
                "Data",
                row_num,
            )
        )

    fallback_year = min(data_dates).year if data_dates else datetime.now().year
    plan_ws = wb["Plan"]
    plan_rows: list[tuple[Any, ...]] = []
    for row_num, row in enumerate(plan_ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        plan_rows.append(
            (
                normalize_marketplace(row[0]),
                parse_month(row[1], fallback_year),
                to_decimal(row[2]),
                to_decimal(row[3]),
                str(file_path),
                "Plan",
                row_num,
            )
        )
    wb.close()
    return dedupe_rows_by_key(daily_rows, 2), dedupe_rows_by_key(plan_rows, 2)


def dedupe_rows_by_key(rows: list[tuple[Any, ...]], key_width: int) -> list[tuple[Any, ...]]:
    deduped: dict[tuple[Any, ...], tuple[Any, ...]] = {}
    for row in rows:
        key = row[:key_width]
        deduped.pop(key, None)
        deduped[key] = row
    return list(deduped.values())


def parse_month_start(value: str | None) -> date:
    if not value:
        return date.today().replace(day=1)
    try:
        return datetime.strptime(value, "%Y-%m").date().replace(day=1)
    except ValueError as exc:
        raise ValueError(f"Month must have YYYY-MM format: {value!r}") from exc


def next_month_start(month_start: date) -> date:
    return date(
        month_start.year + (1 if month_start.month == 12 else 0),
        1 if month_start.month == 12 else month_start.month + 1,
        1,
    )


DDL = """
CREATE TABLE IF NOT EXISTS public.planfact_daily (
    marketplace text NOT NULL,
    report_date date NOT NULL,
    orders_rub numeric(18,2) NOT NULL DEFAULT 0,
    sales_rub numeric(18,2) NOT NULL DEFAULT 0,
    ad_spend_rub numeric(18,2) NOT NULL DEFAULT 0,
    source_file text,
    source_sheet text,
    source_row_num integer,
    imported_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (marketplace, report_date)
);

CREATE TABLE IF NOT EXISTS public.planfact_plan (
    marketplace text NOT NULL,
    plan_month date NOT NULL,
    sales_plan_rub numeric(18,2) NOT NULL DEFAULT 0,
    ad_spend_plan_rub numeric(18,2) NOT NULL DEFAULT 0,
    source_file text,
    source_sheet text,
    source_row_num integer,
    imported_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (marketplace, plan_month)
);
"""


DROP_VIEWS = """
DROP MATERIALIZED VIEW IF EXISTS public.mv_planfact_daily CASCADE;
DROP MATERIALIZED VIEW IF EXISTS public.mv_planfact_monthly CASCADE;
"""


CREATE_VIEWS = """
CREATE MATERIALIZED VIEW public.mv_planfact_daily AS
WITH daily AS (
    SELECT
        d.*,
        date_trunc('month', d.report_date)::date AS plan_month,
        extract(day from d.report_date)::int AS day_of_month,
        extract(day from (date_trunc('month', d.report_date) + interval '1 month - 1 day'))::int AS days_in_month
    FROM public.planfact_daily d
),
joined AS (
    SELECT
        d.marketplace,
        CASE d.marketplace WHEN 'ozon' THEN 'Ozon' WHEN 'wb' THEN 'WB' ELSE d.marketplace END AS marketplace_label,
        d.report_date,
        d.plan_month,
        d.orders_rub,
        d.sales_rub,
        d.ad_spend_rub,
        coalesce(p.sales_plan_rub, 0) AS sales_plan_rub,
        coalesce(p.ad_spend_plan_rub, 0) AS ad_spend_plan_rub,
        round(coalesce(p.sales_plan_rub, 0) / nullif(d.days_in_month, 0), 2) AS sales_plan_daily_rub,
        round(coalesce(p.ad_spend_plan_rub, 0) / nullif(d.days_in_month, 0), 2) AS ad_spend_plan_daily_rub,
        round(coalesce(p.sales_plan_rub, 0) / nullif(d.days_in_month, 0) * d.day_of_month, 2) AS sales_plan_elapsed_rub,
        round(coalesce(p.ad_spend_plan_rub, 0) / nullif(d.days_in_month, 0) * d.day_of_month, 2) AS ad_spend_plan_elapsed_rub
    FROM daily d
    LEFT JOIN public.planfact_plan p
        ON p.marketplace = d.marketplace
        AND p.plan_month = d.plan_month
),
cumulative AS (
    SELECT
        j.*,
        sum(j.orders_rub) OVER (PARTITION BY j.marketplace, j.plan_month ORDER BY j.report_date) AS orders_cum_rub,
        sum(j.sales_rub) OVER (PARTITION BY j.marketplace, j.plan_month ORDER BY j.report_date) AS sales_cum_rub,
        sum(j.ad_spend_rub) OVER (PARTITION BY j.marketplace, j.plan_month ORDER BY j.report_date) AS ad_spend_cum_rub
    FROM joined j
)
SELECT
    *,
    CASE WHEN sales_plan_rub <> 0 THEN round(sales_cum_rub / sales_plan_rub * 100, 2) ELSE 0 END AS sales_month_plan_fact_pct,
    CASE WHEN sales_plan_elapsed_rub <> 0 THEN round(sales_cum_rub / sales_plan_elapsed_rub * 100, 2) ELSE 0 END AS sales_elapsed_plan_fact_pct,
    CASE WHEN ad_spend_plan_rub <> 0 THEN round(ad_spend_cum_rub / ad_spend_plan_rub * 100, 2) ELSE 0 END AS ad_spend_budget_used_pct,
    CASE WHEN ad_spend_plan_elapsed_rub <> 0 THEN round(ad_spend_cum_rub / ad_spend_plan_elapsed_rub * 100, 2) ELSE 0 END AS ad_spend_elapsed_budget_pct,
    CASE WHEN sales_rub <> 0 THEN round(ad_spend_rub / sales_rub * 100, 2) ELSE 0 END AS tacos_pct,
    CASE WHEN sales_cum_rub <> 0 THEN round(ad_spend_cum_rub / sales_cum_rub * 100, 2) ELSE 0 END AS tacos_cum_pct
FROM cumulative;

CREATE INDEX ON public.mv_planfact_daily (marketplace, report_date);
CREATE INDEX ON public.mv_planfact_daily (plan_month);

CREATE MATERIALIZED VIEW public.mv_planfact_monthly AS
WITH fact AS (
    SELECT
        marketplace,
        date_trunc('month', report_date)::date AS plan_month,
        min(report_date) AS date_from,
        max(report_date) AS date_to,
        count(DISTINCT report_date) AS days_with_fact,
        coalesce(sum(orders_rub), 0) AS orders_rub,
        coalesce(sum(sales_rub), 0) AS sales_rub,
        coalesce(sum(ad_spend_rub), 0) AS ad_spend_rub
    FROM public.planfact_daily
    GROUP BY marketplace, date_trunc('month', report_date)::date
)
SELECT
    coalesce(p.marketplace, f.marketplace) AS marketplace,
    CASE coalesce(p.marketplace, f.marketplace) WHEN 'ozon' THEN 'Ozon' WHEN 'wb' THEN 'WB' ELSE coalesce(p.marketplace, f.marketplace) END AS marketplace_label,
    coalesce(p.plan_month, f.plan_month) AS plan_month,
    f.date_from,
    f.date_to,
    coalesce(f.days_with_fact, 0) AS days_with_fact,
    coalesce(p.sales_plan_rub, 0) AS sales_plan_rub,
    coalesce(p.ad_spend_plan_rub, 0) AS ad_spend_plan_rub,
    coalesce(f.orders_rub, 0) AS orders_rub,
    coalesce(f.sales_rub, 0) AS sales_rub,
    coalesce(f.ad_spend_rub, 0) AS ad_spend_rub,
    CASE WHEN coalesce(p.sales_plan_rub, 0) <> 0 THEN round(coalesce(f.sales_rub, 0) / p.sales_plan_rub * 100, 2) ELSE 0 END AS sales_plan_fact_pct,
    CASE WHEN coalesce(p.ad_spend_plan_rub, 0) <> 0 THEN round(coalesce(f.ad_spend_rub, 0) / p.ad_spend_plan_rub * 100, 2) ELSE 0 END AS ad_spend_budget_used_pct,
    CASE WHEN coalesce(p.sales_plan_rub, 0) <> 0 THEN round(coalesce(p.ad_spend_plan_rub, 0) / p.sales_plan_rub * 100, 2) ELSE 0 END AS tacos_plan_pct,
    CASE WHEN coalesce(f.sales_rub, 0) <> 0 THEN round(coalesce(f.ad_spend_rub, 0) / f.sales_rub * 100, 2) ELSE 0 END AS tacos_fact_pct
FROM public.planfact_plan p
FULL OUTER JOIN fact f
    ON f.marketplace = p.marketplace
    AND f.plan_month = p.plan_month;

CREATE INDEX ON public.mv_planfact_monthly (marketplace, plan_month);
"""


def load_to_db(daily_rows: list[tuple[Any, ...]], plan_rows: list[tuple[Any, ...]]) -> None:
    with app.get_conn() as conn, conn.cursor() as cur:
        cur.execute(DDL)
        cur.execute("TRUNCATE public.planfact_daily, public.planfact_plan;")
        if daily_rows:
            execute_values(
                cur,
                """
                INSERT INTO public.planfact_daily (
                    marketplace, report_date, orders_rub, sales_rub, ad_spend_rub,
                    source_file, source_sheet, source_row_num
                ) VALUES %s
                ON CONFLICT (marketplace, report_date) DO UPDATE SET
                    orders_rub = EXCLUDED.orders_rub,
                    sales_rub = EXCLUDED.sales_rub,
                    ad_spend_rub = EXCLUDED.ad_spend_rub,
                    source_file = EXCLUDED.source_file,
                    source_sheet = EXCLUDED.source_sheet,
                    source_row_num = EXCLUDED.source_row_num,
                    imported_at = now()
                """,
                daily_rows,
                page_size=1000,
            )
        if plan_rows:
            execute_values(
                cur,
                """
                INSERT INTO public.planfact_plan (
                    marketplace, plan_month, sales_plan_rub, ad_spend_plan_rub,
                    source_file, source_sheet, source_row_num
                ) VALUES %s
                ON CONFLICT (marketplace, plan_month) DO UPDATE SET
                    sales_plan_rub = EXCLUDED.sales_plan_rub,
                    ad_spend_plan_rub = EXCLUDED.ad_spend_plan_rub,
                    source_file = EXCLUDED.source_file,
                    source_sheet = EXCLUDED.source_sheet,
                    source_row_num = EXCLUDED.source_row_num,
                    imported_at = now()
                """,
                plan_rows,
                page_size=1000,
            )
        cur.execute(DROP_VIEWS)
        cur.execute(CREATE_VIEWS)
        conn.commit()


def replace_wb_month(
    daily_rows: list[tuple[Any, ...]],
    month_start: date,
    ozon_rows: list[tuple[Any, ...]] | None = None,
) -> None:
    """Replace WB month and upsert only fully evidenced Ozon dates atomically."""
    month_start = month_start.replace(day=1)
    month_end = next_month_start(month_start)
    rows = dedupe_rows_by_key(daily_rows, 2)
    if not rows:
        raise RuntimeError(f"Refusing to replace WB {month_start:%Y-%m} with no rows")
    invalid = [
        row[:2]
        for row in rows
        if row[0] != "wb" or not (month_start <= row[1] < month_end)
    ]
    if invalid:
        raise ValueError(f"Scoped WB replacement received out-of-scope rows: {invalid[:5]}")
    ozon_rows = dedupe_rows_by_key(ozon_rows or [], 2)
    invalid_ozon = [
        row[:2]
        for row in ozon_rows
        if row[0] != "ozon" or not (month_start <= row[1] < month_end)
    ]
    if invalid_ozon:
        raise ValueError(f"Scoped Ozon upsert received out-of-scope rows: {invalid_ozon[:5]}")

    with app.get_conn() as conn, conn.cursor() as cur:
        cur.execute(DDL)
        cur.execute(
            """
            DELETE FROM public.planfact_daily
            WHERE marketplace = 'wb'
              AND report_date >= %s
              AND report_date < %s
            """,
            (month_start, month_end),
        )
        execute_values(
            cur,
            """
            INSERT INTO public.planfact_daily (
                marketplace, report_date, orders_rub, sales_rub, ad_spend_rub,
                source_file, source_sheet, source_row_num
            ) VALUES %s
            ON CONFLICT (marketplace, report_date) DO UPDATE SET
                orders_rub = EXCLUDED.orders_rub,
                sales_rub = EXCLUDED.sales_rub,
                ad_spend_rub = EXCLUDED.ad_spend_rub,
                source_file = EXCLUDED.source_file,
                source_sheet = EXCLUDED.source_sheet,
                source_row_num = EXCLUDED.source_row_num,
                imported_at = now()
            """,
            rows,
            page_size=1000,
        )
        if ozon_rows:
            execute_values(
                cur,
                """
                INSERT INTO public.planfact_daily (
                    marketplace, report_date, orders_rub, sales_rub, ad_spend_rub,
                    source_file, source_sheet, source_row_num
                ) VALUES %s
                ON CONFLICT (marketplace, report_date) DO UPDATE SET
                    orders_rub = EXCLUDED.orders_rub,
                    sales_rub = EXCLUDED.sales_rub,
                    ad_spend_rub = EXCLUDED.ad_spend_rub,
                    source_file = EXCLUDED.source_file,
                    source_sheet = EXCLUDED.source_sheet,
                    source_row_num = EXCLUDED.source_row_num,
                    imported_at = now()
                """,
                ozon_rows,
                page_size=1000,
            )
        cur.execute(
            """
            SELECT
                to_regclass('public.mv_planfact_daily') AS daily_view,
                to_regclass('public.mv_planfact_monthly') AS monthly_view
            """
        )
        view_row = cur.fetchone()
        daily_view = view_row.get("daily_view") if isinstance(view_row, dict) else view_row[0]
        monthly_view = view_row.get("monthly_view") if isinstance(view_row, dict) else view_row[1]
        if daily_view and monthly_view:
            cur.execute("REFRESH MATERIALIZED VIEW public.mv_planfact_daily")
            cur.execute("REFRESH MATERIALIZED VIEW public.mv_planfact_monthly")
        else:
            cur.execute(DROP_VIEWS)
            cur.execute(CREATE_VIEWS)
        conn.commit()


def upsert_ozon_month(ozon_rows: list[tuple[Any, ...]], month_start: date) -> None:
    """Refresh only Ozon dates with all primary daily sources present."""
    month_start = month_start.replace(day=1)
    month_end = next_month_start(month_start)
    rows = dedupe_rows_by_key(ozon_rows, 2)
    if not rows:
        raise RuntimeError(f"Ozon source_not_ready for {month_start:%Y-%m}: no fully evidenced daily rows")
    invalid = [
        row[:2]
        for row in rows
        if row[0] != "ozon" or not (month_start <= row[1] < month_end)
    ]
    if invalid:
        raise ValueError(f"Scoped Ozon upsert received out-of-scope rows: {invalid[:5]}")

    with app.get_conn() as conn, conn.cursor() as cur:
        cur.execute(DDL)
        execute_values(
            cur,
            """
            INSERT INTO public.planfact_daily (
                marketplace, report_date, orders_rub, sales_rub, ad_spend_rub,
                source_file, source_sheet, source_row_num
            ) VALUES %s
            ON CONFLICT (marketplace, report_date) DO UPDATE SET
                orders_rub = EXCLUDED.orders_rub,
                sales_rub = EXCLUDED.sales_rub,
                ad_spend_rub = EXCLUDED.ad_spend_rub,
                source_file = EXCLUDED.source_file,
                source_sheet = EXCLUDED.source_sheet,
                source_row_num = EXCLUDED.source_row_num,
                imported_at = now()
            """,
            rows,
            page_size=1000,
        )
        cur.execute("REFRESH MATERIALIZED VIEW public.mv_planfact_daily")
        cur.execute("REFRESH MATERIALIZED VIEW public.mv_planfact_monthly")
        conn.commit()


def main() -> None:
    parser = argparse.ArgumentParser(description="Import Plan/Fact workbook into PostgreSQL.")
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--file", type=Path, default=None)
    parser.add_argument(
        "--legacy-only",
        action="store_true",
        help="Use only the legacy Data sheet; intended for diagnostics and rollback.",
    )
    parser.add_argument(
        "--full-rebuild",
        action="store_true",
        help="Explicitly rebuild WB, Ozon and plan tables from all primary/legacy sources.",
    )
    parser.add_argument(
        "--wb-month",
        default=None,
        help="WB month to replace in YYYY-MM format; defaults to the current month.",
    )
    parser.add_argument(
        "--ozon-only",
        action="store_true",
        help="Refresh only Ozon current-month facts from Spend, funnel and advertising; leaves WB and plans intact.",
    )
    args = parser.parse_args()

    started = time.perf_counter()
    if args.ozon_only:
        if args.full_rebuild or args.legacy_only:
            raise ValueError("--ozon-only cannot be combined with --full-rebuild or --legacy-only")
        month_start = parse_month_start(args.wb_month)
        print(
            f"PLAN: Ozon only {month_start:%Y-%m} | Spend + funnel + advertising | "
            "WB and plans unchanged | incomplete source is not converted to zero.",
            flush=True,
        )
        print("ПРОГРЕСС: 1/3 (33.3%) | Ozon Spend | читаю выручку", flush=True)
        ozon_sales = read_ozon_spend(DEFAULT_OZON_SPEND_DIR)
        print("ПРОГРЕСС: 2/3 (66.7%) | Ozon DB | читаю заказы и расходы рекламы", flush=True)
        ozon_orders, ozon_ad_spend = read_ozon_db(app.get_conn)
        ozon_rows, ozon_stats = build_ozon_month_rows(
            ozon_sales,
            ozon_orders,
            ozon_ad_spend,
            month_start,
            source=f"{DEFAULT_OZON_SPEND_DIR} | db:ozon_funnel+ozon_adv+ozon_media_adv",
        )
        print(
            f"ПРОГРЕСС: 3/3 (100.0%) | Ozon {month_start:%Y-%m} | "
            f"upsert={ozon_stats['rows']} | coverage={ozon_stats['date_from']}..{ozon_stats['date_to']}",
            flush=True,
        )
        upsert_ozon_month(ozon_rows, month_start)
        print(
            f"ИТОГ: Ozon {month_start:%Y-%m} обновлён | строки={ozon_stats['rows']} | "
            f"период {ozon_stats['date_from']} — {ozon_stats['date_to']} | WB unchanged yes | "
            f"plans unchanged yes | errors 0 | partial no | "
            f"elapsed {format_duration(time.perf_counter() - started)}",
            flush=True,
        )
        return

    if not args.full_rebuild and not args.legacy_only:
        month_start = parse_month_start(args.wb_month)
        print(
            f"PLAN: replace WB {month_start:%Y-%m} and upsert fully evidenced Ozon dates | "
            "orders and sales from report.xlsx | expense from Fix | "
            "Ozon facts require Spend + funnel + advertising; plans stay unchanged | "
            "atomic rollback on missing WB Fix days.",
            flush=True,
        )
        print("ПРОГРЕСС: 1/5 (20.0%) | WB report.xlsx | читаю весь выбранный месяц", flush=True)
        wb_report = read_wb_report(DEFAULT_WB_REPORT)
        print("ПРОГРЕСС: 2/5 (40.0%) | WB Fix | проверяю расходы для каждого дня", flush=True)
        wb_fix = read_wb_fix(DEFAULT_WB_FIX_DIR)
        daily_rows, stats = build_wb_month_rows(
            wb_report,
            wb_fix,
            month_start,
            wb_source=f"{DEFAULT_WB_REPORT} | {DEFAULT_WB_FIX_DIR}",
        )
        print("ПРОГРЕСС: 3/5 (60.0%) | Ozon Spend | читаю выручку", flush=True)
        ozon_sales = read_ozon_spend(DEFAULT_OZON_SPEND_DIR)
        print("ПРОГРЕСС: 4/5 (80.0%) | Ozon DB | читаю заказы и расходы рекламы", flush=True)
        ozon_orders, ozon_ad_spend = read_ozon_db(app.get_conn)
        ozon_rows, ozon_stats = build_ozon_month_rows(
            ozon_sales,
            ozon_orders,
            ozon_ad_spend,
            month_start,
            source=f"{DEFAULT_OZON_SPEND_DIR} | db:ozon_funnel+ozon_adv+ozon_media_adv",
        )
        print(
            f"ПРОГРЕСС: 5/5 (100.0%) | План/факт {month_start:%Y-%m} | "
            f"WB: {stats['rows']} дней; Ozon: {ozon_stats['rows']} дней | "
            f"Ozon coverage={ozon_stats['date_from']}..{ozon_stats['date_to']}",
            flush=True,
        )
        replace_wb_month(daily_rows, month_start, ozon_rows)
        print(
            f"ИТОГ: WB {month_start:%Y-%m} заменён ({stats['rows']} дней, "
            f"{stats['date_from']} — {stats['date_to']}) | Ozon upsert={ozon_stats['rows']} дней "
            f"({ozon_stats['date_from']} — {ozon_stats['date_to']}) | plans unchanged yes | "
            f"errors 0 | partial={'yes' if not ozon_rows else 'no'} | "
            f"elapsed {format_duration(time.perf_counter() - started)}",
            flush=True,
        )
        return

    file_path = args.file or find_source_file(args.source_dir)
    print(
        "PLAN: explicit full rebuild | 1 legacy workbook for monthly plans/fallback; "
        "WB report + WB Fix; Ozon funnel/advertising DB + Ozon Spend; "
        "then replace planfact tables and rebuild 2 materialized views.",
        flush=True,
    )
    print(f"Legacy plan/fallback source: {file_path}")

    read_started = time.perf_counter()
    daily_rows, plan_rows = read_rows(file_path)
    primary_stats = None
    if not args.legacy_only:
        daily_rows, primary_stats = load_primary_rows(daily_rows, app.get_conn)
    print(
        f"Read merged fact rows: {len(daily_rows):,} | Plan rows: {len(plan_rows):,} | "
        f"time: {format_duration(time.perf_counter() - read_started)}"
    )
    if primary_stats:
        print(
            "Primary freshness: "
            f"WB {primary_stats['wb_max_date']} ({primary_stats['wb_rows']} rows) | "
            f"Ozon {primary_stats['ozon_max_date']} ({primary_stats['ozon_rows']} rows) | "
            f"Ozon Spend dates {primary_stats['ozon_spend_rows']}"
        )

    steps = ["create tables", "replace raw rows", "refresh materialized views"]
    with tqdm(total=len(steps), unit="step", ascii=True) as progress:
        load_to_db(daily_rows, plan_rows)
        progress.update(len(steps))

    print(
        f"Imported daily rows: {len(daily_rows):,} | plan rows: {len(plan_rows):,} | "
        f"errors: 0 | stopped: no | "
        f"total time: {format_duration(time.perf_counter() - started)}"
    )


if __name__ == "__main__":
    main()
