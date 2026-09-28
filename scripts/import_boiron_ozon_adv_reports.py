# -*- coding: utf-8 -*-

"""Import Boiron Ozon product advertising daily exports and PF metrics."""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import time
from collections import OrderedDict, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import psycopg2
from openpyxl import load_workbook
from psycopg2 import sql
from psycopg2.extras import execute_values


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DASHBOARD_ROOT = PROJECT_ROOT / "ozon_category_dashboard"
sys.path.insert(0, str(DASHBOARD_ROOT))

os.environ.setdefault("DASHBOARD_DB_NAME", os.environ.get("BOIRON_DB_NAME", "boiron"))

import app  # noqa: E402


BOIRON_ADV_DAILY_DIR = Path(
    os.environ.get(
        "BOIRON_ADV_DAILY_DIR",
        r"G:\Общие диски\Kokoc Marketplaces\Clients\Boiron\Отчёты\Анализ РК Daily\Статистика\2026",
    )
)
BOIRON_ADV_WORKBOOK = Path(
    os.environ.get(
        "BOIRON_ADV_WORKBOOK",
        r"G:\Общие диски\Kokoc Marketplaces\Clients\Boiron\Отчёты"
        r"\Отчёты по товарной рекламе\Kokoc Group_Boiron_Анализ рекламных кампаний.xlsb",
    )
)

BRANDS = [
    "Дантинорм",
    "Стодаль",
    "Оциллококцинум",
    "Гомеовокс",
    "Гомеострес",
]
BRAND_ALIASES = {
    "оциллокоцинум": "Оциллококцинум",
    "оциллококцинум": "Оциллококцинум",
}
BATCH_SIZE = 2000

DAILY_HEADER_ALIASES = {
    "sku": ("SKU", "SKU в продвижении"),
    "product_name": ("Название товара", "Название товара в продвижении"),
    "instrument": ("Инструмент",),
    "placement": ("Место размещения",),
    "campaign_id": ("ID кампании",),
    "expense_rub": ("Расход, ₽",),
    "drr_pct": ("ДРР в продвижении, %", "ДРР, %"),
    "orders_amount_rub": ("Продажи в продвижении, ₽", "Продажи, ₽"),
    "orders_qty": ("Продано товаров, шт", "Заказы, шт", "Заказы, шт."),
    "ctr_pct": ("CTR, %",),
    "impressions": ("Показы",),
    "clicks": ("Клики",),
    "added_to_cart": ("Добавления в корзину, шт", "Добавлено в корзину", "В корзину"),
    "cr_pct": ("Конверсия в корзину, %", "CR, %"),
    "cpa_rub": ("Затраты на заказ, ₽", "CPA, ₽"),
    "cpc_rub": ("Средняя стоимость клика, ₽", "CPC, ₽"),
}

UNION_HEADER_ALIASES = {
    "sku": ("SKU в продвижении", "SKU"),
    "product_name": ("Название товара в продвижении", "Название товара"),
    "instrument": ("Инструмент",),
    "placement": ("Место размещения",),
    "campaign_id": ("ID кампании",),
    "union_sku": ("SKU из объединенной карточки",),
    "union_product_name": ("Название товара из объединенной карточки",),
    "orders_amount_rub": ("Продажи в продвижении, ₽", "Продажи, ₽"),
    "orders_qty": ("Продано товаров, шт", "Заказы, шт", "Заказы, шт."),
}

RAW_COLUMNS = [
    "report_date",
    "ozon_marketplace_article",
    "seller_article",
    "sku",
    "product_name",
    "brand_name",
    "instrument",
    "placement",
    "campaign_id",
    "impressions",
    "clicks",
    "ctr_pct",
    "expense_rub",
    "fact_expense_rub",
    "cpc_rub",
    "cpm_rub",
    "added_to_cart",
    "orders_qty",
    "direct_orders_qty",
    "indirect_orders_qty",
    "cr_pct",
    "orders_amount_rub",
    "direct_orders_amount_rub",
    "indirect_orders_amount_rub",
    "cpa_rub",
    "drr_pct",
    "total_orders_qty",
    "total_orders_amount_rub",
    "total_drr_pct",
    "total_cpa_rub",
    "source_file",
    "source_sheet",
    "source_row_num",
]

PLANFACT_COLUMNS = [
    "brand_name",
    "budget_plan_rub",
    "expense_fact_rub",
    "budget_plan_fact_pct",
    "orders_plan_qty",
    "orders_fact_qty",
    "orders_plan_fact_pct",
    "revenue_plan_rub",
    "revenue_fact_rub",
    "revenue_plan_fact_pct",
    "drr_plan_pct",
    "drr_fact_pct",
    "drr_plan_fact_pct",
    "impressions_plan",
    "impressions_fact",
    "impressions_plan_fact_pct",
    "clicks_plan",
    "clicks_fact",
    "clicks_plan_fact_pct",
    "ctr_plan_pct",
    "ctr_fact_pct",
    "ctr_plan_fact_pct",
    "cr_plan_pct",
    "cr_fact_pct",
    "cpo_plan_rub",
    "cpo_fact_rub",
    "cpo_plan_fact_pct",
    "cpc_plan_rub",
    "cpc_fact_rub",
    "cpc_plan_fact_pct",
    "avg_check_plan_rub",
    "avg_check_fact_rub",
    "avg_check_plan_fact_pct",
]


def format_duration(seconds: float) -> str:
    seconds = max(0, int(seconds))
    minutes, secs = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours}h {minutes:02d}m {secs:02d}s"
    if minutes:
        return f"{minutes}m {secs:02d}s"
    return f"{secs}s"


def normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"\s+", " ", text).strip()


def parse_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def parse_decimal(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if not text or text in {"-", "–", "—"}:
        return Decimal("0")
    text = text.rstrip("%").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Unsupported numeric value: {value!r}") from exc


def parse_int(value: Any) -> int:
    return int(parse_decimal(value))


def decimal_or_none(value: Decimal, places: str = "0.0001") -> Decimal | None:
    return value.quantize(Decimal(places)) if value is not None else None


def ratio_pct(numerator: Decimal | int, denominator: Decimal | int) -> Decimal:
    denominator = Decimal(str(denominator))
    if denominator == 0:
        return Decimal("0")
    return (Decimal(str(numerator)) / denominator * Decimal("100")).quantize(Decimal("0.0001"))


def ratio_value(numerator: Decimal | int, denominator: Decimal | int) -> Decimal:
    denominator = Decimal(str(denominator))
    if denominator == 0:
        return Decimal("0")
    return (Decimal(str(numerator)) / denominator).quantize(Decimal("0.0001"))


def cpm_value(expense: Decimal, impressions: int) -> Decimal:
    if not impressions:
        return Decimal("0")
    return (expense / Decimal(impressions) * Decimal("1000")).quantize(Decimal("0.0001"))


def parse_report_date(path: Path, first_cell: Any = None) -> date:
    text = parse_text(first_cell)
    match = re.search(r"(\d{2}\.\d{2}\.\d{4})", text)
    if not match:
        match = re.search(r"(\d{2}\.\d{2}\.\d{4})", path.stem)
    if match:
        return datetime.strptime(match.group(1), "%d.%m.%Y").date()
    match = re.search(r"(\d{4})-(\d{2})-(\d{2})", path.stem)
    if match:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    raise ValueError(f"Cannot parse report date from {path}")


def brand_from_text(value: Any) -> str:
    text = parse_text(value).lower()
    compact = text.replace(" ", "")
    for alias, brand in BRAND_ALIASES.items():
        if alias in compact:
            return brand
    for brand in BRANDS:
        if brand.lower() in text:
            return brand
    return "Без бренда"


def resolve_header_positions(headers: list[Any], aliases: dict[str, tuple[str, ...]], required: tuple[str, ...]) -> dict[str, int]:
    normalized = {normalize_header(header): index for index, header in enumerate(headers) if normalize_header(header)}
    positions: dict[str, int] = {}
    for field, field_aliases in aliases.items():
        for alias in field_aliases:
            if alias in normalized:
                positions[field] = normalized[alias]
                break
    missing: list[str] = []
    for field in required:
        for alias in aliases[field]:
            if alias in normalized:
                break
        else:
            missing.append(field)
    if missing:
        raise ValueError(f"Missing required headers: {', '.join(missing)}")
    return positions


def cell(row: tuple[Any, ...], positions: dict[str, int], field: str) -> Any:
    index = positions.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def reset_sheet_dimensions(sheet: Any) -> None:
    if hasattr(sheet, "reset_dimensions"):
        sheet.reset_dimensions()


def source_file_value(path: Path) -> str:
    try:
        return str(path.relative_to(BOIRON_ADV_DAILY_DIR.parent))
    except ValueError:
        return str(path)


def read_union_rows(workbook: Any, report_date: date) -> list[dict[str, Any]]:
    if "Union" not in workbook.sheetnames:
        return []
    sheet = workbook["Union"]
    reset_sheet_dimensions(sheet)
    rows_iter = sheet.iter_rows(values_only=True)
    next(rows_iter, None)
    headers = next(rows_iter, None)
    if not headers:
        return []
    positions = resolve_header_positions(
        headers,
        UNION_HEADER_ALIASES,
        ("sku", "orders_amount_rub", "orders_qty"),
    )
    union_rows: OrderedDict[tuple[str, str, str], dict[str, Any]] = OrderedDict()
    for source_row_num, row in enumerate(rows_iter, start=3):
        if not any(row):
            continue
        sku = parse_text(cell(row, positions, "sku"))
        union_sku = parse_text(cell(row, positions, "union_sku"))
        if not sku and not union_sku:
            continue
        campaign_id = parse_text(cell(row, positions, "campaign_id"))
        key = (campaign_id, sku, union_sku)
        item = union_rows.setdefault(
            key,
            {
                "sku": sku,
                "union_sku": union_sku,
                "campaign_id": campaign_id,
                "product_name": parse_text(cell(row, positions, "product_name")),
                "instrument": parse_text(cell(row, positions, "instrument")),
                "placement": parse_text(cell(row, positions, "placement")),
                "orders_amount_rub": Decimal("0"),
                "orders_qty": 0,
                "source_row_num": source_row_num,
            },
        )
        item["orders_amount_rub"] += parse_decimal(cell(row, positions, "orders_amount_rub"))
        item["orders_qty"] += parse_int(cell(row, positions, "orders_qty"))
    return list(union_rows.values())

def read_daily_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Statistics" not in workbook.sheetnames:
            raise ValueError(f"Sheet 'Statistics' not found in {path}")
        sheet = workbook["Statistics"]
        reset_sheet_dimensions(sheet)
        rows_iter = sheet.iter_rows(values_only=True)
        first_row = next(rows_iter, None)
        report_date = parse_report_date(path, first_row[0] if first_row else None)
        headers = next(rows_iter, None)
        if not headers:
            return []
        required = (
            "sku",
            "product_name",
            "campaign_id",
            "expense_rub",
            "orders_amount_rub",
            "orders_qty",
            "impressions",
            "clicks",
        )
        positions = resolve_header_positions(headers, DAILY_HEADER_ALIASES, required)
        grouped: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()
        source_rows: dict[tuple[str, str], int] = {}

        def new_item(campaign_id: str, sku: str, product_name: str) -> dict[str, Any]:
            return {
                "report_date": report_date,
                "sku": sku,
                "product_name": product_name,
                "brand_name": brand_from_text(product_name),
                "instrument": set(),
                "placement": set(),
                "campaign_id": campaign_id,
                "impressions": 0,
                "clicks": 0,
                "expense_rub": Decimal("0"),
                "added_to_cart": 0,
                "direct_orders_qty": 0,
                "direct_orders_amount_rub": Decimal("0"),
                "indirect_orders_qty": 0,
                "indirect_orders_amount_rub": Decimal("0"),
                "source_file": source_file_value(path),
                "source_sheet": "Statistics+Union",
            }

        for row_num, row in enumerate(rows_iter, start=3):
            if not any(row):
                continue
            sku = parse_text(cell(row, positions, "sku"))
            if not sku:
                continue
            campaign_id = parse_text(cell(row, positions, "campaign_id"))
            key = (campaign_id, sku)
            product_name = parse_text(cell(row, positions, "product_name"))
            item = grouped.setdefault(key, new_item(campaign_id, sku, product_name))
            if product_name and not item["product_name"]:
                item["product_name"] = product_name
                item["brand_name"] = brand_from_text(product_name)
            for field in ("instrument", "placement"):
                value = parse_text(cell(row, positions, field))
                if value:
                    item[field].add(value)
            item["impressions"] += parse_int(cell(row, positions, "impressions"))
            item["clicks"] += parse_int(cell(row, positions, "clicks"))
            item["expense_rub"] += parse_decimal(cell(row, positions, "expense_rub"))
            item["added_to_cart"] += parse_int(cell(row, positions, "added_to_cart"))
            item["direct_orders_qty"] += parse_int(cell(row, positions, "orders_qty"))
            item["direct_orders_amount_rub"] += parse_decimal(cell(row, positions, "orders_amount_rub"))
            source_rows.setdefault(key, row_num)

        statistics_keys_by_sku: dict[str, list[tuple[str, str]]] = {}
        for key in grouped:
            statistics_keys_by_sku.setdefault(key[1], []).append(key)

        for union in read_union_rows(workbook, report_date):
            promoted_sku = union.get("sku") or ""
            combined_sku = union.get("union_sku") or ""
            campaign_id = union.get("campaign_id") or ""
            if campaign_id:
                target_sku = promoted_sku or combined_sku
                key = (campaign_id, target_sku)
                if key not in grouped:
                    grouped[key] = new_item(campaign_id, target_sku, union.get("product_name") or "")
                    source_rows[key] = int(union.get("source_row_num") or 3) + 100000
                target_keys = [key]
            else:
                target_keys = list(statistics_keys_by_sku.get(promoted_sku, []))
                if not target_keys:
                    target_keys = list(statistics_keys_by_sku.get(combined_sku, []))
                if not target_keys:
                    target_sku = promoted_sku or combined_sku
                    key = ("", target_sku)
                    if key not in grouped:
                        grouped[key] = new_item("", target_sku, union.get("product_name") or "")
                        source_rows[key] = int(union.get("source_row_num") or 3) + 100000
                    target_keys = [key]

            weights = [max(parse_decimal(grouped[key]["expense_rub"]), Decimal("0")) for key in target_keys]
            if sum(weights, Decimal("0")) == 0:
                weights = [Decimal("1") for _ in target_keys]
            total_weight = sum(weights, Decimal("0"))
            orders_amount = parse_decimal(union.get("orders_amount_rub"))
            orders_qty = int(union.get("orders_qty") or 0)

            amount_allocations: list[Decimal] = []
            amount_remaining = orders_amount
            for allocation_index, weight in enumerate(weights):
                if allocation_index == len(weights) - 1:
                    allocation = amount_remaining
                else:
                    allocation = orders_amount * weight / total_weight
                    amount_remaining -= allocation
                amount_allocations.append(allocation)

            raw_qty_allocations = [Decimal(orders_qty) * weight / total_weight for weight in weights]
            qty_allocations = [int(value) for value in raw_qty_allocations]
            qty_remaining = orders_qty - sum(qty_allocations)
            ranked_remainders = sorted(
                range(len(target_keys)),
                key=lambda allocation_index: (
                    -(raw_qty_allocations[allocation_index] - qty_allocations[allocation_index]),
                    str(target_keys[allocation_index]),
                ),
            )
            for allocation_index in ranked_remainders[:qty_remaining]:
                qty_allocations[allocation_index] += 1

            for allocation_index, key in enumerate(target_keys):
                item = grouped[key]
                for field in ("instrument", "placement"):
                    value = union.get(field)
                    if value:
                        item[field].add(value)
                item["indirect_orders_qty"] += qty_allocations[allocation_index]
                item["indirect_orders_amount_rub"] += amount_allocations[allocation_index]

        output: list[dict[str, Any]] = []
        for key, item in grouped.items():
            direct_orders_qty = int(item["direct_orders_qty"])
            direct_orders_amount = parse_decimal(item["direct_orders_amount_rub"])
            indirect_orders_qty = int(item["indirect_orders_qty"])
            indirect_orders_amount = parse_decimal(item["indirect_orders_amount_rub"])
            orders_qty = direct_orders_qty + indirect_orders_qty
            orders_amount = direct_orders_amount + indirect_orders_amount
            expense = item["expense_rub"]
            impressions = int(item["impressions"])
            clicks = int(item["clicks"])
            carts = int(item["added_to_cart"])
            drr_pct = ratio_pct(expense, orders_amount)
            output.append(
                {
                    "report_date": item["report_date"],
                    "ozon_marketplace_article": item["sku"],
                    "seller_article": item["sku"],
                    "sku": item["sku"],
                    "product_name": item["product_name"],
                    "brand_name": item["brand_name"],
                    "instrument": ", ".join(sorted(item["instrument"])),
                    "placement": ", ".join(sorted(item["placement"])),
                    "campaign_id": item["campaign_id"],
                    "impressions": impressions,
                    "clicks": clicks,
                    "ctr_pct": ratio_pct(clicks, impressions),
                    "expense_rub": expense,
                    "fact_expense_rub": expense,
                    "cpc_rub": ratio_value(expense, clicks),
                    "cpm_rub": cpm_value(expense, impressions),
                    "added_to_cart": carts,
                    "orders_qty": orders_qty,
                    "direct_orders_qty": direct_orders_qty,
                    "indirect_orders_qty": indirect_orders_qty,
                    "cr_pct": ratio_pct(orders_qty, clicks),
                    "orders_amount_rub": orders_amount,
                    "direct_orders_amount_rub": direct_orders_amount,
                    "indirect_orders_amount_rub": indirect_orders_amount,
                    "cpa_rub": ratio_value(expense, orders_qty),
                    "drr_pct": drr_pct,
                    "total_orders_qty": orders_qty,
                    "total_orders_amount_rub": orders_amount,
                    "total_drr_pct": drr_pct,
                    "total_cpa_rub": ratio_value(expense, orders_qty),
                    "source_file": item["source_file"],
                    "source_sheet": item["source_sheet"],
                    "source_row_num": source_rows.get(key, len(output) + 3),
                }
            )
        return output
    finally:
        workbook.close()

def find_daily_files(source_dir: Path) -> list[Path]:
    files = sorted(path for path in source_dir.rglob("*.xlsx") if not path.name.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"No .xlsx files found in {source_dir}")
    return files


def ps_single_quote(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def extract_planfact_matrix(workbook_path: Path) -> list[list[Any]]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Boiron PF workbook not found: {workbook_path}")
    script = f"""
$ErrorActionPreference = 'Stop'
[Console]::OutputEncoding = [System.Text.UTF8Encoding]::new($false)
$excel = $null
$workbook = $null
try {{
  $excel = New-Object -ComObject Excel.Application
  $excel.Visible = $false
  $excel.DisplayAlerts = $false
  $workbook = $excel.Workbooks.Open({ps_single_quote(str(workbook_path))}, $null, $true)
  $sheet = $workbook.Worksheets.Item('PF')
  $range = $sheet.Range('C8:AI14')
  $values = $range.Value2
  $rows = @()
  for ($r = 1; $r -le $range.Rows.Count; $r++) {{
    $row = @()
    for ($c = 1; $c -le $range.Columns.Count; $c++) {{
      $row += $values[$r, $c]
    }}
    $rows += ,$row
  }}
  $rows | ConvertTo-Json -Depth 6 -Compress
}} finally {{
  if ($workbook -ne $null) {{ $workbook.Close($false) | Out-Null }}
  if ($excel -ne $null) {{ $excel.Quit() | Out-Null }}
}}
"""
    completed = subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(f"Failed to read PF sheet via Excel COM: {completed.stderr.strip() or completed.stdout.strip()}")
    payload = completed.stdout.strip()
    if not payload:
        return []
    matrix = json.loads(payload)
    if not isinstance(matrix, list):
        raise ValueError("PF extractor returned non-list payload")
    return matrix


def to_float_or_zero(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    return float(text) if text else 0.0


def parse_planfact_matrix(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    if not matrix:
        return []
    rows: list[dict[str, Any]] = []
    for raw in matrix[1:]:
        if isinstance(raw, dict) and "value" in raw:
            raw = raw["value"]
        if not raw:
            continue
        values = list(raw) + [None] * max(0, len(PLANFACT_COLUMNS) - len(raw))
        brand = parse_text(values[0])
        if not brand:
            continue
        row = {"brand_name": brand}
        for index, column in enumerate(PLANFACT_COLUMNS[1:], start=1):
            row[column] = to_float_or_zero(values[index])
        rows.append(row)
    return rows


def ensure_database() -> None:
    config = app.read_db_config()
    target_db = config["database"]
    admin_db = os.environ.get("DASHBOARD_ADMIN_DB_NAME", "postgres")
    admin_config = {**config, "database": admin_db}
    try:
        conn = psycopg2.connect(**admin_config)
    except psycopg2.OperationalError:
        admin_config["database"] = config.get("database")
        conn = psycopg2.connect(**admin_config)
    conn.autocommit = True
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM pg_database WHERE datname = %s", (target_db,))
            if not cur.fetchone():
                cur.execute(sql.SQL("CREATE DATABASE {}").format(sql.Identifier(target_db)))
                print(f"Created PostgreSQL database: {target_db}")
    finally:
        conn.close()


DDL = """
DROP MATERIALIZED VIEW IF EXISTS public.mv_ozon_adv_daily_by_article_category;

CREATE TABLE IF NOT EXISTS public.ozon_adv_daily_raw (
    id bigserial PRIMARY KEY,
    report_date date NOT NULL,
    ozon_marketplace_article text,
    seller_article text,
    sku text,
    product_name text,
    brand_name text,
    instrument text,
    placement text,
    campaign_id text,
    impressions bigint NOT NULL DEFAULT 0,
    clicks bigint NOT NULL DEFAULT 0,
    ctr_pct numeric,
    expense_rub numeric,
    fact_expense_rub numeric,
    cpc_rub numeric,
    cpm_rub numeric,
    added_to_cart bigint NOT NULL DEFAULT 0,
    orders_qty bigint NOT NULL DEFAULT 0,
    direct_orders_qty bigint NOT NULL DEFAULT 0,
    indirect_orders_qty bigint NOT NULL DEFAULT 0,
    cr_pct numeric,
    orders_amount_rub numeric,
    direct_orders_amount_rub numeric,
    indirect_orders_amount_rub numeric,
    cpa_rub numeric,
    drr_pct numeric,
    total_orders_qty bigint NOT NULL DEFAULT 0,
    total_orders_amount_rub numeric,
    total_drr_pct numeric,
    total_cpa_rub numeric,
    source_file text NOT NULL,
    source_sheet text NOT NULL,
    source_row_num integer NOT NULL,
    imported_at timestamp without time zone NOT NULL DEFAULT now(),
    UNIQUE (source_file, source_sheet, source_row_num)
);

ALTER TABLE public.ozon_adv_daily_raw
    ADD COLUMN IF NOT EXISTS direct_orders_qty bigint NOT NULL DEFAULT 0,
    ADD COLUMN IF NOT EXISTS indirect_orders_qty bigint NOT NULL DEFAULT 0,
    ADD COLUMN IF NOT EXISTS direct_orders_amount_rub numeric,
    ADD COLUMN IF NOT EXISTS indirect_orders_amount_rub numeric;

CREATE TABLE IF NOT EXISTS public.ozon_adv_daily_import_files (
    source_file text PRIMARY KEY,
    report_date_from date,
    report_date_to date,
    rows_imported integer NOT NULL DEFAULT 0,
    file_size_bytes bigint,
    file_mtime timestamp without time zone,
    imported_at timestamp without time zone DEFAULT now(),
    status text NOT NULL DEFAULT 'ok',
    error text
);

CREATE TABLE IF NOT EXISTS public.boiron_adv_planfact_brand (
    brand_name text PRIMARY KEY,
    budget_plan_rub numeric,
    expense_fact_rub numeric,
    budget_plan_fact_pct numeric,
    orders_plan_qty numeric,
    orders_fact_qty numeric,
    orders_plan_fact_pct numeric,
    revenue_plan_rub numeric,
    revenue_fact_rub numeric,
    revenue_plan_fact_pct numeric,
    drr_plan_pct numeric,
    drr_fact_pct numeric,
    drr_plan_fact_pct numeric,
    impressions_plan numeric,
    impressions_fact numeric,
    impressions_plan_fact_pct numeric,
    clicks_plan numeric,
    clicks_fact numeric,
    clicks_plan_fact_pct numeric,
    ctr_plan_pct numeric,
    ctr_fact_pct numeric,
    ctr_plan_fact_pct numeric,
    cr_plan_pct numeric,
    cr_fact_pct numeric,
    cpo_plan_rub numeric,
    cpo_fact_rub numeric,
    cpo_plan_fact_pct numeric,
    cpc_plan_rub numeric,
    cpc_fact_rub numeric,
    cpc_plan_fact_pct numeric,
    avg_check_plan_rub numeric,
    avg_check_fact_rub numeric,
    avg_check_plan_fact_pct numeric,
    source_file text,
    source_sheet text,
    imported_at timestamp without time zone NOT NULL DEFAULT now()
);
"""

CREATE_VIEW_SQL = """
CREATE MATERIALIZED VIEW public.mv_ozon_adv_daily_by_article_category AS
SELECT
    report_date,
    ozon_marketplace_article,
    seller_article,
    'boiron_daily_sku' AS match_type,
    NULL::bigint AS product_id,
    sku,
    sku AS product_artikul,
    product_name,
    NULL::bigint AS category_id,
    brand_name AS category_name,
    brand_name,
    instrument,
    placement,
    campaign_id,
    impressions,
    clicks,
    ctr_pct,
    expense_rub,
    fact_expense_rub,
    cpc_rub,
    cpm_rub,
    added_to_cart,
    orders_qty,
    direct_orders_qty,
    indirect_orders_qty,
    cr_pct,
    orders_amount_rub,
    direct_orders_amount_rub,
    indirect_orders_amount_rub,
    cpa_rub,
    drr_pct,
    total_orders_qty,
    total_orders_amount_rub,
    total_drr_pct,
    total_cpa_rub,
    source_file,
    source_sheet,
    source_row_num,
    imported_at
FROM public.ozon_adv_daily_raw;

CREATE INDEX idx_mv_ozon_adv_daily_boiron_date ON public.mv_ozon_adv_daily_by_article_category (report_date);
CREATE INDEX idx_mv_ozon_adv_daily_boiron_brand ON public.mv_ozon_adv_daily_by_article_category (brand_name);
CREATE INDEX idx_mv_ozon_adv_daily_boiron_sku ON public.mv_ozon_adv_daily_by_article_category (sku);
CREATE INDEX idx_mv_ozon_adv_daily_boiron_expense ON public.mv_ozon_adv_daily_by_article_category (expense_rub);
CREATE INDEX idx_mv_ozon_adv_daily_boiron_revenue ON public.mv_ozon_adv_daily_by_article_category (orders_amount_rub);
ANALYZE public.ozon_adv_daily_raw;
ANALYZE public.mv_ozon_adv_daily_by_article_category;
"""


def raw_tuple(row: dict[str, Any]) -> tuple[Any, ...]:
    return tuple(row.get(column) for column in RAW_COLUMNS)


def planfact_tuple(row: dict[str, Any], workbook_path: Path) -> tuple[Any, ...]:
    return tuple(row.get(column) for column in PLANFACT_COLUMNS) + (str(workbook_path), "PF")


def load_to_db(rows: list[dict[str, Any]], planfact_rows: list[dict[str, Any]], workbook_path: Path) -> None:
    with app.get_conn() as conn, conn.cursor() as cur:
        cur.execute(DDL)
        cur.execute("TRUNCATE public.ozon_adv_daily_raw, public.ozon_adv_daily_import_files, public.boiron_adv_planfact_brand;")
        if rows:
            execute_values(
                cur,
                """
                INSERT INTO public.ozon_adv_daily_raw (
                    report_date, ozon_marketplace_article, seller_article, sku, product_name, brand_name,
                    instrument, placement, campaign_id, impressions, clicks, ctr_pct, expense_rub,
                    fact_expense_rub, cpc_rub, cpm_rub, added_to_cart, orders_qty,
                    direct_orders_qty, indirect_orders_qty, cr_pct, orders_amount_rub,
                    direct_orders_amount_rub, indirect_orders_amount_rub,
                    cpa_rub, drr_pct, total_orders_qty, total_orders_amount_rub,
                    total_drr_pct, total_cpa_rub, source_file, source_sheet, source_row_num
                ) VALUES %s
                ON CONFLICT (source_file, source_sheet, source_row_num) DO UPDATE SET
                    report_date = EXCLUDED.report_date,
                    ozon_marketplace_article = EXCLUDED.ozon_marketplace_article,
                    seller_article = EXCLUDED.seller_article,
                    sku = EXCLUDED.sku,
                    product_name = EXCLUDED.product_name,
                    brand_name = EXCLUDED.brand_name,
                    instrument = EXCLUDED.instrument,
                    placement = EXCLUDED.placement,
                    campaign_id = EXCLUDED.campaign_id,
                    impressions = EXCLUDED.impressions,
                    clicks = EXCLUDED.clicks,
                    ctr_pct = EXCLUDED.ctr_pct,
                    expense_rub = EXCLUDED.expense_rub,
                    fact_expense_rub = EXCLUDED.fact_expense_rub,
                    cpc_rub = EXCLUDED.cpc_rub,
                    cpm_rub = EXCLUDED.cpm_rub,
                    added_to_cart = EXCLUDED.added_to_cart,
                    orders_qty = EXCLUDED.orders_qty,
                    direct_orders_qty = EXCLUDED.direct_orders_qty,
                    indirect_orders_qty = EXCLUDED.indirect_orders_qty,
                    cr_pct = EXCLUDED.cr_pct,
                    orders_amount_rub = EXCLUDED.orders_amount_rub,
                    direct_orders_amount_rub = EXCLUDED.direct_orders_amount_rub,
                    indirect_orders_amount_rub = EXCLUDED.indirect_orders_amount_rub,
                    cpa_rub = EXCLUDED.cpa_rub,
                    drr_pct = EXCLUDED.drr_pct,
                    total_orders_qty = EXCLUDED.total_orders_qty,
                    total_orders_amount_rub = EXCLUDED.total_orders_amount_rub,
                    total_drr_pct = EXCLUDED.total_drr_pct,
                    total_cpa_rub = EXCLUDED.total_cpa_rub,
                    imported_at = now()
                """,
                [raw_tuple(row) for row in rows],
                page_size=BATCH_SIZE,
            )
        file_rows = []
        by_file: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in rows:
            by_file[row["source_file"]].append(row)
        for source_file, source_rows in by_file.items():
            stat_path = Path(source_file)
            file_rows.append(
                (
                    source_file,
                    min(row["report_date"] for row in source_rows),
                    max(row["report_date"] for row in source_rows),
                    len(source_rows),
                    None,
                    None,
                    "ok",
                    None,
                )
            )
        if file_rows:
            execute_values(
                cur,
                """
                INSERT INTO public.ozon_adv_daily_import_files (
                    source_file, report_date_from, report_date_to, rows_imported,
                    file_size_bytes, file_mtime, status, error
                ) VALUES %s
                ON CONFLICT (source_file) DO UPDATE SET
                    report_date_from = EXCLUDED.report_date_from,
                    report_date_to = EXCLUDED.report_date_to,
                    rows_imported = EXCLUDED.rows_imported,
                    imported_at = now(),
                    status = EXCLUDED.status,
                    error = EXCLUDED.error
                """,
                file_rows,
                page_size=BATCH_SIZE,
            )
        if planfact_rows:
            execute_values(
                cur,
                """
                INSERT INTO public.boiron_adv_planfact_brand (
                    brand_name, budget_plan_rub, expense_fact_rub, budget_plan_fact_pct,
                    orders_plan_qty, orders_fact_qty, orders_plan_fact_pct,
                    revenue_plan_rub, revenue_fact_rub, revenue_plan_fact_pct,
                    drr_plan_pct, drr_fact_pct, drr_plan_fact_pct,
                    impressions_plan, impressions_fact, impressions_plan_fact_pct,
                    clicks_plan, clicks_fact, clicks_plan_fact_pct,
                    ctr_plan_pct, ctr_fact_pct, ctr_plan_fact_pct,
                    cr_plan_pct, cr_fact_pct,
                    cpo_plan_rub, cpo_fact_rub, cpo_plan_fact_pct,
                    cpc_plan_rub, cpc_fact_rub, cpc_plan_fact_pct,
                    avg_check_plan_rub, avg_check_fact_rub, avg_check_plan_fact_pct,
                    source_file, source_sheet
                ) VALUES %s
                ON CONFLICT (brand_name) DO UPDATE SET
                    budget_plan_rub = EXCLUDED.budget_plan_rub,
                    expense_fact_rub = EXCLUDED.expense_fact_rub,
                    budget_plan_fact_pct = EXCLUDED.budget_plan_fact_pct,
                    orders_plan_qty = EXCLUDED.orders_plan_qty,
                    orders_fact_qty = EXCLUDED.orders_fact_qty,
                    orders_plan_fact_pct = EXCLUDED.orders_plan_fact_pct,
                    revenue_plan_rub = EXCLUDED.revenue_plan_rub,
                    revenue_fact_rub = EXCLUDED.revenue_fact_rub,
                    revenue_plan_fact_pct = EXCLUDED.revenue_plan_fact_pct,
                    drr_plan_pct = EXCLUDED.drr_plan_pct,
                    drr_fact_pct = EXCLUDED.drr_fact_pct,
                    drr_plan_fact_pct = EXCLUDED.drr_plan_fact_pct,
                    impressions_plan = EXCLUDED.impressions_plan,
                    impressions_fact = EXCLUDED.impressions_fact,
                    impressions_plan_fact_pct = EXCLUDED.impressions_plan_fact_pct,
                    clicks_plan = EXCLUDED.clicks_plan,
                    clicks_fact = EXCLUDED.clicks_fact,
                    clicks_plan_fact_pct = EXCLUDED.clicks_plan_fact_pct,
                    ctr_plan_pct = EXCLUDED.ctr_plan_pct,
                    ctr_fact_pct = EXCLUDED.ctr_fact_pct,
                    ctr_plan_fact_pct = EXCLUDED.ctr_plan_fact_pct,
                    cr_plan_pct = EXCLUDED.cr_plan_pct,
                    cr_fact_pct = EXCLUDED.cr_fact_pct,
                    cpo_plan_rub = EXCLUDED.cpo_plan_rub,
                    cpo_fact_rub = EXCLUDED.cpo_fact_rub,
                    cpo_plan_fact_pct = EXCLUDED.cpo_plan_fact_pct,
                    cpc_plan_rub = EXCLUDED.cpc_plan_rub,
                    cpc_fact_rub = EXCLUDED.cpc_fact_rub,
                    cpc_plan_fact_pct = EXCLUDED.cpc_plan_fact_pct,
                    avg_check_plan_rub = EXCLUDED.avg_check_plan_rub,
                    avg_check_fact_rub = EXCLUDED.avg_check_fact_rub,
                    avg_check_plan_fact_pct = EXCLUDED.avg_check_plan_fact_pct,
                    source_file = EXCLUDED.source_file,
                    source_sheet = EXCLUDED.source_sheet,
                    imported_at = now()
                """,
                [planfact_tuple(row, workbook_path) for row in planfact_rows],
                page_size=200,
            )
        cur.execute(CREATE_VIEW_SQL)
        conn.commit()


def main() -> None:
    parser = argparse.ArgumentParser(description="Import Boiron Ozon product advertising exports.")
    parser.add_argument("--source-dir", type=Path, default=BOIRON_ADV_DAILY_DIR)
    parser.add_argument("--workbook", type=Path, default=BOIRON_ADV_WORKBOOK)
    parser.add_argument("--force-reimport", action="store_true", help="Accepted for admin compatibility; Boiron import replaces its own raw tables.")
    parser.add_argument("--resume", action="store_true", help="Accepted for all-client runner compatibility; database writes are idempotent.")
    args = parser.parse_args()

    started = time.perf_counter()
    print(f"Boiron DB: {app.read_db_config()['database']}")
    print(f"Daily source: {args.source_dir}")
    print(f"PF workbook: {args.workbook}")

    ensure_database()

    files = find_daily_files(args.source_dir)
    rows: list[dict[str, Any]] = []
    errors = 0
    print(
        f"ПЛАН: файлов {len(files)}; обработка последовательно; "
        "после файлов — один PF workbook; паузы API и rate-limit не требуются."
    )
    for index, path in enumerate(files, start=1):
        daily_rows = read_daily_rows(path)
        rows.extend(daily_rows)
        elapsed = time.perf_counter() - started
        remaining = elapsed / index * (len(files) - index) if index else 0
        pct = index / len(files) * 100 if files else 100
        print(
            f"ПРОГРЕСС: {index}/{len(files)} ({pct:.1f}%) | {path.name} | "
            f"строк накоплено: {len(rows):,} | ошибок: {errors} | "
            f"elapsed: {format_duration(elapsed)} | ETA: {format_duration(remaining)}"
        )
        print(f"[{index}/{len(files)}] {path.name}: imported {len(daily_rows):,}, errors 0")

    matrix = extract_planfact_matrix(args.workbook)
    planfact_rows = parse_planfact_matrix(matrix)
    print(f"PF rows: {len(planfact_rows):,}")

    load_to_db(rows, planfact_rows, args.workbook)
    print(
        f"ИТОГО: файлов {len(files)}, строк {len(rows):,}, PF-строк {len(planfact_rows):,}, "
        f"ошибок {errors}, пропущено 0, elapsed {format_duration(time.perf_counter() - started)}, статус: завершено."
    )


if __name__ == "__main__":
    main()
