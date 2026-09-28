"""Reviewed TOPTOP source -> exact client catalogue EANs. Dry run unless --apply."""
import argparse
from collections import defaultdict
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sys
import time
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports/unit_economics_20260912'
sys.path.insert(0,str(ROOT/'ozon_category_dashboard'))
import app
from unit_economics_workspace import workspace,barcode_list,import_cogs

def main():
    apply=argparse.ArgumentParser();apply.add_argument('--apply',action='store_true');args=apply.parse_args()
    started=time.monotonic()
    sources=json.loads((OUT/'cost_source_headers.json').read_text(encoding='utf-8'))
    source=Path(sources[0]['file']);inventory=source.parent/'10.08'/'на загрузку'/'Инвентаризация_Авиапарк_31.08.2026.xlsx'
    digest=hashlib.sha256(source.read_bytes()).hexdigest()
    sheet=openpyxl.load_workbook(source,read_only=True,data_only=True).active
    assert sheet.cell(4,3).value=='Период: 07.09.2026'
    assert sheet.cell(8,43).value=='Себестоимость ОПЕРАЦИОННАЯ без НДС'
    print(f'ПЛАН: {sheet.max_row-9} строк прайса + инвентаризация + 2 аккаунта | exact EAN/article/code | apply={args.apply}',flush=True)
    by_article=defaultdict(list);by_code=defaultdict(list);by_ean=defaultdict(list)
    for i,r in enumerate(sheet.iter_rows(min_row=10,values_only=True),10):
        if isinstance(r[42],(int,float)) and r[42]>0:
            item={'amount':str(Decimal(str(r[42]))),'row':i,'article':str(r[4] or '').strip(),'code':str(r[6] or '').strip()}
            if item['article']:by_article[item['article'].upper()].append(item)
            if item['code']:by_code[item['code']].append(item)
            for b in barcode_list(r[7]):by_ean[b].append(item)
        if i%5000==0:print(f'ПРОГРЕСС: {i}/{sheet.max_row} ({i/sheet.max_row:.0%}) | прайс | {len(by_article)} артикулов | elapsed={time.monotonic()-started:.1f}s',flush=True)
    inv=openpyxl.load_workbook(inventory,read_only=True,data_only=True).active
    for i,r in enumerate(inv.iter_rows(min_row=2,values_only=True),2):
        if r[7]!='ОК':continue
        matches=[x for x in by_code.get(str(r[4]),[]) if x['article']==str(r[3])]
        for b in barcode_list(r[0]):
            by_ean[b].extend({**x,'mapping':f'{inventory.name}:A{i},D{i},E{i}'} for x in matches)
    receipt={'source':str(source),'sha256':digest,'cost_kind':'operational_ex_vat','valid_from':'2026-09-07',
             'policy':'Exact source EAN, inventory code+article, or client catalogue article normalized only by case and outer spaces. No size suffix guessing; conflicting EAN amounts excluded. Current scenarios only, not August historical cost.',
             'inventory_sha256':hashlib.sha256(inventory.read_bytes()).hexdigest(),'clients':{}}
    for idx,client in enumerate(('toptop','lera_nena'),1):
        cfg=app.read_db_config();cfg['database']=client
        data=workspace(cfg,client,'2026-08-01','2026-08-31')
        candidates=defaultdict(list)
        for row in data['rows']:
            for b in row['barcodes']:
                candidates[b].extend(by_ean.get(b,[]))
                candidates[b].extend({**x,'mapping':f'{client}:{row["source"]}:case_normalized_article={row["article"]}:EAN={b}'} for x in by_article.get(str(row.get('article') or '').strip().upper(),[]))
        imported=[];conflicts=[]
        for b,items in sorted(candidates.items()):
            amounts={x['amount'] for x in items}
            if len(amounts)>1:conflicts.append(b)
            if len(amounts)!=1:continue
            x=items[0]
            imported.append({'barcode':b,'amount':x['amount'],'valid_from':'2026-09-07',
                'source_ref':f'{source}|TDSheet!AQ{x["row"]}|operational_ex_vat|sha256={digest}|{x.get("mapping","exact source EAN")}'})
        (OUT/f'{client}_reviewed_operating_cogs.json').write_text(json.dumps({'rows':imported},ensure_ascii=False,indent=2),encoding='utf-8')
        result=import_cogs(cfg,client,{'rows':imported}) if args.apply and imported else {'dry_run':True}
        receipt['clients'][client]={'matched_barcodes':len(imported),'conflicts':conflicts,'result':result}
        if args.apply:
            after=workspace(cfg,client,'2026-08-01','2026-08-31')
            receipt['clients'][client]['coverage']=[{'marketplace':s['marketplace'],'matched_rows':s['cogs_rows'],'rows':s['rows']} for s in after['summary']]
            expected={x['barcode']:Decimal(x['amount']) for x in imported}
            for row in after['rows']:
                if row['cogs']:
                    assert all(Decimal(str(row['cogs']['amount']))==expected[b] for b in row['barcodes'] if b in expected)
        print(f'ПРОГРЕСС: {idx}/2 ({idx*50}%) | {client} | EAN={len(imported)} conflicts={len(conflicts)} | {result} | elapsed={time.monotonic()-started:.1f}s',flush=True)
    target=OUT/('cogs_import_receipt.json' if args.apply else 'cogs_import_plan.json')
    target.write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'ИТОГ: {target} | {time.monotonic()-started:.1f}s | без изменения источника',flush=True)

if __name__=='__main__':main()
