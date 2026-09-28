"""Comparable weekly decision report. Read-only source queries, explicit coverage, no inferred zeros."""
from datetime import date, timedelta
from collections import defaultdict
from io import BytesIO
from urllib.parse import parse_qs, urlencode
import math

FUNNEL_FIELDS = ("ordered_amount_rub", "ordered_units", "card_visits", "cart_adds", "bought_units", "bought_amount_rub")
AD_FIELDS = {"adv_expense_rub": "expense_rub", "adv_orders_amount_rub": "orders_amount_rub", "adv_orders": "orders_qty",
             "adv_clicks": "clicks", "adv_impressions": "impressions"}
RATIOS = {"card_visit_to_order_pct": ("ordered_units", "card_visits", 100),
          "card_visit_to_cart_pct": ("cart_adds", "card_visits", 100),
          "cart_to_order_pct": ("ordered_units", "cart_adds", 100),
          "ordered_amount_per_unit_rub": ("ordered_amount_rub", "ordered_units", 1),
          "tacos_pct": ("adv_expense_rub", "ordered_amount_rub", 100),
          "acos_pct": ("adv_expense_rub", "adv_orders_amount_rub", 100)}

def number(value):
    if value is None or value == "" or isinstance(value, bool):
        return None
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (ValueError, TypeError):
        return None

def day(value):
    return date.fromisoformat(str(value)[:10])

def monday(value):
    return value - timedelta(days=value.weekday())

def dates(start, end):
    return [(start + timedelta(days=i)).isoformat() for i in range(max(0, (end - start).days + 1))]

def metric_sum(rows, field):
    values = [number(r.get(field)) for r in rows]
    known = [v for v in values if v is not None]
    missing = sum(v is None for v in values) + sum(int(r.get("_missing_" + field, 0)) for r in rows)
    return (sum(known) if known else None), missing

def aggregate(rows, start, end, ad_rows=()):
    expected = dates(start, end)
    picked = [r for r in rows if str(r["report_date"])[:10] in expected]
    ads = [r for r in ad_rows if str(r["report_date"])[:10] in expected]
    result, coverage = {}, {}
    for field in (*FUNNEL_FIELDS, *AD_FIELDS):
        source = ads if field in AD_FIELDS else picked
        value, missing = metric_sum(source, field)
        known_days = {str(r["report_date"])[:10] for r in source if number(r.get(field)) is not None}
        complete = len(known_days) == len(expected) and bool(expected) and missing == 0
        result[field] = value
        coverage[field] = {"days": len(known_days), "expected": len(expected), "missing_rows": missing,
                           "status": "available" if complete else "partial" if value is not None else "missing"}
    for field, (num, den, scale) in RATIOS.items():
        a, b = result[num], result[den]
        result[field] = a / b * scale if a is not None and b is not None and b > 0 else None
        status = "available" if result[field] is not None and coverage[num]["status"] == coverage[den]["status"] == "available" else "partial" if result[field] is not None else "missing"
        coverage[field] = {"status": status, "days": min(coverage[num]["days"], coverage[den]["days"]), "expected": len(expected)}
    result["coverage"] = coverage
    result["source_days"] = len({str(r["report_date"])[:10] for r in picked})
    return result

def change(current, previous, field):
    a, b = current.get(field), previous.get(field)
    comparable = all(x.get("coverage", {}).get(field, {}).get("status") == "available" for x in (current, previous))
    delta = a - b if comparable and a is not None and b is not None else None
    return {"delta": delta, "pct": delta / abs(b) * 100 if delta is not None and b else None}

def build_report(rows, ad_rows, params, stock_rows=()):
    selected_from = day(params.get("date_from", [""])[0])
    selected_to = day(params.get("date_to", [""])[0])
    latest = max((day(r["report_date"]) for r in rows), default=selected_to)
    cutoff = min(selected_to, latest, date.today() - timedelta(days=1))
    weeks = []
    start = monday(selected_from)
    while start <= selected_to:
        end = start + timedelta(days=6)
        lo, hi = max(start, selected_from), min(end, selected_to, cutoff)
        totals = aggregate(rows, lo, hi, ad_rows)
        full = lo == start and hi == end and totals["coverage"]["ordered_amount_rub"]["status"] == "available"
        weeks.append({"start": start.isoformat(), "end": end.isoformat(), "from": lo.isoformat(),
                      "to": hi.isoformat() if hi >= lo else "", "full": full, "totals": totals})
        start += timedelta(days=7)
    requested = params.get("weekly_week", [""])[0]
    mode = params.get("weekly_mode", ["complete"])[0]
    if requested:
        current_start = monday(day(requested))
        if not monday(selected_from) <= current_start <= monday(selected_to):
            raise ValueError("Выбранная неделя вне периода отчёта")
    elif mode == "current":
        current_start = monday(cutoff)
    else:
        available = [w for w in weeks if w["full"]]
        current_start = day(available[-1]["start"]) if available else monday(cutoff)
    current_end = min(current_start + timedelta(days=6), cutoff, selected_to)
    # Preserve the same weekdays when the user's range starts midweek.
    current_from = max(current_start, selected_from)
    previous_from, previous_to = current_from - timedelta(days=7), current_end - timedelta(days=7)
    current = aggregate(rows, current_from, current_end, ad_rows)
    previous = aggregate(rows, previous_from, previous_to, ad_rows)
    changes = {field: change(current, previous, field) for field in current if field not in ("coverage", "source_days")}
    stock_map = {str(r["sku_key"]): r for r in stock_rows if r.get("stock_date") and day(r["stock_date"]) <= current_end}
    groups = {"sku": defaultdict(list), "category": defaultdict(list)}
    for row in rows:
        d = day(row["report_date"])
        if previous_from <= d <= previous_to or current_from <= d <= current_end:
            groups["sku"][str(row["sku_key"])].append(row)
            groups["category"][str(row.get("category_name") or "Без категории")].append(row)
    entities = {}
    for kind, buckets in groups.items():
        output = []
        for key, group in buckets.items():
            a, b = aggregate(group, current_from, current_end), aggregate(group, previous_from, previous_to)
            # A sparse SKU series is not proof of no sales on missing days. Keep its observed
            # sums but suppress deltas until both windows are present on every expected day.
            ch = change(a, b, "ordered_amount_rub")
            label = str(group[-1].get("sku_label") or key) if kind == "sku" else key
            stock = stock_map.get(key, {}) if kind == "sku" else {}
            stock_qty = number(stock.get("total_stock_qty"))
            signal = "availability" if stock_qty == 0 else "coverage" if ch["delta"] is None else "conversion" if (
                ch["delta"] < 0 and change(a, b, "card_visit_to_order_pct")["delta"] is not None and
                change(a, b, "card_visit_to_order_pct")["delta"] < 0) else "growth" if ch["delta"] > 0 else "traffic"
            output.append({"key": key, "label": label, "category": str(group[-1].get("category_name") or "Без категории"),
                           "current": a, "previous": b, "change": ch,
                           "stock_qty": stock_qty, "stock_date": str(stock.get("stock_date") or ""),
                           "signal": signal,
                           "traffic_change": change(a, b, "card_visits"),
                           "conversion_change": change(a, b, "card_visit_to_order_pct"),
                           "daily": [{"date": str(r["report_date"])[:10], "orders": number(r.get("ordered_amount_rub")),
                                      "units": number(r.get("ordered_units")), "visits": number(r.get("card_visits"))} for r in group] if kind == "sku" else []})
        entities[kind] = sorted(output, key=lambda r: abs(r["change"]["delta"] or 0), reverse=True)
    categories = [r for r in entities["category"] if r["change"]["delta"] is not None]
    contributions = [{"key": r["key"], "label": r["label"], "value": r["change"]["delta"]} for r in categories[:5]]
    total_delta = changes["ordered_amount_rub"]["delta"]
    if total_delta is not None:
        rest = total_delta - sum(r["value"] for r in contributions)
        if abs(rest) > .005:
            contributions.append({"key": "__rest", "label": "Остальные / неполная база", "value": rest})
    else:
        contributions = []
    return {"version": 1, "mode": mode, "cutoff": cutoff.isoformat(), "history_from": selected_from.isoformat(),
            "history_to": selected_to.isoformat(), "current_from": current_from.isoformat(),
            "current_to": current_end.isoformat(), "previous_from": previous_from.isoformat(),
            "previous_to": previous_to.isoformat(), "current_week": current_start.isoformat(),
            "partial_week": current_from != current_start or current_end != current_start + timedelta(days=6),
            "current": current, "previous": previous, "changes": changes,
            "weeks": weeks, "entities": entities, "contributions": contributions,
            "coverage_note": "Покрытие отражает дни и значения в источнике отчёта; полнота загрузки всего ассортимента отдельно не подтверждена.",
            "stock_note": "Остаток — последний доступный снимок не позднее конца выбранной недели. Нулевой остаток не доказывает причину изменения заказов."}

def load_report(parsed, env):
    from psycopg2 import sql
    params = parse_qs(parsed.query)
    end = day(params.get("date_to", [date.today().isoformat()])[0])
    start = day(params.get("date_from", [(end - timedelta(days=34)).isoformat()])[0])
    if start > end or (end - start).days > 730:
        raise ValueError("Выберите корректный период не длиннее двух лет")
    params["date_from"], params["date_to"] = [start.isoformat()], [end.isoformat()]
    # One prior calendar week outside the visible range is required for a fair comparison.
    source_params = {**params, "date_from": [(monday(start) - timedelta(days=7)).isoformat()]}
    source_query = urlencode(source_params, doseq=True)
    marketplace = env["marketplace_from_query"](source_query)
    view = env["funnel_view_for_query"](source_query)
    where, values = env["funnel_filters_from_query"](source_query)
    join = env["mapping_join_for_report"]("funnel", marketplace) if env["has_mapping_filters"](source_query) else sql.SQL("")
    with env["get_conn"]() as conn, conn.cursor() as cur:
        columns = env["get_view_columns"](cur, view)
        select = []
        for field in FUNNEL_FIELDS:
            if field in columns:
                select.extend([sql.SQL("sum(v.{f}) AS {f}").format(f=sql.Identifier(field)),
                               sql.SQL("count(*) FILTER (WHERE v.{f} IS NULL) AS {m}").format(f=sql.Identifier(field), m=sql.Identifier("_missing_" + field))])
            else:
                select.append(sql.SQL("NULL::numeric AS {}").format(sql.Identifier(field)))
        cur.execute(sql.SQL("""
            SELECT v.report_date, coalesce(nullif(trim(v.category_name::text), ''), 'Без категории') category_name,
              coalesce(nullif(trim(v.sku::text), ''), nullif(trim(v.seller_article::text), ''), 'Без SKU') sku_key,
              max(coalesce(nullif(v.product_name, ''), v.sku::text)) sku_label, {metrics}
            FROM public.{view} v {join} {where}
            GROUP BY v.report_date, 2, 3 ORDER BY v.report_date, 3
        """).format(metrics=sql.SQL(", ").join(select), view=sql.Identifier(view), join=join, where=sql.SQL(where)), values)
        rows = env["normalize_rows"](cur.fetchall())
        adv_view = env["adv_view_for_marketplace"](marketplace)
        adv_columns = env["get_view_columns"](cur, adv_view)
        ad_rows = []
        if all(field in adv_columns for field in AD_FIELDS.values()):
            ad_where, ad_values = env["adv_filters_from_query"](source_query)
            ad_join = env["mapping_join_for_report"]("adv", marketplace)
            ad_select = []
            for alias, field in AD_FIELDS.items():
                ad_select.extend([sql.SQL("sum(v.{f}) AS {a}").format(f=sql.Identifier(field), a=sql.Identifier(alias)),
                                  sql.SQL("count(*) FILTER (WHERE v.{f} IS NULL) AS {m}").format(f=sql.Identifier(field), m=sql.Identifier("_missing_" + alias))])
            cur.execute(sql.SQL("SELECT v.report_date, {metrics} FROM public.{view} v {join} {where} GROUP BY v.report_date").format(
                metrics=sql.SQL(", ").join(ad_select), view=sql.Identifier(adv_view), join=ad_join, where=sql.SQL(ad_where)), ad_values)
            ad_rows = env["normalize_rows"](cur.fetchall())
        report = build_report(rows, ad_rows, params)
        cur.execute("SELECT to_regclass('public.inventory_history_daily') IS NOT NULL AS available")
        stock_rows = []
        if cur.fetchone()["available"]:
            cur.execute("""
                WITH latest AS (SELECT max(snapshot_date) d FROM public.inventory_history_daily
                  WHERE marketplace = %s AND snapshot_date <= %s)
                SELECT coalesce(nullif(trim(sku::text), ''), nullif(trim(seller_article::text), '')) sku_key,
                  CASE WHEN count(*) FILTER (WHERE stock_total_qty IS NULL) = 0 THEN sum(stock_total_qty) END total_stock_qty,
                  max(snapshot_date)::date stock_date
                FROM public.inventory_history_daily JOIN latest ON snapshot_date = latest.d
                WHERE marketplace = %s GROUP BY 1
            """, (marketplace, report["current_to"], marketplace))
            stock_rows = env["normalize_rows"](cur.fetchall())
        report = build_report(rows, ad_rows, params, stock_rows)
        report["context"] = {"client": env["current_client_key"](), "marketplace": marketplace,
                             "filters": {k: v for k, v in params.items() if k in ("categories", "article", "product", "inventory_skus", "brand", "sm_gender", "sm_season", "sm_collection")}}
        return report

def export_report(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    meta = wb.active
    meta.title = "Контекст"
    for row in [("Отчёт", "Еженедельная динамика"), ("Текущий период", report["current_from"] + " — " + report["current_to"]),
                ("Сравнение", report["previous_from"] + " — " + report["previous_to"]),
                ("Данные по", report["cutoff"]), ("Покрытие", report["coverage_note"]), ("Остатки", report["stock_note"])]:
        meta.append(row)
    context = report.get("context", {})
    meta.append(("Клиент", context.get("client", "")))
    meta.append(("Площадка", context.get("marketplace", "")))
    for key, values in context.get("filters", {}).items():
        meta.append(("Фильтр: " + key, ", ".join(values)))
    kpi = wb.create_sheet("Показатели")
    kpi.append(["Метрика", "Было", "Стало", "Изменение", "Изменение, %", "Покрытие"])
    labels = {"ordered_amount_rub": "Заказы, ₽", "ordered_units": "Заказано, шт", "card_visits": "Переходы",
              "card_visit_to_order_pct": "Карточка → заказ, %", "ordered_amount_per_unit_rub": "Средняя сумма, ₽/ед.",
              "tacos_pct": "TACoS, %", "acos_pct": "ACoS, %"}
    for field, label in labels.items():
        ch = report["changes"][field]
        kpi.append([label, report["previous"][field], report["current"][field], ch["delta"], ch["pct"], report["current"]["coverage"][field]["status"]])
    trend = wb.create_sheet("Недели")
    trend.append(["Начало", "Конец", "Полная неделя", "Заказы, ₽", "Заказано, шт", "Переходы", "Карточка → заказ, %"])
    for week in report["weeks"]:
        trend.append([week["from"], week["to"], week["full"], *[week["totals"][k] for k in ("ordered_amount_rub", "ordered_units", "card_visits", "card_visit_to_order_pct")]])
    for kind, title in (("category", "Категории"), ("sku", "SKU")):
        sheet = wb.create_sheet(title)
        sheet.append(["Ключ", "Название", "Было, ₽", "Стало, ₽", "Δ, ₽", "Δ, %", "Остаток, шт", "Дата остатка", "Покрытие текущего окна", "Покрытие предыдущего окна", "Трафик, Δ %", "Конверсия, Δ п.п."])
        for r in report["entities"][kind]:
            # Excel text is explicitly typed below to prevent formula injection from product names.
            sheet.append([r["key"], r["label"], r["previous"]["ordered_amount_rub"], r["current"]["ordered_amount_rub"],
                          r["change"]["delta"], r["change"]["pct"], r["stock_qty"], r["stock_date"],
                          r["current"]["coverage"]["ordered_amount_rub"]["status"], r["previous"]["coverage"]["ordered_amount_rub"]["status"],
                          r["traffic_change"]["pct"], r["conversion_change"]["delta"]])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="343B43")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(65, max(16, max(len(str(c.value or "")) for c in col[:40]) + 2))
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
