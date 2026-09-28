"""Persistent, source-backed SEO monitoring projects for PULSE."""

from __future__ import annotations

import json
import csv
import io
import binascii
import hashlib
import math
import os
import re
import threading
import time
import uuid
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from statistics import median
from concurrent.futures import ThreadPoolExecutor, as_completed

import psycopg2
from psycopg2.extras import RealDictCursor, execute_values
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS public.seo_monitoring_projects (
    project_id uuid PRIMARY KEY,
    project_kind text NOT NULL DEFAULT 'monitoring' CHECK (project_kind IN ('monitoring', 'generation')),
    name text NOT NULL,
    marketplace text NOT NULL CHECK (marketplace IN ('ozon', 'wb')),
    status text NOT NULL DEFAULT 'draft',
    date_from date,
    date_to date,
    created_at timestamptz NOT NULL DEFAULT now(),
    updated_at timestamptz NOT NULL DEFAULT now(),
    last_refreshed_at timestamptz,
    last_error text
);
CREATE TABLE IF NOT EXISTS public.seo_monitoring_project_skus (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    product_name text,
    created_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku)
);
CREATE TABLE IF NOT EXISTS public.seo_monitoring_keyword_snapshots (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    snapshot_date date NOT NULL,
    period_from date,
    period_to date,
    sku text NOT NULL,
    product_name text,
    search_query text NOT NULL,
    average_position numeric,
    search_demand numeric,
    traffic numeric,
    cart_adds numeric,
    orders numeric,
    revenue_rub numeric,
    source text NOT NULL,
    raw_payload jsonb,
    imported_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, snapshot_date, sku, search_query, source)
);
CREATE INDEX IF NOT EXISTS idx_seo_monitoring_snapshots_project_date
    ON public.seo_monitoring_keyword_snapshots(project_id, snapshot_date DESC);
CREATE TABLE IF NOT EXISTS public.seo_catalog_position_snapshots (
    marketplace text NOT NULL CHECK (marketplace IN ('ozon', 'wb')),
    snapshot_date date NOT NULL,
    period_from date NOT NULL,
    period_to date NOT NULL,
    sku text NOT NULL,
    product_name text,
    query_count integer NOT NULL DEFAULT 0,
    average_position numeric,
    search_demand numeric,
    source text NOT NULL,
    fetched_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (marketplace, snapshot_date, sku)
);
CREATE INDEX IF NOT EXISTS idx_seo_catalog_positions_marketplace_sku_date
    ON public.seo_catalog_position_snapshots(marketplace, sku, snapshot_date DESC);
CREATE TABLE IF NOT EXISTS public.seo_monitoring_keyword_analysis (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    search_query text NOT NULL,
    semantic_type text NOT NULL,
    modifier_type text NOT NULL,
    core_role text NOT NULL,
    priority_score integer NOT NULL CHECK (priority_score BETWEEN 0 AND 100),
    priority_label text NOT NULL,
    rationale text,
    model text,
    analyzed_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, search_query)
);
CREATE TABLE IF NOT EXISTS public.seo_generation_keyword_analysis (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    search_query text NOT NULL,
    source_group text NOT NULL CHECK (source_group IN ('api', 'mpstats')),
    relevance_decision text NOT NULL CHECK (relevance_decision IN ('keep', 'reject', 'review')),
    clean_query text,
    relevance_reason text,
    semantic_type text NOT NULL DEFAULT 'other',
    frequency_class text CHECK (frequency_class IN ('high', 'mid', 'low')),
    frequency_rank integer,
    priority_score integer NOT NULL DEFAULT 0 CHECK (priority_score BETWEEN 0 AND 100),
    priority_label text NOT NULL DEFAULT 'none' CHECK (priority_label IN ('high', 'medium', 'low', 'none')),
    source_overlap boolean NOT NULL DEFAULT false,
    model text,
    analyzed_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, search_query, source_group)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_keyword_analysis_project_sku
    ON public.seo_generation_keyword_analysis(project_id, sku, source_group, relevance_decision);
CREATE TABLE IF NOT EXISTS public.seo_generation_product_intents (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    search_intent text NOT NULL,
    rationale text,
    confidence numeric(5,4),
    evidence_json jsonb NOT NULL DEFAULT '{}'::jsonb,
    model text,
    generated_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_product_intents_project
    ON public.seo_generation_product_intents(project_id, search_intent);

CREATE TABLE IF NOT EXISTS public.seo_generation_niche_competitors (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    search_intent text NOT NULL,
    niche_id bigint,
    niche_name text,
    competitor_sku text NOT NULL,
    competitor_rank integer NOT NULL,
    product_name text,
    brand text,
    seller text,
    product_url text,
    image_url text,
    price numeric,
    sales numeric,
    revenue numeric,
    rating numeric,
    reviews_count integer,
    days_in_stock integer,
    revenue_potential numeric,
    lost_profit numeric,
    average_position numeric,
    latest_position numeric,
    source_report text NOT NULL DEFAULT 'mpstats_ozon_products_in_search',
    period_from date,
    period_to date,
    raw_payload jsonb NOT NULL DEFAULT '{}'::jsonb,
    collected_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, competitor_sku)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_niche_competitors_intent
    ON public.seo_generation_niche_competitors(project_id, search_intent, competitor_rank);
CREATE TABLE IF NOT EXISTS public.seo_generation_competitor_keywords (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    competitor_sku text NOT NULL,
    competitor_rank integer,
    competitor_name text,
    search_query text NOT NULL,
    search_demand numeric,
    average_position numeric,
    period_from date,
    period_to date,
    source text NOT NULL DEFAULT 'mpstats_ozon_competitor_keywords',
    raw_payload jsonb NOT NULL DEFAULT '{}'::jsonb,
    collected_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, competitor_sku, search_query)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_competitor_keywords_project_sku
    ON public.seo_generation_competitor_keywords(project_id, sku, search_demand DESC);
CREATE TABLE IF NOT EXISTS public.seo_generation_competitor_keyword_analysis (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    search_query text NOT NULL,
    relevance_decision text NOT NULL CHECK (relevance_decision IN ('keep', 'reject', 'review')),
    clean_query text,
    relevance_reason text,
    semantic_type text NOT NULL DEFAULT 'other',
    frequency_class text CHECK (frequency_class IN ('high', 'mid', 'low')),
    frequency_rank integer,
    competitor_coverage integer NOT NULL DEFAULT 0,
    priority_score integer NOT NULL DEFAULT 0 CHECK (priority_score BETWEEN 0 AND 100),
    priority_label text NOT NULL DEFAULT 'none' CHECK (priority_label IN ('high', 'medium', 'low', 'none')),
    model text,
    analyzed_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, search_query)
);
ALTER TABLE public.seo_generation_niche_competitors ADD COLUMN IF NOT EXISTS period_from date;
ALTER TABLE public.seo_generation_niche_competitors ADD COLUMN IF NOT EXISTS period_to date;
ALTER TABLE public.seo_generation_niche_competitors ADD COLUMN IF NOT EXISTS average_position numeric;
ALTER TABLE public.seo_generation_niche_competitors ADD COLUMN IF NOT EXISTS latest_position numeric;
ALTER TABLE public.seo_generation_niche_competitors ADD COLUMN IF NOT EXISTS source_report text NOT NULL DEFAULT 'mpstats_ozon_products_in_search';

CREATE TABLE IF NOT EXISTS public.seo_generation_ai_settings (
    script_key text PRIMARY KEY,
    primary_provider text NOT NULL DEFAULT 'openrouter',
    primary_model text NOT NULL,
    fallback_provider text NOT NULL DEFAULT 'openrouter',
    fallback_model text NOT NULL,
    updated_at timestamptz NOT NULL DEFAULT now()
);
CREATE TABLE IF NOT EXISTS public.seo_generation_operation_locks (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    operation text NOT NULL,
    owner_id uuid NOT NULL,
    started_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, operation)
);
CREATE TABLE IF NOT EXISTS public.seo_generation_competitor_ai_cache (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    search_intent text NOT NULL,
    search_query text NOT NULL,
    relevance_decision text NOT NULL CHECK (relevance_decision IN ('keep', 'reject', 'review')),
    clean_query text,
    relevance_reason text,
    semantic_type text NOT NULL DEFAULT 'other',
    model text,
    analyzed_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, search_intent, search_query)
);
CREATE TABLE IF NOT EXISTS public.seo_generation_customer_messages (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    message_type text NOT NULL CHECK (message_type IN ('review', 'question')),
    source_message_id text NOT NULL,
    message_date date,
    rating numeric(3, 2),
    message_text text,
    pros text,
    cons text,
    answer_text text,
    answered boolean,
    source text NOT NULL,
    period_days integer NOT NULL CHECK (period_days IN (30, 90)),
    raw_payload jsonb NOT NULL DEFAULT '{}'::jsonb,
    collected_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, message_type, source, source_message_id)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_customer_messages_sku
    ON public.seo_generation_customer_messages(project_id, sku, message_type, message_date DESC);
CREATE TABLE IF NOT EXISTS public.seo_generation_customer_message_status (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    message_type text NOT NULL CHECK (message_type IN ('review', 'question')),
    source text NOT NULL,
    period_days integer NOT NULL CHECK (period_days IN (30, 90)),
    status text NOT NULL CHECK (status IN ('ok', 'unavailable', 'error')),
    message_count integer,
    last_error text,
    collected_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, message_type)
);
CREATE TABLE IF NOT EXISTS public.seo_generation_competitor_reviews (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    competitor_sku text NOT NULL,
    competitor_rank integer,
    competitor_name text,
    source_message_id text NOT NULL,
    message_date date,
    rating numeric(3, 2),
    message_text text,
    pros text,
    cons text,
    answer_text text,
    answered boolean,
    source text NOT NULL DEFAULT 'mpstats_ozon_comments',
    period_days integer NOT NULL CHECK (period_days IN (30, 90)),
    raw_payload jsonb NOT NULL DEFAULT '{}'::jsonb,
    collected_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, competitor_sku, source_message_id)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_competitor_reviews_sku
    ON public.seo_generation_competitor_reviews(project_id, sku, competitor_rank, message_date DESC);
CREATE TABLE IF NOT EXISTS public.seo_generation_competitor_review_status (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    competitor_sku text NOT NULL,
    source text NOT NULL DEFAULT 'mpstats_ozon_comments',
    period_days integer NOT NULL CHECK (period_days IN (30, 90)),
    status text NOT NULL CHECK (status IN ('ok', 'unavailable', 'error')),
    review_count integer,
    last_error text,
    collected_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, competitor_sku)
);
CREATE TABLE IF NOT EXISTS public.seo_generation_customer_voice_claims (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    claim_id text NOT NULL,
    claim_text text NOT NULL,
    claim_type text NOT NULL DEFAULT 'other',
    seo_target text NOT NULL DEFAULT 'description',
    verification_status text NOT NULL CHECK (verification_status IN ('safe', 'review')),
    confidence numeric(5, 4),
    own_review_mentions integer NOT NULL DEFAULT 0,
    own_question_mentions integer NOT NULL DEFAULT 0,
    competitor_review_mentions integer NOT NULL DEFAULT 0,
    evidence_count integer NOT NULL DEFAULT 0,
    rationale text,
    evidence_json jsonb NOT NULL DEFAULT '[]'::jsonb,
    model text,
    analyzed_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, claim_id)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_customer_voice_claims_sku
    ON public.seo_generation_customer_voice_claims(project_id, sku, verification_status, confidence DESC);
CREATE TABLE IF NOT EXISTS public.seo_generation_customer_voice_analysis_status (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    status text NOT NULL CHECK (status IN ('ok', 'unavailable', 'error')),
    message_count integer NOT NULL DEFAULT 0,
    claim_count integer,
    last_error text,
    model text,
    analyzed_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku)
);
CREATE TABLE IF NOT EXISTS public.seo_generation_semantic_contexts (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    context_json jsonb NOT NULL DEFAULT '{}'::jsonb,
    context_hash text NOT NULL,
    source_counts jsonb NOT NULL DEFAULT '{}'::jsonb,
    status text NOT NULL DEFAULT 'ok' CHECK (status IN ('ok', 'partial', 'error')),
    last_error text,
    prepared_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_semantic_contexts_status
    ON public.seo_generation_semantic_contexts(project_id, status, prepared_at DESC);
CREATE TABLE IF NOT EXISTS public.seo_generation_content_allocations (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    allocation_json jsonb NOT NULL DEFAULT '{}'::jsonb,
    context_hash text NOT NULL,
    rules_version text NOT NULL,
    prompt_version text NOT NULL,
    status text NOT NULL CHECK (status IN ('ok', 'partial', 'error')),
    last_error text,
    model text,
    generated_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_content_allocations_status
    ON public.seo_generation_content_allocations(project_id, status, generated_at DESC);
CREATE TABLE IF NOT EXISTS public.seo_generation_content_drafts (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    draft_json jsonb NOT NULL DEFAULT '{}'::jsonb,
    audit_json jsonb NOT NULL DEFAULT '{}'::jsonb,
    allocation_hash text NOT NULL,
    rules_version text NOT NULL,
    prompt_version text NOT NULL,
    status text NOT NULL CHECK (status IN ('ok', 'partial', 'error')),
    last_error text,
    model text,
    generated_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_content_drafts_status
    ON public.seo_generation_content_drafts(project_id, status, generated_at DESC);
CREATE TABLE IF NOT EXISTS public.seo_generation_content_reviews (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    review_json jsonb NOT NULL DEFAULT '{}'::jsonb,
    audit_json jsonb NOT NULL DEFAULT '{}'::jsonb,
    draft_hash text NOT NULL,
    rules_version text NOT NULL,
    prompt_version text NOT NULL,
    status text NOT NULL CHECK (status IN ('ok', 'partial', 'error')),
    last_error text,
    model text,
    reviewed_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_content_reviews_status
    ON public.seo_generation_content_reviews(project_id, status, reviewed_at DESC);
CREATE TABLE IF NOT EXISTS public.seo_generation_monitoring_keywords (
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    sku text NOT NULL,
    normalized_query text NOT NULL,
    search_query text NOT NULL,
    placement text NOT NULL CHECK (placement IN ('title', 'description')),
    role text NOT NULL CHECK (role IN ('primary', 'secondary', 'supporting')),
    source_groups jsonb NOT NULL DEFAULT '[]'::jsonb,
    priority_score integer NOT NULL DEFAULT 0 CHECK (priority_score BETWEEN 0 AND 100),
    baseline_position numeric,
    baseline_search_demand numeric,
    active boolean NOT NULL DEFAULT true,
    selected_at timestamptz NOT NULL DEFAULT now(),
    PRIMARY KEY (project_id, sku, normalized_query)
);
CREATE INDEX IF NOT EXISTS idx_seo_generation_monitoring_keywords_active
    ON public.seo_generation_monitoring_keywords(project_id, active, placement, priority_score DESC);
CREATE TABLE IF NOT EXISTS public.seo_generation_jobs (
    job_id uuid PRIMARY KEY,
    project_id uuid NOT NULL REFERENCES public.seo_monitoring_projects(project_id) ON DELETE CASCADE,
    operation text NOT NULL,
    status text NOT NULL CHECK (status IN ('queued','running','stopping','stopped','completed','error')),
    target_skus jsonb NOT NULL DEFAULT '[]'::jsonb,
    total_skus integer NOT NULL DEFAULT 0,
    processed_skus integer NOT NULL DEFAULT 0,
    skipped_skus integer NOT NULL DEFAULT 0,
    analyzed integer NOT NULL DEFAULT 0,
    kept integer NOT NULL DEFAULT 0,
    rejected integer NOT NULL DEFAULT 0,
    review integer NOT NULL DEFAULT 0,
    cache_hits integer NOT NULL DEFAULT 0,
    ai_checked integer NOT NULL DEFAULT 0,
    error_count integer NOT NULL DEFAULT 0,
    current_sku text,
    last_message text,
    last_error text,
    stop_requested boolean NOT NULL DEFAULT false,
    started_at timestamptz,
    updated_at timestamptz NOT NULL DEFAULT now(),
    finished_at timestamptz
);
ALTER TABLE public.seo_generation_jobs DROP CONSTRAINT IF EXISTS seo_generation_jobs_status_check;
ALTER TABLE public.seo_generation_jobs ADD CONSTRAINT seo_generation_jobs_status_check
    CHECK (status IN ('queued','running','stopping','stopped','completed','partial','error'));
CREATE INDEX IF NOT EXISTS idx_seo_generation_jobs_project_operation
    ON public.seo_generation_jobs(project_id, operation, updated_at DESC);
ALTER TABLE public.seo_generation_jobs ADD COLUMN IF NOT EXISTS error_count integer NOT NULL DEFAULT 0;
ALTER TABLE public.seo_generation_jobs ADD COLUMN IF NOT EXISTS cache_hits integer NOT NULL DEFAULT 0;
ALTER TABLE public.seo_generation_jobs ADD COLUMN IF NOT EXISTS ai_checked integer NOT NULL DEFAULT 0;
ALTER TABLE public.seo_generation_jobs ADD COLUMN IF NOT EXISTS progress_json jsonb NOT NULL DEFAULT '{}'::jsonb;
ALTER TABLE public.seo_generation_jobs ADD COLUMN IF NOT EXISTS log_lines jsonb NOT NULL DEFAULT '[]'::jsonb;
ALTER TABLE public.seo_generation_content_drafts ADD COLUMN IF NOT EXISTS audit_json jsonb NOT NULL DEFAULT '{}'::jsonb;
ALTER TABLE public.seo_generation_ai_settings ADD COLUMN IF NOT EXISTS primary_provider text NOT NULL DEFAULT 'openrouter';
ALTER TABLE public.seo_generation_ai_settings ADD COLUMN IF NOT EXISTS fallback_provider text NOT NULL DEFAULT 'openrouter';
ALTER TABLE public.seo_monitoring_projects ADD COLUMN IF NOT EXISTS sku_filters jsonb;
ALTER TABLE public.seo_monitoring_projects ADD COLUMN IF NOT EXISTS project_kind text NOT NULL DEFAULT 'monitoring';
CREATE INDEX IF NOT EXISTS idx_seo_monitoring_projects_kind_updated
    ON public.seo_monitoring_projects(project_kind, updated_at DESC);
"""

SEO_AI_SCRIPT_DEFAULTS = {
    "keyword_relevance": {
        "label": "Очистка и ранжирование ключей",
        "primary_provider": "codex",
        "primary_model": "gpt-5.6-luna",
        "fallback_provider": "openrouter",
        "fallback_model": "deepseek/deepseek-v4-flash",
    },
    "product_intent": {
        "label": "Определение интента товара",
        "primary_provider": "codex",
        "primary_model": "gpt-5.6-luna",
        "fallback_provider": "openrouter",
        "fallback_model": "deepseek/deepseek-v4-flash",
    },
    "competitor_keyword_relevance": {
        "label": "Очистка и ранжирование ключей конкурентов",
        "primary_provider": "codex",
        "primary_model": "gpt-5.6-luna",
        "fallback_provider": "openrouter",
        "fallback_model": "deepseek/deepseek-v4-flash",
    },
    "customer_voice_claims": {
        "label": "Клеймы из отзывов и вопросов",
        "primary_provider": "codex",
        "primary_model": "gpt-5.6-luna",
        "fallback_provider": "openrouter",
        "fallback_model": "deepseek/deepseek-v4-flash",
    },
    "content_allocation": {
        "label": "Разметка ключей и клеймов по полям Ozon",
        "primary_provider": "codex",
        "primary_model": "gpt-5.6-luna",
        "fallback_provider": "openrouter",
        "fallback_model": "deepseek/deepseek-v4-flash",
    },
    "content_generation": {
        "label": "Генерация названия, описания и хештегов",
        "primary_provider": "codex",
        "primary_model": "gpt-5.6-luna",
        "fallback_provider": "openrouter",
        "fallback_model": "deepseek/deepseek-v4-flash",
    },
    "content_expert_review": {
        "label": "Экспертная оценка SEO-текста",
        "primary_provider": "codex",
        "primary_model": "gpt-5.6-terra",
        "fallback_provider": "openrouter",
        "fallback_model": "deepseek/deepseek-v4-flash",
    },
}


CATALOG_ATTRIBUTE_FILTERS = {
    "brand": ("Бренд",),
    "gender": ("Пол",),
    "age": ("Возраст", "Возрастная группа"),
    "collection": ("Коллекция",),
    "style": ("Стиль",),
    "color": ("Название цвета", "Цвет товара", "Цвет"),
    "material": ("Материал",),
    "material_composition": ("Состав материала",),
    "russian_size": ("Российский размер",),
    "manufacturer_size": ("Размер производителя",),
    "target_audience": ("Целевая аудитория",),
}

WB_CATALOG_FILTER_COLUMNS = {
    "brand": "brend",
    "gender": "pol",
    "age": "vozrastnye_ogranicheniya",
    "collection": "kollektsiya",
    "color": "tsvet",
    "material_composition": "sostav",
    "russian_size": "ros_razmer",
    "manufacturer_size": "razmer",
}
CATALOG_FILTER_OPTIONS_CACHE = {}


def _relation_exists(cur, relation):
    cur.execute("SELECT to_regclass(%s) AS relation", (f"public.{relation}",))
    return bool(cur.fetchone()["relation"])


def _trend(current, previous, lower_is_better=False):
    current = _number(current)
    previous = _number(previous)
    if current is None or previous is None:
        return {"direction": "no_data", "pct": None}
    if previous == 0:
        if current == 0:
            return {"direction": "stable", "pct": 0.0}
        return {"direction": "up" if not lower_is_better else "down", "pct": None}
    delta_pct = (current - previous) * 100 / abs(previous)
    if abs(delta_pct) < 5:
        direction = "stable"
    elif delta_pct > 0:
        direction = "down" if lower_is_better else "up"
    else:
        direction = "up" if lower_is_better else "down"
    return {"direction": direction, "pct": round(delta_pct, 1)}


def _decorate_catalog_metrics(row):
    orders_trend = _trend(row.get("orders_7d"), row.get("orders_prev_7d"))
    sales_trend = _trend(row.get("sales_7d_rub"), row.get("sales_prev_7d_rub"))
    position_trend = _trend(row.get("average_position_7d"), row.get("average_position_prev_7d"), True)
    row["orders_trend_direction"] = orders_trend["direction"]
    row["orders_trend_pct"] = orders_trend["pct"]
    row["sales_trend_direction"] = sales_trend["direction"]
    row["sales_trend_pct"] = sales_trend["pct"]
    row["position_trend_direction"] = position_trend["direction"]
    row["position_trend_pct"] = position_trend["pct"]

    stock = _number(row.get("total_stock_qty"))
    position = _number(row.get("average_position_14d"))
    orders_direction = orders_trend["direction"]
    if stock is not None and stock <= 0:
        signal, reason = "no_stock", "Нет остатка: SEO сейчас не является первым действием"
    elif position is None:
        if row.get("position_checked_through") and _number(row.get("position_query_count")) == 0:
            signal, reason = "no_queries", "Ozon проверен: поисковых запросов за период не найдено"
        else:
            signal, reason = "no_data", "Источник поисковой позиции ещё не загружен"
    elif position >= 30 and orders_direction == "down":
        signal, reason = "high", "Позиция хуже 30-й и заказы снижаются к предыдущим 7 дням"
    elif position >= 30:
        signal, reason = "attention", "Средняя поисковая позиция хуже 30-й"
    elif orders_direction == "down":
        signal, reason = "attention", "Заказы снижаются к предыдущим 7 дням"
    else:
        signal, reason = "no_signal", "Явного SEO-сигнала по доступным данным нет"
    row["seo_signal"] = signal
    row["seo_signal_reason"] = reason
    return row


def _conn(config):
    return psycopg2.connect(**config, cursor_factory=RealDictCursor)


SCHEMA_READY = set()
SCHEMA_READY_LOCK = threading.Lock()
READ_REQUIRED_RELATIONS = (
    "public.seo_monitoring_projects",
    "public.seo_monitoring_project_skus",
    "public.seo_monitoring_keyword_snapshots",
)


def ensure_schema(conn):
    """Create the schema once per process and database.

    `SCHEMA_SQL` contains `ALTER TABLE ... ADD COLUMN IF NOT EXISTS`, which takes an
    AccessExclusive lock even when the column exists. Running it on every call
    deadlocked against concurrent snapshot inserts, so it runs once per database.
    """
    key = str(getattr(conn, "dsn", "") or getattr(conn, "info", None) and conn.info.dbname or "default")
    if key in SCHEMA_READY:
        return
    with SCHEMA_READY_LOCK:
        if key in SCHEMA_READY:
            return
        with conn.cursor() as cur:
            cur.execute("SHOW transaction_read_only")
            if cur.fetchone()["transaction_read_only"] == "on":
                cur.execute(
                    "SELECT relation, to_regclass(relation) AS oid FROM unnest(%s::text[]) AS relation",
                    (list(READ_REQUIRED_RELATIONS),),
                )
                missing = [row["relation"] for row in cur.fetchall() if row["oid"] is None]
                if missing:
                    raise RuntimeError(f"SEO read schema is incomplete: {', '.join(missing)}")
                SCHEMA_READY.add(key)
                return
            cur.execute(SCHEMA_SQL)
        # Release AccessExclusive locks before any slow API or AI work starts.
        # Every caller opens a fresh connection and invokes ensure_schema first.
        conn.commit()
        SCHEMA_READY.add(key)


def _clean_text(value, limit=500):
    value = str(value or "").strip()
    return value[:limit]


def _number(value):
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def _first(item, *keys):
    for key in keys:
        if isinstance(item, dict) and item.get(key) not in (None, ""):
            return item.get(key)
    return None


def extract_ozon_items(response):
    current = response if isinstance(response, dict) else {}
    for _ in range(4):
        if isinstance(current, list):
            return current
        # Ozon returns the phrase rows under `queries`; older shapes used `items`.
        for key in ("queries", "items"):
            items = current.get(key) if isinstance(current, dict) else None
            if isinstance(items, list):
                return items
        next_value = current.get("result") or current.get("data") if isinstance(current, dict) else None
        if next_value is current or next_value is None:
            break
        current = next_value
    return []


def normalize_ozon_item(item, fallback_sku=""):
    query = _clean_text(_first(item, "query", "search_query", "phrase"), 1000)
    if not query:
        return None
    sku_value = _first(item, "sku", "skus", "offer_id")
    if isinstance(sku_value, list):
        sku_value = sku_value[0] if sku_value else fallback_sku
    return {
        "sku": _clean_text(sku_value or fallback_sku, 200),
        "product_name": _clean_text(_first(item, "product_name", "name"), 1000) or None,
        "search_query": query,
        "average_position": _number(_first(item, "position", "average_position", "avg_position")),
        "search_demand": _number(_first(item, "unique_search_users", "searches", "search_demand", "query_count")),
        "traffic": _number(_first(item, "unique_view_users", "card_visits", "traffic", "views")),
        "cart_adds": _number(_first(item, "cart_adds", "add_to_cart", "add_to_cart_count")),
        "orders": _number(_first(item, "order_count", "orders", "ordered_units")),
        "revenue_rub": _number(_first(item, "gmv", "revenue", "revenue_rub", "ordered_amount")),
        "raw_payload": item,
    }


def _project_id(value):
    try:
        return str(uuid.UUID(str(value)))
    except (TypeError, ValueError, AttributeError) as exc:
        raise ValueError("Некорректный идентификатор SEO-проекта") from exc


def _project_kind(value):
    kind = _clean_text(value, 20).lower() or "monitoring"
    if kind not in {"monitoring", "generation"}:
        raise ValueError("Тип SEO-проекта должен быть monitoring или generation")
    return kind


def _assert_project_kind(cur, project_id, project_kind, for_update=False):
    suffix = " FOR UPDATE" if for_update else ""
    cur.execute(
        f"SELECT * FROM public.seo_monitoring_projects WHERE project_id=%s AND project_kind=%s{suffix}",
        (project_id, project_kind),
    )
    project = cur.fetchone()
    if not project:
        raise ValueError("SEO-проект не найден в выбранном разделе")
    return project


def _default_project_name(cur, today=None, client_label="", project_kind="monitoring", marketplace=""):
    today = today or date.today()
    if project_kind == "generation":
        marketplace_label = "WB" if marketplace == "wb" else "Ozon"
        return f"{marketplace_label} · {today.strftime('%d.%m.%Y')} · Генерация SEO"
    cur.execute(
        "SELECT count(*)::int AS project_count FROM public.seo_monitoring_projects WHERE project_kind=%s AND created_at >= %s AND created_at < %s",
        (project_kind, today, today + timedelta(days=1)),
    )
    number = int(cur.fetchone()["project_count"] or 0) + 1
    client_part = f" · {client_label}" if client_label else ""
    return f"SEO-проект {today.strftime('%d.%m.%Y')}{client_part} · №{number}"


OZON_QUERY_SORTS_ORDER = ("BY_SEARCHES", "BY_VIEWS", "BY_POSITION", "BY_CONVERSION", "BY_GMV")
OZON_QUERY_SORTS = set(OZON_QUERY_SORTS_ORDER)
OZON_SKUS_PER_REQUEST = 100  # the API accepts up to 1000; 100 keeps a single response small
OZON_PAGE_SIZE = 100  # documented maximum
OZON_MAX_PAGES = 200  # 100 SKU x 15 phrases fits in 15 pages; the cap is a runaway guard
OZON_PAGE_RETRIES = 4  # 429 answers are retried with a growing pause
OZON_RETRY_BACKOFF_SECONDS = 1.5
OZON_PAGE_PAUSE_SECONDS = 0.15  # ~6 requests/s, far below the 50 RPS per-cabinet limit


SKU_FILTER_KEYS = (
    "filters",
    "category",
    "categories",
    "category_filters",
    "sku_filters",
    "category_sort",
    "sku_sort",
)


def _sku_filters_json(value):
    """Keep only the known builder-state keys and cap the stored payload size."""
    if not isinstance(value, dict):
        return None
    state = {key: value[key] for key in SKU_FILTER_KEYS if key in value and value[key] not in (None, "", {}, [])}
    if not state:
        return None
    dumped = json.dumps(state, ensure_ascii=False)
    if len(dumped) > 20000:
        raise ValueError("Слишком большой набор фильтров отбора SKU")
    return dumped


def create_project(config, payload):
    name = _clean_text(payload.get("name"), 160)
    client_label = _clean_text(payload.get("client_label"), 80)
    project_kind = _project_kind(payload.get("project_kind"))
    marketplace = _clean_text(payload.get("marketplace"), 20).lower()
    if marketplace not in {"ozon", "wb"}:
        raise ValueError("Маркетплейс должен быть Ozon или WB")
    raw_skus = payload.get("skus") or []
    sku_rows = []
    seen = set()
    for value in raw_skus:
        item = value if isinstance(value, dict) else {"sku": value}
        sku = _clean_text(item.get("sku"), 200)
        if sku and sku not in seen:
            seen.add(sku)
            sku_rows.append((sku, _clean_text(item.get("product_name"), 1000) or None))
    if not sku_rows:
        raise ValueError("Добавьте хотя бы один SKU")
    project_id = str(uuid.uuid4())
    date_from = payload.get("date_from") or None
    date_to = payload.get("date_to") or None
    sku_filters = _sku_filters_json(payload.get("sku_filters"))
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            if not name:
                name = _default_project_name(cur, client_label=client_label, project_kind=project_kind, marketplace=marketplace)
            cur.execute(
                """INSERT INTO public.seo_monitoring_projects
                   (project_id, project_kind, name, marketplace, status, date_from, date_to, sku_filters)
                   VALUES (%s, %s, %s, %s, 'active', %s, %s, %s)""",
                (project_id, project_kind, name, marketplace, date_from, date_to, sku_filters),
            )
            execute_values(
                cur,
                "INSERT INTO public.seo_monitoring_project_skus (project_id, sku, product_name) VALUES %s",
                [(project_id, sku, product_name) for sku, product_name in sku_rows],
            )
    return {"ok": True, "project_id": project_id, "name": name, "sku_count": len(sku_rows)}


def rename_project(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    name = _clean_text(payload.get("name"), 160)
    if not name:
        raise ValueError("Название проекта не может быть пустым")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE public.seo_monitoring_projects SET name=%s, updated_at=now() WHERE project_id=%s AND project_kind=%s RETURNING name",
                (name, project_id, project_kind),
            )
            row = cur.fetchone()
            if not row:
                raise ValueError("SEO-проект не найден")
    return {"ok": True, "project_id": project_id, "name": row["name"]}


def update_project_skus(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    rows = []
    seen = set()
    for value in payload.get("skus") or []:
        item = value if isinstance(value, dict) else {"sku": value}
        sku = _clean_text(item.get("sku"), 200)
        if sku and sku not in seen:
            seen.add(sku)
            rows.append((sku, _clean_text(item.get("product_name"), 1000) or None))
    if not rows:
        raise ValueError("Добавьте хотя бы один SKU")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                "DELETE FROM public.seo_monitoring_project_skus WHERE project_id=%s AND NOT (sku = ANY(%s))",
                (project_id, [sku for sku, _ in rows]),
            )
            removed = cur.rowcount or 0
            execute_values(
                cur,
                """INSERT INTO public.seo_monitoring_project_skus (project_id, sku, product_name) VALUES %s
                   ON CONFLICT (project_id, sku) DO UPDATE SET product_name = COALESCE(EXCLUDED.product_name, public.seo_monitoring_project_skus.product_name)""",
                [(project_id, sku, product_name) for sku, product_name in rows],
            )
            marketplace = _clean_text(payload.get("marketplace"), 20).lower()
            if marketplace in {"ozon", "wb"}:
                cur.execute(
                    "UPDATE public.seo_monitoring_projects SET marketplace=%s WHERE project_id=%s",
                    (marketplace, project_id),
                )
            if "sku_filters" in payload:
                cur.execute(
                    "UPDATE public.seo_monitoring_projects SET sku_filters=%s, updated_at=now() WHERE project_id=%s",
                    (_sku_filters_json(payload.get("sku_filters")), project_id),
                )
            else:
                cur.execute(
                    "UPDATE public.seo_monitoring_projects SET updated_at=now() WHERE project_id=%s",
                    (project_id,),
                )
    return {"ok": True, "project_id": project_id, "sku_count": len(rows), "removed_count": removed}


def delete_project(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                "DELETE FROM public.seo_monitoring_projects WHERE project_id=%s AND project_kind=%s RETURNING name",
                (project_id, project_kind),
            )
            row = cur.fetchone()
            if not row:
                raise ValueError("SEO-проект не найден")
    return {"ok": True, "project_id": project_id, "name": row["name"]}


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _serialize(row):
    return {key: _json_value(value) for key, value in dict(row).items()}


def list_projects(config, marketplace=None, project_kind="monitoring"):
    project_kind = _project_kind(project_kind)
    values = [project_kind]
    conditions = ["p.project_kind = %s"]
    if marketplace in {"ozon", "wb"}:
        conditions.append("p.marketplace = %s")
        values.append(marketplace)
    where = "WHERE " + " AND ".join(conditions)
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT p.*,
                       coalesce(s.sku_count, 0)::int AS sku_count,
                       coalesce(k.keyword_count, 0)::int AS keyword_count,
                       k.average_position, k.traffic, k.orders
                FROM public.seo_monitoring_projects p
                LEFT JOIN LATERAL (
                    SELECT count(*) AS sku_count
                    FROM public.seo_monitoring_project_skus ps WHERE ps.project_id = p.project_id
                ) s ON true
                LEFT JOIN LATERAL (
                    SELECT count(DISTINCT ks.search_query) AS keyword_count,
                           round(avg(ks.average_position), 1) AS average_position,
                           sum(ks.traffic) AS traffic, sum(ks.orders) AS orders
                    FROM public.seo_monitoring_keyword_snapshots ks
                    WHERE ks.project_id = p.project_id
                      AND ks.snapshot_date = (
                          SELECT max(latest.snapshot_date)
                          FROM public.seo_monitoring_keyword_snapshots latest
                          WHERE latest.project_id = p.project_id
                      )
                ) k ON true
                {where}
                ORDER BY p.updated_at DESC, p.created_at DESC
                """,
                values,
            )
            rows = [_serialize(row) for row in cur.fetchall()]
    for row in rows:
        traffic = _number(row.get("traffic"))
        orders = _number(row.get("orders"))
        row["conversion_pct"] = round(orders * 100 / traffic, 2) if traffic and orders is not None else None
        row["has_data"] = bool(row.get("keyword_count"))
    return {"ok": True, "rows": rows}


def _catalog_attribute_filter_sql(marketplace, selected_filters, values, alias="products"):
    filters = []
    if marketplace == "wb":
        for name, column in WB_CATALOG_FILTER_COLUMNS.items():
            selected = selected_filters.get(name)
            if not selected:
                continue
            filters.append(
                f"EXISTS (SELECT 1 FROM public.products native_catalog "
                f"WHERE native_catalog.artikul_wb::text = {alias}.sku::text "
                f"AND to_jsonb(native_catalog) ->> '{column}' = %s)"
            )
            values.append(selected)
        return filters

    for name, attribute_names in CATALOG_ATTRIBUTE_FILTERS.items():
        selected = selected_filters.get(name)
        if not selected:
            continue
        filters.append(
            f"""EXISTS (
                SELECT 1
                FROM public.ozon_cat_products native_product
                JOIN public.ozon_cat_product_attributes native_attr
                  ON native_attr.product_id = native_product.product_id
                LEFT JOIN public.ozon_cat_category_attributes category_attr
                  ON category_attr.category_id = native_product.category_id
                 AND category_attr.attribute_id = native_attr.attribute_id
                LEFT JOIN public.ozon_cat_common_attributes common_attr
                  ON common_attr.attribute_id = native_attr.attribute_id
                CROSS JOIN LATERAL regexp_split_to_table(coalesce(native_attr.value_text, ''), ';') split_value
                WHERE native_product.sku::text = {alias}.sku::text
                  AND coalesce(category_attr.attribute_name, common_attr.attribute_name) = ANY(%s)
                  AND trim(split_value) = %s
            )"""
        )
        values.extend([list(attribute_names), selected])
    return filters


def _catalog_filter_options(cur, marketplace, cache_key=""):
    cache_id = (cache_key, marketplace)
    cached = CATALOG_FILTER_OPTIONS_CACHE.get(cache_id)
    if cached and time.monotonic() - cached[0] < 300:
        return cached[1]
    options = {name: [] for name in CATALOG_ATTRIBUTE_FILTERS}
    if marketplace == "wb":
        if not _relation_exists(cur, "products"):
            return options
        # WB client schemas are not fully uniform: some catalogs do not have
        # every optional assortment attribute (for example ``pol``). Reading
        # the row as JSON keeps a missing attribute as NULL instead of making
        # the whole SEO catalog fail with an undefined-column error.
        select_parts = []
        for name, column in WB_CATALOG_FILTER_COLUMNS.items():
            expression = f"nullif(to_jsonb(native_catalog) ->> '{column}', '')"
            select_parts.append(
                f"array_remove(array_agg(DISTINCT {expression} ORDER BY {expression}), NULL) AS {name}"
            )
        cur.execute(f"SELECT {', '.join(select_parts)} FROM public.products native_catalog")
        row = dict(cur.fetchone())
        for name in WB_CATALOG_FILTER_COLUMNS:
            options[name] = [value for value in (row.get(name) or []) if str(value).strip()][:400]
        CATALOG_FILTER_OPTIONS_CACHE[cache_id] = (time.monotonic(), options)
        return options

    required = (
        _relation_exists(cur, "ozon_cat_products")
        and _relation_exists(cur, "ozon_cat_product_attributes")
        and _relation_exists(cur, "ozon_cat_category_attributes")
        and _relation_exists(cur, "ozon_cat_common_attributes")
    )
    if not required:
        return options
    name_to_filter = {
        attribute_name: filter_name
        for filter_name, attribute_names in CATALOG_ATTRIBUTE_FILTERS.items()
        for attribute_name in attribute_names
    }
    cur.execute(
        """
        WITH values_by_filter AS (
            SELECT coalesce(category_attr.attribute_name, common_attr.attribute_name) AS attribute_name,
                   nullif(trim(split_value), '') AS value
            FROM public.ozon_cat_products native_product
            JOIN public.ozon_cat_product_attributes native_attr
              ON native_attr.product_id = native_product.product_id
            LEFT JOIN public.ozon_cat_category_attributes category_attr
              ON category_attr.category_id = native_product.category_id
             AND category_attr.attribute_id = native_attr.attribute_id
            LEFT JOIN public.ozon_cat_common_attributes common_attr
              ON common_attr.attribute_id = native_attr.attribute_id
            CROSS JOIN LATERAL regexp_split_to_table(coalesce(native_attr.value_text, ''), ';') split_value
            WHERE coalesce(category_attr.attribute_name, common_attr.attribute_name) = ANY(%s)
        )
        SELECT attribute_name, array_agg(DISTINCT value ORDER BY value) AS values
        FROM values_by_filter
        WHERE value IS NOT NULL
        GROUP BY attribute_name
        """,
        [list(name_to_filter)],
    )
    for row in cur.fetchall():
        filter_name = name_to_filter.get(row["attribute_name"])
        if filter_name:
            options[filter_name] = list(dict.fromkeys([*options[filter_name], *(row.get("values") or [])]))[:400]
    CATALOG_FILTER_OPTIONS_CACHE[cache_id] = (time.monotonic(), options)
    return options


def _prepare_catalog_metrics(cur, marketplace):
    funnel_view = f"mv_{marketplace}_funnel_daily_by_article_category"
    if not _relation_exists(cur, funnel_view):
        cur.execute(
            """CREATE TEMP TABLE seo_candidate_metrics (
                sku text, funnel_date_to date, orders_7d numeric, orders_prev_7d numeric, orders_14d numeric,
                sales_7d_rub numeric, sales_prev_7d_rub numeric, sales_14d_rub numeric,
                average_position_7d numeric, average_position_prev_7d numeric, average_position_14d numeric,
                position_date_to date, position_checked_through date, position_query_count numeric
            ) ON COMMIT DROP"""
        )
        return {"funnel_date_to": None, "position_date_to": None}
    cur.execute(
        f"""
        CREATE TEMP TABLE seo_candidate_metrics ON COMMIT DROP AS
        WITH latest AS (
            SELECT max(report_date) AS date_to FROM public.{funnel_view}
        )
        SELECT f.sku::text AS sku,
               max(latest.date_to) AS funnel_date_to,
               sum(f.ordered_units) FILTER (WHERE f.report_date > latest.date_to - 7) AS orders_7d,
               sum(f.ordered_units) FILTER (WHERE f.report_date BETWEEN latest.date_to - 13 AND latest.date_to - 7) AS orders_prev_7d,
               sum(f.ordered_units) FILTER (WHERE f.report_date BETWEEN latest.date_to - 13 AND latest.date_to) AS orders_14d,
               sum(f.ordered_amount_rub) FILTER (WHERE f.report_date > latest.date_to - 7) AS sales_7d_rub,
               sum(f.ordered_amount_rub) FILTER (WHERE f.report_date BETWEEN latest.date_to - 13 AND latest.date_to - 7) AS sales_prev_7d_rub,
               sum(f.ordered_amount_rub) FILTER (WHERE f.report_date BETWEEN latest.date_to - 13 AND latest.date_to) AS sales_14d_rub,
               avg(nullif(f.search_catalog_position, 0)) FILTER (WHERE f.report_date > latest.date_to - 7) AS average_position_7d,
               avg(nullif(f.search_catalog_position, 0)) FILTER (WHERE f.report_date BETWEEN latest.date_to - 13 AND latest.date_to - 7) AS average_position_prev_7d,
               avg(nullif(f.search_catalog_position, 0)) FILTER (WHERE f.report_date BETWEEN latest.date_to - 13 AND latest.date_to) AS average_position_14d,
               max(latest.date_to) FILTER (WHERE f.search_catalog_position > 0) AS position_date_to,
               NULL::date AS position_checked_through,
               NULL::numeric AS position_query_count
        FROM public.{funnel_view} f
        CROSS JOIN latest
        WHERE f.report_date BETWEEN latest.date_to - 13 AND latest.date_to
        GROUP BY f.sku
        """
    )
    if marketplace == "wb" and _relation_exists(cur, "mv_wb_search_product_daily"):
        cur.execute(
            """
            WITH latest AS (
                SELECT max(report_date) AS date_to FROM public.mv_wb_search_product_daily
            ), positions AS (
                SELECT p.wb_sku::text AS sku,
                       max(latest.date_to) AS position_date_to,
                       avg(nullif(p.average_position, 0)) FILTER (WHERE p.report_date > latest.date_to - 7) AS average_position_7d,
                       avg(nullif(p.average_position, 0)) FILTER (WHERE p.report_date BETWEEN latest.date_to - 13 AND latest.date_to - 7) AS average_position_prev_7d,
                       avg(nullif(p.average_position, 0)) FILTER (WHERE p.report_date BETWEEN latest.date_to - 13 AND latest.date_to) AS average_position_14d
                FROM public.mv_wb_search_product_daily p
                CROSS JOIN latest
                WHERE p.report_date BETWEEN latest.date_to - 13 AND latest.date_to
                GROUP BY p.wb_sku
            )
            UPDATE seo_candidate_metrics metrics
            SET position_date_to = positions.position_date_to,
                average_position_7d = positions.average_position_7d,
                average_position_prev_7d = positions.average_position_prev_7d,
                average_position_14d = positions.average_position_14d
            FROM positions
            WHERE positions.sku = metrics.sku
            """
        )
    if marketplace == "ozon" and _relation_exists(cur, "seo_catalog_position_snapshots"):
        cur.execute(
            """
            WITH latest AS (
                SELECT max(snapshot_date) AS date_to
                FROM public.seo_catalog_position_snapshots
                WHERE marketplace = 'ozon'
            ), positions AS (
                SELECT p.sku,
                       max(p.snapshot_date) FILTER (WHERE p.average_position IS NOT NULL) AS position_date_to,
                       max(p.snapshot_date) AS position_checked_through,
                       sum(p.query_count) FILTER (WHERE p.snapshot_date BETWEEN latest.date_to - 13 AND latest.date_to) AS position_query_count,
                       avg(p.average_position) FILTER (WHERE p.snapshot_date > latest.date_to - 7) AS average_position_7d,
                       avg(p.average_position) FILTER (WHERE p.snapshot_date BETWEEN latest.date_to - 13 AND latest.date_to - 7) AS average_position_prev_7d,
                       avg(p.average_position) FILTER (WHERE p.snapshot_date BETWEEN latest.date_to - 13 AND latest.date_to) AS average_position_14d
                FROM public.seo_catalog_position_snapshots p
                CROSS JOIN latest
                WHERE p.marketplace = 'ozon'
                  AND p.snapshot_date BETWEEN latest.date_to - 13 AND latest.date_to
                GROUP BY p.sku
            )
            UPDATE seo_candidate_metrics metrics
            SET position_date_to = positions.position_date_to,
                position_checked_through = positions.position_checked_through,
                position_query_count = positions.position_query_count,
                average_position_7d = positions.average_position_7d,
                average_position_prev_7d = positions.average_position_prev_7d,
                average_position_14d = positions.average_position_14d
            FROM positions
            WHERE positions.sku = metrics.sku
            """
        )
    cur.execute("CREATE INDEX ON seo_candidate_metrics (sku)")
    cur.execute(
        """SELECT max(funnel_date_to) AS funnel_date_to,
                  max(position_date_to) AS position_date_to,
                  max(position_checked_through) AS position_checked_through
           FROM seo_candidate_metrics"""
    )
    return _serialize(cur.fetchone())


def _summarize_ozon_catalog_positions(items, requested_skus):
    grouped = {str(sku): [] for sku in requested_skus}
    for item in items:
        normalized = normalize_ozon_item(item)
        if not normalized or normalized["sku"] not in grouped:
            continue
        grouped[normalized["sku"]].append(normalized)
    result = []
    for sku, rows in grouped.items():
        positioned = [row for row in rows if row.get("average_position") is not None]
        weighted = [
            (row["average_position"], row.get("search_demand"))
            for row in positioned if (row.get("search_demand") or 0) > 0
        ]
        if weighted:
            weight = sum(float(value) for _, value in weighted)
            average_position = sum(float(position) * float(value) for position, value in weighted) / weight
        elif positioned:
            average_position = sum(float(row["average_position"]) for row in positioned) / len(positioned)
        else:
            average_position = None
        result.append({
            "sku": sku,
            "product_name": next((row.get("product_name") for row in rows if row.get("product_name")), None),
            "query_count": len(rows),
            "average_position": round(average_position, 2) if average_position is not None else None,
            "search_demand": sum(float(row.get("search_demand") or 0) for row in rows) if rows else None,
        })
    return result


def sync_ozon_catalog_positions(config, ozon_fetcher, date_from, date_to, batch_size=100, progress=None):
    """Refresh one source-backed Ozon search-position snapshot for the full catalog."""
    period_from = date.fromisoformat(str(date_from))
    period_to = date.fromisoformat(str(date_to))
    batch_size = max(1, min(int(batch_size or 100), 1000))
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            if not _relation_exists(cur, "mv_ozon_abc_product_stock_base"):
                raise RuntimeError("Ozon-каталог не загружен")
            cur.execute(
                """SELECT artikul_wb::text AS sku, max(naimenovanie) AS product_name
                   FROM public.mv_ozon_abc_product_stock_base
                   WHERE nullif(artikul_wb::text, '') IS NOT NULL
                   GROUP BY artikul_wb::text ORDER BY artikul_wb::text"""
            )
            catalog_rows = list(cur.fetchall())
        total = len(catalog_rows)
        batches = [catalog_rows[index:index + batch_size] for index in range(0, total, batch_size)]
        if progress:
            progress({"event": "plan", "total": total, "batches": len(batches), "date_from": str(period_from), "date_to": str(period_to)})
        imported = 0
        positioned = 0
        errors = []
        started = time.monotonic()
        for index, batch in enumerate(batches, start=1):
            sku_values = [str(row["sku"]) for row in batch]
            try:
                response = ozon_fetcher({
                    "skus": sku_values,
                    "date_from": str(period_from),
                    "date_to": str(period_to),
                    "limit_by_sku": 1,
                    "page": 0,
                    "page_size": 100,
                    "sort_by": "BY_SEARCHES",
                    "sort_dir": "DESCENDING",
                })
                summaries = _summarize_ozon_catalog_positions(
                    extract_ozon_items(response.get("response") if isinstance(response, dict) else response),
                    sku_values,
                )
                names = {str(row["sku"]): row.get("product_name") for row in batch}
                values = [(
                    "ozon", period_to, period_from, period_to, row["sku"],
                    row.get("product_name") or names.get(row["sku"]), row["query_count"],
                    row["average_position"], row["search_demand"],
                    "ozon_seller_api_product_queries_details",
                ) for row in summaries]
                with conn.cursor() as cur:
                    execute_values(
                        cur,
                        """INSERT INTO public.seo_catalog_position_snapshots
                           (marketplace, snapshot_date, period_from, period_to, sku, product_name,
                            query_count, average_position, search_demand, source)
                           VALUES %s
                           ON CONFLICT (marketplace, snapshot_date, sku) DO UPDATE SET
                             period_from=excluded.period_from, period_to=excluded.period_to,
                             product_name=excluded.product_name, query_count=excluded.query_count,
                             average_position=excluded.average_position, search_demand=excluded.search_demand,
                             source=excluded.source, fetched_at=now()""",
                        values,
                    )
                conn.commit()
                imported += len(values)
                positioned += sum(1 for row in summaries if row["average_position"] is not None)
            except Exception as exc:
                conn.rollback()
                errors.append({"batch": index, "error": str(exc)[:500]})
            if progress:
                elapsed = max(time.monotonic() - started, 0.001)
                done = min(index * batch_size, total)
                rate = done / elapsed
                progress({
                    "event": "progress", "batch": index, "batches": len(batches),
                    "current": done, "total": total, "imported": imported,
                    "positioned": positioned, "errors": len(errors),
                    "elapsed": elapsed, "eta": (total - done) / rate if rate else None,
                })
        return {
            "ok": not errors,
            "partial": bool(errors),
            "total_skus": total,
            "imported_skus": imported,
            "positioned_skus": positioned,
            "missing_query_skus": max(imported - positioned, 0),
            "errors": errors,
            "date_from": str(period_from),
            "date_to": str(period_to),
            "elapsed_seconds": round(time.monotonic() - started, 2),
        }


def project_candidates(
    config, marketplace, query="", limit=100, category="", subcategory="", page=1,
    gj_model="", assortment_bia="", tg="", tg_plus="", cg="", season="", selection=False,
    brand="", gender="", age="", collection="", style="", color="", material="",
    material_composition="", russian_size="", manufacturer_size="", target_audience="",
    availability="", rating_min="", catalog_category="", sort_col="product_name", sort_dir="asc",
    column_filters="", categories="", project_id="",
):
    """Return the current product catalog from the indexed stock base."""
    if marketplace not in {"ozon", "wb"}:
        raise ValueError("Маркетплейс должен быть Ozon или WB")
    prefix = "ozon" if marketplace == "ozon" else "wb"
    stock_view = f"mv_{prefix}_abc_product_stock_base"
    clean_query = _clean_text(query, 200)
    clean_category = _clean_text(category, 300)
    clean_catalog_category = _clean_text(catalog_category, 300)
    clean_subcategory = _clean_text(subcategory, 300)
    selection = str(selection or "").strip().lower() in {"1", "true", "yes"}
    limit = max(10, min(int(limit or 50), 100))
    page = max(1, int(page or 1))
    filters = []
    values = []
    category_filters = []
    category_values = []
    if clean_query:
        filters.append("(coalesce(candidates.product_name, '') ILIKE %s OR candidates.sku ILIKE %s)")
        values.extend([f"%{clean_query}%", f"%{clean_query}%"])
        category_filters.append("(coalesce(products.product_name, '') ILIKE %s OR products.sku ILIKE %s)")
        category_values.extend([f"%{clean_query}%", f"%{clean_query}%"])
    if clean_category:
        if clean_category == "Без категории":
            filters.append("category_name IS NULL")
        else:
            filters.append("category_name = %s")
            values.append(clean_category)
    try:
        decoded_categories = json.loads(categories) if isinstance(categories, str) and categories else categories
    except (TypeError, ValueError, json.JSONDecodeError):
        decoded_categories = []
    selected_categories = list(dict.fromkeys(
        clean for value in (decoded_categories if isinstance(decoded_categories, list) else [])
        if (clean := _clean_text(value, 300))
    ))
    if selected_categories:
        filters.append(f"category_name IN ({', '.join(['%s'] * len(selected_categories))})")
        values.extend(selected_categories)
    if clean_catalog_category:
        if clean_catalog_category == "Без категории":
            filters.append("category_name IS NULL")
            category_filters.append("category_name IS NULL")
        else:
            filters.append("category_name = %s")
            values.append(clean_catalog_category)
            category_filters.append("category_name = %s")
            category_values.append(clean_catalog_category)
    if clean_subcategory:
        filters.append("subcategory_name = %s")
        values.append(clean_subcategory)
        category_filters.append("subcategory_name = %s")
        category_values.append(clean_subcategory)
    extra_filters = {
        "gj_model": _clean_text(gj_model, 300),
        "assortment_bia": _clean_text(assortment_bia, 300),
        "tg": _clean_text(tg, 300),
        "tg_plus": _clean_text(tg_plus, 300),
        "cg": _clean_text(cg, 300),
        "season": _clean_text(season, 300),
    }
    for column, filter_value in extra_filters.items():
        if filter_value:
            operator = "ILIKE" if column == "gj_model" else "="
            prepared_value = f"%{filter_value}%" if column == "gj_model" else filter_value
            filters.append(f"{column} {operator} %s")
            values.append(prepared_value)
            category_filters.append(f"{column} {operator} %s")
            category_values.append(prepared_value)
    catalog_attribute_filters = {
        "brand": _clean_text(brand, 300),
        "gender": _clean_text(gender, 300),
        "age": _clean_text(age, 300),
        "collection": _clean_text(collection, 300),
        "style": _clean_text(style, 300),
        "color": _clean_text(color, 300),
        "material": _clean_text(material, 300),
        "material_composition": _clean_text(material_composition, 300),
        "russian_size": _clean_text(russian_size, 300),
        "manufacturer_size": _clean_text(manufacturer_size, 300),
        "target_audience": _clean_text(target_audience, 300),
    }
    clean_project_id = _project_id(project_id) if str(project_id or "").strip() else ""
    if clean_project_id:
        project_scope_sql = (
            "{alias}.sku::text IN (SELECT project_skus.sku FROM public.seo_monitoring_project_skus project_skus"
            " WHERE project_skus.project_id = %s)"
        )
        filters.append(project_scope_sql.format(alias="candidates"))
        values.append(clean_project_id)
        category_filters.append(project_scope_sql.format(alias="products"))
        category_values.append(clean_project_id)

    # The SKU/count queries wrap the source in a subquery aliased `candidates`,
    # while the category query keeps the `products` alias.
    filters.extend(_catalog_attribute_filter_sql(marketplace, catalog_attribute_filters, values, alias="candidates"))
    category_filters.extend(_catalog_attribute_filter_sql(marketplace, catalog_attribute_filters, category_values))
    clean_availability = _clean_text(availability, 30)
    if clean_availability in {"in_stock", "out_of_stock"}:
        availability_sql = "coalesce(total_stock_qty, 0) > 0" if clean_availability == "in_stock" else "coalesce(total_stock_qty, 0) <= 0"
        filters.append(availability_sql)
        category_filters.append(availability_sql)
    clean_rating_min = _number(rating_min)
    if clean_rating_min is not None:
        rating_sql = "CASE WHEN replace(reyting_po_otzyvam::text, ',', '.') ~ '^-?[0-9]+(\\.[0-9]+)?$' THEN replace(reyting_po_otzyvam::text, ',', '.')::numeric END >= %s"
        filters.append(rating_sql)
        values.append(clean_rating_min)
        category_filters.append(rating_sql)
        category_values.append(clean_rating_min)
    table_columns = {
        "product_name": ("coalesce(product_name, '')", "text"),
        "sku": ("coalesce(sku, '')", "text"),
        "gj_model": ("coalesce(gj_model, '')", "text"),
        "total_stock_qty": ("total_stock_qty", "number"),
        "orders_14d": ("orders_14d", "number"),
        "sales_14d_rub": ("sales_14d_rub", "number"),
        "average_position_14d": ("average_position_14d", "number"),
        "orders_trend_pct": ("orders_trend_pct", "number"),
        "seo_signal": ("coalesce(seo_signal, '')", "text"),
        "reyting_kartochki": ("CASE WHEN replace(reyting_kartochki::text, ',', '.') ~ '^-?[0-9]+(\\.[0-9]+)?$' THEN replace(reyting_kartochki::text, ',', '.')::numeric END", "number"),
        "reyting_po_otzyvam": ("CASE WHEN replace(reyting_po_otzyvam::text, ',', '.') ~ '^-?[0-9]+(\\.[0-9]+)?$' THEN replace(reyting_po_otzyvam::text, ',', '.')::numeric END", "number"),
        "api_keyword_count": ("api_keyword_count", "number"),
        "mpstats_keyword_count": ("mpstats_keyword_count", "number"),
        "niche_name": ("coalesce(niche_name, '')", "text"),
        "search_intent": ("coalesce(search_intent, '')", "text"),
        "niche_competitor_count": ("niche_competitor_count", "number"),
        "competitor_keyword_count": ("competitor_keyword_count", "number"),
        "review_count": ("review_count", "number"),
        "question_count": ("question_count", "number"),
        "competitor_review_count": ("competitor_review_count", "number"),
        "customer_voice_claim_count": ("customer_voice_claim_count", "number"),
    }
    # Column filters apply in selection mode too, so "select all by filter"
    # returns exactly the rows the table shows.
    if column_filters:
        try:
            decoded_filters = json.loads(column_filters) if isinstance(column_filters, str) else column_filters
        except (TypeError, ValueError, json.JSONDecodeError):
            decoded_filters = {}
        for key, item in (decoded_filters.items() if isinstance(decoded_filters, dict) else []):
            expression, value_type = table_columns.get(key, (None, None))
            operator = str(item.get("op") or "").strip() if isinstance(item, dict) else ""
            raw_value = item.get("value") if isinstance(item, dict) else None
            if not expression or raw_value is None or raw_value == "":
                continue
            if key == "seo_signal":
                allowed_signals = {"high", "attention", "no_signal", "no_stock", "no_queries", "no_data"}
                selected_signals = raw_value if isinstance(raw_value, list) else [raw_value]
                selected_signals = list(dict.fromkeys(
                    clean for value in selected_signals
                    if (clean := _clean_text(value, 30)) in allowed_signals
                ))
                if selected_signals:
                    filters.append(f"{expression} IN ({', '.join(['%s'] * len(selected_signals))})")
                    values.extend(selected_signals)
                continue
            if value_type == "number":
                numeric_value = _number(raw_value)
                sql_operator = {"eq": "=", "neq": "<>", "gt": ">", "gte": ">=", "lt": "<", "lte": "<="}.get(operator)
                if numeric_value is None or not sql_operator:
                    continue
                filters.append(f"{expression} {sql_operator} %s")
                values.append(numeric_value)
            else:
                clean_value = _clean_text(raw_value, 300)
                if not clean_value:
                    continue
                if operator == "contains":
                    filters.append(f"{expression} ILIKE %s")
                    values.append(f"%{clean_value}%")
                elif operator == "not_contains":
                    filters.append(f"{expression} NOT ILIKE %s")
                    values.append(f"%{clean_value}%")
                elif operator in {"eq", "neq"}:
                    filters.append(f"lower({expression}) {'=' if operator == 'eq' else '<>'} lower(%s)")
                    values.append(clean_value)
    where = "WHERE " + " AND ".join(filters) if filters else ""
    category_where = "WHERE " + " AND ".join(category_filters) if category_filters else ""
    source_sql = f"""
        SELECT artikul_wb::text AS sku,
               naimenovanie AS product_name,
               category_name,
               subcategory_name,
               coalesce(nullif(subcategory_name, ''), nullif(category_name, '')) AS niche_name,
               total_stock_qty,
               reyting_kartochki,
               reyting_po_otzyvam,
               gj_model,
               assortment_bia,
               tg,
               tg_plus,
               cg,
               season
        FROM public.{stock_view}
    """
    with _conn(config) as conn, conn.cursor() as cur:
        if not _relation_exists(cur, stock_view):
            if selection:
                return {
                    "ok": True, "rows": [], "total": 0,
                    "selection_overflow": False, "selection_limit": None,
                }
            empty_options = {
                key: [] for key in (
                    "assortment_bia", "tg", "tg_plus", "cg", "season",
                    *CATALOG_ATTRIBUTE_FILTERS.keys(),
                )
            }
            empty_options["availability"] = ["in_stock", "out_of_stock"]
            return {
                "ok": True, "rows": [], "total": 0, "page": 1, "page_size": limit,
                "total_pages": 1, "category_rows": [], "category_values": [],
                "subcategories_by_category": {}, "filter_options": empty_options,
                "metric_sources": {"funnel_date_to": None, "position_date_to": None},
                "catalog_available": False,
                "catalog_missing_reason": "Каталог площадки ещё не загружен для выбранного аккаунта",
            }
        # Selection mode needs the metric columns too: the saved filters and the
        # table sort may reference orders, sales, position or the SEO signal.
        metrics_meta = _prepare_catalog_metrics(cur, marketplace)
        metric_columns = """, metrics.orders_7d, metrics.orders_prev_7d, metrics.orders_14d,
                metrics.sales_7d_rub, metrics.sales_prev_7d_rub, metrics.sales_14d_rub,
                metrics.average_position_7d, metrics.average_position_prev_7d, metrics.average_position_14d,
                metrics.funnel_date_to, metrics.position_date_to,
                metrics.position_checked_through, metrics.position_query_count,
                CASE WHEN metrics.orders_prev_7d IS NULL OR metrics.orders_prev_7d = 0 THEN NULL
                     ELSE round((metrics.orders_7d - metrics.orders_prev_7d) * 100.0 / metrics.orders_prev_7d, 1) END AS orders_trend_pct,
                CASE WHEN coalesce(products.total_stock_qty, 0) <= 0 THEN 'no_stock'
                     WHEN metrics.average_position_14d IS NULL
                          AND metrics.position_checked_through IS NOT NULL
                          AND coalesce(metrics.position_query_count, 0) = 0 THEN 'no_queries'
                     WHEN metrics.average_position_14d IS NULL THEN 'no_data'
                     WHEN metrics.average_position_14d > 30 AND metrics.orders_prev_7d > 0
                          AND (metrics.orders_7d - metrics.orders_prev_7d) * 100.0 / metrics.orders_prev_7d <= -5 THEN 'high'
                     WHEN metrics.average_position_14d > 30 OR (metrics.orders_prev_7d > 0
                          AND (metrics.orders_7d - metrics.orders_prev_7d) * 100.0 / metrics.orders_prev_7d <= -5) THEN 'attention'
                     ELSE 'no_signal' END AS seo_signal"""
        metric_join = "LEFT JOIN seo_candidate_metrics metrics ON metrics.sku = products.sku"
        keyword_metric_columns = """, keyword_metrics.api_keyword_count,
                keyword_metrics.mpstats_keyword_count,
                product_intent.search_intent,
                competitor_metrics.niche_competitor_count,
                competitor_keyword_metrics.competitor_keyword_count,
                customer_message_metrics.review_count,
                customer_message_metrics.question_count,
                customer_message_metrics.review_status,
                customer_message_metrics.question_status,
                customer_message_metrics.review_error,
                customer_message_metrics.question_error,
                customer_message_metrics.customer_message_period_days,
                competitor_review_metrics.competitor_review_count,
                competitor_review_metrics.competitor_review_status,
                competitor_review_metrics.competitor_review_error,
                competitor_review_metrics.competitor_review_period_days,
                customer_voice_metrics.customer_voice_claim_count,
                customer_voice_metrics.customer_voice_status,
                customer_voice_metrics.customer_voice_error,
                customer_voice_metrics.customer_voice_model"""
        keyword_metric_join = ""
        intent_metric_join = ""
        competitor_metric_join = ""
        competitor_keyword_metric_join = ""
        customer_message_metric_join = ""
        competitor_review_metric_join = ""
        customer_voice_metric_join = ""
        candidate_values = []
        if clean_project_id:
            keyword_metric_join = """LEFT JOIN (
                SELECT sku,
                       CASE WHEN count(*) FILTER (WHERE source IN ('ozon_seller_api_product_queries_details', 'wb_search_queries_daily')) > 0
                            THEN count(DISTINCT search_query) FILTER (WHERE source IN ('ozon_seller_api_product_queries_details', 'wb_search_queries_daily'))::int END AS api_keyword_count,
                       CASE WHEN count(*) FILTER (WHERE source LIKE 'mpstats_%%') > 0
                            THEN count(DISTINCT search_query) FILTER (WHERE source LIKE 'mpstats_%%')::int END AS mpstats_keyword_count
                FROM public.seo_monitoring_keyword_snapshots
                WHERE project_id = %s
                GROUP BY sku
            ) keyword_metrics ON keyword_metrics.sku = products.sku"""
            candidate_values.append(clean_project_id)
            intent_metric_join = """LEFT JOIN public.seo_generation_product_intents product_intent
                ON product_intent.project_id = %s AND product_intent.sku = products.sku"""
            candidate_values.append(clean_project_id)
            competitor_metric_join = """LEFT JOIN (
                SELECT sku, count(*)::integer AS niche_competitor_count
                FROM public.seo_generation_niche_competitors
                WHERE project_id = %s
                GROUP BY sku
            ) competitor_metrics ON competitor_metrics.sku = products.sku"""
            candidate_values.append(clean_project_id)
            competitor_keyword_metric_join = """LEFT JOIN (
                SELECT sku, count(DISTINCT search_query)::integer AS competitor_keyword_count
                FROM public.seo_generation_competitor_keywords
                WHERE project_id = %s
                GROUP BY sku
            ) competitor_keyword_metrics ON competitor_keyword_metrics.sku = products.sku"""
            candidate_values.append(clean_project_id)
            customer_message_metric_join = """LEFT JOIN (
                SELECT sku,
                       max(message_count) FILTER (WHERE message_type='review' AND status='ok')::integer AS review_count,
                       max(message_count) FILTER (WHERE message_type='question' AND status='ok')::integer AS question_count,
                       max(status) FILTER (WHERE message_type='review') AS review_status,
                       max(status) FILTER (WHERE message_type='question') AS question_status,
                       max(last_error) FILTER (WHERE message_type='review') AS review_error,
                       max(last_error) FILTER (WHERE message_type='question') AS question_error,
                       max(period_days) FILTER (WHERE status='ok')::integer AS customer_message_period_days
                FROM public.seo_generation_customer_message_status
                WHERE project_id = %s
                GROUP BY sku
            ) customer_message_metrics ON customer_message_metrics.sku = products.sku"""
            candidate_values.append(clean_project_id)
            competitor_review_metric_join = """LEFT JOIN (
                SELECT sku,
                       sum(review_count) FILTER (WHERE status='ok')::integer AS competitor_review_count,
                       CASE WHEN bool_or(status='error') THEN 'error'
                            WHEN bool_or(status='unavailable') THEN 'unavailable'
                            WHEN bool_or(status='ok') THEN 'ok' END AS competitor_review_status,
                       max(last_error) FILTER (WHERE status<>'ok') AS competitor_review_error,
                       max(period_days) FILTER (WHERE status='ok')::integer AS competitor_review_period_days
                FROM public.seo_generation_competitor_review_status
                WHERE project_id = %s
                GROUP BY sku
            ) competitor_review_metrics ON competitor_review_metrics.sku = products.sku"""
            candidate_values.append(clean_project_id)
            customer_voice_metric_join = """LEFT JOIN (
                SELECT sku, claim_count AS customer_voice_claim_count,
                       status AS customer_voice_status, last_error AS customer_voice_error,
                       model AS customer_voice_model
                FROM public.seo_generation_customer_voice_analysis_status
                WHERE project_id = %s
            ) customer_voice_metrics ON customer_voice_metrics.sku = products.sku"""
            candidate_values.append(clean_project_id)
        else:
            keyword_metric_columns = """, NULL::integer AS api_keyword_count,
                NULL::integer AS mpstats_keyword_count,
                NULL::text AS search_intent,
                NULL::integer AS niche_competitor_count,
                NULL::integer AS competitor_keyword_count,
                NULL::integer AS review_count,
                NULL::integer AS question_count,
                NULL::text AS review_status,
                NULL::text AS question_status,
                NULL::text AS review_error,
                NULL::text AS question_error,
                NULL::integer AS customer_message_period_days,
                NULL::integer AS competitor_review_count,
                NULL::text AS competitor_review_status,
                NULL::text AS competitor_review_error,
                NULL::integer AS competitor_review_period_days,
                NULL::integer AS customer_voice_claim_count,
                NULL::text AS customer_voice_status,
                NULL::text AS customer_voice_error,
                NULL::text AS customer_voice_model"""
        candidate_sql = f"""SELECT products.* {metric_columns} {keyword_metric_columns}
            FROM ({source_sql}) products
            {metric_join}
            {keyword_metric_join}
            {intent_metric_join}
            {competitor_metric_join}
            {competitor_keyword_metric_join}
            {customer_message_metric_join}
            {competitor_review_metric_join}
            {customer_voice_metric_join}"""
        cur.execute(
            f"""SELECT count(*) AS total FROM ({candidate_sql}) candidates {where}""",
            candidate_values + values,
        )
        total = int(cur.fetchone()["total"])
        total_pages = 1 if selection else max(1, (total + limit - 1) // limit)
        page = min(page, total_pages)
        clean_sort_col = sort_col if sort_col in table_columns else "product_name"
        clean_sort_dir = "desc" if str(sort_dir).lower() == "desc" else "asc"
        sort_expression = table_columns[clean_sort_col][0]
        pagination_sql = "" if selection else "LIMIT %s OFFSET %s"
        pagination_values = [] if selection else [limit, (page - 1) * limit]
        cur.execute(
            f"""SELECT candidates.*
                FROM ({candidate_sql}) candidates
                {where}
                ORDER BY {sort_expression} {clean_sort_dir} NULLS LAST, product_name NULLS LAST, sku
                {pagination_sql}""",
            candidate_values + values + pagination_values,
        )
        rows = [_serialize(row) for row in cur.fetchall()]
        if selection:
            return {
                "ok": True,
                "rows": rows,
                "total": total,
                "selection_overflow": False,
                "selection_limit": None,
            }
        cur.execute(
            f"""SELECT coalesce(products.category_name, 'Без категории') AS category_name,
                       count(*)::int AS sku_count,
                       coalesce(sum(products.total_stock_qty), 0) AS total_stock_qty,
                       round(avg(CASE WHEN replace(products.reyting_kartochki::text, ',', '.') ~ '^-?[0-9]+(\\.[0-9]+)?$'
                           THEN replace(products.reyting_kartochki::text, ',', '.')::numeric END), 1) AS average_card_rating,
                       round(avg(CASE WHEN replace(products.reyting_po_otzyvam::text, ',', '.') ~ '^-?[0-9]+(\\.[0-9]+)?$'
                           THEN replace(products.reyting_po_otzyvam::text, ',', '.')::numeric END), 1) AS average_review_rating,
                       array_agg(DISTINCT products.subcategory_name ORDER BY products.subcategory_name)
                           FILTER (WHERE products.subcategory_name IS NOT NULL) AS subcategories,
                       sum(metrics.orders_7d) AS orders_7d,
                       sum(metrics.orders_prev_7d) AS orders_prev_7d,
                       sum(metrics.orders_14d) AS orders_14d,
                       sum(metrics.sales_7d_rub) AS sales_7d_rub,
                       sum(metrics.sales_prev_7d_rub) AS sales_prev_7d_rub,
                       sum(metrics.sales_14d_rub) AS sales_14d_rub,
                       round(avg(metrics.average_position_7d), 1) AS average_position_7d,
                       round(avg(metrics.average_position_prev_7d), 1) AS average_position_prev_7d,
                       round(avg(metrics.average_position_14d), 1) AS average_position_14d,
                       max(metrics.funnel_date_to) AS funnel_date_to,
                       max(metrics.position_date_to) AS position_date_to
                FROM ({source_sql}) products
                LEFT JOIN seo_candidate_metrics metrics ON metrics.sku = products.sku
                {category_where}
                GROUP BY coalesce(products.category_name, 'Без категории')
                ORDER BY sku_count DESC, category_name""",
            category_values,
        )
        category_rows = [_serialize(row) for row in cur.fetchall()]
        cur.execute(
            f"""SELECT
                array_agg(DISTINCT assortment_bia ORDER BY assortment_bia) FILTER (WHERE assortment_bia IS NOT NULL AND assortment_bia <> '') AS assortment_bia,
                array_agg(DISTINCT tg ORDER BY tg) FILTER (WHERE tg IS NOT NULL AND tg <> '') AS tg,
                array_agg(DISTINCT tg_plus ORDER BY tg_plus) FILTER (WHERE tg_plus IS NOT NULL AND tg_plus <> '') AS tg_plus,
                array_agg(DISTINCT cg ORDER BY cg) FILTER (WHERE cg IS NOT NULL AND cg <> '') AS cg,
                array_agg(DISTINCT season ORDER BY season) FILTER (WHERE season IS NOT NULL AND season <> '') AS season
                FROM ({source_sql}) products"""
        )
        option_values = _serialize(cur.fetchone())
        native_options = _catalog_filter_options(cur, marketplace, str(config.get("dbname") or config.get("database") or ""))
    rows = [_decorate_catalog_metrics(row) for row in rows]
    category_rows = [_decorate_catalog_metrics(row) for row in category_rows]
    filter_options = {key: value or [] for key, value in option_values.items()}
    filter_options.update(native_options)
    filter_options["availability"] = ["in_stock", "out_of_stock"]
    return {
        "ok": True, "rows": rows, "total": total, "page": page, "page_size": limit,
        "total_pages": total_pages, "category_rows": category_rows,
        "category_values": [row["category_name"] for row in category_rows],
        "subcategories_by_category": {row["category_name"]: row.get("subcategories") or [] for row in category_rows},
        "filter_options": filter_options,
        "metric_sources": metrics_meta, "sort_col": clean_sort_col, "sort_dir": clean_sort_dir,
    }


def build_project_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SKU"
    sheet.append(["sku", "product_name", "comment"])
    sheet.append([None, None, None])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF4747")
        cell.alignment = Alignment(horizontal="center")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 48
    sheet.column_dimensions["C"].width = 42
    instructions = workbook.create_sheet("Инструкция")
    instructions.append(["Поле", "Правило"])
    instructions.append(["sku", "Обязательное. Один SKU в строке. Сохраняйте как текст, чтобы не потерять цифры."])
    instructions.append(["product_name", "Необязательное. Если пусто, PULSE подставит название из каталога, когда найдёт SKU."])
    instructions.append(["comment", "Необязательное служебное поле; в проект не загружается."])
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 92
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "seo_monitoring_sku_template.xlsx"


SEO_CONTENT_EXPORT_HEADERS = (
    "SKU",
    "Модель",
    "Название",
    "Наименование было",
    "Наименование стало",
    "Описание было",
    "Описание стало",
    "Хештеги были",
    "Хештеги стали",
    "Ключи использованы",
    "Характеристики использованы",
    "Приоритетные ключи для дальнейшего мониторинга",
)


def _excel_safe_text(value, missing="—"):
    """Keep missing data explicit and prevent spreadsheet formula injection."""
    if value is None:
        return missing
    if isinstance(value, (list, tuple, set)):
        value = "\n".join(str(item).strip() for item in value if str(item).strip())
    value = str(value).strip()
    if not value:
        return missing
    if value[0] in "=+-@":
        value = "'" + value
    return value


def _export_hashtags(value):
    if isinstance(value, str):
        values = re.findall(r"#[\wа-яёА-ЯЁ-]+", value, flags=re.UNICODE)
    else:
        values = value if isinstance(value, list) else []
    return _excel_safe_text(list(dict.fromkeys(_clean_text(item, 160) for item in values if _clean_text(item, 160))))


def _export_keyword_usage(draft):
    rows = []
    for item in (draft or {}).get("keyword_usage_audit") or []:
        if not isinstance(item, dict) or not item.get("passed"):
            continue
        query = _clean_text(item.get("query"), 500)
        if not query:
            continue
        coverage = "точное" if item.get("exact_occurrences") else "морфология"
        fields = list(dict.fromkeys(
            _clean_text(location.get("field"), 80)
            for location in item.get("locations") or []
            if isinstance(location, dict) and _clean_text(location.get("field"), 80)
        ))
        field_text = ", ".join(fields) or _clean_text(item.get("placement"), 80) or "поле не указано"
        rows.append(f"{query} — {coverage}; {field_text}")
    return _excel_safe_text(rows)


def _export_characteristics(draft):
    rows = []
    for item in (draft or {}).get("used_characteristics") or []:
        if not isinstance(item, dict):
            continue
        name = _clean_text(item.get("name"), 160)
        value = _clean_text(item.get("value"), 700)
        if not name or not value:
            continue
        placement = _clean_text(item.get("placement"), 80)
        rows.append(f"{name}: {value}" + (f" — {placement}" if placement else ""))
    return _excel_safe_text(rows)


def _export_monitoring_keywords(rows):
    values = []
    for item in rows or []:
        if not isinstance(item, dict):
            continue
        query = _clean_text(item.get("search_query") or item.get("query"), 500)
        if not query:
            continue
        role = _clean_text(item.get("role"), 40)
        placement = _clean_text(item.get("placement"), 40)
        score = item.get("priority_score")
        meta = ", ".join(part for part in (placement, role, f"приоритет {score}" if score is not None else "") if part)
        values.append(f"{query}" + (f" — {meta}" if meta else ""))
    return _excel_safe_text(values)


def _content_export_values(row):
    context = row.get("context_json") or {}
    product = context.get("product") or {}
    draft = row.get("draft_json") or {}
    original_title = product.get("title") or row.get("project_product_name")
    original_description = product.get("description")
    original_hashtags = product.get("hashtags") or context.get("hashtags")
    model = row.get("gj_model") or draft.get("required_source_model") or _source_model_identifier(context)
    return (
        _excel_safe_text(row.get("sku")),
        _excel_safe_text(model),
        _excel_safe_text(row.get("project_product_name") or original_title),
        _excel_safe_text(original_title),
        _excel_safe_text(draft.get("title")),
        _excel_safe_text(original_description),
        _excel_safe_text(draft.get("description")),
        _export_hashtags(original_hashtags),
        _export_hashtags(draft.get("hashtags")),
        _export_keyword_usage(draft),
        _export_characteristics(draft),
        _export_monitoring_keywords(row.get("monitoring_keywords")),
    )


def build_project_content_export(rows, project_name="SEO-проект"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SEO-экспорт"
    sheet.append(SEO_CONTENT_EXPORT_HEADERS)
    for row in rows:
        sheet.append(_content_export_values(row))

    red = PatternFill("solid", fgColor="FF4747")
    white_bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=9, color="26332F")
    divider = Side(style="thin", color="DFE8DA")
    for cell in sheet[1]:
        cell.fill = red
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=divider)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=divider)
        sheet.row_dimensions[cell.row].height = 72
    widths = (16, 24, 44, 48, 48, 72, 72, 34, 34, 62, 58, 62)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(sheet.max_row, 1)}"
    sheet.row_dimensions[1].height = 32
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.oddHeader.center.text = _excel_safe_text(project_name, "SEO-проект")
    sheet.oddFooter.right.text = "Страница &P из &N"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_project_content(config, payload):
    """Export all project SKU and their source-backed generated content; never publish."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    if project_kind != "generation":
        raise ValueError("Экспорт SEO-контента доступен только в проектах генерации")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            marketplace = project.get("marketplace")
            stock_view = f"mv_{marketplace}_abc_product_stock_base"
            stock_join = (
                f"LEFT JOIN public.{stock_view} catalog ON catalog.artikul_wb::text=s.sku"
                if _relation_exists(cur, stock_view) else
                "LEFT JOIN (SELECT NULL::text AS artikul_wb, NULL::text AS gj_model) catalog ON false"
            )
            cur.execute(
                f"""SELECT s.sku, s.product_name AS project_product_name, catalog.gj_model,
                            c.context_json, d.draft_json,
                            coalesce(m.monitoring_keywords, '[]'::jsonb) AS monitoring_keywords
                     FROM public.seo_monitoring_project_skus s
                     {stock_join}
                     LEFT JOIN public.seo_generation_semantic_contexts c
                       ON c.project_id=s.project_id AND c.sku=s.sku
                     LEFT JOIN public.seo_generation_content_drafts d
                       ON d.project_id=s.project_id AND d.sku=s.sku
                     LEFT JOIN LATERAL (
                       SELECT jsonb_agg(jsonb_build_object(
                         'search_query', mk.search_query, 'placement', mk.placement,
                         'role', mk.role, 'priority_score', mk.priority_score
                       ) ORDER BY mk.priority_score DESC, mk.search_query) AS monitoring_keywords
                       FROM public.seo_generation_monitoring_keywords mk
                       WHERE mk.project_id=s.project_id AND mk.sku=s.sku AND mk.active
                     ) m ON true
                     WHERE s.project_id=%s
                     ORDER BY s.product_name NULLS LAST, s.sku""",
                (project_id,),
            )
            rows = [dict(row) for row in cur.fetchall()]
    body = build_project_content_export(rows, project.get("name") or "SEO-проект")
    slug = re.sub(r"[^a-zA-Z0-9_-]+", "_", _clean_text(project.get("name"), 80)).strip("_").lower() or "seo_project"
    return body, f"{slug}_{date.today().strftime('%Y%m%d')}.xlsx", len(rows)


def _template_rows(filename, file_bytes):
    suffix = str(filename or "").lower().rsplit(".", 1)[-1]
    if suffix == "xlsx":
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = workbook["SKU"] if "SKU" in workbook.sheetnames else workbook.active
        values = list(sheet.iter_rows(values_only=True))
    elif suffix == "csv":
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = file_bytes.decode("cp1251")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        values = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise ValueError("Поддерживаются только XLSX и CSV")
    if not values:
        raise ValueError("Файл пуст")
    headers = [_clean_text(value, 100).lower() for value in values[0]]
    aliases = {"sku", "артикул", "артикул sku", "marketplace_sku", "nm_id", "nmid"}
    sku_index = next((index for index, value in enumerate(headers) if value in aliases), None)
    if sku_index is None:
        raise ValueError("В шаблоне не найден обязательный столбец sku")
    name_aliases = {"product_name", "название", "наименование", "товар"}
    name_index = next((index for index, value in enumerate(headers) if value in name_aliases), None)
    parsed, duplicates, invalid = [], [], []
    seen = set()
    for row_number, row in enumerate(values[1:], start=2):
        raw_sku = row[sku_index] if sku_index < len(row) else None
        if isinstance(raw_sku, float) and raw_sku.is_integer():
            raw_sku = int(raw_sku)
        sku = _clean_text(raw_sku, 200)
        if not sku:
            if any(_clean_text(value, 100) for value in row):
                invalid.append({"row": row_number, "error": "SKU не заполнен"})
            continue
        if sku in seen:
            duplicates.append({"row": row_number, "sku": sku})
            continue
        seen.add(sku)
        product_name = _clean_text(row[name_index] if name_index is not None and name_index < len(row) else None, 1000)
        parsed.append({"sku": sku, "product_name": product_name or None, "row": row_number})
    if not parsed:
        raise ValueError("В файле нет заполненных SKU")
    return parsed, duplicates, invalid


def parse_project_template(config, payload):
    import base64

    filename = _clean_text(payload.get("filename"), 300)
    encoded = str(payload.get("file_base64") or "")
    if not filename or not encoded:
        raise ValueError("Выберите XLSX или CSV файл")
    try:
        file_bytes = base64.b64decode(encoded, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise ValueError("Не удалось прочитать файл") from exc
    if len(file_bytes) > 8 * 1024 * 1024:
        raise ValueError("Размер файла не должен превышать 8 МБ")
    rows, duplicates, invalid = _template_rows(filename, file_bytes)
    marketplace = _clean_text(payload.get("marketplace"), 20).lower()
    if marketplace not in {"ozon", "wb"}:
        raise ValueError("Маркетплейс должен быть Ozon или WB")
    prefix = "ozon" if marketplace == "ozon" else "wb"
    sku_values = [row["sku"] for row in rows]
    found = {}
    with _conn(config) as conn, conn.cursor() as cur:
        cur.execute(
            f"""SELECT artikul_wb::text AS sku, naimenovanie AS product_name
                FROM public.mv_{prefix}_abc_product_stock_base
                WHERE artikul_wb::text = ANY(%s)""",
            (sku_values,),
        )
        found = {str(row["sku"]): row.get("product_name") for row in cur.fetchall()}
    for row in rows:
        row["found"] = row["sku"] in found
        row["product_name"] = row.get("product_name") or found.get(row["sku"])
    return {
        "ok": True, "rows": rows, "row_count": len(rows), "found_count": sum(1 for row in rows if row["found"]),
        "not_found_count": sum(1 for row in rows if not row["found"]),
        "duplicates": duplicates, "invalid": invalid,
    }


def _classify(rows):
    traffic_values = [float(row["traffic"]) for row in rows if row.get("traffic") is not None and float(row["traffic"]) > 0]
    demand_values = [float(row["search_demand"]) for row in rows if row.get("search_demand") is not None and float(row["search_demand"]) > 0]
    conversions = []
    for row in rows:
        traffic = _number(row.get("traffic"))
        orders = _number(row.get("orders"))
        conversion = orders * 100 / traffic if traffic and orders is not None else None
        row["conversion_pct"] = round(conversion, 2) if conversion is not None else None
        if conversion and conversion > 0:
            conversions.append(conversion)
    traffic_mid = median(traffic_values) if traffic_values else None
    demand_mid = median(demand_values) if demand_values else None
    conversion_mid = median(conversions) if conversions else None
    for row in rows:
        traffic = _number(row.get("traffic")) or 0
        demand = _number(row.get("search_demand")) or 0
        orders = _number(row.get("orders")) or 0
        conversion = row.get("conversion_pct")
        if orders > 0 and ((traffic_mid is not None and traffic >= traffic_mid) or (conversion_mid is not None and conversion >= conversion_mid)):
            row["segment"] = "priority"
            row["segment_reason"] = "Есть заказы и трафик или конверсия не ниже медианы проекта"
        elif (traffic_mid is not None and traffic >= traffic_mid) or (demand_mid is not None and demand >= demand_mid):
            row["segment"] = "core"
            row["segment_reason"] = "Трафик или спрос не ниже медианы проекта"
        else:
            row["segment"] = "tail"
            row["segment_reason"] = "Низкочастотный запрос внутри текущего снимка"
    return rows


def project_detail(config, project_id, project_kind="monitoring"):
    project_id = _project_id(project_id)
    project_kind = _project_kind(project_kind)
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                "SELECT sku, product_name FROM public.seo_monitoring_project_skus WHERE project_id=%s ORDER BY product_name NULLS LAST, sku",
                (project_id,),
            )
            skus = [_serialize(row) for row in cur.fetchall()]
            cur.execute(
                """
                WITH latest AS (
                    SELECT max(snapshot_date) AS snapshot_date
                    FROM public.seo_monitoring_keyword_snapshots WHERE project_id=%s
                )
                SELECT k.sku, max(k.product_name) AS product_name, k.search_query,
                       round(avg(k.average_position), 1) AS average_position,
                       sum(k.search_demand) AS search_demand,
                       sum(k.traffic) AS traffic,
                       sum(k.cart_adds) AS cart_adds,
                       sum(k.orders) AS orders,
                       sum(k.revenue_rub) AS revenue_rub,
                       max(k.source) AS source,
                       max(k.snapshot_date) AS snapshot_date
                FROM public.seo_monitoring_keyword_snapshots k, latest
                WHERE k.project_id=%s AND k.snapshot_date=latest.snapshot_date
                GROUP BY k.sku, k.search_query
                ORDER BY sum(k.traffic) DESC NULLS LAST, sum(k.search_demand) DESC NULLS LAST, k.search_query
                LIMIT 1000
                """,
                (project_id, project_id),
            )
            keywords = _classify([_serialize(row) for row in cur.fetchall()])
            cur.execute(
                """
                SELECT snapshot_date,
                       count(DISTINCT search_query)::int AS keyword_count,
                       round(avg(average_position), 1) AS average_position,
                       sum(traffic) AS traffic,
                       sum(orders) AS orders
                FROM public.seo_monitoring_keyword_snapshots
                WHERE project_id=%s
                GROUP BY snapshot_date ORDER BY snapshot_date
                """,
                (project_id,),
            )
            dynamics = [_serialize(row) for row in cur.fetchall()]
    project = _serialize(project)
    traffic = sum(_number(row.get("traffic")) or 0 for row in keywords) if keywords else None
    orders = sum(_number(row.get("orders")) or 0 for row in keywords) if keywords else None
    positions = [_number(row.get("average_position")) for row in keywords if row.get("average_position") is not None]
    summary = {
        "sku_count": len(skus),
        "keyword_count": len({row["search_query"] for row in keywords}),
        "average_position": round(sum(positions) / len(positions), 1) if positions else None,
        "traffic": traffic,
        "orders": orders,
        "conversion_pct": round(orders * 100 / traffic, 2) if traffic and orders is not None else None,
        "priority_count": sum(1 for row in keywords if row.get("segment") == "priority"),
        "core_count": sum(1 for row in keywords if row.get("segment") == "core"),
        "tail_count": sum(1 for row in keywords if row.get("segment") == "tail"),
    }
    return {"ok": True, "project": project, "skus": skus, "summary": summary, "keywords": keywords, "dynamics": dynamics}


KEYWORD_SEGMENT_LABELS = {"priority": "Приоритет", "core": "Ядро", "tail": "Хвост"}
KEYWORD_FREQUENCY_LABELS = {"high": "ВЧ", "mid": "СЧ", "low": "НЧ"}

SEO_PRODUCT_FAMILIES = {
    "рубашка": ("рубаш",), "блузка": ("блуз",), "футболка": ("футбол",),
    "лонгслив": ("лонгслив",), "майка": ("майк",), "топ": (" топ", "топ "),
    "платье": ("плать",), "юбка": ("юбк",), "брюки": ("брюк",),
    "джинсы": ("джинс",), "шорты": ("шорт",), "легинсы": ("легин", "лосин"),
    "свитшот": ("свитшот",), "худи": ("худи",), "толстовка": ("толстов",),
    "джемпер": ("джемпер",), "свитер": ("свитер",), "кардиган": ("кардиган",),
    "жакет": ("жакет",), "пиджак": ("пиджак",), "куртка": ("куртк",),
    "пальто": ("пальто",), "жилет": ("жилет",), "костюм": ("костюм",),
    "комбинезон": ("комбинез",), "белье": ("бель", "трус", "бюст"),
    "носки": ("носк",), "колготки": ("колгот",), "пижама": ("пижам",),
    "обувь": ("кроссов", "кед", "ботин", "сапог", "туфл", "сандал", "тапоч"),
    "сумка": ("сумк", "рюкзак"), "головной убор": ("шапк", "кепк", "панам", "шляп"),
}
SEO_RELEVANCE_STOPWORDS = {
    "для", "или", "как", "это", "при", "под", "над", "без", "про", "все", "товар",
    "купить", "озон", "ozon", "интернет", "магазин", "женский", "женская", "женские",
    "мужской", "мужская", "мужские", "детский", "детская", "детские", "gloria", "jeans",
}
SEO_GLOBAL_NEGATIVE_DICT_PATH = Path(__file__).with_name("config") / "seo_global_negative_keywords.json"


def _load_global_negative_dictionary(path=SEO_GLOBAL_NEGATIVE_DICT_PATH):
    """Load only the curated repository dictionary; AI history is never a hard gate."""
    try:
        payload = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, TypeError):
        return {"version": "unavailable", "allowlist_exact": set(), "categories": []}
    allowlist = {_seo_normalize(value) for value in payload.get("allowlist_exact", []) if _seo_normalize(value)}
    categories = []
    for raw in payload.get("categories", []):
        if not isinstance(raw, dict):
            continue
        category = {
            "id": _clean_text(raw.get("id"), 80) or "uncategorized",
            "label": _clean_text(raw.get("label"), 160) or "глобальный минус-словарь",
        }
        for field in ("exact", "tokens", "phrases"):
            category[field] = tuple(dict.fromkeys(
                normalized for normalized in (_seo_normalize(value) for value in raw.get(field, [])) if normalized
            ))
        if any(category[field] for field in ("exact", "tokens", "phrases")):
            categories.append(category)
    return {
        "version": _clean_text(payload.get("version"), 80) or "unknown",
        "allowlist_exact": allowlist,
        "categories": categories,
    }


def _global_negative_match(query, negative_dictionary=None):
    dictionary = negative_dictionary if negative_dictionary is not None else _load_global_negative_dictionary()
    normalized = _seo_normalize(query)
    if not normalized or normalized in dictionary.get("allowlist_exact", set()):
        return None
    tokens = set(normalized.split())
    padded = f" {normalized} "
    for category in dictionary.get("categories", []):
        for term in category.get("exact", ()):
            if normalized == term:
                return category, term
        for term in category.get("tokens", ()):
            if term in tokens:
                return category, term
        for term in category.get("phrases", ()):
            if f" {term} " in padded:
                return category, term
    return None


def _seo_normalize(value):
    return re.sub(r"\s+", " ", re.sub(r"[^0-9a-zа-яё]+", " ", str(value or "").lower().replace("ё", "е"))).strip()


def _seo_families(value):
    padded = f" {_seo_normalize(value)} "
    return {family for family, stems in SEO_PRODUCT_FAMILIES.items() if any(stem in padded for stem in stems)}


def _seo_audience(value):
    text = _seo_normalize(value)
    result = set()
    if re.search(r"\b(жен|женск|женщин|девуш)", text): result.add("female")
    if re.search(r"\b(муж|мужск|мужчин)", text): result.add("male")
    if re.search(r"\b(девоч|мальчик|дет|подрост|школь)", text): result.add("child")
    if re.search(r"\b(взросл)", text): result.add("adult")
    return result


def _seo_meaningful_tokens(value):
    return {
        token for token in _seo_normalize(value).split()
        if len(token) >= 4 and token not in SEO_RELEVANCE_STOPWORDS and not re.fullmatch(r"[a-z]{1,5}\d+", token)
    }


def _seo_semantic_type(query, product_context):
    text = _seo_normalize(query)
    if "gloria jeans" in text or "глория джинс" in text:
        return "brand"
    if _seo_families(query):
        return "product"
    if _seo_audience(query):
        return "audience"
    if _seo_meaningful_tokens(query) & _seo_meaningful_tokens(product_context):
        return "attribute"
    return "other"


def _local_keyword_relevance(query, product_context, negative_dictionary=None):
    """Conservative hard gates adapted from the SEO bot relevance checker."""
    normalized = _seo_normalize(query)
    if len(normalized) < 2 or not re.search(r"[a-zа-яё]", normalized):
        return "reject", "Пустой или цифровой запрос"
    if re.search(r"https?\b|www\b|[a-z]+\d{5,}|(.)\1{5,}", normalized):
        return "reject", "Технический или кодовый шум"
    negative_match = _global_negative_match(query, negative_dictionary)
    if negative_match:
        category, term = negative_match
        return "reject", f"Глобальный минус-словарь · {category['label']}: {term}"
    product_audience = _seo_audience(product_context)
    query_audience = _seo_audience(query)
    if ({"female", "male"} <= product_audience | query_audience
            and (("female" in product_audience and "male" in query_audience)
                 or ("male" in product_audience and "female" in query_audience))):
        return "reject", "Конфликт пола с карточкой товара"
    if (("child" in product_audience and "adult" in query_audience)
            or ("adult" in product_audience and "child" in query_audience)):
        return "reject", "Конфликт возраста с карточкой товара"
    product_families = _seo_families(product_context)
    query_families = _seo_families(query)
    if product_families and query_families:
        if product_families & query_families:
            return "keep", "Совпадает товарная группа карточки"
        return "reject", "Другая товарная группа"
    if "gloria jeans" in normalized or "глория джинс" in normalized:
        return "keep", "Релевантный брендовый запрос"
    shared = _seo_meaningful_tokens(query) & _seo_meaningful_tokens(product_context)
    if shared:
        return "keep", "Совпадает характеристика карточки: " + ", ".join(sorted(shared)[:3])
    return "review", "Нужна смысловая проверка по карточке"


def _apply_source_frequency_ranks(rows):
    grouped = {}
    for row in rows:
        grouped.setdefault((row["sku"], row["source_group"]), []).append(row)
    for items in grouped.values():
        kept = [row for row in items if row["relevance_decision"] == "keep"]
        kept.sort(key=lambda row: (-(float(row.get("frequency_value") or 0)), row["search_query"]))
        total = sum(max(0.0, float(row.get("frequency_value") or 0)) for row in kept)
        cumulative = 0.0
        dense_rank = 0
        previous_value = None
        for row in kept:
            value = max(0.0, float(row.get("frequency_value") or 0))
            if previous_value is None or value != previous_value:
                dense_rank += 1
                previous_value = value
            before = cumulative / total if total else 1.0
            row["frequency_class"] = "high" if before < .80 else ("mid" if before < .95 else "low")
            row["frequency_rank"] = dense_rank
            cumulative += value
        for row in items:
            if row["relevance_decision"] != "keep":
                row["frequency_class"] = None
                row["frequency_rank"] = None


def _apply_keyword_priorities(rows):
    sources_by_key = {}
    for row in rows:
        if row["relevance_decision"] == "keep":
            sources_by_key.setdefault((row["sku"], row["search_query"]), set()).add(row["source_group"])
    for row in rows:
        if row["relevance_decision"] != "keep":
            row.update(priority_score=0, priority_label="none", source_overlap=False)
            continue
        overlap = len(sources_by_key.get((row["sku"], row["search_query"]), ())) > 1
        score = {"high": 55, "mid": 35, "low": 15}.get(row.get("frequency_class"), 10)
        if overlap: score += 20
        if float(row.get("orders") or 0) > 0: score += 15
        traffic = float(row.get("traffic") or 0)
        if traffic > 0: score += min(10, max(2, int(math.log10(traffic + 1) * 3)))
        score = min(100, score)
        row["source_overlap"] = overlap
        row["priority_score"] = score
        row["priority_label"] = "high" if score >= 70 else ("medium" if score >= 45 else "low")


def _keyword_scope(payload):
    """Filters shared by the keyword dashboards and the keyword table."""
    project_id = _project_id(payload.get("project_id"))
    source_group = _clean_text(payload.get("source_group"), 20).lower()
    date_from = _clean_text(payload.get("date_from"), 10) or None
    date_to = _clean_text(payload.get("date_to"), 10) or None
    filters = ["project_id=%s"]
    if source_group != "mpstats":
        filters.append("period_from = period_to")
    values = [project_id]
    if date_from:
        filters.append("snapshot_date >= %s")
        values.append(date_from)
    if date_to:
        filters.append("snapshot_date <= %s")
        values.append(date_to)
    query = _clean_text(payload.get("query"), 200)
    if query:
        filters.append("search_query ILIKE %s")
        values.append(f"%{query}%")
    sku = _clean_text(payload.get("sku"), 200)
    if sku:
        filters.append("sku ILIKE %s")
        values.append(f"%{sku}%")
    if source_group == "mpstats":
        filters.append("source LIKE 'mpstats_%%'")
    elif source_group == "api":
        filters.append("source NOT LIKE 'mpstats_%%'")
    column_filters = payload.get("column_filters")
    try:
        decoded_filters = json.loads(column_filters) if isinstance(column_filters, str) and column_filters else column_filters
    except (TypeError, ValueError, json.JSONDecodeError):
        decoded_filters = {}
    filter_columns = {
        "snapshot_date": ("snapshot_date", "date"),
        "search_query": ("search_query", "text"),
        "average_position": ("average_position", "number"),
        "search_demand": ("search_demand", "number"),
        "traffic": ("traffic", "number"),
        "orders": ("orders", "number"),
        "revenue_rub": ("revenue_rub", "number"),
        "source": ("source", "source"),
    }
    for key, item in (decoded_filters.items() if isinstance(decoded_filters, dict) else []):
        expression, value_type = filter_columns.get(key, (None, None))
        operator = str(item.get("op") or "").strip() if isinstance(item, dict) else ""
        raw_value = item.get("value") if isinstance(item, dict) else None
        if not expression or raw_value in (None, ""):
            continue
        if value_type == "number":
            numeric_value = _number(raw_value)
            sql_operator = {"eq": "=", "neq": "<>", "gt": ">", "gte": ">=", "lt": "<", "lte": "<="}.get(operator)
            if numeric_value is not None and sql_operator:
                filters.append(f"{expression} {sql_operator} %s")
                values.append(numeric_value)
        elif value_type == "date":
            clean_value = _clean_text(raw_value, 10)
            if clean_value:
                filters.append(f"{expression} = %s")
                values.append(clean_value)
        elif value_type == "source":
            clean_value = _clean_text(raw_value, 20).lower()
            if clean_value == "mpstats":
                filters.append("source LIKE 'mpstats_%%'")
            elif clean_value == "seller_api":
                filters.append("source NOT LIKE 'mpstats_%%'")
        else:
            clean_value = _clean_text(raw_value, 300)
            sql_operator = {"contains": "ILIKE", "not_contains": "NOT ILIKE", "eq": "=", "neq": "<>"}.get(operator)
            if clean_value and sql_operator:
                filters.append(f"{expression} {sql_operator} %s")
                values.append(f"%{clean_value}%" if operator in {"contains", "not_contains"} else clean_value)
    return project_id, " AND ".join(filters), values


# Per query and day first: `search_demand` repeats for every SKU of the same phrase,
# so it is taken once per day, while traffic/orders/GMV are summed across SKUs.
KEYWORD_AGGREGATE_SQL = """
    WITH base AS (
        SELECT * FROM public.seo_monitoring_keyword_snapshots WHERE {where}
    ), per_day AS (
        SELECT search_query, snapshot_date,
               max(search_demand) AS demand,
               sum(traffic) AS traffic,
               sum(orders) AS orders,
               sum(revenue_rub) AS gmv,
               avg(nullif(average_position, 0)) AS position,
               count(DISTINCT sku) AS skus
        FROM base GROUP BY search_query, snapshot_date
    ), per_query AS (
        SELECT search_query,
               sum(demand) AS demand,
               sum(traffic) AS traffic,
               sum(orders) AS orders,
               sum(gmv) AS gmv,
               round(avg(position), 1) AS average_position,
               max(skus)::int AS sku_count,
               count(*)::int AS day_count,
               CASE WHEN sum(traffic) > 0 THEN round(sum(orders) * 100.0 / sum(traffic), 2) END AS conversion_pct
        FROM per_day GROUP BY search_query
    )
"""


def _keyword_thresholds(cur, where, values):
    """Cut-offs for the whole selection, not for a single day or SKU."""
    cur.execute(
        KEYWORD_AGGREGATE_SQL.format(where=where)
        + """SELECT percentile_cont(0.80) WITHIN GROUP (ORDER BY traffic) AS traffic_p80,
                    percentile_cont(0.50) WITHIN GROUP (ORDER BY traffic) AS traffic_p50,
                    percentile_cont(0.66) WITHIN GROUP (ORDER BY traffic) AS traffic_p66,
                    percentile_cont(0.33) WITHIN GROUP (ORDER BY traffic) AS traffic_p33,
                    percentile_cont(0.50) WITHIN GROUP (ORDER BY conversion_pct)
                        FILTER (WHERE conversion_pct > 0) AS conversion_mid,
                    count(*)::int AS query_count
             FROM per_query""",
        values,
    )
    row = _serialize(cur.fetchone() or {})
    return {key: _number(value) if key != "query_count" else value for key, value in row.items()}


def _keyword_class_sql():
    """Segment and frequency class of a phrase, from its rank inside the whole selection.

    Ranks beat absolute cut-offs here: period traffic has many equal values, so a
    `traffic >= percentile` rule pushed almost every phrase into one class.
    """
    return """
        CASE
            WHEN per_query.orders > 0
                 AND (per_query.conversion_pct IS NULL OR per_query.conversion_pct >= %s) THEN 'priority'
            WHEN ntile(100) OVER (ORDER BY per_query.traffic NULLS FIRST) >= 40 THEN 'core'
            ELSE 'tail'
        END AS segment,
        CASE
            WHEN ntile(100) OVER (ORDER BY per_query.traffic NULLS FIRST) >= 80 THEN 'high'
            WHEN ntile(100) OVER (ORDER BY per_query.traffic NULLS FIRST) >= 40 THEN 'mid'
            ELSE 'low'
        END AS frequency
    """


def _class_values(thresholds):
    return [thresholds.get("conversion_mid") or 0]


def project_keyword_stats(config, payload):
    """Dashboards over the collected phrases: totals, daily series and class breakdown."""
    project_id, where, values = _keyword_scope(payload)
    project_kind = _project_kind(payload.get("project_kind"))
    segment = _clean_text(payload.get("segment"), 20)
    frequency = _clean_text(payload.get("frequency"), 20)
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                """SELECT min(snapshot_date) AS first_day, max(snapshot_date) AS last_day,
                          count(DISTINCT snapshot_date)::int AS day_count
                   FROM public.seo_monitoring_keyword_snapshots
                   WHERE project_id=%s AND period_from = period_to""",
                (project_id,),
            )
            available = _serialize(cur.fetchone() or {})
            thresholds = _keyword_thresholds(cur, where, values)
            if not thresholds.get("query_count"):
                return {"ok": True, "available": available, "summary": {}, "series": [], "segments": [], "frequency": [], "top": [], "thresholds": thresholds}
            class_values = _class_values(thresholds)
            having = []
            class_filter_values = []
            if segment in KEYWORD_SEGMENT_LABELS:
                having.append("classified.segment = %s")
                class_filter_values.append(segment)
            if frequency in KEYWORD_FREQUENCY_LABELS:
                having.append("classified.frequency = %s")
                class_filter_values.append(frequency)
            class_where = ("WHERE " + " AND ".join(having)) if having else ""
            classified_sql = (
                KEYWORD_AGGREGATE_SQL.format(where=where)
                + f", classified AS (SELECT per_query.*, {_keyword_class_sql()} FROM per_query) "
            )
            cur.execute(
                classified_sql
                + f"""SELECT count(*)::int AS query_count, sum(traffic) AS traffic, sum(demand) AS demand,
                             sum(orders) AS orders, sum(gmv) AS gmv, round(avg(average_position), 1) AS average_position,
                             sum(sku_count)::int AS sku_links
                      FROM classified {class_where}""",
                values + class_values + class_filter_values,
            )
            summary = _serialize(cur.fetchone() or {})
            summary["conversion_pct"] = (
                round(float(summary["orders"]) * 100 / float(summary["traffic"]), 2)
                if summary.get("traffic") and summary.get("orders") is not None else None
            )
            cur.execute(
                classified_sql
                + f"""SELECT segment, count(*)::int AS query_count, sum(traffic) AS traffic,
                             sum(orders) AS orders, sum(gmv) AS gmv
                      FROM classified {class_where} GROUP BY segment""",
                values + class_values + class_filter_values,
            )
            segments = [_serialize(row) for row in cur.fetchall()]
            cur.execute(
                classified_sql
                + f"""SELECT frequency, count(*)::int AS query_count, sum(traffic) AS traffic,
                             sum(orders) AS orders, sum(gmv) AS gmv
                      FROM classified {class_where} GROUP BY frequency""",
                values + class_values + class_filter_values,
            )
            frequency_rows = [_serialize(row) for row in cur.fetchall()]
            cur.execute(
                classified_sql
                + f"""SELECT search_query, segment, frequency, traffic, demand, orders, gmv,
                             average_position, conversion_pct, sku_count, day_count
                      FROM classified {class_where}
                      ORDER BY traffic DESC NULLS LAST LIMIT 15""",
                values + class_values + class_filter_values,
            )
            top = [_serialize(row) for row in cur.fetchall()]
            cur.execute(
                KEYWORD_AGGREGATE_SQL.format(where=where)
                + """SELECT snapshot_date AS day, sum(traffic) AS traffic, sum(orders) AS orders,
                           count(DISTINCT search_query)::int AS query_count
                    FROM base GROUP BY snapshot_date ORDER BY snapshot_date""",
                values,
            )
            series = [_serialize(row) for row in cur.fetchall()]
    return {
        "ok": True, "available": available, "summary": summary, "series": series,
        "segments": segments, "frequency": frequency_rows, "top": top, "thresholds": thresholds,
        "segment_labels": KEYWORD_SEGMENT_LABELS, "frequency_labels": KEYWORD_FREQUENCY_LABELS,
    }


def project_keywords(config, payload):
    """Collected phrases as rows, ranked by the totals of the whole selected period."""
    project_id, where, values = _keyword_scope(payload)
    project_kind = _project_kind(payload.get("project_kind"))
    segment = _clean_text(payload.get("segment"), 20)
    frequency = _clean_text(payload.get("frequency"), 20)
    try:
        limit = max(10, min(int(payload.get("limit") or 50), 500))
        page = max(1, int(payload.get("page") or 1))
    except (TypeError, ValueError) as exc:
        raise ValueError("limit и page должны быть числами") from exc
    sort_columns = {
        "traffic": "rows.traffic", "search_demand": "rows.search_demand", "average_position": "rows.average_position",
        "orders": "rows.orders", "revenue_rub": "rows.revenue_rub", "search_query": "rows.search_query",
        "sku": "rows.sku", "snapshot_date": "rows.snapshot_date", "source": "rows.source",
        "relevance_decision": "rows.relevance_decision", "frequency_class": "rows.frequency_class",
        "priority_score": "rows.priority_score",
    }
    sort_col = sort_columns.get(str(payload.get("sort_col") or "traffic"), "rows.traffic")
    sort_dir = "ASC" if str(payload.get("sort_dir") or "desc").lower() == "asc" else "DESC"
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                """SELECT snapshot_date, count(*)::int AS row_count, count(DISTINCT sku)::int AS sku_count
                   FROM public.seo_monitoring_keyword_snapshots
                   WHERE project_id=%s AND period_from = period_to
                   GROUP BY snapshot_date ORDER BY snapshot_date DESC""",
                (project_id,),
            )
            snapshots = [_serialize(row) for row in cur.fetchall()]
            if not snapshots:
                return {"ok": True, "rows": [], "snapshots": [], "total": 0, "page": 1, "total_pages": 1}
            thresholds = _keyword_thresholds(cur, where, values)
            class_values = _class_values(thresholds)
            class_filters = []
            class_filter_values = []
            if segment in KEYWORD_SEGMENT_LABELS:
                class_filters.append("classified.segment = %s")
                class_filter_values.append(segment)
            if frequency in KEYWORD_FREQUENCY_LABELS:
                class_filters.append("classified.frequency = %s")
                class_filter_values.append(frequency)
            class_where = ("WHERE " + " AND ".join(class_filters)) if class_filters else ""
            joined_sql = (
                KEYWORD_AGGREGATE_SQL.format(where=where)
                + f", classified AS (SELECT per_query.*, {_keyword_class_sql()} FROM per_query), "
                + f"""rows AS (
                        SELECT base.snapshot_date, base.sku, base.product_name, base.search_query, base.source,
                               base.average_position, base.search_demand, base.traffic, base.orders,
                               base.revenue_rub, classified.segment,
                               coalesce(generation_analysis.frequency_class, classified.frequency) AS frequency,
                               classified.traffic AS period_traffic, classified.orders AS period_orders,
                               classified.conversion_pct AS period_conversion,
                               coalesce(generation_analysis.semantic_type, analysis.semantic_type) AS semantic_type,
                               analysis.modifier_type, analysis.core_role,
                               coalesce(generation_analysis.priority_score, analysis.priority_score) AS priority_score,
                               coalesce(generation_analysis.priority_label, analysis.priority_label) AS priority_label,
                               coalesce(generation_analysis.relevance_reason, analysis.rationale) AS rationale,
                               generation_analysis.relevance_decision, generation_analysis.clean_query,
                               generation_analysis.frequency_rank, generation_analysis.source_overlap,
                               coalesce(generation_analysis.model, analysis.model) AS analysis_model,
                               coalesce(generation_analysis.analyzed_at, analysis.analyzed_at) AS analyzed_at
                        FROM base JOIN classified ON classified.search_query = base.search_query
                        LEFT JOIN public.seo_monitoring_keyword_analysis analysis
                          ON analysis.project_id = base.project_id AND analysis.search_query = base.search_query
                        LEFT JOIN public.seo_generation_keyword_analysis generation_analysis
                          ON generation_analysis.project_id = base.project_id
                         AND generation_analysis.sku = base.sku
                         AND generation_analysis.search_query = base.search_query
                         AND generation_analysis.source_group = CASE WHEN base.source LIKE 'mpstats_%%' THEN 'mpstats' ELSE 'api' END
                        {class_where}
                     ) """
            )
            query_values = values + class_values + class_filter_values
            row_filters = []
            row_filter_values = []
            relevance = _clean_text(payload.get("relevance"), 20).lower()
            if relevance in {"keep", "reject", "review"}:
                row_filters.append("rows.relevance_decision = %s")
                row_filter_values.append(relevance)
            elif relevance == "analyzed":
                row_filters.append("rows.relevance_decision IS NOT NULL")
            elif relevance == "clean":
                row_filters.append("(rows.relevance_decision = 'keep' OR rows.relevance_decision IS NULL)")
            priority_label = _clean_text(payload.get("priority_label"), 20).lower()
            if priority_label in {"high", "medium", "low", "none"}:
                row_filters.append("rows.priority_label = %s")
                row_filter_values.append(priority_label)
            try:
                analysis_filters = json.loads(payload.get("column_filters")) if isinstance(payload.get("column_filters"), str) and payload.get("column_filters") else (payload.get("column_filters") or {})
            except (TypeError, ValueError, json.JSONDecodeError):
                analysis_filters = {}
            for key, allowed_values in {
                "relevance_decision": {"keep", "reject", "review"},
                "frequency_class": {"high", "mid", "low"},
                "priority_label": {"high", "medium", "low", "none"},
            }.items():
                item = analysis_filters.get(key) if isinstance(analysis_filters, dict) else None
                clean_value = _clean_text(item.get("value"), 20).lower() if isinstance(item, dict) else ""
                if clean_value in allowed_values:
                    row_filters.append(f"rows.{key} = %s")
                    row_filter_values.append(clean_value)
            priority_filter = analysis_filters.get("priority_score") if isinstance(analysis_filters, dict) else None
            if isinstance(priority_filter, dict):
                numeric_value = _number(priority_filter.get("value"))
                sql_operator = {"eq": "=", "neq": "<>", "gt": ">", "gte": ">=", "lt": "<", "lte": "<="}.get(str(priority_filter.get("op") or ""))
                if numeric_value is not None and sql_operator:
                    row_filters.append(f"rows.priority_score {sql_operator} %s")
                    row_filter_values.append(numeric_value)
            row_where = (" WHERE " + " AND ".join(row_filters)) if row_filters else ""
            cur.execute(joined_sql + f"SELECT count(*)::int AS total FROM rows{row_where}", query_values + row_filter_values)
            total = int(cur.fetchone()["total"])
            total_pages = max(1, (total + limit - 1) // limit)
            page = min(page, total_pages)
            cur.execute(
                joined_sql
                + f"""SELECT * FROM rows{row_where} ORDER BY {sort_col} {sort_dir} NULLS LAST, rows.search_query
                      LIMIT %s OFFSET %s""",
                query_values + row_filter_values + [limit, (page - 1) * limit],
            )
            rows = [_serialize(row) for row in cur.fetchall()]
            for row in rows:
                row["segment_reason"] = (
                    f"Итоги периода: трафик {int(_number(row.get('period_traffic')) or 0)}, "
                    f"заказы {int(_number(row.get('period_orders')) or 0)}, "
                    f"конверсия {row.get('period_conversion') if row.get('period_conversion') is not None else '—'}%"
                )
    return {
        "ok": True, "rows": rows, "snapshots": snapshots, "total": total,
        "page": page, "total_pages": total_pages, "thresholds": thresholds,
        "segment_labels": KEYWORD_SEGMENT_LABELS, "frequency_labels": KEYWORD_FREQUENCY_LABELS,
    }


def analyze_project_keywords(config, payload, analyze_batch):
    """Clean and rank keywords at SKU x source grain, preserving every raw snapshot."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    try:
        ai_limit = max(0, min(int(payload.get("ai_limit") or 1000), 3000))
    except (TypeError, ValueError) as exc:
        raise ValueError("ai_limit должен быть числом") from exc
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if len(requested_skus) > 10:
        raise ValueError("Один пакет анализа может содержать не более 10 SKU")
    with _conn(config) as conn:
        ensure_schema(conn)
        model_settings = _ai_script_models(conn, "keyword_relevance")
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            if requested_skus:
                cur.execute(
                    "SELECT sku FROM public.seo_monitoring_project_skus WHERE project_id=%s AND sku=ANY(%s) ORDER BY sku",
                    (project_id, requested_skus),
                )
            else:
                cur.execute(
                    "SELECT sku FROM public.seo_monitoring_project_skus WHERE project_id=%s ORDER BY sku LIMIT 10",
                    (project_id,),
                )
            selected_skus = [str(row["sku"]) for row in cur.fetchall()]
            if not selected_skus:
                return {"ok": True, "analyzed": 0, "total": 0, "message": "В проекте нет выбранных SKU"}
            cur.execute(
                """SELECT sku, search_query,
                          CASE WHEN source LIKE 'mpstats_%%' THEN 'mpstats' ELSE 'api' END AS source_group,
                          max(product_name) AS product_name,
                          round(avg(search_demand), 2) AS frequency_value,
                          sum(coalesce(traffic, 0)) AS traffic,
                          sum(coalesce(orders, 0)) AS orders,
                          round(avg(nullif(average_position, 0)), 1) AS average_position
                   FROM public.seo_monitoring_keyword_snapshots
                   WHERE project_id=%s AND sku=ANY(%s)
                   GROUP BY sku, search_query, CASE WHEN source LIKE 'mpstats_%%' THEN 'mpstats' ELSE 'api' END
                   ORDER BY sku, source_group, frequency_value DESC NULLS LAST, search_query""",
                (project_id, selected_skus),
            )
            source_rows = [_serialize(row) for row in cur.fetchall()]
            context_by_sku = {}
            stock_view = f"mv_{project['marketplace']}_abc_product_stock_base"
            if _relation_exists(cur, stock_view):
                cur.execute(
                    f"""SELECT artikul_wb::text AS sku, naimenovanie AS product_name,
                               category_name, subcategory_name, gj_model, tg, tg_plus, cg, season
                        FROM public.{stock_view} WHERE artikul_wb::text=ANY(%s)""",
                    (selected_skus,),
                )
                context_by_sku = {str(row["sku"]): _serialize(row) for row in cur.fetchall()}
        if not source_rows:
            return {"ok": True, "analyzed": 0, "total": 0, "message": "В проекте пока нет собранных ключей"}

        pending_ai = []
        negative_dictionary = _load_global_negative_dictionary()
        for index, row in enumerate(source_rows):
            context = context_by_sku.get(row["sku"], {})
            product_context = " | ".join(str(value) for value in (
                context.get("product_name") or row.get("product_name"), context.get("category_name"),
                context.get("subcategory_name"), context.get("gj_model"), context.get("tg"),
                context.get("tg_plus"), context.get("cg"), context.get("season"),
            ) if value not in (None, ""))
            decision, reason = _local_keyword_relevance(row["search_query"], product_context, negative_dictionary)
            local_model = (
                f"local:global-negative-dict:{negative_dictionary['version']}"
                if reason.startswith("Глобальный минус-словарь")
                else "local:seo-bot-rules-v1"
            )
            row.update(
                analysis_id=str(index), product_context=product_context[:1200],
                relevance_decision=decision, clean_query=_seo_normalize(row["search_query"]) if decision == "keep" else "",
                relevance_reason=reason, semantic_type=_seo_semantic_type(row["search_query"], product_context),
                analysis_model=local_model,
            )
            if decision == "review":
                pending_ai.append(row)

        ai_rows = pending_ai[:ai_limit]
        model = "local:seo-bot-rules-v1"
        if ai_rows and analyze_batch:
            model_payload = analyze_batch(ai_rows, model_settings)
            model = _clean_text(model_payload.get("model"), 120) or model
            allowed_semantics = {"brand", "category", "product", "attribute", "audience", "use_case", "problem", "competitor", "other"}
            results = {str(item.get("analysis_id") or item.get("id") or ""): item for item in (model_payload.get("items") or [])}
            for row in ai_rows:
                item = results.get(row["analysis_id"])
                if not item:
                    continue
                decision = str(item.get("relevance_decision") or item.get("decision") or "review").lower()
                if decision not in {"keep", "reject", "review"}:
                    decision = "review"
                row["relevance_decision"] = decision
                row["clean_query"] = _seo_normalize(item.get("clean_query") or row["search_query"]) if decision == "keep" else ""
                row["relevance_reason"] = _clean_text(item.get("rationale") or item.get("reason"), 500) or row["relevance_reason"]
                semantic = str(item.get("semantic_type") or row["semantic_type"]).lower()
                row["semantic_type"] = semantic if semantic in allowed_semantics else "other"
                row["analysis_model"] = model

        _apply_source_frequency_ranks(source_rows)
        _apply_keyword_priorities(source_rows)
        db_rows = [(
            project_id, row["sku"], row["search_query"], row["source_group"], row["relevance_decision"],
            row.get("clean_query") or None, _clean_text(row.get("relevance_reason"), 500), row.get("semantic_type") or "other",
            row.get("frequency_class"), row.get("frequency_rank"), int(row.get("priority_score") or 0),
            row.get("priority_label") or "none", bool(row.get("source_overlap")), row.get("analysis_model") or model,
        ) for row in source_rows]
        if db_rows:
            with conn.cursor() as cur:
                execute_values(cur, """INSERT INTO public.seo_generation_keyword_analysis
                    (project_id, sku, search_query, source_group, relevance_decision, clean_query,
                     relevance_reason, semantic_type, frequency_class, frequency_rank, priority_score,
                     priority_label, source_overlap, model) VALUES %s
                    ON CONFLICT (project_id, sku, search_query, source_group) DO UPDATE SET
                      relevance_decision=EXCLUDED.relevance_decision, clean_query=EXCLUDED.clean_query,
                      relevance_reason=EXCLUDED.relevance_reason, semantic_type=EXCLUDED.semantic_type,
                      frequency_class=EXCLUDED.frequency_class, frequency_rank=EXCLUDED.frequency_rank,
                      priority_score=EXCLUDED.priority_score, priority_label=EXCLUDED.priority_label,
                      source_overlap=EXCLUDED.source_overlap, model=EXCLUDED.model, analyzed_at=now()""", db_rows)
    conn.commit()
    counts = {
        decision: sum(row["relevance_decision"] == decision for row in source_rows)
        for decision in ("keep", "reject", "review")
    }
    return {
        "ok": True,
        "analyzed": len(source_rows),
        "total": len(source_rows),
        "sku_count": len(selected_skus),
        "kept": counts["keep"],
        "rejected": counts["reject"],
        "review": counts["review"],
        "ai_checked": len(ai_rows),
        "priority": sum(row.get("priority_label") == "high" for row in source_rows),
        "model": model,
    }


def _valid_ai_model(value):
    model = str(value or "").strip()
    if not model or len(model) > 160 or not re.fullmatch(r"[A-Za-z0-9._:/-]+", model):
        raise ValueError("Укажите корректный идентификатор модели")
    return model


def _valid_ai_provider(value):
    provider = str(value or "").strip().lower()
    if provider not in {"local", "codex", "claude", "openrouter"}:
        raise ValueError("Неизвестный провайдер модели")
    return provider


def _ai_settings_from_conn(conn):
    with conn.cursor() as cur:
        cur.execute("SELECT script_key, primary_provider, primary_model, fallback_provider, fallback_model, updated_at FROM public.seo_generation_ai_settings")
        saved = {row["script_key"]: row for row in cur.fetchall()}
    rows = []
    for key, defaults in SEO_AI_SCRIPT_DEFAULTS.items():
        row = saved.get(key) or {}
        rows.append({
            "script_key": key,
            "label": defaults["label"],
            "primary_provider": row.get("primary_provider") or defaults["primary_provider"],
            "primary_model": row.get("primary_model") or defaults["primary_model"],
            "fallback_provider": row.get("fallback_provider") or defaults["fallback_provider"],
            "fallback_model": row.get("fallback_model") or defaults["fallback_model"],
            "updated_at": row.get("updated_at"),
        })
    return rows


def get_ai_settings(config):
    with _conn(config) as conn:
        ensure_schema(conn)
        return {"ok": True, "rows": [_serialize(row) for row in _ai_settings_from_conn(conn)]}


def save_ai_settings(config, payload):
    supplied = payload.get("rows")
    if not isinstance(supplied, list):
        raise ValueError("Настройки моделей должны быть списком")
    values = []
    seen = set()
    for item in supplied:
        key = str((item or {}).get("script_key") or "").strip()
        if key not in SEO_AI_SCRIPT_DEFAULTS or key in seen:
            raise ValueError("Неизвестный AI-сценарий")
        seen.add(key)
        values.append((key, _valid_ai_provider(item.get("primary_provider")), _valid_ai_model(item.get("primary_model")),
                       _valid_ai_provider(item.get("fallback_provider")), _valid_ai_model(item.get("fallback_model"))))
    if seen != set(SEO_AI_SCRIPT_DEFAULTS):
        raise ValueError("Передайте настройки для всех AI-сценариев")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            execute_values(cur, """INSERT INTO public.seo_generation_ai_settings
                (script_key, primary_provider, primary_model, fallback_provider, fallback_model) VALUES %s
                ON CONFLICT (script_key) DO UPDATE SET primary_provider=EXCLUDED.primary_provider,
                  primary_model=EXCLUDED.primary_model, fallback_provider=EXCLUDED.fallback_provider,
                  fallback_model=EXCLUDED.fallback_model, updated_at=now()""", values)
        conn.commit()
        return {"ok": True, "rows": [_serialize(row) for row in _ai_settings_from_conn(conn)]}


def _ai_script_models(conn, script_key):
    return next(row for row in _ai_settings_from_conn(conn) if row["script_key"] == script_key)


SEO_INTENT_COLORS = {
    "белый", "белая", "белые", "черный", "черная", "черные", "красный", "красная", "красные",
    "синий", "синяя", "синие", "зеленый", "зеленая", "зеленые", "розовый", "розовая", "розовые",
    "серый", "серая", "серые", "бежевый", "бежевая", "бежевые", "голубой", "голубая", "голубые",
}


def _json_text(value, limit=1200):
    parts = []
    def visit(item):
        if len(" ".join(parts)) >= limit:
            return
        if isinstance(item, dict):
            for key, nested in item.items():
                if str(key).lower() in {"src", "srcmobile", "href", "url", "images", "photos"}:
                    continue
                visit(nested)
        elif isinstance(item, list):
            for nested in item[:40]:
                visit(nested)
        elif isinstance(item, (str, int, float)):
            text = re.sub(r"\s+", " ", str(item)).strip()
            if text and not text.startswith(("http://", "https://")):
                parts.append(text)
    try:
        visit(json.loads(value) if isinstance(value, str) and value.lstrip().startswith(("{", "[")) else value)
    except (TypeError, ValueError, json.JSONDecodeError):
        visit(value)
    return " ".join(parts)[:limit]


def _normalize_product_intent(value, sku=""):
    text = _seo_normalize(value)
    text = text.replace("gloria jeans", " ").replace("глория джинс", " ")
    tokens = []
    for token in text.split():
        if token == str(sku).lower() or token in SEO_INTENT_COLORS or re.fullmatch(r"\d+[a-zа-я-]*", token):
            continue
        if token in {"товар", "озон", "ozon", "купить", "бренд"}:
            continue
        tokens.append(token)
    return " ".join(tokens[:5]).strip()


def _fallback_product_intent(row):
    source = " | ".join(str(row.get(key) or "") for key in ("product_name", "type_name", "category_name", "attributes_text"))
    families = _seo_families(source)
    product_type = sorted(families)[0] if families else ""
    if not product_type:
        candidate = _seo_normalize(row.get("type_name") or row.get("category_name") or row.get("product_name"))
        candidate = re.sub(r"\b(товары|товар|одежда|обувь|аксессуары|разное)\b", " ", candidate)
        product_type = " ".join(token for token in candidate.split() if len(token) > 2)[:60]
        product_type = " ".join(product_type.split()[:2])
    if not product_type:
        return ""
    audience_text = _seo_normalize(source)
    if re.search(r"\bдевоч", audience_text):
        return _normalize_product_intent(f"{product_type} для девочки", row.get("sku"))
    if re.search(r"\bмальчик", audience_text):
        return _normalize_product_intent(f"{product_type} для мальчика", row.get("sku"))
    audience = _seo_audience(source)
    if "female" in audience or "male" in audience:
        plural = product_type in {"брюки", "джинсы", "шорты", "легинсы", "носки", "колготки"}
        feminine = product_type.endswith(("а", "я", "ка"))
        if "female" in audience:
            adjective = "женские" if plural else ("женская" if feminine else "женский")
        else:
            adjective = "мужские" if plural else ("мужская" if feminine else "мужской")
        return _normalize_product_intent(f"{product_type} {adjective}", row.get("sku"))
    return _normalize_product_intent(product_type, row.get("sku"))


def generate_project_intents(config, payload, analyze_batch):
    """Infer one concise, auditable search container for each selected project SKU."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if len(requested_skus) > 10:
        raise ValueError("Один пакет интентов может содержать не более 10 SKU")
    with _conn(config) as conn:
        ensure_schema(conn)
        model_settings = _ai_script_models(conn, "product_intent")
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            sku_filter = "AND ps.sku=ANY(%s)" if requested_skus else ""
            values = [project_id, requested_skus] if requested_skus else [project_id]
            cur.execute(
                f"""SELECT ps.sku, ps.product_name FROM public.seo_monitoring_project_skus ps
                    WHERE ps.project_id=%s {sku_filter} ORDER BY ps.sku LIMIT 10""", values,
            )
            rows = [dict(item) for item in cur.fetchall()]
            if not rows:
                return {"ok": True, "generated": 0, "message": "SKU не найдены"}
            skus = [str(row["sku"]) for row in rows]
            by_sku = {str(row["sku"]): row for row in rows}
            if project["marketplace"] == "ozon" and _relation_exists(cur, "ozon_cat_products"):
                cur.execute(
                    """SELECT sku::text AS sku, nazvanie_tovara AS product_name, annotatsiya AS description,
                              rich_kontent_json, category_name, tip AS type_name, product_id, category_id
                       FROM public.ozon_cat_products WHERE sku::text=ANY(%s)""", (skus,),
                )
                for item in cur.fetchall():
                    row = by_sku[str(item["sku"])]
                    row.update(dict(item))
                    row["description"] = " ".join(filter(None, [
                        _clean_text(item.get("description"), 1200), _json_text(item.get("rich_kontent_json"), 800)
                    ]))[:1800]
                cur.execute(
                    """SELECT p.sku::text AS sku, coalesce(ca.attribute_name, co.attribute_name) AS name,
                              pa.value_text AS value
                       FROM public.ozon_cat_products p
                       JOIN public.ozon_cat_product_attributes pa ON pa.product_id=p.product_id
                       LEFT JOIN public.ozon_cat_category_attributes ca
                         ON ca.category_id=p.category_id AND ca.attribute_id=pa.attribute_id
                       LEFT JOIN public.ozon_cat_common_attributes co ON co.attribute_id=pa.attribute_id
                       WHERE p.sku::text=ANY(%s) AND pa.value_text IS NOT NULL""", (skus,),
                )
                attributes = {}
                for item in cur.fetchall():
                    name, value = _clean_text(item.get("name"), 120), _clean_text(item.get("value"), 250)
                    if name and value and not re.search(r"размер|штрих|код маркиров|таблиц", name, re.I):
                        attributes.setdefault(str(item["sku"]), []).append(f"{name}: {value}")
                        audience = by_sku[str(item["sku"])].setdefault("audience", {})
                        normalized_name = re.sub(r"[^a-zа-яё ]+", " ", name.lower()).strip()
                        if normalized_name == "пол":
                            audience["gender"] = value
                            audience.setdefault("evidence", []).append(f"характеристика Пол: {value}")
                        elif "возраст" in normalized_name or normalized_name == "целевая аудитория":
                            audience["age_group"] = value
                            audience.setdefault("evidence", []).append(f"характеристика {name}: {value}")
                for sku, items in attributes.items():
                    by_sku[sku]["attributes_text"] = " | ".join(items[:18])[:1800]
            cur.execute(
                """SELECT sku, source_group, clean_query, priority_score, frequency_class,
                          row_number() OVER (PARTITION BY sku, source_group ORDER BY priority_score DESC, frequency_rank NULLS LAST) AS rn
                   FROM public.seo_generation_keyword_analysis
                   WHERE project_id=%s AND sku=ANY(%s) AND relevance_decision='keep'""", (project_id, skus),
            )
            for item in cur.fetchall():
                if int(item.get("rn") or 99) <= 6:
                    by_sku[str(item["sku"])].setdefault("keywords", []).append({
                        "query": item.get("clean_query"), "source": item.get("source_group"),
                        "priority": item.get("priority_score"), "frequency": item.get("frequency_class"),
                    })
        prepared = []
        for row in rows:
            row.setdefault("description", "")
            row.setdefault("attributes_text", "")
            row.setdefault("category_name", "")
            row.setdefault("type_name", "")
            row.setdefault("keywords", [])
            audience = row.setdefault("audience", {})
            if not audience.get("age_group"):
                audience_text = _seo_normalize(" ".join(filter(None, [row.get("product_name"), row.get("description")])))
                age_hints = (
                    (r"\bподрост", "подростки"), (r"\bноворожден", "новорождённые"),
                    (r"\bмалыш", "малыши"), (r"\bдетск", "дети"),
                )
                for pattern, age_group in age_hints:
                    if re.search(pattern, audience_text):
                        audience["age_group"] = age_group
                        audience.setdefault("evidence", []).append("явное указание в названии или описании")
                        break
            prepared.append(row)
        ai_payload = analyze_batch(prepared, project["marketplace"], model_settings) if analyze_batch else {"items": [], "model": ""}
        results = {str(item.get("sku") or ""): item for item in (ai_payload.get("items") or []) if isinstance(item, dict)}
        model = _clean_text(ai_payload.get("model"), 120) or "local:intent-fallback-v1"
        db_rows = []
        ai_count = fallback_count = 0
        for row in prepared:
            item = results.get(str(row["sku"]), {})
            intent = _normalize_product_intent(item.get("intent") or item.get("search_intent"), row["sku"])
            row_model = model
            if intent:
                ai_count += 1
            else:
                intent = _fallback_product_intent(row)
                row_model = "local:intent-fallback-v1"
                fallback_count += 1
            if not intent:
                continue
            confidence = _number(item.get("confidence"))
            confidence = min(1, max(0, confidence)) if confidence is not None else (0.55 if row_model.startswith("local:") else 0.8)
            evidence = {
                "product_name": _clean_text(row.get("product_name"), 300),
                "category": _clean_text(row.get("category_name"), 200),
                "type": _clean_text(row.get("type_name"), 120),
                "audience": row.get("audience") or {},
                "keywords": row.get("keywords", [])[:12],
            }
            db_rows.append((project_id, row["sku"], intent, _clean_text(item.get("rationale"), 500), confidence,
                            json.dumps(evidence, ensure_ascii=False), row_model))
        if db_rows:
            with conn.cursor() as cur:
                execute_values(cur, """INSERT INTO public.seo_generation_product_intents
                    (project_id, sku, search_intent, rationale, confidence, evidence_json, model) VALUES %s
                    ON CONFLICT (project_id, sku) DO UPDATE SET search_intent=EXCLUDED.search_intent,
                      rationale=EXCLUDED.rationale, confidence=EXCLUDED.confidence,
                      evidence_json=EXCLUDED.evidence_json, model=EXCLUDED.model, generated_at=now()""", db_rows)
            conn.commit()
    return {"ok": True, "generated": len(db_rows), "ai_generated": ai_count,
            "fallback_generated": fallback_count, "model": model}


def _insert_snapshot_rows(cur, project_id, snapshot_date, period_from, period_to, source, rows):
    if not rows:
        return 0
    values = []
    seen = set()
    for row in rows:
        key = (str(row.get("sku")), str(row.get("search_query")))
        if key in seen:  # one phrase can arrive from several sorts in the same run
            continue
        seen.add(key)
        values.append((
            project_id, snapshot_date, period_from, period_to, row.get("sku"), row.get("product_name"),
            row.get("search_query"), row.get("average_position"), row.get("search_demand"), row.get("traffic"),
            row.get("cart_adds"), row.get("orders"), row.get("revenue_rub"), source,
            json.dumps(row.get("raw_payload") or {}, ensure_ascii=False, default=str),
        ))
    execute_values(
        cur,
        """INSERT INTO public.seo_monitoring_keyword_snapshots
           (project_id, snapshot_date, period_from, period_to, sku, product_name, search_query,
            average_position, search_demand, traffic, cart_adds, orders, revenue_rub, source, raw_payload)
           VALUES %s
           ON CONFLICT (project_id, snapshot_date, sku, search_query, source) DO UPDATE SET
             product_name=excluded.product_name, average_position=excluded.average_position,
             search_demand=excluded.search_demand, traffic=excluded.traffic, cart_adds=excluded.cart_adds,
             orders=excluded.orders, revenue_rub=excluded.revenue_rub, raw_payload=excluded.raw_payload,
             imported_at=now()""",
        values,
    )
    return len(values)


OZON_ANALYTICS_LAG_DAYS = 2  # Ozon calculates query analytics 1-2 days after the fact
OZON_DAILY_HISTORY_DAYS = 30  # inside the last month any interval works; deeper data is weekly only
MPSTATS_BATCH_LIMIT = 25
MPSTATS_TOKEN_SOURCES = {
    "ozon": Path(r"G:\Общие диски\Kokoc Marketplaces\Парсеры\Парсер позиций по ключам\Ozon\ozon_parser.py"),
    "wb": Path(r"G:\Общие диски\Kokoc Marketplaces\Парсеры\Парсер позиций по ключам\WB\mpstats_parser.py"),
}
MPSTATS_API_BASES = {
    "ozon": "https://mpstats.io/api/analytics/v1/oz/items",
    "wb": "https://mpstats.io/api/analytics/v1/wb/items",
}


def _load_mpstats_token(marketplace):
    marketplace = str(marketplace or "").strip().lower()
    token = os.environ.get(f"MPSTATS_{marketplace.upper()}_TOKEN") or os.environ.get("MPSTATS_TOKEN")
    if token:
        return token.strip()
    source_env = os.environ.get(f"MPSTATS_{marketplace.upper()}_TOKEN_SOURCE")
    source = Path(source_env) if source_env else MPSTATS_TOKEN_SOURCES.get(marketplace)
    if not source or not source.is_file():
        raise RuntimeError("Интеграция MPStats не настроена для выбранной площадки")
    text = source.read_text(encoding="utf-8", errors="ignore")
    match = re.search(r"^\s*TOKEN\s*=\s*([\"'])(.+?)\1", text, flags=re.MULTILINE)
    if not match:
        raise RuntimeError("В источнике интеграции MPStats не найден токен")
    return match.group(2)


def _mpstats_request_json(url, token, params, payload=None, retries=2):
    full_url = f"{url}?{urllib.parse.urlencode(params)}"
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8") if payload is not None else None
    request = urllib.request.Request(
        full_url,
        data=body,
        headers={"X-Mpstats-TOKEN": token, "Content-Type": "application/json"},
        method="POST" if payload is not None else "GET",
    )
    for attempt in range(retries + 1):
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                return json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            if exc.code == 404:
                return []
            retryable = exc.code == 429 or 500 <= exc.code <= 599
            if not retryable or attempt >= retries:
                raise RuntimeError(f"MPStats HTTP {exc.code}") from exc
            retry_after = exc.headers.get("Retry-After") if exc.headers else None
            try:
                pause = min(20.0, max(1.0, float(retry_after)))
            except (TypeError, ValueError):
                pause = min(20.0, 3.0 * (attempt + 1))
            time.sleep(pause)
        except (TimeoutError, urllib.error.URLError) as exc:
            if attempt >= retries:
                raise RuntimeError("MPStats не ответил вовремя") from exc
            time.sleep(min(10.0, 2.0 * (attempt + 1)))
    return []


def _customer_message_date(value):
    """Parse provider timestamps without inventing a date for missing values."""
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        try:
            return date.fromisoformat(text[:10])
        except ValueError:
            return None


def _normalize_mpstats_customer_review(sku, raw, period_from, period_to):
    """Normalize one MPStats Ozon comment; return None outside the requested period."""
    if not isinstance(raw, dict):
        return None
    review_date = _customer_message_date(
        raw.get("date") or raw.get("createdDate") or raw.get("created_at") or raw.get("published_at")
    )
    if review_date is None or review_date < period_from or review_date > period_to:
        return None
    text = _clean_text(raw.get("text") or raw.get("comment"), 20000)
    pros = _clean_text(raw.get("pros") or raw.get("advantages"), 10000)
    cons = _clean_text(raw.get("cons") or raw.get("defects"), 10000)
    rating = _number(raw.get("valuation") or raw.get("productValuation") or raw.get("rating"))
    source_id = _clean_text(raw.get("id") or raw.get("review_id"), 500)
    if not source_id:
        source_id = hashlib.sha256(
            "|".join((str(sku), review_date.isoformat(), str(rating or ""), text, pros, cons)).encode("utf-8")
        ).hexdigest()
    answer = raw.get("answer")
    answer_text = _clean_text(answer.get("text") if isinstance(answer, dict) else answer, 10000)
    return {
        "sku": str(sku), "message_type": "review", "source_message_id": source_id,
        "message_date": review_date, "rating": rating, "message_text": text,
        "pros": pros, "cons": cons, "answer_text": answer_text,
        "answered": bool(answer_text), "source": "mpstats_ozon_comments", "raw_payload": raw,
    }


def fetch_mpstats_product_reviews(sku, period_from, period_to, token=None):
    """Fetch exact-SKU Ozon reviews from the existing MPStats integration."""
    token = token or _load_mpstats_token("ozon")
    payload = _mpstats_request_json(
        f"{MPSTATS_API_BASES['ozon']}/{urllib.parse.quote(str(sku), safe='')}/comments",
        token,
        {},
    )
    if isinstance(payload, dict):
        rows = payload.get("comments") or payload.get("items") or payload.get("data") or []
    else:
        rows = payload
    if not isinstance(rows, list):
        rows = []
    normalized = [
        item for row in rows
        if (item := _normalize_mpstats_customer_review(sku, row, period_from, period_to)) is not None
    ]
    return normalized


def fetch_mpstats_project_keywords(marketplace, sku, date_from, date_to):
    marketplace = str(marketplace or "").strip().lower()
    if marketplace not in MPSTATS_API_BASES:
        raise ValueError("MPStats поддерживает только Ozon и WB")
    token = _load_mpstats_token(marketplace)
    url = f"{MPSTATS_API_BASES[marketplace]}/{urllib.parse.quote(str(sku), safe='')}/keywords"
    params = {"d1": str(date_from), "d2": str(date_to)}
    payload = None
    if marketplace == "wb":
        params.update({
            "fbs": 0, "positionType": "position", "dateSortDirection": "desc",
            "startRow": 0, "endRow": 5000,
        })
        payload = {"startRow": 0, "endRow": 5000, "filterModel": {}, "sortModel": []}
    response = _mpstats_request_json(url, token, params, payload)
    if isinstance(response, list):
        return response
    if not isinstance(response, dict):
        return []
    data = response.get("data")
    if isinstance(data, dict) and isinstance(data.get("words"), list):
        return data["words"]
    return response.get("words") if isinstance(response.get("words"), list) else []


def _normalize_mpstats_keyword_rows(marketplace, sku_row, words):
    rows = []
    marketplace = str(marketplace or "").strip().lower()
    for word in words or []:
        if not isinstance(word, dict):
            continue
        query = _clean_text(word.get("query") or word.get("query_cluster"), 1000)
        if not query:
            continue
        demand = word.get("count") if marketplace == "ozon" else word.get("wb_count")
        rows.append({
            "sku": sku_row["sku"],
            "product_name": sku_row.get("product_name"),
            "search_query": query,
            "average_position": _number(word.get("avg_position")),
            "search_demand": _number(demand),
            "traffic": None,
            "cart_adds": None,
            "orders": None,
            "revenue_rub": None,
            "raw_payload": {"provider": "mpstats", "marketplace": marketplace, "keyword": word},
        })
    return rows


def _normalize_mpstats_competitor(item, rank, niche_id=None, niche_name=None, marketplace="ozon"):
    marketplace = str(marketplace or "ozon").strip().lower()
    positions = [_number(value) for value in (item.get("positions") or [])]
    positions = [value for value in positions if value is not None and value > 0]
    item_position = _number(item.get("position") or item.get("avg_position") or item.get("average_position"))
    competitor_sku = _clean_text(item.get("id") or item.get("nmId") or item.get("nm_id"), 120)
    product_url = _clean_text(item.get("url"), 1000)
    if marketplace == "wb" and competitor_sku and not product_url:
        product_url = f"https://www.wildberries.ru/catalog/{competitor_sku}/detail.aspx"
    source_report = f"mpstats_{marketplace}_products_in_search"
    return {
        "competitor_sku": competitor_sku,
        "competitor_rank": int(rank),
        "product_name": _clean_text(item.get("name"), 500),
        "brand": _clean_text(item.get("brand") or item.get("brand_name"), 250),
        "seller": _clean_text(item.get("seller") or item.get("seller_name"), 250),
        "product_url": product_url,
        "image_url": _clean_text(item.get("thumb_middle") or item.get("thumb") or item.get("image"), 1000),
        "price": _number(item.get("final_price") or item.get("price")),
        "sales": _number(item.get("sales") or item.get("orders")),
        "revenue": _number(item.get("revenue")),
        "rating": _number(item.get("rating")),
        "reviews_count": int(_number(item.get("comments") or item.get("reviews")) or 0),
        "days_in_stock": int(_number(item.get("days_in_stock")) or 0),
        "revenue_potential": _number(item.get("revenue_potential")),
        "lost_profit": _number(item.get("lost_profit")),
        "average_position": round(sum(positions) / len(positions), 2) if positions else item_position,
        "latest_position": positions[-1] if positions else item_position,
        "niche_id": int(niche_id or item.get("niche_id")) if (niche_id or item.get("niche_id")) else None,
        "niche_name": _clean_text(niche_name or item.get("niche") or item.get("category"), 500),
        "source_report": source_report,
        "raw_payload": item,
    }


def fetch_mpstats_niche_competitors(marketplace, sku, date_from, date_to, own_skus=None, limit=10, search_intent=None):
    """Return Ozon or WB competitors from MPStats products-in-search for the product intent."""
    marketplace = str(marketplace or "").strip().lower()
    if marketplace not in {"ozon", "wb"}:
        raise ValueError("Сбор конкурентов MPStats не подключён для выбранной площадки")
    token = _load_mpstats_token(marketplace)
    intent = _clean_text(search_intent, 250)
    if not intent:
        raise ValueError(f"Для SKU {sku} не определён интент")
    if marketplace == "wb":
        page_size = max(50, min(200, int(limit) * 10))
        page = _mpstats_request_json(
            "https://mpstats.io/api/analytics/v1/wb/search/items",
            token,
            {"path": intent, "type": "json", "d1": str(date_from), "d2": str(date_to),
             "fbs": 0, "currency": "RUB", "startRow": 0, "endRow": page_size},
            {"startRow": 0, "endRow": page_size, "filterModel": {}, "sortModel": [],
             "fields": [], "filters": []},
        )
        source_rows = (page.get("data") or []) if isinstance(page, dict) else []
        result_count = page.get("total") if isinstance(page, dict) else None
    else:
        page = _mpstats_request_json(
            f"https://mpstats.io/api/oz/keywords/{urllib.parse.quote(intent, safe='')}/serp", token,
            {"d1": str(date_from), "d2": str(date_to), "isFbs": "false"},
        )
        source_rows = (page.get("items") or []) if isinstance(page, dict) else []
        result_count = page.get("result_count") if isinstance(page, dict) else None
    excluded_skus = {str(value) for value in (own_skus or [])}
    candidates = []
    for source_rank, item in enumerate(source_rows, start=1):
        competitor_sku = str(item.get("id") or item.get("nmId") or item.get("nm_id") or "").strip()
        brand = _seo_normalize(item.get("brand") or item.get("brand_name"))
        if not competitor_sku or competitor_sku in excluded_skus or "gloria jeans" in brand or "глория джинс" in brand:
            continue
        normalized = _normalize_mpstats_competitor(item, source_rank, marketplace=marketplace)
        if normalized.get("average_position") is None and marketplace == "wb":
            normalized["average_position"] = source_rank
            normalized["latest_position"] = source_rank
        candidates.append(normalized)
    candidates.sort(key=lambda row: (
        row.get("average_position") is None,
        row.get("average_position") if row.get("average_position") is not None else 10**9,
        -float(row.get("revenue") or 0),
    ))
    rows = candidates[:max(1, int(limit))]
    for rank, row in enumerate(rows, start=1):
        row["competitor_rank"] = rank
    return {"search_intent": intent, "source_report": f"mpstats_{marketplace}_products_in_search",
            "items": rows, "result_count": result_count}


def collect_project_niche_competitors(config, payload, fetcher=None):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if not requested_skus:
        raise ValueError("SKU для сбора конкурентов не переданы")
    if len(requested_skus) > MPSTATS_BATCH_LIMIT:
        raise ValueError(f"За один запрос можно обработать не более {MPSTATS_BATCH_LIMIT} SKU")
    force = str(payload.get("force") or "").lower() in {"1", "true", "yes"}
    fetch = fetcher or fetch_mpstats_niche_competitors
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            if project["marketplace"] not in {"ozon", "wb"}:
                raise ValueError("Сбор конкурентов MPStats не подключён для выбранной площадки")
            default_competitor_limit = 5 if project["marketplace"] == "wb" else 10
            competitor_limit = max(1, min(10, int(payload.get("competitor_limit") or default_competitor_limit)))
            cur.execute("SELECT sku FROM public.seo_monitoring_project_skus WHERE project_id=%s", (project_id,))
            own_skus = {str(row["sku"]) for row in cur.fetchall()}
            cur.execute(
                """SELECT ps.sku, ps.product_name, pi.search_intent
                   FROM public.seo_monitoring_project_skus ps
                   LEFT JOIN public.seo_generation_product_intents pi
                     ON pi.project_id=ps.project_id AND pi.sku=ps.sku
                   WHERE ps.project_id=%s AND ps.sku=ANY(%s) ORDER BY ps.sku""",
                (project_id, requested_skus),
            )
            sku_rows = [dict(row) for row in cur.fetchall()]
            if not sku_rows:
                raise ValueError("Переданные SKU не входят в проект")
            date_from = _as_date(payload.get("date_from") or project.get("date_from") or date.today())
            date_to = _as_date(payload.get("date_to") or project.get("date_to") or date.today())
            intents = list({row.get("search_intent") for row in sku_rows if row.get("search_intent")})
            cached = {}
            if intents and not force:
                cur.execute(
                    """SELECT * FROM public.seo_generation_niche_competitors
                       WHERE project_id=%s AND search_intent=ANY(%s) AND period_from=%s AND period_to=%s
                       ORDER BY search_intent, sku, competitor_rank""",
                    (project_id, intents, date_from, date_to),
                )
                cache_source = {}
                for item in cur.fetchall():
                    intent, source_sku = item["search_intent"], item["sku"]
                    cache_source.setdefault(intent, source_sku)
                    if cache_source[intent] == source_sku:
                        cached.setdefault(intent, []).append(dict(item))
    errors, collected, reused, requests = [], 0, 0, 0
    intent_results = {intent: rows[:competitor_limit] for intent, rows in cached.items() if rows}
    for row in sku_rows:
        intent = _clean_text(row.get("search_intent"), 200)
        if not intent:
            errors.append({"sku": row["sku"], "error": "Интент не определён"})
            continue
        source_rows = intent_results.get(intent)
        if source_rows is None:
            try:
                result = fetch(project["marketplace"], row["sku"], date_from, date_to, own_skus=own_skus,
                               limit=competitor_limit, search_intent=intent)
                requests += 1
                source_rows = result.get("items") or []
                if source_rows:
                    intent_results[intent] = source_rows
            except Exception as exc:
                errors.append({"sku": row["sku"], "intent": intent, "error": str(exc)[:500]})
                continue
        else:
            reused += 1
        if not source_rows:
            errors.append({"sku": row["sku"], "intent": intent, "error": "MPStats не вернул конкурентов"})
            continue
        values = []
        for rank, item in enumerate(source_rows[:competitor_limit], start=1):
            raw = item.get("raw_payload") or {}
            values.append((project_id, row["sku"], intent, item.get("niche_id"), item.get("niche_name"),
                           item.get("competitor_sku"), rank, item.get("product_name"), item.get("brand"),
                           item.get("seller"), item.get("product_url"), item.get("image_url"), item.get("price"),
                           item.get("sales"), item.get("revenue"), item.get("rating"), item.get("reviews_count"),
                           item.get("days_in_stock"), item.get("revenue_potential"), item.get("lost_profit"),
                           item.get("average_position"), item.get("latest_position"),
                           item.get("source_report") or f"mpstats_{project['marketplace']}_products_in_search",
                           date_from, date_to, json.dumps(raw, ensure_ascii=False, default=str)))
        with _conn(config) as conn:
            ensure_schema(conn)
            with conn.cursor() as cur:
                _assert_project_kind(cur, project_id, project_kind, for_update=True)
                cur.execute("DELETE FROM public.seo_generation_niche_competitors WHERE project_id=%s AND sku=%s", (project_id, row["sku"]))
                execute_values(cur, """INSERT INTO public.seo_generation_niche_competitors
                    (project_id, sku, search_intent, niche_id, niche_name, competitor_sku, competitor_rank,
                     product_name, brand, seller, product_url, image_url, price, sales, revenue, rating,
                     reviews_count, days_in_stock, revenue_potential, lost_profit, average_position,
                     latest_position, source_report, period_from, period_to, raw_payload)
                    VALUES %s""", values)
        collected += len(values)
    return {"ok": True, "project_id": project_id, "sku_count": len(sku_rows), "row_count": collected,
            "requests": requests, "reused": reused, "errors": errors,
            "competitor_limit": competitor_limit,
            "date_from": date_from.isoformat(), "date_to": date_to.isoformat()}


def project_niche_competitors(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                """SELECT competitors.competitor_rank, competitors.competitor_sku, competitors.product_name,
                          competitors.brand, competitors.seller, competitors.product_url, competitors.image_url,
                          competitors.price, competitors.sales, competitors.revenue, competitors.rating,
                          competitors.reviews_count, competitors.days_in_stock, competitors.revenue_potential,
                          competitors.lost_profit, competitors.average_position, competitors.latest_position,
                          competitors.source_report, competitors.niche_id, competitors.niche_name,
                          competitors.search_intent, competitors.period_from, competitors.period_to,
                          competitors.collected_at, review_status.review_count AS collected_review_count,
                          review_status.status AS review_status, review_status.last_error AS review_error,
                          review_status.period_days AS review_period_days,
                          review_status.collected_at AS reviews_collected_at
                   FROM public.seo_generation_niche_competitors competitors
                   LEFT JOIN public.seo_generation_competitor_review_status review_status
                     ON review_status.project_id=competitors.project_id
                    AND review_status.sku=competitors.sku
                    AND review_status.competitor_sku=competitors.competitor_sku
                   WHERE competitors.project_id=%s AND competitors.sku=%s
                   ORDER BY competitors.competitor_rank LIMIT 10""",
                (project_id, sku),
            )
            rows = [_serialize(row) for row in cur.fetchall()]
    return {"ok": True, "project_id": project_id, "sku": sku, "rows": rows,
            "count": len(rows), "intent": rows[0]["search_intent"] if rows else None,
            "niche_id": rows[0]["niche_id"] if rows else None,
            "niche_name": rows[0]["niche_name"] if rows else None,
            "source_report": rows[0]["source_report"] if rows else f"mpstats_{project['marketplace']}_products_in_search",
            "date_from": rows[0]["period_from"] if rows else None,
            "date_to": rows[0]["period_to"] if rows else None}


def collect_project_competitor_keywords(config, payload, fetcher=None):
    """Collect MPStats item keywords for the top niche competitors of selected own SKU."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if not requested_skus:
        raise ValueError("SKU для сбора ключей конкурентов не переданы")
    if len(requested_skus) > 5:
        raise ValueError("Один пакет может содержать не более 5 SKU проекта")
    force = str(payload.get("force") or "").lower() in {"1", "true", "yes"}
    fetch = fetcher or fetch_mpstats_project_keywords
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            if project["marketplace"] not in {"ozon", "wb"}:
                raise ValueError("Ключи конкурентов MPStats не подключены для выбранной площадки")
            default_keyword_limit = 50 if project["marketplace"] == "wb" else 0
            keyword_limit = max(0, min(500, int(payload.get("keyword_limit_per_competitor") or default_keyword_limit)))
            cur.execute(
                """SELECT sku, competitor_sku, competitor_rank, product_name
                   FROM public.seo_generation_niche_competitors
                   WHERE project_id=%s AND sku=ANY(%s)
                   ORDER BY sku, competitor_rank""", (project_id, requested_skus),
            )
            competitors = [dict(row) for row in cur.fetchall()]
            if not competitors:
                raise ValueError("Сначала соберите конкурентов по нише")
            period_from = _as_date(payload.get("date_from") or project.get("date_from") or date.today())
            period_to = _as_date(payload.get("date_to") or project.get("date_to") or date.today())

    by_competitor = {}
    for row in competitors:
        by_competitor.setdefault(str(row["competitor_sku"]), []).append(row)
    cache, requests, reused, empty, errors = {}, 0, 0, [], []
    if not force:
        with _conn(config) as conn, conn.cursor() as cur:
            ensure_schema(conn)
            cur.execute(
                """SELECT DISTINCT ON (competitor_sku, search_query)
                          competitor_sku, search_query, search_demand, average_position, raw_payload
                   FROM public.seo_generation_competitor_keywords
                   WHERE project_id=%s AND competitor_sku=ANY(%s) AND period_from=%s AND period_to=%s
                   ORDER BY competitor_sku, search_query, collected_at DESC""",
                (project_id, list(by_competitor), period_from, period_to),
            )
            for item in cur.fetchall():
                cache.setdefault(str(item["competitor_sku"]), []).append(dict(item))
    values = []
    for competitor_sku, links in by_competitor.items():
        normalized = cache.get(competitor_sku)
        if normalized:
            reused += 1
        else:
            try:
                words = fetch(project["marketplace"], competitor_sku, period_from, period_to)
                requests += 1
                normalized = _normalize_mpstats_keyword_rows(
                    project["marketplace"], {"sku": competitor_sku, "product_name": links[0].get("product_name")}, words,
                )
            except Exception as exc:
                errors.append({"competitor_sku": competitor_sku, "error": str(exc)[:500]})
                continue
        if not normalized:
            empty.append(competitor_sku)
            continue
        normalized = sorted(
            normalized,
            key=lambda item: (
                -float(item.get("search_demand") or 0),
                item.get("average_position") is None,
                float(item.get("average_position") or 10**9),
                str(item.get("search_query") or ""),
            ),
        )
        if keyword_limit:
            normalized = normalized[:keyword_limit]
        for link in links:
            for item in normalized:
                raw = item.get("raw_payload") or {}
                values.append((
                    project_id, link["sku"], competitor_sku, link.get("competitor_rank"),
                    link.get("product_name"), item["search_query"], item.get("search_demand"),
                    item.get("average_position"), period_from, period_to,
                    json.dumps(raw, ensure_ascii=False, default=str),
                ))
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind, for_update=True)
            cur.execute(
                "DELETE FROM public.seo_generation_competitor_keywords WHERE project_id=%s AND sku=ANY(%s)",
                (project_id, requested_skus),
            )
            if values:
                execute_values(cur, """INSERT INTO public.seo_generation_competitor_keywords
                    (project_id, sku, competitor_sku, competitor_rank, competitor_name, search_query,
                     search_demand, average_position, period_from, period_to, raw_payload) VALUES %s""", values)
    return {"ok": True, "project_id": project_id, "sku_count": len(set(row["sku"] for row in competitors)),
            "competitor_count": len(by_competitor), "row_count": len(values), "requests": requests,
            "reused": reused, "empty": empty, "errors": errors,
            "keyword_limit_per_competitor": keyword_limit or None,
            "date_from": period_from.isoformat(), "date_to": period_to.isoformat()}


def _analyze_project_competitor_keywords_unlocked(config, payload, analyze_batch):
    """Clean and rank competitor queries against the own product context."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if not requested_skus or len(requested_skus) > 10:
        raise ValueError("Передайте от 1 до 10 SKU проекта")
    with _conn(config) as conn:
        ensure_schema(conn)
        settings = _ai_script_models(conn, "competitor_keyword_relevance")
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                """SELECT k.sku, k.search_query, max(ps.product_name) AS product_name,
                          max(pi.search_intent) AS search_intent,
                          max(k.search_demand) AS frequency_value,
                          min(nullif(k.average_position, 0)) AS average_position,
                          count(DISTINCT k.competitor_sku)::integer AS competitor_coverage
                   FROM public.seo_generation_competitor_keywords k
                   JOIN public.seo_monitoring_project_skus ps ON ps.project_id=k.project_id AND ps.sku=k.sku
                   LEFT JOIN public.seo_generation_product_intents pi ON pi.project_id=k.project_id AND pi.sku=k.sku
                   WHERE k.project_id=%s AND k.sku=ANY(%s)
                   GROUP BY k.sku, k.search_query ORDER BY k.sku, frequency_value DESC NULLS LAST""",
                (project_id, requested_skus),
            )
            rows = [_serialize(row) for row in cur.fetchall()]
    if not rows:
        return {"ok": True, "analyzed": 0, "message": "Ключи конкурентов ещё не собраны"}
    negative_dictionary = _load_global_negative_dictionary()
    pending = []
    for index, row in enumerate(rows):
        context = " | ".join(value for value in (row.get("product_name"), row.get("search_intent")) if value)
        decision, reason = _local_keyword_relevance(row["search_query"], context, negative_dictionary)
        row.update(analysis_id=str(index), product_context=context[:1200], source_group="competitor_mpstats",
                   relevance_decision=decision, clean_query=_seo_normalize(row["search_query"]) if decision == "keep" else "",
                   relevance_reason=reason, semantic_type=_seo_semantic_type(row["search_query"], context),
                   analysis_model=f"local:competitor-rules:{negative_dictionary['version']}")
        if decision == "review":
            pending.append(row)
    cache_hits = 0
    if pending:
        intents = list({str(row.get("search_intent") or "") for row in pending})
        queries = list({str(row.get("search_query") or "") for row in pending})
        with _conn(config) as conn, conn.cursor() as cur:
            cur.execute(
                """SELECT search_intent, search_query, relevance_decision, clean_query,
                          relevance_reason, semantic_type, model
                   FROM public.seo_generation_competitor_ai_cache
                   WHERE project_id=%s AND search_intent=ANY(%s) AND search_query=ANY(%s)""",
                (project_id, intents, queries),
            )
            cached = {(str(item["search_intent"] or ""), str(item["search_query"] or "")): item for item in cur.fetchall()}
        uncached = []
        for row in pending:
            item = cached.get((str(row.get("search_intent") or ""), str(row.get("search_query") or "")))
            if not item:
                uncached.append(row)
                continue
            row.update(relevance_decision=item["relevance_decision"], clean_query=item.get("clean_query") or "",
                       relevance_reason=item.get("relevance_reason"), semantic_type=item.get("semantic_type") or "other",
                       analysis_model=item.get("model") or "cache")
            cache_hits += 1
        pending = uncached
    # SEO Bot's proven pattern: send each semantic key to AI only once and fan
    # the decision back out to all SKU rows.  The cache key includes intent, so
    # the same phrase may still be classified differently for different goods.
    pending_groups = {}
    for row in pending:
        key = (str(row.get("search_intent") or ""), _seo_normalize(row.get("search_query")))
        pending_groups.setdefault(key, []).append(row)
    ai_rows = []
    for index, group in enumerate(pending_groups.values()):
        representative = dict(group[0])
        representative["analysis_id"] = str(index)
        ai_rows.append(representative)

    model = "local:competitor-rules"
    if ai_rows and analyze_batch:
        result = analyze_batch(ai_rows, settings)
        model = _clean_text(result.get("model"), 160) or model
        items = {str(item.get("analysis_id") or item.get("id") or ""): item for item in (result.get("items") or [])}
        for representative, group in zip(ai_rows, pending_groups.values()):
            item = items.get(representative["analysis_id"])
            if not item:
                raise RuntimeError(
                    f"ИИ не вернул решение для запроса «{representative['search_query']}»"
                )
            decision = str(item.get("relevance_decision") or item.get("decision") or "review").lower()
            decision = decision if decision in {"keep", "reject", "review"} else "review"
            clean_query = _seo_normalize(item.get("clean_query") or representative["search_query"]) if decision == "keep" else ""
            reason = _clean_text(item.get("rationale") or item.get("reason"), 500) or "Смысловая проверка SEO-бота"
            for row in group:
                row["relevance_decision"] = decision
                row["clean_query"] = clean_query
                row["relevance_reason"] = reason
                row["analysis_model"] = model
        cache_values = [(project_id, key[0], group[0]["search_query"],
                         group[0]["relevance_decision"], group[0].get("clean_query") or None,
                         group[0].get("relevance_reason"), group[0].get("semantic_type") or "other",
                         group[0].get("analysis_model") or model)
                        for key, group in pending_groups.items()]
        with _conn(config) as conn, conn.cursor() as cur:
            execute_values(cur, """INSERT INTO public.seo_generation_competitor_ai_cache
                (project_id, search_intent, search_query, relevance_decision, clean_query,
                 relevance_reason, semantic_type, model) VALUES %s
                ON CONFLICT(project_id, search_intent, search_query) DO UPDATE SET
                  relevance_decision=excluded.relevance_decision, clean_query=excluded.clean_query,
                  relevance_reason=excluded.relevance_reason, semantic_type=excluded.semantic_type,
                  model=excluded.model, analyzed_at=now()""", cache_values)
    for sku in {row["sku"] for row in rows}:
        group = [row for row in rows if row["sku"] == sku and row["relevance_decision"] == "keep"]
        group.sort(key=lambda row: (-(float(row.get("frequency_value") or 0)), -int(row.get("competitor_coverage") or 0), row["search_query"]))
        total = max(1, len(group))
        for rank, row in enumerate(group, 1):
            row["frequency_rank"] = rank
            row["frequency_class"] = "high" if rank <= max(1, math.ceil(total * .2)) else ("mid" if rank <= max(2, math.ceil(total * .6)) else "low")
            score = min(100, round(45 + 35 * (1 - (rank - 1) / total) + min(20, int(row.get("competitor_coverage") or 0) * 2)))
            row["priority_score"] = score
            row["priority_label"] = "high" if score >= 80 else ("medium" if score >= 60 else "low")
    for row in rows:
        if row.get("relevance_decision") != "keep":
            row.update(frequency_rank=None, frequency_class=None, priority_score=0, priority_label="none")
    values = [(project_id, row["sku"], row["search_query"], row["relevance_decision"], row.get("clean_query") or None,
               row.get("relevance_reason"), row.get("semantic_type") or "other", row.get("frequency_class"),
               row.get("frequency_rank"), int(row.get("competitor_coverage") or 0), int(row.get("priority_score") or 0),
               row.get("priority_label") or "none", row.get("analysis_model") or model) for row in rows]
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute("DELETE FROM public.seo_generation_competitor_keyword_analysis WHERE project_id=%s AND sku=ANY(%s)", (project_id, requested_skus))
            execute_values(cur, """INSERT INTO public.seo_generation_competitor_keyword_analysis
                (project_id, sku, search_query, relevance_decision, clean_query, relevance_reason, semantic_type,
                 frequency_class, frequency_rank, competitor_coverage, priority_score, priority_label, model) VALUES %s""", values)
    return {"ok": True, "project_id": project_id, "analyzed": len(rows),
            "kept": sum(row["relevance_decision"] == "keep" for row in rows),
            "rejected": sum(row["relevance_decision"] == "reject" for row in rows),
            "review": sum(row["relevance_decision"] == "review" for row in rows), "model": model,
            "cache_hits": cache_hits, "ai_checked": len(ai_rows),
            "deduplicated": max(0, len(pending) - len(ai_rows))}


def analyze_project_competitor_keywords(config, payload, analyze_batch):
    """Run one checkpointed competitor-keyword batch with a project-wide server lock."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if not requested_skus or len(requested_skus) > 10:
        raise ValueError("Передайте от 1 до 10 SKU проекта")
    operation = "competitor_keyword_analysis"
    owner_id = str(uuid.uuid4())
    force = str(payload.get("force") or "").lower() in {"1", "true", "yes"}
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                """DELETE FROM public.seo_generation_operation_locks
                   WHERE project_id=%s AND operation=%s AND started_at < now() - interval '30 minutes'""",
                (project_id, operation),
            )
            cur.execute(
                """INSERT INTO public.seo_generation_operation_locks(project_id, operation, owner_id)
                   VALUES (%s,%s,%s) ON CONFLICT DO NOTHING RETURNING owner_id""",
                (project_id, operation, owner_id),
            )
            if not cur.fetchone():
                raise ValueError("Очистка ключей конкурентов уже выполняется для этого проекта")
    try:
        completed = set()
        if not force:
            with _conn(config) as conn, conn.cursor() as cur:
                cur.execute(
                    """SELECT a.sku
                       FROM public.seo_generation_competitor_keyword_analysis a
                       WHERE a.project_id=%s AND a.sku=ANY(%s)
                       GROUP BY a.sku
                       HAVING count(*) = (
                           SELECT count(DISTINCT k.search_query)
                           FROM public.seo_generation_competitor_keywords k
                           WHERE k.project_id=%s AND k.sku=a.sku
                       )""",
                    (project_id, requested_skus, project_id),
                )
                completed = {str(row["sku"]) for row in cur.fetchall()}
        pending_skus = [sku for sku in requested_skus if sku not in completed]
        if not pending_skus:
            return {"ok": True, "project_id": project_id, "analyzed": 0, "kept": 0,
                    "rejected": 0, "review": 0, "skipped_skus": len(completed),
                    "message": "Батч уже обработан; использован сохранённый checkpoint"}
        run_payload = dict(payload)
        run_payload["skus"] = pending_skus
        result = _analyze_project_competitor_keywords_unlocked(config, run_payload, analyze_batch)
        return {**result, "skipped_skus": len(completed), "processed_skus": len(pending_skus)}
    finally:
        with _conn(config) as conn, conn.cursor() as cur:
            cur.execute(
                """DELETE FROM public.seo_generation_operation_locks
                   WHERE project_id=%s AND operation=%s AND owner_id=%s""",
                (project_id, operation, owner_id),
            )


def _job_row(row):
    raw_started = row.get("started_at") if row and hasattr(row, "get") else None
    result = _serialize(row) if row else None
    if not result:
        return None
    progress = result.get("progress_json") if isinstance(result.get("progress_json"), dict) else {}
    total = int(result.get("total_skus") or 0); done = int(result.get("processed_skus") or 0) + int(result.get("skipped_skus") or 0)
    if result.get("operation") == "full_seo_pipeline" and progress:
        done = int(progress.get("current_stage_done") or 0)
        result["current_stage"] = progress.get("current_stage")
        result["current_stage_done"] = done
        result["current_stage_total"] = int(progress.get("current_stage_total") or total)
        result["overall_percent"] = float(progress.get("overall_percent") or 0)
        result["stages"] = progress.get("stages") or []
        result["models"] = progress.get("models") or []
        result["publication"] = "not_requested"
    result["done_skus"] = min(total, done); result["percent"] = round(100 * done / total, 1) if total else 100
    started = raw_started or result.get("started_at")
    if isinstance(started, str):
        try:
            started = datetime.fromisoformat(started.replace("Z", "+00:00"))
        except ValueError:
            started = None
    if started is not None and started.tzinfo is None:
        started = started.replace(tzinfo=timezone.utc)
    if started and result.get("status") in {"queued", "running", "stopping"}:
        elapsed = max(1, (datetime.now(timezone.utc) - started).total_seconds())
        overall = float(result.get("overall_percent") or 0)
        if result.get("operation") == "full_seo_pipeline" and 0 < overall < 100:
            result["eta_seconds"] = round(elapsed / overall * (100 - overall))
        elif done and total > done:
            result["eta_seconds"] = round(elapsed / done * (total - done))
        else:
            result["eta_seconds"] = None
    else:
        result["eta_seconds"] = None
    result["is_active"] = result.get("status") in {"queued", "running", "stopping"}
    return result


def _finalize_stale_competitor_jobs(cur, project_id):
    """Expose a lost in-process worker honestly and allow checkpoint resume."""
    cur.execute(
        """UPDATE public.seo_generation_jobs
           SET status='error', finished_at=coalesce(finished_at,now()), updated_at=now(),
               last_message='Ран прерван: процесс сервера остановлен; повторный запуск продолжит с checkpoint',
               last_error=coalesce(last_error,'Потерян heartbeat серверного worker более 20 минут назад')
           WHERE project_id=%s AND operation='competitor_keyword_analysis'
             AND status IN ('queued','running','stopping')
             AND updated_at < now() - interval '20 minutes'""",
        (project_id,),
    )


def start_competitor_analysis_job(config, payload):
    """Create one durable job instead of one browser request per SKU batch."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            if not requested:
                cur.execute(
                    "SELECT DISTINCT sku FROM public.seo_generation_competitor_keywords WHERE project_id=%s ORDER BY sku",
                    (project_id,),
                )
                requested = [str(row["sku"]) for row in cur.fetchall()]
            if not requested:
                raise ValueError("Ключи конкурентов ещё не собраны")
            _finalize_stale_competitor_jobs(cur, project_id)
            cur.execute(
                "SELECT DISTINCT sku FROM public.seo_generation_competitor_keywords WHERE project_id=%s AND sku=ANY(%s)",
                (project_id, requested),
            )
            available = {str(row["sku"]) for row in cur.fetchall()}
            requested = [sku for sku in requested if sku in available]
            if not requested:
                raise ValueError("Для выбранных SKU ключи конкурентов не собраны")
            cur.execute("""SELECT * FROM public.seo_generation_jobs WHERE project_id=%s AND operation='competitor_keyword_analysis'
                           AND status IN ('queued','running','stopping') ORDER BY updated_at DESC LIMIT 1""", (project_id,))
            existing = cur.fetchone()
            if existing:
                return {"ok": True, "job": _job_row(existing), "reused": True}
            job_id = str(uuid.uuid4())
            cur.execute("""INSERT INTO public.seo_generation_jobs
                (job_id, project_id, operation, status, target_skus, total_skus, last_message)
                VALUES (%s,%s,'competitor_keyword_analysis','queued',%s::jsonb,%s,'Задание поставлено в очередь') RETURNING *""",
                (job_id, project_id, json.dumps(requested, ensure_ascii=False), len(requested)))
            return {"ok": True, "job": _job_row(cur.fetchone()), "reused": False}


def competitor_analysis_job_status(config, payload):
    project_id = _project_id(payload.get("project_id"))
    job_id = _clean_text(payload.get("job_id"), 80)
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _finalize_stale_competitor_jobs(cur, project_id)
            if job_id:
                cur.execute("SELECT * FROM public.seo_generation_jobs WHERE job_id=%s AND project_id=%s", (job_id, project_id))
            else:
                cur.execute("""SELECT * FROM public.seo_generation_jobs WHERE project_id=%s AND operation='competitor_keyword_analysis'
                               ORDER BY updated_at DESC LIMIT 1""", (project_id,))
            return {"ok": True, "job": _job_row(cur.fetchone())}


def stop_competitor_analysis_job(config, payload):
    project_id = _project_id(payload.get("project_id"))
    job_id = _clean_text(payload.get("job_id"), 80)
    if not job_id:
        raise ValueError("job_id не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute("""UPDATE public.seo_generation_jobs SET stop_requested=true, status='stopping',
                           last_message='Запрошена остановка', updated_at=now()
                           WHERE job_id=%s AND project_id=%s AND status IN ('queued','running') RETURNING *""", (job_id, project_id))
            return {"ok": True, "job": _job_row(cur.fetchone())}


def run_competitor_analysis_job(config, job_id, analyze_batch):
    """Run the SEO Bot pipeline server-side with DB checkpoints between 10-SKU batches."""
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                """UPDATE public.seo_generation_jobs
                   SET status='running', stop_requested=false,
                       started_at=coalesce(started_at,now()), finished_at=NULL,
                       last_message='Подготовка словаря, кэша и checkpoint', updated_at=now()
                   WHERE job_id=%s RETURNING *""",
                (job_id,),
            )
            job = cur.fetchone()
    if not job:
        return
    started_at = job.get("started_at") or datetime.now(timezone.utc)
    if isinstance(started_at, str):
        started_at = datetime.fromisoformat(started_at.replace("Z", "+00:00"))
    if started_at.tzinfo is None:
        started_at = started_at.replace(tzinfo=timezone.utc)
    skus = [str(value) for value in (job.get("target_skus") or [])]
    project_id = str(job["project_id"])
    operation = "competitor_keyword_analysis"
    owner_id = str(job_id)
    lock_acquired = False
    try:
        with _conn(config) as conn, conn.cursor() as cur:
            cur.execute(
                """DELETE FROM public.seo_generation_operation_locks
                   WHERE project_id=%s AND operation=%s AND started_at < now() - interval '30 minutes'""",
                (project_id, operation),
            )
            cur.execute(
                """INSERT INTO public.seo_generation_operation_locks(project_id, operation, owner_id)
                   VALUES (%s,%s,%s) ON CONFLICT DO NOTHING RETURNING owner_id""",
                (project_id, operation, owner_id),
            )
            lock_acquired = bool(cur.fetchone())
        if not lock_acquired:
            raise RuntimeError("Очистка ключей конкурентов уже выполняется для этого проекта")

        with _conn(config) as conn, conn.cursor() as cur:
            cur.execute(
                """SELECT a.sku
                   FROM public.seo_generation_competitor_keyword_analysis a
                   WHERE a.project_id=%s AND a.sku=ANY(%s)
                   GROUP BY a.sku
                   HAVING count(*) = (
                       SELECT count(DISTINCT k.search_query)
                       FROM public.seo_generation_competitor_keywords k
                       WHERE k.project_id=%s AND k.sku=a.sku
                   )""",
                (project_id, skus, project_id),
            )
            completed = {str(row["sku"]) for row in cur.fetchall()}
            cur.execute(
                """SELECT k.sku, coalesce(max(i.search_intent),'') AS search_intent
                   FROM public.seo_generation_competitor_keywords k
                   LEFT JOIN public.seo_generation_product_intents i
                     ON i.project_id=k.project_id AND i.sku=k.sku
                   WHERE k.project_id=%s AND k.sku=ANY(%s)
                   GROUP BY k.sku ORDER BY search_intent, k.sku""",
                (project_id, skus),
            )
            ordered_skus = [str(row["sku"]) for row in cur.fetchall() if str(row["sku"]) not in completed]
            cur.execute(
                """SELECT count(*)::integer AS analyzed,
                          count(*) FILTER (WHERE relevance_decision='keep')::integer AS kept,
                          count(*) FILTER (WHERE relevance_decision='reject')::integer AS rejected,
                          count(*) FILTER (WHERE relevance_decision='review')::integer AS review
                   FROM public.seo_generation_competitor_keyword_analysis
                   WHERE project_id=%s AND sku=ANY(%s)""",
                (project_id, skus),
            )
            durable = dict(cur.fetchone())

        totals = {
            "analyzed": int(durable.get("analyzed") or 0),
            "kept": int(durable.get("kept") or 0),
            "rejected": int(durable.get("rejected") or 0),
            "review": int(durable.get("review") or 0),
            "processed_skus": 0,
            "skipped_skus": len(completed),
            "cache_hits": 0,
            "ai_checked": 0,
            "error_count": 0,
        }
        batch_size = 10
        batches = [ordered_skus[index:index + batch_size] for index in range(0, len(ordered_skus), batch_size)]
        with _conn(config) as conn, conn.cursor() as cur:
            cur.execute(
                """UPDATE public.seo_generation_jobs SET skipped_skus=%s, analyzed=%s, kept=%s,
                          rejected=%s, review=%s, last_message=%s, updated_at=now()
                   WHERE job_id=%s""",
                (totals["skipped_skus"], totals["analyzed"], totals["kept"], totals["rejected"],
                 totals["review"],
                 f"ПЛАН: {len(skus)} SKU · checkpoint {len(completed)} · {len(batches)} батчей по {batch_size}",
                 job_id),
            )

        for batch_index, batch in enumerate(batches, 1):
            with _conn(config) as conn, conn.cursor() as cur:
                cur.execute("SELECT stop_requested FROM public.seo_generation_jobs WHERE job_id=%s", (job_id,))
                state = cur.fetchone()
                if state and state["stop_requested"]:
                    cur.execute(
                        """UPDATE public.seo_generation_jobs SET status='stopped', finished_at=now(),
                                  updated_at=now(), last_message='Остановлено пользователем'
                           WHERE job_id=%s""",
                        (job_id,),
                    )
                    return
            batch_failed = False
            try:
                result = _analyze_project_competitor_keywords_unlocked(
                    config,
                    {"project_id": project_id, "project_kind": "generation", "skus": batch},
                    analyze_batch,
                )
                for key in ("analyzed", "kept", "rejected", "review", "cache_hits", "ai_checked"):
                    totals[key] += int(result.get(key) or 0)
                totals["processed_skus"] += len(batch)
                message = (
                    f"ПРОГРЕСС: батч {batch_index}/{len(batches)} · SKU {totals['processed_skus'] + totals['skipped_skus']}/{len(skus)} · "
                    f"проверено {totals['analyzed']} · кэш {totals['cache_hits']} · ИИ {totals['ai_checked']}"
                )
                last_error = None
            except Exception as exc:
                totals["error_count"] += len(batch)
                message = f"ОШИБКА: батч {batch_index}/{len(batches)} · SKU {', '.join(batch[:3])}"
                last_error = str(exc)[:1000]
                batch_failed = True
            with _conn(config) as conn, conn.cursor() as cur:
                cur.execute("""UPDATE public.seo_generation_jobs SET processed_skus=%s, skipped_skus=%s, analyzed=%s,
                    kept=%s, rejected=%s, review=%s, cache_hits=%s, ai_checked=%s, error_count=%s,
                    current_sku=%s, last_message=%s, last_error=%s, updated_at=now() WHERE job_id=%s""",
                    (totals["processed_skus"], totals["skipped_skus"], totals["analyzed"], totals["kept"],
                     totals["rejected"], totals["review"], totals["cache_hits"], totals["ai_checked"],
                     totals["error_count"], batch[-1], message, last_error, job_id))
                if batch_failed:
                    cur.execute(
                        """UPDATE public.seo_generation_jobs
                           SET status='error', finished_at=now(), updated_at=now(),
                               last_message='Остановлено на первой ошибке; повторный запуск продолжит с checkpoint'
                           WHERE job_id=%s""",
                        (job_id,),
                    )
            if batch_failed:
                return
        status = "completed" if totals["error_count"] == 0 else "error"
        with _conn(config) as conn, conn.cursor() as cur:
            cur.execute(
                """UPDATE public.seo_generation_jobs SET status=%s, finished_at=now(), updated_at=now(),
                          last_message=%s WHERE job_id=%s""",
                (status, "Завершено" if status == "completed" else "Завершено с ошибками; повторный запуск обработает только пропущенные SKU", job_id),
            )
    except Exception as exc:
        with _conn(config) as conn, conn.cursor() as cur:
            cur.execute(
                """UPDATE public.seo_generation_jobs SET status='error', error_count=error_count+1,
                          last_error=%s, last_message='Аварийная остановка', finished_at=now(), updated_at=now()
                   WHERE job_id=%s""",
                (str(exc)[:1000], job_id),
            )
    finally:
        if lock_acquired:
            with _conn(config) as conn, conn.cursor() as cur:
                cur.execute(
                    """DELETE FROM public.seo_generation_operation_locks
                       WHERE project_id=%s AND operation=%s AND owner_id=%s""",
                    (project_id, operation, owner_id),
                )


FULL_SEO_RUN_STAGES = (
    {"id": "seller_keywords", "label": "Запросы", "batch_size": 1000},
    {"id": "mpstats", "label": "MPStats", "batch_size": 25},
    {"id": "keyword_analysis", "label": "Очистить ключи", "batch_size": 10},
    {"id": "intents", "label": "Интенты", "batch_size": 10},
    {"id": "competitors", "label": "Конкуренты", "batch_size": 25},
    {"id": "competitor_keywords", "label": "Ключи конкурентов", "batch_size": 5},
    {"id": "competitor_analysis", "label": "Очистить конкурентов", "batch_size": 10},
    {"id": "customer_messages", "label": "Отзывы", "batch_size": 10},
    {"id": "customer_voice", "label": "SEO-клеймы", "batch_size": 5},
    {"id": "semantic_context", "label": "SEO-контекст", "batch_size": 20},
    {"id": "allocation", "label": "SEO-разметка", "batch_size": 1},
    {"id": "draft", "label": "SEO-тексты", "batch_size": 1},
    {"id": "review", "label": "Эксперт", "batch_size": 1},
)


def _full_run_progress(total_skus):
    return {
        "version": 1,
        "current_stage": None,
        "current_stage_done": 0,
        "current_stage_total": int(total_skus),
        "overall_percent": 0,
        "models": [],
        "stages": [{
            "id": item["id"], "label": item["label"], "status": "pending",
            "done": 0, "total": int(total_skus), "errors": 0,
        } for item in FULL_SEO_RUN_STAGES],
    }


def _full_run_append_log(log_lines, line):
    rows = list(log_lines or [])
    rows.append(f"{datetime.now().strftime('%H:%M:%S')} | {_clean_text(line, 1000)}")
    return rows[-240:]


def _finalize_stale_full_run_jobs(cur, project_id):
    cur.execute(
        """UPDATE public.seo_generation_jobs
           SET status='error', finished_at=coalesce(finished_at,now()), updated_at=now(),
               last_message='Ран прерван: серверный worker потерян; запустите полный ран повторно',
               last_error=coalesce(last_error,'Нет heartbeat полного рана более 30 минут')
           WHERE project_id=%s AND operation='full_seo_pipeline'
             AND status IN ('queued','running','stopping')
             AND updated_at < now() - interval '30 minutes'""",
        (project_id,),
    )


def start_full_seo_run_job(config, payload):
    """Create one durable all-stage job for every SKU currently in the project."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                "SELECT sku FROM public.seo_monitoring_project_skus WHERE project_id=%s ORDER BY sku",
                (project_id,),
            )
            skus = [str(row["sku"]) for row in cur.fetchall()]
            if not skus:
                raise ValueError("В проекте нет выбранных SKU")
            _finalize_stale_full_run_jobs(cur, project_id)
            cur.execute(
                """SELECT * FROM public.seo_generation_jobs
                   WHERE project_id=%s AND operation='full_seo_pipeline'
                     AND status IN ('queued','running','stopping')
                   ORDER BY updated_at DESC LIMIT 1""",
                (project_id,),
            )
            existing = cur.fetchone()
            if existing:
                return {"ok": True, "job": _job_row(existing), "reused": True}
            job_id = str(uuid.uuid4())
            progress = _full_run_progress(len(skus))
            log_lines = [_full_run_append_log([], (
                f"ПЛАН: {len(skus)} SKU · {len(FULL_SEO_RUN_STAGES)} этапов · "
                "checkpoint после каждого пакета · публикация в WB не выполняется"
            ))[-1]]
            cur.execute(
                """INSERT INTO public.seo_generation_jobs
                    (job_id, project_id, operation, status, target_skus, total_skus,
                     last_message, progress_json, log_lines)
                   VALUES (%s,%s,'full_seo_pipeline','queued',%s::jsonb,%s,%s,%s::jsonb,%s::jsonb)
                   RETURNING *""",
                (job_id, project_id, json.dumps(skus, ensure_ascii=False), len(skus),
                 "Полный ран поставлен в очередь", json.dumps(progress, ensure_ascii=False),
                 json.dumps(log_lines, ensure_ascii=False)),
            )
            return {
                "ok": True, "job": _job_row(cur.fetchone()), "reused": False,
                "marketplace": project.get("marketplace"),
            }


def full_seo_run_job_status(config, payload):
    project_id = _project_id(payload.get("project_id"))
    job_id = _clean_text(payload.get("job_id"), 80)
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _finalize_stale_full_run_jobs(cur, project_id)
            if job_id:
                cur.execute(
                    "SELECT * FROM public.seo_generation_jobs WHERE job_id=%s AND project_id=%s AND operation='full_seo_pipeline'",
                    (job_id, project_id),
                )
            else:
                cur.execute(
                    """SELECT * FROM public.seo_generation_jobs
                       WHERE project_id=%s AND operation='full_seo_pipeline'
                       ORDER BY updated_at DESC LIMIT 1""",
                    (project_id,),
                )
            return {"ok": True, "job": _job_row(cur.fetchone())}


def stop_full_seo_run_job(config, payload):
    project_id = _project_id(payload.get("project_id"))
    job_id = _clean_text(payload.get("job_id"), 80)
    if not job_id:
        raise ValueError("job_id не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                """UPDATE public.seo_generation_jobs
                   SET stop_requested=true, status='stopping',
                       last_message='Запрошена безопасная остановка после текущего пакета', updated_at=now()
                   WHERE job_id=%s AND project_id=%s AND operation='full_seo_pipeline'
                     AND status IN ('queued','running') RETURNING *""",
                (job_id, project_id),
            )
            return {"ok": True, "job": _job_row(cur.fetchone())}


def _full_run_result_errors(result):
    if not isinstance(result, dict):
        return 0
    errors = result.get("errors")
    if isinstance(errors, list):
        return len(errors)
    if isinstance(errors, int):
        return max(0, errors)
    return int(result.get("error_count") or 0)


def _full_run_result_models(result):
    if not isinstance(result, dict):
        return []
    values = []
    for key in ("model", "models", "analysis_model", "review_model"):
        value = result.get(key)
        if isinstance(value, list):
            values.extend(value)
        elif value:
            values.extend(str(value).split(","))
    return [value.strip() for value in values if str(value).strip()]


def _full_run_started_at(job):
    """Return the durable job start as an aware UTC datetime for ETA math."""
    started_at = job.get("started_at") or datetime.now(timezone.utc)
    if isinstance(started_at, str):
        started_at = datetime.fromisoformat(started_at.replace("Z", "+00:00"))
    if started_at.tzinfo is None:
        started_at = started_at.replace(tzinfo=timezone.utc)
    return started_at.astimezone(timezone.utc)


def run_full_seo_run_job(config, job_id, stage_handlers):
    """Run the full pipeline server-side and checkpoint every bounded SKU package."""
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                """UPDATE public.seo_generation_jobs
                   SET status='running', stop_requested=false, started_at=coalesce(started_at,now()),
                       finished_at=NULL, last_message='Подготовка полного рана', updated_at=now()
                   WHERE job_id=%s AND operation='full_seo_pipeline' RETURNING *""",
                (job_id,),
            )
            job = cur.fetchone()
    if not job:
        return
    started_at = _full_run_started_at(job)
    skus = [str(value) for value in (job.get("target_skus") or [])]
    project_id = str(job["project_id"])
    progress = job.get("progress_json") if isinstance(job.get("progress_json"), dict) else _full_run_progress(len(skus))
    logs = list(job.get("log_lines") or [])
    stage_rows = progress.get("stages") or _full_run_progress(len(skus))["stages"]
    total_units = max(1, len(skus) * len(stage_rows))
    completed_units = sum(int(row.get("done") or 0) for row in stage_rows if row.get("status") in {"completed", "partial"})
    total_errors = int(job.get("error_count") or 0)
    try:
        for stage_index, stage_def in enumerate(FULL_SEO_RUN_STAGES):
            stage = stage_rows[stage_index]
            if stage.get("status") in {"completed", "partial"} and int(stage.get("done") or 0) >= len(skus):
                continue
            handler = stage_handlers.get(stage_def["id"])
            if not handler:
                raise RuntimeError(f"Обработчик этапа {stage_def['id']} не настроен")
            stage.update(status="running", started_at=datetime.now(timezone.utc).isoformat(), done=0, errors=0)
            progress.update(current_stage=stage_def["id"], current_stage_done=0, current_stage_total=len(skus))
            logs = _full_run_append_log(logs, f"ЭТАП {stage_index + 1}/{len(stage_rows)}: {stage_def['label']} · {len(skus)} SKU")
            batches = [skus[index:index + stage_def["batch_size"]] for index in range(0, len(skus), stage_def["batch_size"])]
            stage_models = set()
            for batch_index, batch in enumerate(batches, 1):
                with _conn(config) as conn, conn.cursor() as cur:
                    cur.execute("SELECT stop_requested FROM public.seo_generation_jobs WHERE job_id=%s", (job_id,))
                    state = cur.fetchone()
                if state and state["stop_requested"]:
                    stage["status"] = "stopped"
                    logs = _full_run_append_log(logs, "ОСТАНОВЛЕНО: пользователь запросил остановку")
                    with _conn(config) as conn, conn.cursor() as cur:
                        cur.execute(
                            """UPDATE public.seo_generation_jobs SET status='stopped', finished_at=now(),
                               progress_json=%s::jsonb, log_lines=%s::jsonb,
                               last_message='Остановлено пользователем', updated_at=now() WHERE job_id=%s""",
                            (json.dumps(progress, ensure_ascii=False), json.dumps(logs, ensure_ascii=False), job_id),
                        )
                    return
                try:
                    result = handler({
                        "project_id": project_id, "project_kind": "generation", "skus": batch,
                        "stage_id": stage_def["id"], "batch_index": batch_index,
                        "batch_total": len(batches),
                    }) or {}
                    batch_errors = _full_run_result_errors(result)
                    stage["errors"] = int(stage.get("errors") or 0) + batch_errors
                    total_errors += batch_errors
                    stage_models.update(_full_run_result_models(result))
                    detail = f"ошибки {batch_errors}" if batch_errors else "готово"
                except Exception as exc:
                    batch_errors = len(batch)
                    stage["errors"] = int(stage.get("errors") or 0) + batch_errors
                    total_errors += batch_errors
                    detail = f"ошибка: {_clean_text(str(exc), 400)}"
                stage["done"] = min(len(skus), int(stage.get("done") or 0) + len(batch))
                progress["current_stage_done"] = stage["done"]
                completed_before = sum(int(row.get("done") or 0) for row in stage_rows[:stage_index])
                progress["overall_percent"] = round(100 * (completed_before + stage["done"]) / total_units, 1)
                if stage_models:
                    progress["models"] = sorted(set(progress.get("models") or []) | stage_models)
                elapsed = max(1, (datetime.now(timezone.utc) - started_at).total_seconds())
                completed_now = max(1, completed_before + stage["done"])
                eta = round(elapsed / completed_now * max(0, total_units - completed_now))
                message = (
                    f"ПРОГРЕСС: {stage_def['label']} · {stage['done']}/{len(skus)} "
                    f"({round(100 * stage['done'] / len(skus), 1)}%) · {detail} · ETA {eta} сек."
                )
                logs = _full_run_append_log(logs, message)
                with _conn(config) as conn, conn.cursor() as cur:
                    cur.execute(
                        """UPDATE public.seo_generation_jobs SET processed_skus=%s, error_count=%s,
                           current_sku=%s, last_message=%s, last_error=%s,
                           progress_json=%s::jsonb, log_lines=%s::jsonb, updated_at=now()
                           WHERE job_id=%s""",
                        (stage["done"], total_errors, batch[-1], message,
                         None if not batch_errors else detail,
                         json.dumps(progress, ensure_ascii=False), json.dumps(logs, ensure_ascii=False), job_id),
                    )
            stage["status"] = "partial" if stage.get("errors") else "completed"
            stage["finished_at"] = datetime.now(timezone.utc).isoformat()
            completed_units = sum(int(row.get("done") or 0) for row in stage_rows[:stage_index + 1])
            progress["overall_percent"] = round(100 * completed_units / total_units, 1)
        progress.update(current_stage=None, current_stage_done=len(skus), current_stage_total=len(skus), overall_percent=100)
        final_status = "completed" if total_errors == 0 else "partial"
        final_message = "Полный ран завершён" if total_errors == 0 else f"Полный ран завершён частично: ошибок источников {total_errors}"
        logs = _full_run_append_log(logs, f"ИТОГ: {len(skus)} SKU · ошибок источников {total_errors} · публикация в WB не выполнялась")
        with _conn(config) as conn, conn.cursor() as cur:
            cur.execute(
                """UPDATE public.seo_generation_jobs SET status=%s, processed_skus=%s,
                   error_count=%s, current_sku=NULL, last_message=%s, finished_at=now(),
                   progress_json=%s::jsonb, log_lines=%s::jsonb, updated_at=now() WHERE job_id=%s""",
                (final_status, len(skus), total_errors, final_message,
                 json.dumps(progress, ensure_ascii=False), json.dumps(logs, ensure_ascii=False), job_id),
            )
            cur.execute(
                """UPDATE public.seo_monitoring_projects
                   SET status=%s, last_error=%s, updated_at=now()
                   WHERE project_id=%s""",
                ("ready" if total_errors == 0 else "partial",
                 None if total_errors == 0 else final_message, project_id),
            )
    except Exception as exc:
        logs = _full_run_append_log(logs, f"АВАРИЙНАЯ ОСТАНОВКА: {_clean_text(str(exc), 500)}")
        with _conn(config) as conn, conn.cursor() as cur:
            cur.execute(
                """UPDATE public.seo_generation_jobs SET status='error', error_count=error_count+1,
                   last_error=%s, last_message='Аварийная остановка полного рана', finished_at=now(),
                   progress_json=%s::jsonb, log_lines=%s::jsonb, updated_at=now() WHERE job_id=%s""",
                (str(exc)[:1000], json.dumps(progress, ensure_ascii=False),
                 json.dumps(logs, ensure_ascii=False), job_id),
            )
            cur.execute(
                """UPDATE public.seo_monitoring_projects
                   SET status='error', last_error=%s, updated_at=now()
                   WHERE project_id=%s""",
                (str(exc)[:1000], project_id),
            )


def project_competitor_keywords(config, payload):
    project_id = _project_id(payload.get("project_id")); project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku: raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute("""SELECT k.search_query, max(k.search_demand) AS search_demand,
                       min(nullif(k.average_position,0)) AS best_position,
                       count(DISTINCT k.competitor_sku)::integer AS competitor_coverage,
                       string_agg(DISTINCT k.competitor_sku, ', ') AS competitor_skus,
                       a.relevance_decision, a.clean_query, a.frequency_class, a.frequency_rank,
                       a.priority_score, a.priority_label, a.relevance_reason
                FROM public.seo_generation_competitor_keywords k
                LEFT JOIN public.seo_generation_competitor_keyword_analysis a
                  ON a.project_id=k.project_id AND a.sku=k.sku AND a.search_query=k.search_query
                WHERE k.project_id=%s AND k.sku=%s
                GROUP BY k.search_query, a.relevance_decision, a.clean_query, a.frequency_class,
                         a.frequency_rank, a.priority_score, a.priority_label, a.relevance_reason
                ORDER BY coalesce(a.priority_score,0) DESC, search_demand DESC NULLS LAST, k.search_query""", (project_id, sku))
            rows = [_serialize(row) for row in cur.fetchall()]
    return {"ok": True, "project_id": project_id, "sku": sku, "rows": rows, "count": len(rows)}


def _normalize_ozon_customer_question(raw, allowed_skus, period_from, period_to):
    if not isinstance(raw, dict):
        return None
    sku = str(raw.get("sku") or "").strip()
    question_id = _clean_text(raw.get("id"), 500)
    question_date = _customer_message_date(raw.get("published_at") or raw.get("date"))
    if not sku or sku not in allowed_skus or not question_id or question_date is None:
        return None
    if question_date < period_from or question_date > period_to:
        return None
    answers_count = int(_number(raw.get("answers_count")) or 0)
    status = str(raw.get("status") or "").strip().upper()
    return {
        "sku": sku, "message_type": "question", "source_message_id": question_id,
        "message_date": question_date, "rating": None,
        "message_text": _clean_text(raw.get("text"), 20000), "pros": "", "cons": "",
        "answer_text": "", "answered": answers_count > 0 or status == "PROCESSED",
        "source": "ozon_seller_api_questions", "raw_payload": raw,
    }


def _dedupe_customer_message_rows(rows):
    """Collapse duplicate provider rows before one PostgreSQL UPSERT statement."""
    unique = {}
    for row in rows or []:
        key = (
            str(row.get("sku") or ""), str(row.get("message_type") or ""),
            str(row.get("source") or ""), str(row.get("source_message_id") or ""),
        )
        if all(key):
            unique[key] = row
    return list(unique.values())


def _dedupe_competitor_review_rows(rows):
    """Collapse duplicate competitor reviews by their table primary key."""
    unique = {}
    for row in rows or []:
        key = (
            str(row.get("sku") or ""), str(row.get("competitor_sku") or ""),
            str(row.get("source_message_id") or ""),
        )
        if all(key):
            unique[key] = row
    return list(unique.values())


def collect_project_questions(config, payload, question_fetcher=None):
    """Collect only own-SKU Ozon questions; competitor FAQ is unavailable for Ozon."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if not requested_skus or len(requested_skus) > 50:
        raise ValueError("Передайте от 1 до 50 SKU проекта")
    try:
        period_days = int(payload.get("period_days") or 90)
    except (TypeError, ValueError):
        period_days = 90
    if period_days not in {30, 90}:
        raise ValueError("Период вопросов должен быть 30 или 90 дней")
    period_to = date.today()
    period_from = period_to - timedelta(days=period_days - 1)
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            if project.get("marketplace") != "ozon":
                raise ValueError("Сбор вопросов сейчас подключён для Ozon")
            cur.execute(
                "SELECT sku FROM public.seo_monitoring_project_skus WHERE project_id=%s AND sku=ANY(%s)",
                (project_id, requested_skus),
            )
            project_skus = {str(row["sku"]) for row in cur.fetchall()}
    missing = [sku for sku in requested_skus if sku not in project_skus]
    if missing:
        raise ValueError(f"SKU не входят в проект: {', '.join(missing[:5])}")

    rows, error = [], None
    if question_fetcher is None:
        error = "Интеграция вопросов Ozon не подключена"
    else:
        try:
            source_rows = question_fetcher(requested_skus, period_from, period_to)
            allowed_skus = set(requested_skus)
            rows = [
                item for raw in (source_rows or [])
                if (item := _normalize_ozon_customer_question(raw, allowed_skus, period_from, period_to)) is not None
            ]
            rows = _dedupe_customer_message_rows(rows)
        except Exception as exc:
            error = str(exc)[:500]
    counts = {sku: 0 for sku in requested_skus}
    for row in rows:
        counts[row["sku"]] = counts.get(row["sku"], 0) + 1
    status = "ok" if error is None else (
        "unavailable" if "403" in error or "subscription" in error.lower() else "error"
    )
    status_values = [(
        project_id, sku, "question", "ozon_seller_api_questions", period_days, status,
        counts.get(sku, 0) if error is None else None, error,
    ) for sku in requested_skus]
    message_values = [(
        project_id, row["sku"], "question", row["source_message_id"], row["message_date"],
        None, row.get("message_text"), "", "", row.get("answer_text"), row.get("answered"),
        row["source"], period_days, json.dumps(row.get("raw_payload") or {}, ensure_ascii=False, default=str),
    ) for row in rows]
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            if error is None:
                cur.execute(
                    "DELETE FROM public.seo_generation_customer_messages WHERE project_id=%s AND message_type='question' AND sku=ANY(%s)",
                    (project_id, requested_skus),
                )
            if message_values:
                execute_values(cur, """INSERT INTO public.seo_generation_customer_messages
                    (project_id, sku, message_type, source_message_id, message_date, rating, message_text,
                     pros, cons, answer_text, answered, source, period_days, raw_payload) VALUES %s
                    ON CONFLICT(project_id, sku, message_type, source, source_message_id) DO UPDATE SET
                      message_date=excluded.message_date, message_text=excluded.message_text,
                      answer_text=excluded.answer_text, answered=excluded.answered,
                      period_days=excluded.period_days, raw_payload=excluded.raw_payload,
                      collected_at=now()""", message_values)
            execute_values(cur, """INSERT INTO public.seo_generation_customer_message_status
                (project_id, sku, message_type, source, period_days, status, message_count, last_error) VALUES %s
                ON CONFLICT(project_id, sku, message_type) DO UPDATE SET
                  source=excluded.source, period_days=excluded.period_days, status=excluded.status,
                  message_count=excluded.message_count, last_error=excluded.last_error, collected_at=now()""",
                status_values)
    return {
        "ok": error is None, "partial": error is not None, "project_id": project_id,
        "sku_count": len(requested_skus), "question_count": len(rows),
        "question_skus": len(requested_skus) if error is None else 0,
        "period_days": period_days, "date_from": period_from.isoformat(), "date_to": period_to.isoformat(),
        "competitor_questions_status": "unavailable",
        "competitor_questions_reason": "MPStats Ozon FAQ для товаров-конкурентов недоступен",
        "errors": ([{"type": "question", "error": error}] if error else []),
    }


def _collect_wb_customer_messages(config, project_id, requested_skus, period_days, period_from, period_to):
    """Copy already synchronized official WB feedback/questions into the SEO project store."""
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            has_reviews = _relation_exists(cur, "marketplace_reviews")
            has_questions = _relation_exists(cur, "marketplace_questions")
            cur.execute(
                """SELECT sku, competitor_sku, competitor_rank, product_name
                   FROM public.seo_generation_niche_competitors
                   WHERE project_id=%s AND sku=ANY(%s) AND competitor_rank <= 10
                   ORDER BY sku, competitor_rank""", (project_id, requested_skus),
            )
            competitor_links = [dict(row) for row in cur.fetchall()]
            review_rows, question_rows, competitor_rows = [], [], []
            if has_reviews:
                cur.execute(
                    """SELECT product_id AS sku, source_review_id AS source_message_id,
                              review_date AS message_date, rating, review_text AS message_text,
                              pros, cons, answer_text, answered, source, '{}'::jsonb AS raw_payload
                       FROM public.marketplace_reviews
                       WHERE marketplace='wb' AND product_id=ANY(%s)
                         AND review_date BETWEEN %s AND %s""", (requested_skus, period_from, period_to),
                )
                review_rows = [{**dict(row), "sku": str(row["sku"]), "message_type": "review"} for row in cur.fetchall()]
                competitor_ids = sorted({str(row["competitor_sku"]) for row in competitor_links})
                if competitor_ids:
                    cur.execute(
                        """SELECT product_id AS competitor_sku, source_review_id AS source_message_id,
                                  review_date AS message_date, rating, review_text AS message_text,
                                  pros, cons, answer_text, answered, source, '{}'::jsonb AS raw_payload
                           FROM public.marketplace_reviews
                           WHERE marketplace='wb' AND product_id=ANY(%s)
                             AND review_date BETWEEN %s AND %s""", (competitor_ids, period_from, period_to),
                    )
                    by_competitor = {}
                    for row in cur.fetchall():
                        by_competitor.setdefault(str(row["competitor_sku"]), []).append(dict(row))
                    for link in competitor_links:
                        for row in by_competitor.get(str(link["competitor_sku"]), []):
                            competitor_rows.append({
                                **row, "sku": str(link["sku"]), "competitor_sku": str(link["competitor_sku"]),
                                "competitor_rank": link.get("competitor_rank"), "competitor_name": link.get("product_name"),
                            })
            if has_questions:
                cur.execute(
                    """SELECT product_id AS sku, source_question_id AS source_message_id,
                              question_date AS message_date, NULL::numeric AS rating,
                              question_text AS message_text, NULL::text AS pros, NULL::text AS cons,
                              answer_text, answered, source, raw_payload
                       FROM public.marketplace_questions
                       WHERE marketplace='wb' AND product_id=ANY(%s)
                         AND question_date BETWEEN %s AND %s""", (requested_skus, period_from, period_to),
                )
                question_rows = [{**dict(row), "sku": str(row["sku"]), "message_type": "question"} for row in cur.fetchall()]
            message_values = [(
                project_id, row["sku"], row["message_type"], row["source_message_id"], row["message_date"],
                row.get("rating"), row.get("message_text"), row.get("pros"), row.get("cons"), row.get("answer_text"),
                row.get("answered"), row.get("source") or "wb_seller_api", period_days,
                json.dumps(row.get("raw_payload") or {}, ensure_ascii=False, default=str),
            ) for row in review_rows + question_rows]
            if message_values:
                execute_values(cur, """INSERT INTO public.seo_generation_customer_messages
                    (project_id, sku, message_type, source_message_id, message_date, rating, message_text,
                     pros, cons, answer_text, answered, source, period_days, raw_payload) VALUES %s
                    ON CONFLICT(project_id, sku, message_type, source, source_message_id) DO UPDATE SET
                      message_date=excluded.message_date, rating=excluded.rating, message_text=excluded.message_text,
                      pros=excluded.pros, cons=excluded.cons, answer_text=excluded.answer_text,
                      answered=excluded.answered, period_days=excluded.period_days,
                      raw_payload=excluded.raw_payload, collected_at=now()""", message_values)
            review_counts = {sku: 0 for sku in requested_skus}
            question_counts = {sku: 0 for sku in requested_skus}
            for row in review_rows: review_counts[row["sku"]] += 1
            for row in question_rows: question_counts[row["sku"]] += 1
            missing_reason = "Сначала синхронизируйте официальный WB API отзывов и вопросов для аккаунта"
            statuses = []
            for sku in requested_skus:
                statuses.extend([
                    (project_id, sku, "review", "wb_seller_api", period_days,
                     "ok" if has_reviews else "unavailable", review_counts[sku] if has_reviews else None,
                     None if has_reviews else missing_reason),
                    (project_id, sku, "question", "wb_seller_api", period_days,
                     "ok" if has_questions else "unavailable", question_counts[sku] if has_questions else None,
                     None if has_questions else missing_reason),
                ])
            execute_values(cur, """INSERT INTO public.seo_generation_customer_message_status
                (project_id, sku, message_type, source, period_days, status, message_count, last_error) VALUES %s
                ON CONFLICT(project_id, sku, message_type) DO UPDATE SET source=excluded.source,
                  period_days=excluded.period_days, status=excluded.status, message_count=excluded.message_count,
                  last_error=excluded.last_error, collected_at=now()""", statuses)
            competitor_values = [(
                project_id, row["sku"], row["competitor_sku"], row.get("competitor_rank"), row.get("competitor_name"),
                row["source_message_id"], row["message_date"], row.get("rating"), row.get("message_text"),
                row.get("pros"), row.get("cons"), row.get("answer_text"), row.get("answered"),
                row.get("source") or "wb_seller_api", period_days,
                json.dumps(row.get("raw_payload") or {}, ensure_ascii=False, default=str),
            ) for row in competitor_rows]
            if competitor_values:
                execute_values(cur, """INSERT INTO public.seo_generation_competitor_reviews
                    (project_id, sku, competitor_sku, competitor_rank, competitor_name, source_message_id,
                     message_date, rating, message_text, pros, cons, answer_text, answered, source,
                     period_days, raw_payload) VALUES %s
                    ON CONFLICT(project_id, sku, competitor_sku, source_message_id) DO UPDATE SET
                      message_date=excluded.message_date, rating=excluded.rating, message_text=excluded.message_text,
                      pros=excluded.pros, cons=excluded.cons, answer_text=excluded.answer_text,
                      answered=excluded.answered, period_days=excluded.period_days,
                      raw_payload=excluded.raw_payload, collected_at=now()""", competitor_values)
        conn.commit()
    errors = []
    if not has_reviews: errors.append({"type": "review", "status": "unavailable", "error": missing_reason})
    if not has_questions: errors.append({"type": "question", "status": "unavailable", "error": missing_reason})
    if not competitor_links: errors.append({"type": "competitor_review", "status": "unavailable", "error": "Сначала соберите конкурентов по нише"})
    return {
        # Source availability is a data-quality state, not a transport/runtime
        # failure. The caller must be able to persist and render `unavailable`
        # without treating the completed collection check as a failed request.
        "ok": True, "partial": bool(errors), "project_id": project_id,
        "sku_count": len(requested_skus), "period_days": period_days,
        "date_from": period_from.isoformat(), "date_to": period_to.isoformat(),
        "review_count": len(review_rows), "question_count": len(question_rows),
        "competitor_review_count": len(competitor_rows), "competitor_count": len({row['competitor_sku'] for row in competitor_links}),
        "sources": {"reviews": "WB Seller API", "questions": "WB Seller API", "competitor_reviews": "WB Seller API"},
        "errors": errors,
    }


def collect_project_customer_messages(config, payload, question_fetcher=None, review_fetcher=None):
    """Collect own-SKU reviews/questions and top-10 competitor reviews for a bounded batch."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if not requested_skus or len(requested_skus) > 10:
        raise ValueError("Передайте от 1 до 10 SKU проекта")
    try:
        period_days = int(payload.get("period_days") or 90)
    except (TypeError, ValueError):
        period_days = 90
    if period_days not in {30, 90}:
        raise ValueError("Период отзывов и вопросов должен быть 30 или 90 дней")
    period_to = date.today()
    period_from = period_to - timedelta(days=period_days - 1)
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            marketplace = project.get("marketplace")
            cur.execute(
                "SELECT sku FROM public.seo_monitoring_project_skus WHERE project_id=%s AND sku=ANY(%s)",
                (project_id, requested_skus),
            )
            project_skus = {str(row["sku"]) for row in cur.fetchall()}
            cur.execute(
                """SELECT sku, competitor_sku, competitor_rank, product_name
                   FROM public.seo_generation_niche_competitors
                   WHERE project_id=%s AND sku=ANY(%s) AND competitor_rank <= 10
                   ORDER BY sku, competitor_rank""",
                (project_id, requested_skus),
            )
            competitor_links = [dict(row) for row in cur.fetchall()]
    missing = [sku for sku in requested_skus if sku not in project_skus]
    if missing:
        raise ValueError(f"SKU не входят в проект: {', '.join(missing[:5])}")
    if marketplace == "wb":
        return _collect_wb_customer_messages(
            config, project_id, requested_skus, period_days, period_from, period_to
        )

    fetch_reviews = review_fetcher or fetch_mpstats_product_reviews
    review_rows, review_errors, successful_review_skus = [], [], set()
    token = _load_mpstats_token("ozon")
    with ThreadPoolExecutor(max_workers=min(4, len(requested_skus))) as pool:
        futures = {
            pool.submit(fetch_reviews, sku, period_from, period_to, token): sku
            for sku in requested_skus
        }
        for future in as_completed(futures):
            sku = futures[future]
            try:
                rows = future.result()
                review_rows.extend(rows)
                successful_review_skus.add(sku)
            except Exception as exc:
                review_errors.append({"sku": sku, "type": "review", "error": str(exc)[:500]})

    links_by_competitor = {}
    for link in competitor_links:
        links_by_competitor.setdefault(str(link["competitor_sku"]), []).append(link)
    competitor_review_rows, competitor_review_errors = [], []
    successful_competitors = set()
    fetched_competitor_rows = {}
    if links_by_competitor:
        with ThreadPoolExecutor(max_workers=min(6, len(links_by_competitor))) as pool:
            futures = {
                pool.submit(fetch_reviews, competitor_sku, period_from, period_to, token): competitor_sku
                for competitor_sku in links_by_competitor
            }
            for future in as_completed(futures):
                competitor_sku = futures[future]
                try:
                    fetched_competitor_rows[competitor_sku] = future.result()
                    successful_competitors.add(competitor_sku)
                except Exception as exc:
                    competitor_review_errors.append({
                        "type": "competitor_review", "competitor_sku": competitor_sku,
                        "skus": sorted({str(item["sku"]) for item in links_by_competitor[competitor_sku]}),
                        "error": str(exc)[:500],
                    })
        for competitor_sku, source_rows in fetched_competitor_rows.items():
            for link in links_by_competitor[competitor_sku]:
                for row in source_rows:
                    competitor_review_rows.append({
                        **row, "sku": str(link["sku"]), "competitor_sku": competitor_sku,
                        "competitor_rank": link.get("competitor_rank"),
                        "competitor_name": link.get("product_name"),
                    })

    question_rows, question_error = [], None
    if question_fetcher is None:
        question_error = "Интеграция вопросов Ozon не подключена"
    else:
        try:
            source_questions = question_fetcher(requested_skus, period_from, period_to)
            allowed_skus = set(requested_skus)
            question_rows = [
                item for raw in (source_questions or [])
                if (item := _normalize_ozon_customer_question(raw, allowed_skus, period_from, period_to)) is not None
            ]
        except Exception as exc:
            question_error = str(exc)[:500]

    review_rows = _dedupe_customer_message_rows(review_rows)
    question_rows = _dedupe_customer_message_rows(question_rows)
    competitor_review_rows = _dedupe_competitor_review_rows(competitor_review_rows)

    review_counts = {sku: 0 for sku in successful_review_skus}
    for row in review_rows:
        review_counts[row["sku"]] = review_counts.get(row["sku"], 0) + 1
    question_counts = {sku: 0 for sku in requested_skus}
    for row in question_rows:
        question_counts[row["sku"]] = question_counts.get(row["sku"], 0) + 1

    competitor_review_counts = {}
    for row in competitor_review_rows:
        key = (row["sku"], row["competitor_sku"])
        competitor_review_counts[key] = competitor_review_counts.get(key, 0) + 1
    competitor_status_values = []
    for competitor_sku, links in links_by_competitor.items():
        for link in links:
            own_sku = str(link["sku"])
            if competitor_sku in successful_competitors:
                competitor_status_values.append((
                    project_id, own_sku, competitor_sku, "mpstats_ozon_comments", period_days, "ok",
                    competitor_review_counts.get((own_sku, competitor_sku), 0), None,
                ))
            else:
                error = next((
                    item["error"] for item in competitor_review_errors
                    if item["competitor_sku"] == competitor_sku
                ), "MPStats недоступен")
                competitor_status_values.append((
                    project_id, own_sku, competitor_sku, "mpstats_ozon_comments", period_days,
                    "error", None, error,
                ))

    status_values = []
    for sku in requested_skus:
        if sku in successful_review_skus:
            status_values.append((project_id, sku, "review", "mpstats_ozon_comments", period_days, "ok", review_counts.get(sku, 0), None))
        else:
            error = next((item["error"] for item in review_errors if item["sku"] == sku), "MPStats недоступен")
            status_values.append((project_id, sku, "review", "mpstats_ozon_comments", period_days, "error", None, error))
        if question_error is None:
            status_values.append((project_id, sku, "question", "ozon_seller_api_questions", period_days, "ok", question_counts.get(sku, 0), None))
        else:
            status = "unavailable" if "403" in question_error or "subscription" in question_error.lower() else "error"
            status_values.append((project_id, sku, "question", "ozon_seller_api_questions", period_days, status, None, question_error))

    message_values = [(
        project_id, row["sku"], row["message_type"], row["source_message_id"], row["message_date"],
        row.get("rating"), row.get("message_text"), row.get("pros"), row.get("cons"),
        row.get("answer_text"), row.get("answered"), row["source"], period_days,
        json.dumps(row.get("raw_payload") or {}, ensure_ascii=False, default=str),
    ) for row in review_rows + question_rows]
    competitor_review_values = [(
        project_id, row["sku"], row["competitor_sku"], row.get("competitor_rank"),
        row.get("competitor_name"), row["source_message_id"], row["message_date"], row.get("rating"),
        row.get("message_text"), row.get("pros"), row.get("cons"), row.get("answer_text"),
        row.get("answered"), row["source"], period_days,
        json.dumps(row.get("raw_payload") or {}, ensure_ascii=False, default=str),
    ) for row in competitor_review_rows]
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            if successful_review_skus:
                cur.execute(
                    "DELETE FROM public.seo_generation_customer_messages WHERE project_id=%s AND message_type='review' AND sku=ANY(%s)",
                    (project_id, list(successful_review_skus)),
                )
            if question_error is None:
                cur.execute(
                    "DELETE FROM public.seo_generation_customer_messages WHERE project_id=%s AND message_type='question' AND sku=ANY(%s)",
                    (project_id, requested_skus),
                )
            if message_values:
                execute_values(cur, """INSERT INTO public.seo_generation_customer_messages
                    (project_id, sku, message_type, source_message_id, message_date, rating, message_text,
                     pros, cons, answer_text, answered, source, period_days, raw_payload) VALUES %s
                    ON CONFLICT(project_id, sku, message_type, source, source_message_id) DO UPDATE SET
                      message_date=excluded.message_date, rating=excluded.rating, message_text=excluded.message_text,
                      pros=excluded.pros, cons=excluded.cons, answer_text=excluded.answer_text,
                      answered=excluded.answered, period_days=excluded.period_days,
                      raw_payload=excluded.raw_payload, collected_at=now()""", message_values)
            execute_values(cur, """INSERT INTO public.seo_generation_customer_message_status
                (project_id, sku, message_type, source, period_days, status, message_count, last_error) VALUES %s
                ON CONFLICT(project_id, sku, message_type) DO UPDATE SET
                  source=excluded.source, period_days=excluded.period_days, status=excluded.status,
                  message_count=excluded.message_count, last_error=excluded.last_error, collected_at=now()""", status_values)
            for competitor_sku in successful_competitors:
                own_skus = sorted({str(item["sku"]) for item in links_by_competitor[competitor_sku]})
                cur.execute(
                    """DELETE FROM public.seo_generation_competitor_reviews
                       WHERE project_id=%s AND competitor_sku=%s AND sku=ANY(%s)""",
                    (project_id, competitor_sku, own_skus),
                )
            if competitor_review_values:
                execute_values(cur, """INSERT INTO public.seo_generation_competitor_reviews
                    (project_id, sku, competitor_sku, competitor_rank, competitor_name, source_message_id,
                     message_date, rating, message_text, pros, cons, answer_text, answered, source,
                     period_days, raw_payload) VALUES %s
                    ON CONFLICT(project_id, sku, competitor_sku, source_message_id) DO UPDATE SET
                      competitor_rank=excluded.competitor_rank, competitor_name=excluded.competitor_name,
                      message_date=excluded.message_date, rating=excluded.rating,
                      message_text=excluded.message_text, pros=excluded.pros, cons=excluded.cons,
                      answer_text=excluded.answer_text, answered=excluded.answered,
                      period_days=excluded.period_days, raw_payload=excluded.raw_payload,
                      collected_at=now()""", competitor_review_values)
            if competitor_status_values:
                execute_values(cur, """INSERT INTO public.seo_generation_competitor_review_status
                    (project_id, sku, competitor_sku, source, period_days, status, review_count, last_error)
                    VALUES %s
                    ON CONFLICT(project_id, sku, competitor_sku) DO UPDATE SET
                      source=excluded.source, period_days=excluded.period_days, status=excluded.status,
                      review_count=excluded.review_count, last_error=excluded.last_error,
                      collected_at=now()""", competitor_status_values)
    missing_competitor_skus = sorted(set(requested_skus) - {str(item["sku"]) for item in competitor_links})
    missing_competitor_errors = [{
        "sku": sku, "type": "competitor_review", "status": "unavailable",
        "error": "Сначала соберите конкурентов по нише",
    } for sku in missing_competitor_skus]
    errors = review_errors + competitor_review_errors + missing_competitor_errors + (
        [{"type": "question", "error": question_error}] if question_error else []
    )
    technical_errors = [item for item in errors if str(item.get("status") or "").lower() != "unavailable"]
    return {
        "ok": not technical_errors, "partial": bool(errors), "project_id": project_id,
        "sku_count": len(requested_skus), "period_days": period_days,
        "date_from": period_from.isoformat(), "date_to": period_to.isoformat(),
        "review_count": len(review_rows), "question_count": len(question_rows),
        "competitor_review_count": len(competitor_review_rows),
        "competitor_count": len(links_by_competitor),
        "competitor_review_requests": len(links_by_competitor),
        "competitor_review_skus": len({str(item["sku"]) for item in competitor_links}),
        "review_skus": len(successful_review_skus),
        "question_skus": len(requested_skus) if question_error is None else 0,
        "sources": {"reviews": "MPStats Ozon", "questions": "Ozon Seller API",
                    "competitor_reviews": "MPStats Ozon"},
        "errors": errors,
    }


def project_customer_messages(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    message_type = _clean_text(payload.get("message_type"), 20)
    if not sku:
        raise ValueError("SKU не передан")
    if message_type not in {"", "review", "question"}:
        raise ValueError("message_type должен быть review или question")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                """SELECT message_type, source, period_days, status, message_count, last_error, collected_at
                   FROM public.seo_generation_customer_message_status
                   WHERE project_id=%s AND sku=%s ORDER BY message_type""",
                (project_id, sku),
            )
            statuses = [_serialize(row) for row in cur.fetchall()]
            type_sql = " AND message_type=%s" if message_type else ""
            values = [project_id, sku] + ([message_type] if message_type else [])
            cur.execute(
                f"SELECT count(*) AS total FROM public.seo_generation_customer_messages WHERE project_id=%s AND sku=%s{type_sql}",
                values,
            )
            total = int(cur.fetchone()["total"])
            cur.execute(
                f"""SELECT message_type, source_message_id, message_date, rating, message_text, pros, cons,
                            answer_text, answered, source, period_days, collected_at
                     FROM public.seo_generation_customer_messages
                     WHERE project_id=%s AND sku=%s{type_sql}
                     ORDER BY message_date DESC NULLS LAST, source_message_id DESC LIMIT 500""",
                values,
            )
            rows = [_serialize(row) for row in cur.fetchall()]
    return {"ok": True, "project_id": project_id, "sku": sku, "message_type": message_type or None,
            "rows": rows, "count": total, "statuses": statuses, "truncated": total > len(rows)}


def project_competitor_reviews(config, payload):
    """Return collected review texts for the top competitors of one project SKU."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    competitor_sku = _clean_text(payload.get("competitor_sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    competitor_sql = " AND competitor_sku=%s" if competitor_sku else ""
    values = [project_id, sku] + ([competitor_sku] if competitor_sku else [])
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                f"""SELECT competitor_sku, source, period_days, status, review_count, last_error, collected_at
                     FROM public.seo_generation_competitor_review_status
                     WHERE project_id=%s AND sku=%s{competitor_sql}
                     ORDER BY competitor_sku""",
                values,
            )
            statuses = [_serialize(row) for row in cur.fetchall()]
            cur.execute(
                f"""SELECT count(*) AS total FROM public.seo_generation_competitor_reviews
                     WHERE project_id=%s AND sku=%s{competitor_sql}""",
                values,
            )
            total = int(cur.fetchone()["total"])
            cur.execute(
                f"""SELECT competitor_sku, competitor_rank, competitor_name, source_message_id,
                            message_date, rating, message_text, pros, cons, answer_text, answered,
                            source, period_days, collected_at
                     FROM public.seo_generation_competitor_reviews
                     WHERE project_id=%s AND sku=%s{competitor_sql}
                     ORDER BY competitor_rank NULLS LAST, competitor_sku,
                              message_date DESC NULLS LAST, source_message_id DESC LIMIT 500""",
                values,
            )
            rows = [_serialize(row) for row in cur.fetchall()]
    return {
        "ok": True, "project_id": project_id, "sku": sku,
        "competitor_sku": competitor_sku or None, "rows": rows, "count": total,
        "statuses": statuses, "truncated": total > len(rows),
    }


def _customer_voice_product_contexts(cur, project_id, skus, marketplace="ozon"):
    contexts = {str(sku): {"sku": str(sku), "product_name": "", "search_intent": "",
                          "description": "", "attributes": ""} for sku in skus}
    cur.execute(
        """SELECT ps.sku, ps.product_name, pi.search_intent
           FROM public.seo_monitoring_project_skus ps
           LEFT JOIN public.seo_generation_product_intents pi
             ON pi.project_id=ps.project_id AND pi.sku=ps.sku
           WHERE ps.project_id=%s AND ps.sku=ANY(%s)""",
        (project_id, skus),
    )
    for row in cur.fetchall():
        contexts[str(row["sku"])].update({
            "product_name": _clean_text(row.get("product_name"), 350),
            "search_intent": _clean_text(row.get("search_intent"), 180),
        })
    if marketplace == "wb":
        products = _semantic_product_contexts(cur, project_id, skus, marketplace)
        for sku, product in products.items():
            contexts[sku].update({
                "product_name": product.get("title") or contexts[sku].get("product_name") or "",
                "description": product.get("description") or "",
                "category": product.get("category") or "",
                "type": product.get("product_type") or "",
                "attributes": " | ".join(
                    f"{row.get('name')}: {row.get('value')}"
                    for row in (product.get("seo_characteristics") or [])
                )[:1800],
            })
        return contexts
    if _relation_exists(cur, "ozon_cat_products"):
        cur.execute(
            """SELECT sku::text AS sku, annotatsiya, rich_kontent_json, category_name, tip,
                      product_id, category_id
               FROM public.ozon_cat_products WHERE sku::text=ANY(%s)""",
            (skus,),
        )
        product_rows = {str(row["sku"]): dict(row) for row in cur.fetchall()}
        for sku, row in product_rows.items():
            contexts[sku].update({
                "description": " ".join(filter(None, [
                    _clean_text(row.get("annotatsiya"), 1000),
                    _json_text(row.get("rich_kontent_json"), 600),
                ]))[:1400],
                "category": _clean_text(row.get("category_name"), 220),
                "type": _clean_text(row.get("tip"), 160),
            })
        if product_rows:
            cur.execute(
                """SELECT p.sku::text AS sku, coalesce(ca.attribute_name, co.attribute_name) AS name,
                          pa.value_text AS value
                   FROM public.ozon_cat_products p
                   JOIN public.ozon_cat_product_attributes pa ON pa.product_id=p.product_id
                   LEFT JOIN public.ozon_cat_category_attributes ca
                     ON ca.category_id=p.category_id AND ca.attribute_id=pa.attribute_id
                   LEFT JOIN public.ozon_cat_common_attributes co ON co.attribute_id=pa.attribute_id
                   WHERE p.sku::text=ANY(%s) AND pa.value_text IS NOT NULL""",
                (skus,),
            )
            attributes = {}
            for row in cur.fetchall():
                name, value = _clean_text(row.get("name"), 120), _clean_text(row.get("value"), 220)
                if name and value and not re.search(r"штрих|код маркиров", name, re.I):
                    attributes.setdefault(str(row["sku"]), []).append(f"{name}: {value}")
            for sku, values in attributes.items():
                contexts[sku]["attributes"] = " | ".join(values[:20])[:1800]
    return contexts


def _customer_voice_evidence_rows(cur, project_id, skus):
    by_sku = {str(sku): [] for sku in skus}
    cur.execute(
        """SELECT sku, message_type, source_message_id, message_date, rating, message_text,
                  pros, cons, answer_text, source
           FROM public.seo_generation_customer_messages
           WHERE project_id=%s AND sku=ANY(%s)
           ORDER BY sku, message_type, message_date DESC NULLS LAST""",
        (project_id, skus),
    )
    own_limits = {sku: {"review": 0, "question": 0} for sku in by_sku}
    for row in cur.fetchall():
        sku, kind = str(row["sku"]), str(row["message_type"])
        limit = 30
        if own_limits[sku].get(kind, 0) >= limit:
            continue
        text = " | ".join(filter(None, [
            _clean_text(row.get("message_text"), 500),
            f"Плюсы: {_clean_text(row.get('pros'), 250)}" if row.get("pros") else "",
            f"Минусы: {_clean_text(row.get('cons'), 250)}" if row.get("cons") else "",
            f"Ответ: {_clean_text(row.get('answer_text'), 250)}" if row.get("answer_text") else "",
        ]))[:900]
        if not text:
            continue
        evidence_type = "own_question" if kind == "question" else "own_review"
        evidence_id = f"{evidence_type}:{row['source_message_id']}"
        by_sku[sku].append({
            "evidence_id": evidence_id, "source_type": evidence_type, "text": text,
            "date": row.get("message_date"), "rating": row.get("rating"),
        })
        own_limits[sku][kind] = own_limits[sku].get(kind, 0) + 1
    cur.execute(
        """SELECT sku, competitor_sku, competitor_rank, source_message_id, message_date, rating,
                  message_text, pros, cons, answer_text
           FROM public.seo_generation_competitor_reviews
           WHERE project_id=%s AND sku=ANY(%s)
           ORDER BY sku, competitor_rank, message_date DESC NULLS LAST""",
        (project_id, skus),
    )
    competitor_limits = {}
    for row in cur.fetchall():
        sku, competitor_sku = str(row["sku"]), str(row["competitor_sku"])
        key = (sku, competitor_sku)
        if competitor_limits.get(key, 0) >= 8:
            continue
        text = " | ".join(filter(None, [
            _clean_text(row.get("message_text"), 500),
            f"Плюсы: {_clean_text(row.get('pros'), 250)}" if row.get("pros") else "",
            f"Минусы: {_clean_text(row.get('cons'), 250)}" if row.get("cons") else "",
            f"Ответ: {_clean_text(row.get('answer_text'), 250)}" if row.get("answer_text") else "",
        ]))[:900]
        if not text:
            continue
        evidence_id = f"competitor_review:{competitor_sku}:{row['source_message_id']}"
        by_sku[sku].append({
            "evidence_id": evidence_id, "source_type": "competitor_review", "text": text,
            "competitor_sku": competitor_sku, "competitor_rank": row.get("competitor_rank"),
            "date": row.get("message_date"), "rating": row.get("rating"),
        })
        competitor_limits[key] = competitor_limits.get(key, 0) + 1
    return by_sku


_CUSTOMER_VOICE_CLAIM_STOPWORDS = {
    "для", "это", "как", "что", "или", "при", "под", "над", "без", "очень", "можно",
    "нужно", "надо", "товар", "модель", "изделие", "карточка", "указать", "уточнить",
}


def _customer_voice_claim_limits(evidence_count):
    """Return a compact synthesis range instead of one claim per source message."""
    count = max(0, int(evidence_count or 0))
    if count <= 1:
        return {"min": 1, "max": 1}
    if count <= 3:
        return {"min": 1, "max": 2}
    if count <= 7:
        return {"min": 2, "max": 4}
    if count <= 19:
        return {"min": 3, "max": 5}
    return {"min": 3, "max": 7}


def _customer_voice_claim_tokens(value):
    return {
        token for token in re.findall(r"[a-zа-яё0-9]+", _seo_normalize(value), re.I)
        if len(token) >= 3 and token not in _CUSTOMER_VOICE_CLAIM_STOPWORDS
    }


def _customer_voice_claim_similarity(left, right):
    left_text, right_text = _seo_normalize(left), _seo_normalize(right)
    if not left_text or not right_text:
        return 0.0
    if left_text == right_text:
        return 1.0
    if min(len(left_text), len(right_text)) >= 16 and (left_text in right_text or right_text in left_text):
        return 0.95
    left_tokens, right_tokens = _customer_voice_claim_tokens(left_text), _customer_voice_claim_tokens(right_text)
    if not left_tokens or not right_tokens:
        return 0.0
    return len(left_tokens & right_tokens) / len(left_tokens | right_tokens)


def analyze_project_customer_voice(config, payload, analyze_batch):
    """Extract source-backed SEO claims from own/competitor customer voice."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if not requested_skus or len(requested_skus) > 5:
        raise ValueError("Передайте от 1 до 5 SKU проекта")
    with _conn(config) as conn:
        ensure_schema(conn)
        settings = _ai_script_models(conn, "customer_voice_claims")
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            marketplace = project.get("marketplace")
            cur.execute(
                "SELECT sku FROM public.seo_monitoring_project_skus WHERE project_id=%s AND sku=ANY(%s)",
                (project_id, requested_skus),
            )
            found = {str(row["sku"]) for row in cur.fetchall()}
            missing = [sku for sku in requested_skus if sku not in found]
            if missing:
                raise ValueError(f"SKU не входят в проект: {', '.join(missing[:5])}")
            contexts = _customer_voice_product_contexts(cur, project_id, requested_skus, marketplace)
            evidence_by_sku = _customer_voice_evidence_rows(cur, project_id, requested_skus)

    prepared, unavailable = [], []
    for sku in requested_skus:
        source_evidence = evidence_by_sku.get(sku) or []
        evidence = (
            [item for item in source_evidence if item["source_type"] == "own_review"][:20]
            + [item for item in source_evidence if item["source_type"] == "own_question"][:20]
            + [item for item in source_evidence if item["source_type"] == "competitor_review"][:40]
        )
        if not evidence:
            unavailable.append(sku)
            continue
        source_counts = {
            "own_reviews": sum(item["source_type"] == "own_review" for item in evidence),
            "own_questions": sum(item["source_type"] == "own_question" for item in evidence),
            "competitor_reviews": sum(item["source_type"] == "competitor_review" for item in evidence),
        }
        prepared.append({
            "sku": sku, "product_context": contexts.get(sku) or {}, "evidence": evidence,
            "source_counts": source_counts, "claim_limits": _customer_voice_claim_limits(len(evidence)),
        })
    result, ai_error = {"items": [], "model": ""}, None
    if prepared:
        try:
            result = analyze_batch(prepared, settings) if analyze_batch else result
        except Exception as exc:
            ai_error = str(exc)[:500]
    result_by_sku = {
        str(item.get("sku") or ""): item for item in (result.get("items") or []) if isinstance(item, dict)
    }
    model = _clean_text(result.get("model"), 160)
    db_rows, status_values, errors = [], [], []
    for row in prepared:
        sku = row["sku"]
        evidence_map = {item["evidence_id"]: item for item in row["evidence"]}
        item = result_by_sku.get(sku) or {}
        claims = item.get("claims") if isinstance(item.get("claims"), list) else []
        claim_candidates = []
        claim_limit = int((row.get("claim_limits") or {}).get("max") or 7)
        require_repeated_theme = len(row["evidence"]) >= 6
        for claim in claims[:20]:
            if not isinstance(claim, dict):
                continue
            text = _clean_text(claim.get("claim") or claim.get("claim_text"), 400)
            if not text:
                continue
            refs = list(dict.fromkeys(
                str(value) for value in (claim.get("source_refs") or [])
                if str(value) in evidence_map or str(value) == "context:product"
            ))
            voice_refs = [ref for ref in refs if ref != "context:product"]
            if not voice_refs:
                continue
            evidence = [evidence_map[ref] for ref in voice_refs]
            own_reviews = sum(item["source_type"] == "own_review" for item in evidence)
            own_questions = sum(item["source_type"] == "own_question" for item in evidence)
            competitor_reviews = sum(item["source_type"] == "competitor_review" for item in evidence)
            requested_status = str(claim.get("verification_status") or "review").strip().lower()
            verification_status = "safe" if (
                requested_status == "safe" and "context:product" in refs and (own_reviews or competitor_reviews)
            ) else "review"
            confidence = _number(claim.get("confidence"))
            confidence = min(1, max(0, confidence)) if confidence is not None else None
            claim_type = _clean_text(claim.get("claim_type"), 40) or "other"
            seo_target = _clean_text(claim.get("seo_target"), 40) or "description"
            if require_repeated_theme and len(voice_refs) < 2 and claim_type != "faq":
                continue
            claim_id = hashlib.sha256(
                f"{claim_type}|{seo_target}|{_seo_normalize(text)}".encode("utf-8")
            ).hexdigest()[:24]
            evidence_json = [{
                "evidence_id": source["evidence_id"], "source_type": source["source_type"],
                "text": _clean_text(source.get("text"), 500), "date": source.get("date"),
                "rating": source.get("rating"), "competitor_sku": source.get("competitor_sku"),
            } for source in evidence]
            source_diversity = sum(bool(value) for value in (own_reviews, own_questions, competitor_reviews))
            cross_market_signal = bool((own_reviews or own_questions) and competitor_reviews)
            candidate_row = (
                project_id, sku, claim_id, text, claim_type, seo_target, verification_status,
                confidence, own_reviews, own_questions, competitor_reviews, len(evidence),
                _clean_text(claim.get("rationale"), 500),
                json.dumps(evidence_json, ensure_ascii=False, default=str), model,
            )
            claim_candidates.append({
                "claim_id": claim_id, "text": text, "row": candidate_row,
                "rank": (
                    int(cross_market_signal), int(bool(own_reviews or own_questions)), source_diversity,
                    len(evidence), float(confidence or 0), int(verification_status == "safe"),
                ),
            })
        claim_candidates.sort(key=lambda value: value["rank"], reverse=True)
        accepted_candidates = []
        for candidate in claim_candidates:
            if any(_customer_voice_claim_similarity(candidate["text"], saved["text"]) >= 0.72
                   for saved in accepted_candidates):
                continue
            accepted_candidates.append(candidate)
            if len(accepted_candidates) >= claim_limit:
                break
        db_rows.extend(candidate["row"] for candidate in accepted_candidates)
        accepted = [candidate["claim_id"] for candidate in accepted_candidates]
        status = "error" if ai_error else "ok"
        last_error = ai_error or None
        status_values.append((project_id, sku, status, len(row["evidence"]), len(accepted) if not ai_error else None,
                              last_error, model or None))
        if ai_error:
            errors.append({"sku": sku, "type": "customer_voice_ai", "error": ai_error})
    for sku in unavailable:
        status_values.append((project_id, sku, "unavailable", 0, None,
                              "Сначала соберите отзывы или вопросы", None))
        errors.append({"sku": sku, "type": "customer_voice", "status": "unavailable",
                       "error": "Сначала соберите отзывы или вопросы"})
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            successful_skus = [value[1] for value in status_values if value[2] == "ok"]
            if successful_skus:
                cur.execute(
                    "DELETE FROM public.seo_generation_customer_voice_claims WHERE project_id=%s AND sku=ANY(%s)",
                    (project_id, successful_skus),
                )
            if db_rows:
                execute_values(cur, """INSERT INTO public.seo_generation_customer_voice_claims
                    (project_id, sku, claim_id, claim_text, claim_type, seo_target, verification_status,
                     confidence, own_review_mentions, own_question_mentions, competitor_review_mentions,
                     evidence_count, rationale, evidence_json, model) VALUES %s""", db_rows)
            execute_values(cur, """INSERT INTO public.seo_generation_customer_voice_analysis_status
                (project_id, sku, status, message_count, claim_count, last_error, model) VALUES %s
                ON CONFLICT(project_id, sku) DO UPDATE SET status=excluded.status,
                  message_count=excluded.message_count, claim_count=excluded.claim_count,
                  last_error=excluded.last_error, model=excluded.model, analyzed_at=now()""", status_values)
    technical_errors = [item for item in errors if str(item.get("status") or "").lower() != "unavailable"]
    return {
        "ok": not technical_errors, "partial": bool(errors), "project_id": project_id,
        "sku_count": len(requested_skus), "analyzed_skus": len(prepared) if not ai_error else 0,
        "message_count": sum(len(row["evidence"]) for row in prepared),
        "claim_count": len(db_rows), "model": model, "errors": errors,
    }


def project_customer_voice_claims(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                """SELECT status, message_count, claim_count, last_error, model, analyzed_at
                   FROM public.seo_generation_customer_voice_analysis_status
                   WHERE project_id=%s AND sku=%s""", (project_id, sku),
            )
            status_row = cur.fetchone()
            status = _serialize(status_row) if status_row else None
            cur.execute(
                """SELECT claim_id, claim_text, claim_type, seo_target, verification_status, confidence,
                          own_review_mentions, own_question_mentions, competitor_review_mentions,
                          evidence_count, rationale, evidence_json, model, analyzed_at
                   FROM public.seo_generation_customer_voice_claims
                   WHERE project_id=%s AND sku=%s
                   ORDER BY verification_status, confidence DESC NULLS LAST, evidence_count DESC, claim_text""",
                (project_id, sku),
            )
            rows = [_serialize(row) for row in cur.fetchall()]
    return {"ok": True, "project_id": project_id, "sku": sku, "status": status,
            "rows": rows, "count": len(rows)}


SEO_CONTEXT_VERSION = "seo-semantic-context-v1"
SEO_CONTEXT_OWN_KEYWORD_LIMIT = 80
SEO_CONTEXT_COMPETITOR_KEYWORD_LIMIT = 80
SEO_CONTEXT_CLAIM_LIMIT = 12
SEO_CONTEXT_ATTRIBUTE_LIMIT = 30
SEO_CONTENT_ALLOCATION_RULES_VERSION = "ozon-product-content-rules-v2"
SEO_CONTENT_ALLOCATION_PROMPT_VERSION = "seo-content-allocation-v2"
SEO_CONTENT_DRAFT_RULES_VERSION = "ozon-seo-content-draft-rules-v10-dynamic-structure"
SEO_CONTENT_DRAFT_PROMPT_VERSION = "seo-content-draft-v11-dynamic-structure"
SEO_CONTENT_REVIEW_RULES_VERSION = "ozon-seo-content-expert-review-rules-v5"
SEO_CONTENT_REVIEW_PROMPT_VERSION = "seo-content-expert-review-v4"


def _content_stage_versions(marketplace, stage):
    """Keep stored audit versions marketplace-specific without duplicating the validators."""
    prefix = "wb" if marketplace == "wb" else "ozon"
    versions = {
        "allocation": (f"{prefix}-product-content-rules-v2", "seo-content-allocation-v2"),
        "draft": (f"{prefix}-seo-content-draft-rules-v10-dynamic-structure", "seo-content-draft-v11-dynamic-structure"),
        "review": (f"{prefix}-seo-content-expert-review-rules-v5", "seo-content-expert-review-v4"),
    }
    return versions[stage]
SEO_CONTENT_DESCRIPTION_TARGET_MIN = 1300
SEO_CONTENT_DESCRIPTION_TARGET_MAX = 1800
SEO_CONTENT_DESCRIPTION_SOFT_MAX = 2200
SEO_CONTENT_DESCRIPTION_SECTION_KINDS = {
    "product_identity": "Товар и назначение",
    "material_and_construction": "Материал и конструкция",
    "fit_dimensions_compatibility": "Размер, посадка или совместимость",
    "usage_scenarios": "Сценарии использования",
    "functional_details": "Функции и детали",
    "care_operation_storage": "Уход, эксплуатация или хранение",
    "style_audience_context": "Стиль и аудитория",
    "source_backed_benefits": "Подтверждённые преимущества",
}
SEO_SOURCE_MODEL_ATTRIBUTE_RE = re.compile(r"^(?:название\s+|код\s+)?модел(?:ь|и)?$", re.I)
SEO_SOURCE_MODEL_TOKEN_RE = re.compile(
    r"^(?=.{3,40}$)(?=.*[A-Za-zА-Яа-яЁё])(?=.*\d)[A-Za-zА-Яа-яЁё0-9._/-]+$"
)
SEO_CONTEXT_IMPORTANT_ATTRIBUTE_RE = re.compile(
    r"пол|гендер|возраст|аудитор|бренд|марка|категор|тип товара|вид товара|предмет|назначен|"
    r"материал|состав|цвет|оттенок|сезон|стиль|узор|рисунок|принт|посадк|силуэт|"
    r"рукав|воротник|длина|рост|размер|комплект|особенност|застеж|карман|форма", re.I,
)
SEO_CONTEXT_TECHNICAL_ATTRIBUTE_RE = re.compile(
    r"фото|изображ|медиа|картин|ссылк|url|rich|контент|json|jison|widget|"
    r"таблиц.*размер|размерн.*таблиц|упаков|габарит|вес (?:товара|брутто|нетто)|"
    r"штрихкод|артикул|идентификатор|код маркиров", re.I,
)


def _source_model_identifier(context):
    """Return the exact source model identifier that must close the generated title."""
    product = (context or {}).get("product") or {}
    for row in product.get("seo_characteristics") or []:
        name = _clean_text((row or {}).get("name"), 160)
        value = _clean_text((row or {}).get("value"), 160)
        if value and SEO_SOURCE_MODEL_ATTRIBUTE_RE.fullmatch(name):
            return value
    source_title = _clean_text(product.get("title"), 500)
    for raw_token in source_title.split():
        token = raw_token.strip(".,;:()[]{}")
        if SEO_SOURCE_MODEL_TOKEN_RE.fullmatch(token):
            return token
    return ""


def _source_warranty_claims(context):
    """Extract exact source-backed warranty sentences; absence stays absence."""
    description = _clean_text(((context or {}).get("product") or {}).get("description"), 12000)
    if not description:
        return []
    claims = []
    for sentence in re.split(r"(?<=[.!?])\s+|[\r\n]+", description):
        sentence = _clean_text(sentence, 1000)
        if sentence and re.search(r"\bгарант", sentence, re.I) and sentence not in claims:
            claims.append(sentence)
    return claims


def _enforce_source_preservation(title, description, context):
    """Deterministically preserve source model suffix and warranty without inventing facts."""
    model = _source_model_identifier(context)
    warranties = _source_warranty_claims(context)
    if model and not title.casefold().endswith(model.casefold()):
        separator = " · "
        available = 200 - len(separator) - len(model)
        if available <= 0:
            raise ValueError("Исходное название модели не помещается в лимит названия Ozon")
        title = f"{title[:available].rstrip(' ·,;:-')}{separator}{model}"
    for claim in warranties:
        if _seo_normalize(claim) not in _seo_normalize(description):
            description = f"{description.rstrip()} {claim}".strip()
    description = _clean_text(description, 4000)
    if model and not title.casefold().endswith(model.casefold()):
        raise ValueError("Название должно заканчиваться исходным названием модели")
    for claim in warranties:
        if _seo_normalize(claim) not in _seo_normalize(description):
            raise ValueError(f"В описании не сохранён исходный гарантийный клейм: {claim}")
    return title, description, model, warranties


def _build_title_structure(title, main_intent, title_characteristics, required_model):
    """Describe the actual generic title composition without imposing a category template."""
    title_folded = str(title or "").casefold()
    components = []
    seen = set()

    def add(kind, label, value, source, required=False):
        clean_value = _clean_text(value, 500)
        key = (kind, _seo_normalize(clean_value))
        if not clean_value or not key[1] or key in seen:
            return
        seen.add(key)
        position = title_folded.find(clean_value.casefold())
        normalized_value = _seo_normalize(clean_value)
        normalized_title = _seo_normalize(title)
        value_bases = {
            _seo_russian_stem(token) for token in normalized_value.split()
            if len(token) >= 3 and _seo_russian_stem(token)
        }
        title_bases = {
            _seo_russian_stem(token) for token in normalized_title.split()
            if len(token) >= 3 and _seo_russian_stem(token)
        }
        found = position >= 0 or normalized_value in normalized_title or bool(value_bases and value_bases <= title_bases)
        row = {
            "kind": kind,
            "label": label,
            "value": clean_value,
            "source": source,
            "required": bool(required),
            "found": bool(found),
            "position": position if position >= 0 else None,
        }
        components.append(row)

    add("primary_intent", "Основной интент", main_intent, "product.intent", True)
    for row in title_characteristics or []:
        add(
            "characteristic",
            _clean_text(row.get("name"), 160) or "Различающая характеристика",
            row.get("value"),
            "SEO-разметка характеристики",
        )
    add("model_suffix", "Модель", required_model, "Исходная карточка", bool(required_model))
    components.sort(key=lambda row: (
        row.get("position") is None,
        row.get("position") if row.get("position") is not None else 999999,
        row.get("kind") or "",
    ))
    return {
        "dynamic": True,
        "policy": "основной интент → уместные различающие факты → обязательная модель в конце",
        "components": [row for row in components if row["found"]],
        "omitted_selected": [row for row in components if not row["found"]],
    }


def _semantic_context_attribute(name, value):
    """Return a compact SEO-relevant characteristic or None for technical/raw fields."""
    clean_name = _clean_text(name, 160)
    clean_value = _clean_text(value, 500)
    if not clean_name or not clean_value:
        return None
    if not SEO_CONTEXT_IMPORTANT_ATTRIBUTE_RE.search(clean_name):
        return None
    if SEO_CONTEXT_TECHNICAL_ATTRIBUTE_RE.search(clean_name):
        return None
    if re.match(r"^https?://", clean_value, re.I) or clean_value.startswith(("{", "[")):
        return None
    return {"name": clean_name, "value": clean_value}


def _semantic_keyword_rank(row):
    labels = {"high": 3, "medium": 2, "low": 1, "none": 0}
    frequencies = {"high": 3, "mid": 2, "low": 1, None: 0}
    return (
        labels.get(str(row.get("priority_label") or "none"), 0),
        int(_number(row.get("priority_score")) or 0),
        frequencies.get(row.get("frequency_class"), 0),
        float(_number(row.get("search_demand")) or 0),
        float(_number(row.get("traffic")) or 0),
        -float(_number(row.get("best_position") or row.get("average_position")) or 999999),
    )


def _select_semantic_keywords(rows, limit):
    """Deduplicate analyzed phrases while retaining source coverage and ranking evidence."""
    selected = {}
    for raw in rows or []:
        if str(raw.get("relevance_decision") or "") != "keep":
            continue
        query = _clean_text(raw.get("clean_query") or raw.get("search_query"), 300)
        normalized = _seo_normalize(query)
        if not normalized:
            continue
        item = {
            "query": query,
            "sources": [str(raw.get("source_group") or "competitors")],
            "semantic_type": _clean_text(raw.get("semantic_type"), 60) or "other",
            "frequency_class": raw.get("frequency_class"),
            "priority_label": raw.get("priority_label") or "none",
            "priority_score": int(_number(raw.get("priority_score")) or 0),
            "search_demand": _number(raw.get("search_demand")),
            "traffic": _number(raw.get("traffic")),
            "orders": _number(raw.get("orders")),
            "best_position": _number(raw.get("best_position") or raw.get("average_position")),
            "competitor_coverage": int(_number(raw.get("competitor_coverage")) or 0),
            "reason": _clean_text(raw.get("relevance_reason"), 300),
        }
        current = selected.get(normalized)
        if current is None:
            selected[normalized] = item
            continue
        current["sources"] = sorted(set(current["sources"] + item["sources"]))
        current["search_demand"] = max(filter(lambda value: value is not None, [current.get("search_demand"), item.get("search_demand")]), default=None)
        current["traffic"] = max(filter(lambda value: value is not None, [current.get("traffic"), item.get("traffic")]), default=None)
        current["orders"] = max(filter(lambda value: value is not None, [current.get("orders"), item.get("orders")]), default=None)
        current["competitor_coverage"] = max(current.get("competitor_coverage") or 0, item.get("competitor_coverage") or 0)
        positions = [value for value in (current.get("best_position"), item.get("best_position")) if value is not None]
        current["best_position"] = min(positions) if positions else None
        if _semantic_keyword_rank(item) > _semantic_keyword_rank(current):
            sources = current["sources"]
            selected[normalized] = {**item, "sources": sources}
    ranked = sorted(selected.values(), key=_semantic_keyword_rank, reverse=True)
    return ranked[:limit]


def _build_semantic_context(sku, product, own_rows, competitor_rows, claims, competitors, marketplace="ozon"):
    own_selected = _select_semantic_keywords(own_rows, SEO_CONTEXT_OWN_KEYWORD_LIMIT)
    competitor_selected = _select_semantic_keywords(competitor_rows, SEO_CONTEXT_COMPETITOR_KEYWORD_LIMIT)
    ranked_claims = sorted(
        claims or [],
        key=lambda row: (
            int(str(row.get("verification_status") or "") == "safe"),
            float(_number(row.get("confidence")) or 0),
            int(_number(row.get("evidence_count")) or 0),
        ), reverse=True,
    )[:SEO_CONTEXT_CLAIM_LIMIT]
    source_counts = {
        "all_characteristics": int(product.get("all_characteristic_count") or 0),
        "seo_characteristics": len(product.get("seo_characteristics") or []),
        "own_keywords_available": len(own_rows or []),
        "own_keywords_selected": len(own_selected),
        "competitor_keywords_available": len(competitor_rows or []),
        "competitor_keywords_selected": len(competitor_selected),
        "claims_available": len(claims or []),
        "claims_selected": len(ranked_claims),
        "competitors": len(competitors or []),
    }
    warnings = []
    if not product.get("title"):
        warnings.append("Название товара недоступно")
    if not product.get("description"):
        warnings.append("Описание товара недоступно")
    if not product.get("seo_characteristics"):
        warnings.append("SEO-чувствительные характеристики не найдены")
    if not own_rows:
        warnings.append("Очищенные свои ключи отсутствуют")
    if not competitor_rows:
        warnings.append("Очищенные ключи конкурентов отсутствуют")
    if not claims:
        warnings.append("SEO-клеймы из отзывов и вопросов отсутствуют")
    context = {
        "context_version": SEO_CONTEXT_VERSION,
        "marketplace": marketplace,
        "sku": str(sku),
        "product": {
            "title": product.get("title") or "",
            "description": product.get("description") or "",
            "category": product.get("category") or "",
            "product_type": product.get("product_type") or "",
            "intent": product.get("intent") or "",
            "seo_characteristics": product.get("seo_characteristics") or [],
        },
        "own_keywords": own_selected,
        "competitor_keywords": competitor_selected,
        "customer_voice_claims": [{
            "claim": row.get("claim_text"),
            "target": row.get("seo_target"),
            "verification_status": row.get("verification_status"),
            "confidence": _number(row.get("confidence")),
            "evidence_count": int(_number(row.get("evidence_count")) or 0),
            "rationale": row.get("rationale"),
        } for row in ranked_claims],
        "competitors": [{
            "sku": row.get("competitor_sku"), "rank": row.get("competitor_rank"),
            "name": row.get("product_name"), "brand": row.get("brand"),
        } for row in (competitors or [])[:10]],
        "coverage": source_counts,
        "warnings": warnings,
    }
    stable_json = json.dumps(context, ensure_ascii=False, sort_keys=True, default=str)
    return context, source_counts, hashlib.sha256(stable_json.encode("utf-8")).hexdigest(), ("partial" if warnings else "ok")


def _semantic_product_contexts(cur, project_id, skus, marketplace="ozon"):
    contexts = {str(sku): {"seo_characteristics": [], "all_characteristic_count": 0} for sku in skus}
    cur.execute(
        """SELECT ps.sku, ps.product_name, i.search_intent
           FROM public.seo_monitoring_project_skus ps
           LEFT JOIN public.seo_generation_product_intents i
             ON i.project_id=ps.project_id AND i.sku=ps.sku
           WHERE ps.project_id=%s AND ps.sku=ANY(%s)""", (project_id, skus),
    )
    for row in cur.fetchall():
        contexts[str(row["sku"])].update({"title": row.get("product_name") or "", "intent": row.get("search_intent") or ""})
    if marketplace == "wb":
        if not _relation_exists(cur, "products"):
            return contexts
        cur.execute(
            """SELECT product_id, artikul_wb::text AS sku, naimenovanie, opisanie,
                      seller_category_name, brend, pol, razmer, ros_razmer, sostav,
                      tsvet, kollektsiya, dekorativnye_elementy, uhod_za_veschami,
                      vozrastnye_ogranicheniya, artikul_prodavtsa
               FROM public.products WHERE artikul_wb::text=ANY(%s)""", (skus,),
        )
        product_ids = {}
        direct_attributes = {
            "brend": "Бренд", "pol": "Пол", "razmer": "Размер производителя",
            "ros_razmer": "Российский размер", "sostav": "Состав",
            "tsvet": "Цвет", "kollektsiya": "Коллекция",
            "dekorativnye_elementy": "Декоративные элементы",
            "uhod_za_veschami": "Уход", "vozrastnye_ogranicheniya": "Возрастные ограничения",
            "artikul_prodavtsa": "Модель",
        }
        for row in cur.fetchall():
            sku = str(row["sku"])
            product_ids[row["product_id"]] = sku
            contexts[sku].update({
                "title": row.get("naimenovanie") or contexts[sku].get("title") or "",
                "description": _clean_text(row.get("opisanie"), 5000),
                "category": _clean_text(row.get("seller_category_name"), 300),
                "product_type": _clean_text(row.get("seller_category_name"), 200),
            })
            for column, name in direct_attributes.items():
                item = _semantic_context_attribute(name, row.get(column))
                if item:
                    contexts[sku]["all_characteristic_count"] += 1
                    if item not in contexts[sku]["seo_characteristics"]:
                        contexts[sku]["seo_characteristics"].append(item)
        if product_ids and _relation_exists(cur, "product_attributes") and _relation_exists(cur, "category_attributes"):
            cur.execute(
                """SELECT pa.product_id, ca.attribute_name AS name, pa.value_text AS value
                   FROM public.product_attributes pa
                   LEFT JOIN public.products p ON p.product_id=pa.product_id
                   LEFT JOIN public.category_attributes ca
                     ON ca.category_id=p.category_id AND ca.attribute_id=pa.attribute_id
                   WHERE pa.product_id=ANY(%s) AND pa.value_text IS NOT NULL
                   ORDER BY pa.product_id, lower(ca.attribute_name)""", (list(product_ids),),
            )
            for row in cur.fetchall():
                sku = product_ids.get(row.get("product_id"))
                if not sku:
                    continue
                contexts[sku]["all_characteristic_count"] += 1
                item = _semantic_context_attribute(row.get("name"), row.get("value"))
                if item and item not in contexts[sku]["seo_characteristics"] and len(contexts[sku]["seo_characteristics"]) < SEO_CONTEXT_ATTRIBUTE_LIMIT:
                    contexts[sku]["seo_characteristics"].append(item)
        return contexts
    if not _relation_exists(cur, "ozon_cat_products"):
        return contexts
    cur.execute(
        """SELECT DISTINCT ON (sku) sku::text AS sku, product_id, category_id,
                  nazvanie_tovara, annotatsiya, category_name, tip
           FROM public.ozon_cat_products WHERE sku::text=ANY(%s)
           ORDER BY sku, updated_at DESC NULLS LAST, imported_at DESC NULLS LAST, product_id DESC""", (skus,),
    )
    product_ids = {}
    for row in cur.fetchall():
        sku = str(row["sku"])
        product_ids[row["product_id"]] = sku
        contexts[sku].update({
            "title": row.get("nazvanie_tovara") or contexts[sku].get("title") or "",
            "description": _clean_text(row.get("annotatsiya"), 4000),
            "category": _clean_text(row.get("category_name"), 300),
            "product_type": _clean_text(row.get("tip"), 200),
        })
    if not product_ids:
        return contexts
    cur.execute(
        """SELECT pa.product_id, coalesce(ca.attribute_name, co.attribute_name) AS name, pa.value_text AS value
           FROM public.ozon_cat_product_attributes pa
           LEFT JOIN public.ozon_cat_products p ON p.product_id=pa.product_id
           LEFT JOIN public.ozon_cat_category_attributes ca
             ON ca.category_id=p.category_id AND ca.attribute_id=pa.attribute_id
           LEFT JOIN public.ozon_cat_common_attributes co ON co.attribute_id=pa.attribute_id
           WHERE pa.product_id=ANY(%s) AND pa.value_text IS NOT NULL
           ORDER BY pa.product_id, lower(coalesce(ca.attribute_name, co.attribute_name))""", (list(product_ids),),
    )
    for row in cur.fetchall():
        sku = product_ids.get(row.get("product_id"))
        if not sku:
            continue
        contexts[sku]["all_characteristic_count"] += 1
        item = _semantic_context_attribute(row.get("name"), row.get("value"))
        if item and item not in contexts[sku]["seo_characteristics"] and len(contexts[sku]["seo_characteristics"]) < SEO_CONTEXT_ATTRIBUTE_LIMIT:
            contexts[sku]["seo_characteristics"].append(item)
    return contexts


def prepare_project_semantic_context(config, payload):
    """Build and persist a deterministic, auditable input package for later AI semantic-core generation."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = list(dict.fromkeys(
        _clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)
    ))
    if not requested_skus or len(requested_skus) > 20:
        raise ValueError("Передайте от 1 до 20 SKU проекта")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            marketplace = project.get("marketplace")
            cur.execute("SELECT sku FROM public.seo_monitoring_project_skus WHERE project_id=%s AND sku=ANY(%s)", (project_id, requested_skus))
            found = {str(row["sku"]) for row in cur.fetchall()}
            missing = [sku for sku in requested_skus if sku not in found]
            if missing:
                raise ValueError(f"SKU не входят в проект: {', '.join(missing[:5])}")
            products = _semantic_product_contexts(cur, project_id, requested_skus, marketplace)
            cur.execute(
                """WITH metrics AS (
                     SELECT sku, search_query,
                            CASE WHEN source LIKE 'mpstats_%%' THEN 'mpstats' ELSE 'api' END AS source_group,
                            max(search_demand) AS search_demand, sum(traffic) AS traffic,
                            sum(orders) AS orders, min(nullif(average_position,0)) AS average_position
                     FROM public.seo_monitoring_keyword_snapshots
                     WHERE project_id=%s AND sku=ANY(%s)
                     GROUP BY sku, search_query, CASE WHEN source LIKE 'mpstats_%%' THEN 'mpstats' ELSE 'api' END)
                   SELECT a.*, m.search_demand, m.traffic, m.orders, m.average_position
                   FROM public.seo_generation_keyword_analysis a
                   LEFT JOIN metrics m ON m.sku=a.sku AND m.search_query=a.search_query AND m.source_group=a.source_group
                   WHERE a.project_id=%s AND a.sku=ANY(%s)""",
                (project_id, requested_skus, project_id, requested_skus),
            )
            own_by_sku = {sku: [] for sku in requested_skus}
            for row in cur.fetchall(): own_by_sku[str(row["sku"])].append(dict(row))
            cur.execute(
                """SELECT a.*, k.search_demand, k.best_position, k.competitor_coverage
                   FROM public.seo_generation_competitor_keyword_analysis a
                   LEFT JOIN (
                     SELECT project_id, sku, search_query, max(search_demand) AS search_demand,
                            min(nullif(average_position,0)) AS best_position,
                            count(DISTINCT competitor_sku)::int AS competitor_coverage
                     FROM public.seo_generation_competitor_keywords
                     WHERE project_id=%s AND sku=ANY(%s)
                     GROUP BY project_id, sku, search_query
                   ) k ON k.project_id=a.project_id AND k.sku=a.sku AND k.search_query=a.search_query
                   WHERE a.project_id=%s AND a.sku=ANY(%s)""",
                (project_id, requested_skus, project_id, requested_skus),
            )
            competitor_keys = {sku: [] for sku in requested_skus}
            for row in cur.fetchall(): competitor_keys[str(row["sku"])].append(dict(row))
            cur.execute("SELECT * FROM public.seo_generation_customer_voice_claims WHERE project_id=%s AND sku=ANY(%s)", (project_id, requested_skus))
            claims_by_sku = {sku: [] for sku in requested_skus}
            for row in cur.fetchall(): claims_by_sku[str(row["sku"])].append(dict(row))
            cur.execute("""SELECT sku, competitor_sku, competitor_rank, product_name, brand
                           FROM public.seo_generation_niche_competitors
                           WHERE project_id=%s AND sku=ANY(%s) ORDER BY sku, competitor_rank""", (project_id, requested_skus))
            competitors_by_sku = {sku: [] for sku in requested_skus}
            for row in cur.fetchall(): competitors_by_sku[str(row["sku"])].append(dict(row))
            stored = []
            for sku in requested_skus:
                context, counts, context_hash, status = _build_semantic_context(
                    sku, products.get(sku) or {}, own_by_sku[sku], competitor_keys[sku], claims_by_sku[sku],
                    competitors_by_sku[sku], marketplace
                )
                stored.append((project_id, sku, json.dumps(context, ensure_ascii=False, default=str), context_hash,
                               json.dumps(counts, ensure_ascii=False), status, None))
            execute_values(cur, """INSERT INTO public.seo_generation_semantic_contexts
                (project_id, sku, context_json, context_hash, source_counts, status, last_error) VALUES %s
                ON CONFLICT(project_id, sku) DO UPDATE SET context_json=excluded.context_json,
                  context_hash=excluded.context_hash, source_counts=excluded.source_counts,
                  status=excluded.status, last_error=excluded.last_error, prepared_at=now()""", stored)
    partial = sum(value[5] == "partial" for value in stored)
    return {"ok": True, "project_id": project_id, "prepared": len(stored), "partial": partial,
            "complete": len(stored) - partial, "errors": []}


def project_semantic_context(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute("""SELECT context_json, context_hash, source_counts, status, last_error, prepared_at
                           FROM public.seo_generation_semantic_contexts WHERE project_id=%s AND sku=%s""", (project_id, sku))
            row = cur.fetchone()
    if not row:
        return {"ok": True, "project_id": project_id, "sku": sku, "available": False, "status": None}
    result = _serialize(row)
    return {"ok": True, "project_id": project_id, "sku": sku, "available": True, **result}


def _allocation_candidate_maps(context):
    keywords = {}
    for group in ("own_keywords", "competitor_keywords"):
        for row in context.get(group) or []:
            query = _clean_text(row.get("query"), 300)
            normalized = _seo_normalize(query)
            if not normalized:
                continue
            current = keywords.setdefault(normalized, {
                "query": query, "sources": [], "priority_score": 0,
                "search_demand": None, "best_position": None,
            })
            current["sources"] = sorted(set(current["sources"] + list(row.get("sources") or []) +
                                               (["competitors"] if group == "competitor_keywords" else [])))
            current["priority_score"] = max(current["priority_score"], int(_number(row.get("priority_score")) or 0))
            current["search_demand"] = max(
                [value for value in (current.get("search_demand"), _number(row.get("search_demand"))) if value is not None],
                default=None,
            )
            positions = [value for value in (current.get("best_position"), _number(row.get("best_position"))) if value is not None]
            current["best_position"] = min(positions) if positions else None
    product = context.get("product") or {}
    main_intent = _clean_text(product.get("intent"), 300)
    normalized_intent = _seo_normalize(main_intent)
    if normalized_intent:
        intent_candidate = keywords.setdefault(normalized_intent, {
            "query": main_intent,
            "sources": ["product_intent"],
            "priority_score": 100,
            "search_demand": None,
            "best_position": None,
        })
        intent_candidate["sources"] = sorted(set(list(intent_candidate.get("sources") or []) + ["product_intent"]))
        intent_candidate["priority_score"] = max(int(intent_candidate.get("priority_score") or 0), 100)
    claims = {}
    for row in context.get("customer_voice_claims") or []:
        claim = _clean_text(row.get("claim"), 500)
        normalized = _seo_normalize(claim)
        if normalized:
            claims[normalized] = {**row, "claim": claim}
    characteristics = {}
    for row in product.get("seo_characteristics") or []:
        name = _clean_text(row.get("name"), 300)
        value = _clean_text(row.get("value"), 1000)
        key = (_seo_normalize(name), _seo_normalize(value))
        if key[0] and key[1]:
            characteristics[key] = {"name": name, "value": value}
    return keywords, claims, characteristics


def _validate_content_allocation(context, raw_allocation):
    """Validate AI selection against the deterministic input; invented facts never reach persistence."""
    if not isinstance(raw_allocation, dict):
        raise ValueError("AI вернул SEO-разметку неподдерживаемого типа")
    candidate_keywords, candidate_claims, candidate_characteristics = _allocation_candidate_maps(context)
    selected_keywords, seen = [], set()
    errors = []
    for placement, field in (("title", "title_keywords"), ("description", "description_keywords")):
        rows = raw_allocation.get(field) or []
        if not isinstance(rows, list):
            errors.append(f"{field}: ожидается список")
            continue
        limit = 8 if placement == "title" else 20
        for raw in rows[:limit]:
            raw = raw if isinstance(raw, dict) else {"query": raw}
            normalized = _seo_normalize(raw.get("query"))
            candidate = candidate_keywords.get(normalized)
            if not candidate:
                errors.append(f"Ключ отсутствует в SEO-контексте: {_clean_text(raw.get('query'), 120)}")
                continue
            if normalized in seen:
                continue
            role = str(raw.get("role") or ("primary" if not selected_keywords else "secondary")).lower()
            if role not in {"primary", "secondary", "supporting"}:
                role = "secondary"
            selected_keywords.append({
                **candidate, "placement": placement, "role": role,
                "reason": _clean_text(raw.get("reason"), 500), "monitor": True,
            })
            seen.add(normalized)

    # The product intent is the only exact phrase reserved for the title.  AI may
    # rank/select supporting queries, but it cannot promote several exact phrases
    # into the title or omit the canonical intent altogether.
    main_intent = _clean_text((context.get("product") or {}).get("intent"), 300)
    normalized_intent = _seo_normalize(main_intent)
    intent_candidate = candidate_keywords.get(normalized_intent) if normalized_intent else None
    if intent_candidate:
        intent_row = next(
            (row for row in selected_keywords if _seo_normalize(row.get("query")) == normalized_intent),
            None,
        )
        if intent_row is None:
            intent_row = {
                **intent_candidate,
                "placement": "title",
                "role": "primary",
                "reason": "Канонический основной интент товара",
                "monitor": True,
            }
            selected_keywords.insert(0, intent_row)
            seen.add(normalized_intent)
        for row in selected_keywords:
            is_intent = _seo_normalize(row.get("query")) == normalized_intent
            row["placement"] = "title" if is_intent else "description"
            row["exact_required"] = bool(is_intent)
            if is_intent:
                row["role"] = "primary"
            elif row.get("role") == "primary":
                row["role"] = "secondary"
    else:
        title_rows = [row for row in selected_keywords if row.get("placement") == "title"]
        if title_rows:
            title_row = max(title_rows, key=lambda row: int(row.get("priority_score") or 0))
            for row in selected_keywords:
                is_title = row is title_row
                row["placement"] = "title" if is_title else "description"
                row["exact_required"] = bool(is_title)
    selected_claims = []
    for raw in (raw_allocation.get("description_claims") or [])[:12]:
        raw = raw if isinstance(raw, dict) else {"claim": raw}
        normalized = _seo_normalize(raw.get("claim"))
        candidate = candidate_claims.get(normalized)
        if not candidate:
            errors.append(f"Клейм отсутствует в SEO-контексте: {_clean_text(raw.get('claim'), 120)}")
            continue
        selected_claims.append({
            **candidate,
            "reason": _clean_text(raw.get("reason"), 500) or _clean_text(candidate.get("rationale"), 500),
        })
    selected_characteristics, characteristic_seen = [], set()
    for placement, field in (("title", "title_characteristics"), ("description", "description_characteristics")):
        rows = raw_allocation.get(field) or []
        if not isinstance(rows, list):
            errors.append(f"{field}: ожидается список")
            continue
        limit = 6 if placement == "title" else 12
        for raw in rows[:limit]:
            if not isinstance(raw, dict):
                errors.append(f"{field}: характеристика должна содержать name и value")
                continue
            key = (_seo_normalize(raw.get("name")), _seo_normalize(raw.get("value")))
            candidate = candidate_characteristics.get(key)
            if not candidate:
                errors.append(
                    "Характеристика отсутствует в SEO-контексте: "
                    f"{_clean_text(raw.get('name'), 80)} = {_clean_text(raw.get('value'), 120)}"
                )
                continue
            if key in characteristic_seen:
                continue
            selected_characteristics.append({
                **candidate, "placement": placement,
                "reason": _clean_text(raw.get("reason"), 500),
            })
            characteristic_seen.add(key)
    title_outline = _clean_text(raw_allocation.get("title_outline"), 200)
    description_outline = _clean_text(raw_allocation.get("description_outline"), 1200)
    if not selected_keywords:
        raise ValueError("AI не выбрал ни одного проверяемого ключа")
    if candidate_characteristics and not selected_characteristics:
        raise ValueError("AI не выбрал ни одной проверяемой SEO-характеристики")
    warnings = list(context.get("warnings") or [])
    warnings.extend(errors)
    return {
        "allocation_version": SEO_CONTENT_ALLOCATION_PROMPT_VERSION,
        "rules_version": SEO_CONTENT_ALLOCATION_RULES_VERSION,
        "sku": str(context.get("sku") or ""),
        "title_outline": title_outline,
        "description_outline": description_outline,
        "title_keywords": [row for row in selected_keywords if row["placement"] == "title"],
        "description_keywords": [row for row in selected_keywords if row["placement"] == "description"],
        "description_claims": selected_claims,
        "title_characteristics": [row for row in selected_characteristics if row["placement"] == "title"],
        "description_characteristics": [row for row in selected_characteristics if row["placement"] == "description"],
        "characteristics_available": len(candidate_characteristics),
        "characteristics_selected": len(selected_characteristics),
        "excluded": [
            {"text": _clean_text(row.get("text"), 300), "reason": _clean_text(row.get("reason"), 500)}
            for row in (raw_allocation.get("excluded") or [])[:30] if isinstance(row, dict) and _clean_text(row.get("text"), 300)
        ],
        "monitoring_keywords": selected_keywords,
        "warnings": list(dict.fromkeys(value for value in warnings if value)),
    }


def generate_project_content_allocation(config, payload, analyze_one):
    """Create a fact-checked field allocation and atomically replace its monitoring marks."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            marketplace = project.get("marketplace")
            rules_version, prompt_version = _content_stage_versions(marketplace, "allocation")
            cur.execute("""SELECT context_json, context_hash, status FROM public.seo_generation_semantic_contexts
                           WHERE project_id=%s AND sku=%s""", (project_id, sku))
            context_row = cur.fetchone()
            if not context_row:
                raise ValueError("Сначала подготовьте SEO-контекст для SKU")
            context = context_row.get("context_json") or {}
            settings = _ai_script_models(conn, "content_allocation")
        ai_payload = analyze_one(context, settings)
        raw = ai_payload.get("allocation") if isinstance(ai_payload, dict) else None
        allocation = _validate_content_allocation(context, raw)
        allocation["rules_version"] = rules_version
        allocation["allocation_version"] = prompt_version
        model = _clean_text((ai_payload or {}).get("model"), 200)
        status = "partial" if allocation["warnings"] else "ok"
        with conn.cursor() as cur:
            cur.execute("""INSERT INTO public.seo_generation_content_allocations
                (project_id, sku, allocation_json, context_hash, rules_version, prompt_version, status, last_error, model)
                VALUES (%s,%s,%s::jsonb,%s,%s,%s,%s,NULL,%s)
                ON CONFLICT(project_id, sku) DO UPDATE SET allocation_json=excluded.allocation_json,
                  context_hash=excluded.context_hash, rules_version=excluded.rules_version,
                  prompt_version=excluded.prompt_version, status=excluded.status, last_error=NULL,
                  model=excluded.model, generated_at=now()""",
                (project_id, sku, json.dumps(allocation, ensure_ascii=False, default=str), context_row["context_hash"],
                 rules_version, prompt_version, status, model))
            cur.execute("UPDATE public.seo_generation_monitoring_keywords SET active=false WHERE project_id=%s AND sku=%s", (project_id, sku))
            rows = [(project_id, sku, _seo_normalize(item["query"]), item["query"], item["placement"], item["role"],
                     json.dumps(item.get("sources") or [], ensure_ascii=False), int(item.get("priority_score") or 0),
                     item.get("best_position"), item.get("search_demand"), True) for item in allocation["monitoring_keywords"]]
            if rows:
                execute_values(cur, """INSERT INTO public.seo_generation_monitoring_keywords
                    (project_id, sku, normalized_query, search_query, placement, role, source_groups,
                     priority_score, baseline_position, baseline_search_demand, active) VALUES %s
                    ON CONFLICT(project_id, sku, normalized_query) DO UPDATE SET search_query=excluded.search_query,
                      placement=excluded.placement, role=excluded.role, source_groups=excluded.source_groups,
                      priority_score=excluded.priority_score, baseline_position=excluded.baseline_position,
                      baseline_search_demand=excluded.baseline_search_demand, active=true, selected_at=now()""", rows)
        conn.commit()
    return {"ok": True, "project_id": project_id, "sku": sku, "status": status,
            "model": model, "selected": len(allocation["monitoring_keywords"]), "allocation": allocation}


def project_content_allocation(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute("""SELECT allocation_json, context_hash, rules_version, prompt_version, status,
                                  last_error, model, generated_at
                           FROM public.seo_generation_content_allocations WHERE project_id=%s AND sku=%s""",
                        (project_id, sku))
            row = cur.fetchone()
            cur.execute("""SELECT search_query, placement, role, source_groups, priority_score,
                                  baseline_position, baseline_search_demand, active, selected_at
                           FROM public.seo_generation_monitoring_keywords
                           WHERE project_id=%s AND sku=%s AND active=true
                           ORDER BY placement, priority_score DESC, search_query""", (project_id, sku))
            monitoring = [_serialize(item) for item in cur.fetchall()]
    if not row:
        return {"ok": True, "project_id": project_id, "sku": sku, "available": False, "monitoring_keywords": []}
    return {"ok": True, "project_id": project_id, "sku": sku, "available": True,
            **_serialize(row), "monitoring_keywords": monitoring}


def _allocation_hash(allocation):
    stable = json.dumps(allocation or {}, ensure_ascii=False, sort_keys=True, default=str, separators=(",", ":"))
    return hashlib.sha256(stable.encode("utf-8")).hexdigest()


def _keyword_priority(row):
    role_rank = {"primary": 3, "secondary": 2, "supporting": 1}
    try:
        priority_score = float(row.get("priority_score") or 0)
    except (TypeError, ValueError):
        priority_score = 0.0
    try:
        search_demand = float(row.get("search_demand") or row.get("baseline_search_demand") or 0)
    except (TypeError, ValueError):
        search_demand = 0.0
    return (-role_rank.get(str(row.get("role") or "").lower(), 0), -priority_score, -search_demand,
            _seo_normalize(row.get("query")))


def _keyword_prefix_check(keyword, text, prefix_chars, *, require_near_start=False):
    query = _clean_text((keyword or {}).get("query"), 300)
    query_tokens = list(dict.fromkeys(token for token in _seo_normalize(query).split() if len(token) > 1))
    prefix_tokens = _seo_normalize(_clean_text(text, prefix_chars)[:prefix_chars]).split()
    prefix_set = set(prefix_tokens)
    matched = [token for token in query_tokens if token in prefix_set]
    coverage = round(len(matched) / len(query_tokens), 3) if query_tokens else 0.0
    positions = [prefix_tokens.index(token) for token in matched if token in prefix_set]
    first_position = min(positions) if positions else None
    passed = coverage >= 0.8 and (not require_near_start or (first_position is not None and first_position <= 3))
    return {
        "query": query,
        "role": keyword.get("role") or "",
        "prefix_chars": prefix_chars,
        "token_coverage": coverage,
        "matched_tokens": matched,
        "first_token_position": first_position,
        "passed": passed,
    }


def _exact_phrase_count(text, query):
    """Count a selected search phrase as written, ignoring only case and whitespace."""
    parts = [part for part in re.split(r"\s+", _clean_text(query, 300)) if part]
    if not parts:
        return 0
    pattern = r"(?<!\w)" + r"\s+".join(re.escape(part) for part in parts) + r"(?!\w)"
    return len(re.findall(pattern, _clean_text(text, 5000), flags=re.IGNORECASE))


SEO_KEYWORD_STOP_WORDS = {
    "а", "без", "в", "во", "для", "до", "и", "из", "к", "ко", "на", "не", "о", "об", "от",
    "по", "под", "при", "про", "с", "со", "у", "или", "как", "что", "это", "этот", "эта", "эти",
}


def _seo_russian_stem(value):
    """Small deterministic Russian Porter stemmer for audit, not linguistic generation."""
    word = _seo_normalize(value).replace("ё", "е")
    if len(word) <= 3:
        return word
    match = re.search(r"[аеиоуыэюя]", word)
    if not match:
        return word
    start = match.end()
    prefix, rv = word[:start], word[start:]

    def strip(pattern):
        nonlocal rv
        changed = re.sub(pattern + r"$", "", rv, count=1)
        if changed == rv:
            return False
        rv = changed
        return True

    if not strip(r"(?:ив|ивши|ившись|ыв|ывши|ывшись)"):
        strip(r"(?<=[ая])(?:в|вши|вшись)")
    strip(r"(?:ся|сь)")
    if not strip(
        r"(?:ее|ие|ые|ое|ими|ыми|ей|ий|ый|ой|ем|им|ым|ом|его|ого|ему|ому|их|ых|ую|юю|ая|яя|ою|ею)"
    ):
        if not strip(
            r"(?:ила|ыла|ена|ейте|уйте|ите|или|ыли|ей|уй|ил|ыл|им|ым|ен|ило|ыло|ено|ят|ует|уют|ит|ыт|ены|ить|ыть|ишь|ую|ю)"
        ):
            strip(r"(?:а|ев|ов|ие|ье|е|иями|ями|ами|еи|ии|и|ией|ей|ой|ий|й|иям|ям|ием|ем|ам|ом|о|у|ах|иях|ях|ы|ь|ию|ью|ю|ия|ья|я)")
    strip(r"и")
    strip(r"(?:ейше|ейш)")
    rv = re.sub(r"нн$", "н", rv)
    strip(r"ь") or strip(r"ей") or strip(r"н")
    return prefix + rv


def _seo_keyword_units(query):
    tokens = re.findall(r"[0-9A-Za-zА-Яа-яЁё]+", _seo_normalize(query))
    content_tokens = [token for token in tokens if token not in SEO_KEYWORD_STOP_WORDS and len(token) > 1]
    bases = []
    for token in content_tokens:
        base = _seo_russian_stem(token)
        if base and base not in bases:
            bases.append(base)
    links = [[bases[index], bases[index + 1]] for index in range(len(bases) - 1)]
    return {"tokens": content_tokens, "bases": bases, "links": links}


def _seo_field_morphology(text, query):
    units = _seo_keyword_units(query)
    field_tokens = re.findall(r"[0-9A-Za-zА-Яа-яЁё]+", _seo_normalize(text).replace("_", " "))
    field_bases = [_seo_russian_stem(token) for token in field_tokens]
    positions = {}
    forms = {}
    for index, (token, base) in enumerate(zip(field_tokens, field_bases)):
        positions.setdefault(base, []).append(index)
        forms.setdefault(base, [])
        if token not in forms[base]:
            forms[base].append(token)
    matched = [base for base in units["bases"] if base in positions]
    matched_links = []
    for left, right in units["links"]:
        if any(0 < right_pos - left_pos <= 4 for left_pos in positions.get(left, []) for right_pos in positions.get(right, [])):
            matched_links.append([left, right])
    base_coverage = round(len(matched) / len(units["bases"]), 3) if units["bases"] else 0.0
    link_coverage = round(len(matched_links) / len(units["links"]), 3) if units["links"] else (1.0 if matched else 0.0)
    covered = bool(units["bases"]) and base_coverage == 1.0 and link_coverage >= 0.5
    return {
        "exact_occurrences": _exact_phrase_count(text, query),
        "matched_bases": matched,
        "matched_forms": {base: forms.get(base, []) for base in matched},
        "base_coverage": base_coverage,
        "matched_links": matched_links,
        "link_coverage": link_coverage,
        "covered": covered,
    }


def _build_keyword_strategy(keyword_candidates, main_intent=""):
    ordered = sorted(keyword_candidates.values(), key=_keyword_priority)
    rows = []
    normalized_intent = _seo_normalize(main_intent)
    if normalized_intent:
        intent_units = _seo_keyword_units(main_intent)
        intent_candidate = next(
            (row for row in ordered if _seo_normalize(row.get("query")) == normalized_intent),
            {},
        )
        rows.append({
            "query": _clean_text(main_intent, 300),
            "placement": "title",
            "role": "primary",
            "priority_score": intent_candidate.get("priority_score"),
            "source_type": "product_intent",
            "exact_required": True,
            "recommended_field": "title",
            "tokens": intent_units["tokens"],
            "morphological_bases": intent_units["bases"],
            "semantic_links": intent_units["links"],
        })
    for row in ordered:
        if normalized_intent and _seo_normalize(row.get("query")) == normalized_intent:
            continue
        units = _seo_keyword_units(row.get("query"))
        exact_required = not normalized_intent and not rows
        recommended_field = "title" if exact_required else "description"
        rows.append({
            "query": row.get("query") or "",
            "placement": row.get("placement") or "description",
            "role": row.get("role") or "",
            "priority_score": row.get("priority_score"),
            "source_type": "search_query",
            "exact_required": exact_required,
            "recommended_field": recommended_field,
            "tokens": units["tokens"],
            "morphological_bases": units["bases"],
            "semantic_links": units["links"],
        })
    return rows


def _audit_keyword_strategy(keyword_candidates, title, description, hashtags, main_intent=""):
    strategy = _build_keyword_strategy(keyword_candidates, main_intent)
    fields = {"title": title, "description": description, "hashtags": " ".join(hashtags or [])}
    rows, all_bases, used_bases = [], set(), set()
    for row in strategy:
        field_matches = {name: _seo_field_morphology(text, row["query"]) for name, text in fields.items()}
        exact_occurrences = sum(item["exact_occurrences"] for item in field_matches.values())
        matched_bases = []
        matched_forms = {}
        locations = []
        morph_covered = False
        for name, match in field_matches.items():
            if match["exact_occurrences"]:
                locations.append({"field": name, "type": "exact", "count": match["exact_occurrences"]})
            elif match["covered"]:
                locations.append({"field": name, "type": "morphological", "count": 1})
            morph_covered = morph_covered or match["covered"]
            for base in match["matched_bases"]:
                if base not in matched_bases:
                    matched_bases.append(base)
                matched_forms.setdefault(base, [])
                for form in match["matched_forms"].get(base, []):
                    if form not in matched_forms[base]:
                        matched_forms[base].append(form)
        if not row["exact_required"]:
            all_bases.update(row["morphological_bases"])
            used_bases.update(matched_bases)
        base_coverage = round(len(matched_bases) / len(row["morphological_bases"]), 3) if row["morphological_bases"] else 0.0
        coverage_type = "exact" if exact_occurrences else ("morphological" if morph_covered else ("partial" if matched_bases else "absent"))
        passed = exact_occurrences > 0 if row["exact_required"] else coverage_type in {"exact", "morphological"}
        rows.append({
            **row,
            "coverage_type": coverage_type,
            "passed": passed,
            "exact_occurrences": exact_occurrences,
            "matched_bases": matched_bases,
            "unmatched_bases": [base for base in row["morphological_bases"] if base not in matched_bases],
            "matched_forms": matched_forms,
            "base_coverage": base_coverage,
            "locations": locations,
            "field_matches": field_matches,
        })
    exact_rows = [row for row in rows if row["exact_required"]]
    support_rows = [row for row in rows if not row["exact_required"]]
    summary = {
        "main_intent": next((row["query"] for row in exact_rows), ""),
        "exact_required": len(exact_rows),
        "exact_passed": sum(1 for row in exact_rows if row["passed"]),
        "support_queries": len(support_rows),
        "support_covered": sum(1 for row in support_rows if row["passed"]),
        "unique_bases": len(all_bases),
        "unique_bases_used": len(all_bases & used_bases),
        "base_coverage": round(len(all_bases & used_bases) / len(all_bases), 3) if all_bases else 0.0,
    }
    summary["passed"] = summary["exact_passed"] == summary["exact_required"] and summary["base_coverage"] >= 0.8
    return rows, summary


def _hashtag_value(source):
    value = _clean_text(source, 120).lstrip("#").replace(" ", "_")
    value = re.sub(r"[^0-9A-Za-zА-Яа-яЁё_-]+", "", value).strip("_-")
    return f"#{value}" if value else ""


def _build_quality_hashtags(context, used_keywords, raw_hashtags):
    product = (context or {}).get("product") or {}
    product_type_tokens = set(_seo_normalize(product.get("product_type")).split())
    intent_tokens = _seo_normalize(product.get("intent")).split()
    # Product type from marketplace cards may be only a broad catalogue label
    # (for example, "Верхняя одежда").  In that case use the canonical intent
    # as a second source-backed identity signal.  A concrete product type stays
    # authoritative so unrelated queries cannot enter merely by sharing an
    # adjective with the intent.
    generic_type_tokens = {"товар", "товары", "одежда", "верхняя", "аксессуары", "аксессуар"}
    concrete_type_tokens = product_type_tokens - generic_type_tokens
    identity_tokens = product_type_tokens if concrete_type_tokens else product_type_tokens | set(intent_tokens[:2])
    details = []
    seen = set()
    token_sets = []

    def add(source, source_type, *, query="", role="", priority_score=None):
        hashtag = _hashtag_value(source)
        normalized_tokens = tuple(_seo_normalize(hashtag.lstrip("#").replace("_", " ")).split())
        token_set = set(normalized_tokens)
        if not hashtag or len(token_set) < 2 or len(token_set) > 6:
            return
        if identity_tokens and not token_set.intersection(identity_tokens):
            return
        key = hashtag.lower()
        if key in seen:
            return
        if any(len(token_set & prior) / max(len(token_set | prior), 1) >= 0.85 for prior in token_sets):
            return
        seen.add(key)
        token_sets.append(token_set)
        details.append({
            "hashtag": hashtag,
            "source_type": source_type,
            "source_query": query,
            "role": role,
            "priority_score": priority_score,
        })

    for row in sorted(used_keywords, key=_keyword_priority):
        add(row.get("query"), "keyword", query=row.get("query") or "", role=row.get("role") or "",
            priority_score=row.get("priority_score"))
        if len(details) >= 8:
            break

    return [row["hashtag"] for row in details], details


def _validate_content_draft(context, allocation, raw_draft):
    """Accept only a compact draft whose evidence references exist in the saved allocation."""
    if not isinstance(raw_draft, dict):
        raise ValueError("AI вернул SEO-текст неподдерживаемого типа")
    title_limit = 60 if str(context.get("marketplace") or "").lower() == "wb" else 200
    title = _clean_text(raw_draft.get("title"), 500)
    raw_sections = raw_draft.get("description_sections") or []
    description_sections = []
    seen_section_kinds = set()
    if raw_sections and not isinstance(raw_sections, list):
        raise ValueError("description_sections должен быть массивом смысловых блоков")
    for raw_section in raw_sections[:6]:
        if not isinstance(raw_section, dict):
            raise ValueError("Каждый смысловой блок описания должен быть объектом")
        kind = _clean_text(raw_section.get("kind"), 80)
        if kind not in SEO_CONTENT_DESCRIPTION_SECTION_KINDS:
            raise ValueError(f"Неизвестный тип смыслового блока описания: {kind or 'пусто'}")
        if kind in seen_section_kinds:
            raise ValueError(f"Смысловой блок описания продублирован: {kind}")
        heading = _clean_text(raw_section.get("heading"), 100)
        text = _clean_text(raw_section.get("text"), 1400)
        if not text:
            continue
        description_sections.append({
            "kind": kind,
            "label": SEO_CONTENT_DESCRIPTION_SECTION_KINDS[kind],
            "heading": heading,
            "text": text,
            "char_count": len(text),
        })
        seen_section_kinds.add(kind)
    if description_sections:
        # Section headings are audit metadata. They must never leak into the customer-facing copy.
        description = "\n\n".join(row["text"] for row in description_sections)
        description = _clean_text(description, 4000)
    else:
        description = _clean_text(raw_draft.get("description"), 4000)
    if not title:
        raise ValueError("AI не сформировал название")
    if not description:
        raise ValueError("AI не сформировал описание")
    title, description, required_model, required_warranties = _enforce_source_preservation(
        title, description, context
    )
    if len(title) > title_limit:
        raise ValueError(f"Название превышает лимит {title_limit} символов для выбранного маркетплейса")

    keyword_candidates = {}
    for placement, field in (("title", "title_keywords"), ("description", "description_keywords")):
        for row in allocation.get(field) or []:
            query = _clean_text(row.get("query"), 300)
            normalized = _seo_normalize(query)
            if normalized:
                keyword_candidates[normalized] = {**row, "query": query, "placement": placement}
    claim_candidates = {}
    for row in allocation.get("description_claims") or []:
        claim = _clean_text(row.get("claim"), 500)
        normalized = _seo_normalize(claim)
        if normalized:
            claim_candidates[normalized] = {**row, "claim": claim}
    for claim in required_warranties:
        normalized = _seo_normalize(claim)
        if normalized:
            claim_candidates[normalized] = {
                "claim": claim,
                "verification_status": "source",
                "confidence": 1.0,
                "source": "original_description",
            }
    characteristic_candidates = {}
    for placement, field in (("title", "title_characteristics"), ("description", "description_characteristics")):
        for row in allocation.get(field) or []:
            key = (_seo_normalize(row.get("name")), _seo_normalize(row.get("value")))
            if key[0] and key[1]:
                characteristic_candidates[key] = {
                    **row, "name": _clean_text(row.get("name"), 300),
                    "value": _clean_text(row.get("value"), 1000), "placement": placement,
                }

    used_keywords, seen_keywords = [], set()
    declared_but_absent_keywords = []
    combined_copy = f"{title}\n{description}"
    declared_keyword_norms = set()
    for raw in raw_draft.get("used_keywords") or []:
        raw = raw if isinstance(raw, dict) else {"query": raw}
        normalized = _seo_normalize(raw.get("query"))
        candidate = keyword_candidates.get(normalized)
        if not candidate:
            raise ValueError(f"SEO-текст ссылается на невыбранный ключ: {_clean_text(raw.get('query'), 120)}")
        declared_keyword_norms.add(normalized)

    # The model declaration is advisory. The canonical audit is derived from the saved copy itself.
    for normalized, candidate in keyword_candidates.items():
        occurrence_count = _exact_phrase_count(combined_copy, candidate.get("query"))
        if occurrence_count > 0 and normalized not in seen_keywords:
            candidate = {**candidate, "exact_occurrences": occurrence_count}
            used_keywords.append(candidate)
            seen_keywords.add(normalized)
    main_intent = _clean_text(((context or {}).get("product") or {}).get("intent"), 300)
    keyword_strategy = _build_keyword_strategy(keyword_candidates, main_intent)
    required_exact_norms = {
        _seo_normalize(row.get("query")) for row in keyword_strategy
        if row.get("exact_required")
    }
    declared_but_absent_keywords = [
        keyword_candidates[normalized].get("query") or ""
        for normalized in declared_keyword_norms
        if normalized not in seen_keywords and normalized in required_exact_norms
    ]
    ordered_candidates = sorted(keyword_candidates.values(), key=_keyword_priority)
    title_used = sorted((row for row in used_keywords if row.get("placement") == "title"), key=_keyword_priority)
    description_used = sorted((row for row in used_keywords if row.get("placement") == "description"), key=_keyword_priority)
    title_target = next((row for row in keyword_strategy if row.get("exact_required")), {})
    title_keyword_check = _keyword_prefix_check(title_target, title, 80, require_near_start=True)
    description_keyword_check = None
    available_keyword_count = len(keyword_candidates)
    natural_keyword_target = min(1, available_keyword_count)
    duplicate_exact_phrases = [
        row.get("query") for row in used_keywords if int(row.get("exact_occurrences") or 0) > 1
    ]
    exact_phrase_occurrences = sum(int(row.get("exact_occurrences") or 0) for row in used_keywords)
    # Do not exaggerate density on short evidence-limited drafts; the guardrail is calibrated per 1000 chars.
    exact_phrase_density = round(exact_phrase_occurrences * 1000 / max(len(combined_copy), 1000), 2)
    seo_writing_checks = {
        "primary_title_keyword_near_start": title_keyword_check,
        "description_opening_keyword": description_keyword_check,
        "keyword_usage": {
            "used": len(used_keywords),
            "available": available_keyword_count,
            "natural_target": natural_keyword_target,
            "passed": len(used_keywords) >= natural_keyword_target,
        },
        "keyword_stuffing": {
            "exact_occurrences": exact_phrase_occurrences,
            "exact_occurrences_per_1000_chars": exact_phrase_density,
            "maximum_occurrences_per_1000_chars": 6.0,
            "duplicate_queries": duplicate_exact_phrases,
            "passed": not duplicate_exact_phrases and exact_phrase_density <= 6.0,
        },
        "policy_note": (
            "Один основной интент проверяется как точная фраза в названии; поисковые запросы — по морфологическим основам и "
            "смысловым связкам. Проверка не является обещанием позиции в поиске Ozon."
        ),
    }

    used_claims, seen_claims = [], set()
    for raw in raw_draft.get("used_claims") or []:
        claim = raw.get("claim") if isinstance(raw, dict) else raw
        normalized = _seo_normalize(claim)
        candidate = claim_candidates.get(normalized)
        if not candidate:
            raise ValueError(f"SEO-текст ссылается на невыбранный клейм: {_clean_text(claim, 120)}")
        if str(candidate.get("verification_status") or "").lower() not in {"safe", "source"}:
            raise ValueError(f"Непроверенный клейм нельзя переносить в SEO-текст: {candidate['claim']}")
        if normalized not in seen_claims:
            used_claims.append(candidate)
            seen_claims.add(normalized)

    used_characteristics, seen_characteristics = [], set()
    for raw in raw_draft.get("used_characteristics") or []:
        if not isinstance(raw, dict):
            raise ValueError("Использованная характеристика должна содержать name и value")
        key = (_seo_normalize(raw.get("name")), _seo_normalize(raw.get("value")))
        candidate = characteristic_candidates.get(key)
        if not candidate:
            raise ValueError(
                "SEO-текст ссылается на невыбранную характеристику: "
                f"{_clean_text(raw.get('name'), 80)} = {_clean_text(raw.get('value'), 120)}"
            )
        if key not in seen_characteristics:
            used_characteristics.append(candidate)
            seen_characteristics.add(key)
    if characteristic_candidates and not used_characteristics:
        raise ValueError("AI не указал использованные характеристики из SEO-разметки")

    hashtags, hashtag_details = _build_quality_hashtags(
        context, list(keyword_candidates.values()), raw_draft.get("hashtags") or []
    )
    keyword_usage_audit, keyword_coverage = _audit_keyword_strategy(
        keyword_candidates, title, description, hashtags, main_intent
    )
    seo_writing_checks["keyword_coverage"] = keyword_coverage
    title_structure = _build_title_structure(
        title, main_intent, allocation.get("title_characteristics") or [], required_model
    )

    warnings = list(allocation.get("warnings") or [])
    if declared_but_absent_keywords:
        warnings.append(
            "AI заявил не найденные в тексте точные фразы: " + ", ".join(declared_but_absent_keywords[:5])
        )
    if duplicate_exact_phrases:
        warnings.append(
            "Точные SEO-фразы повторены и требуют редакторской проверки: " + ", ".join(duplicate_exact_phrases[:5])
        )
    if exact_phrase_density > 6.0:
        warnings.append(
            f"Плотность точных SEO-фраз {exact_phrase_density} на 1000 символов выше безопасного ориентира 6,0"
        )
    if not description_sections:
        warnings.append("AI не вернул проверяемую структуру описания; сохранён совместимый сплошной текст")
    elif len(description_sections) < 3:
        warnings.append("Описание содержит меньше трёх содержательных блоков")
    if len(description) < SEO_CONTENT_DESCRIPTION_TARGET_MIN:
        warnings.append(
            f"Описание короче целевого диапазона {SEO_CONTENT_DESCRIPTION_TARGET_MIN}–"
            f"{SEO_CONTENT_DESCRIPTION_TARGET_MAX} символов; факты не дополнялись выдумками"
        )
    elif len(description) > SEO_CONTENT_DESCRIPTION_SOFT_MAX:
        warnings.append(f"Описание длиннее мягкого лимита {SEO_CONTENT_DESCRIPTION_SOFT_MAX} символов")
    if not used_claims:
        warnings.append("Проверенные клеймы в текст не включены")
    hashtag_target = min(8, len(keyword_candidates))
    if len(hashtags) < hashtag_target:
        warnings.append(
            f"Сформировано {len(hashtags)} качественных хештегов из целевых {hashtag_target}; "
            "недостающие общие или повторяющиеся варианты отброшены"
        )
    if not title_keyword_check["passed"]:
        warnings.append("Слова главного ключа не собраны близко к началу названия")
    if description_keyword_check and not description_keyword_check["passed"]:
        warnings.append("В открывающем блоке описания не найден основной ключ описания")
    if keyword_coverage["exact_passed"] < keyword_coverage["exact_required"]:
        warnings.append(
            f"Точно включено {keyword_coverage['exact_passed']} из {keyword_coverage['exact_required']} "
            "главных ключей"
        )
    if keyword_coverage["base_coverage"] < 0.8:
        warnings.append(
            f"Покрыто {round(keyword_coverage['base_coverage'] * 100)}% уникальных морфологических основ; "
            "целевой ориентир — не менее 80%"
        )
    return {
        "draft_version": SEO_CONTENT_DRAFT_PROMPT_VERSION,
        "rules_version": SEO_CONTENT_DRAFT_RULES_VERSION,
        "sku": str(context.get("sku") or ""),
        "title": title,
        "description": description,
        "description_char_count": len(description),
        "description_target_chars": {
            "min": SEO_CONTENT_DESCRIPTION_TARGET_MIN,
            "max": SEO_CONTENT_DESCRIPTION_TARGET_MAX,
            "soft_max": SEO_CONTENT_DESCRIPTION_SOFT_MAX,
        },
        "description_sections": description_sections,
        "title_structure": title_structure,
        "hashtags": hashtags,
        "primary_intent": main_intent,
        "hashtag_details": hashtag_details,
        "used_keywords": used_keywords,
        "keyword_usage_audit": keyword_usage_audit,
        "seo_writing_checks": seo_writing_checks,
        "used_claims": used_claims,
        "used_characteristics": used_characteristics,
        "required_source_model": required_model,
        "required_source_warranty_claims": required_warranties,
        "monitoring_keywords": allocation.get("monitoring_keywords") or [],
        "warnings": list(dict.fromkeys(value for value in warnings if value)),
        "publication_status": "draft_only",
    }


def generate_project_content_draft(config, payload, analyze_one):
    """Generate and persist a source-backed title/description/internal-tag draft without publishing it."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            marketplace = project.get("marketplace")
            rules_version, prompt_version = _content_stage_versions(marketplace, "draft")
            cur.execute(
                """SELECT c.context_json, c.context_hash, a.allocation_json, a.context_hash AS allocation_context_hash
                   FROM public.seo_generation_semantic_contexts c
                   JOIN public.seo_generation_content_allocations a ON a.project_id=c.project_id AND a.sku=c.sku
                   WHERE c.project_id=%s AND c.sku=%s""", (project_id, sku),
            )
            source = cur.fetchone()
            if not source:
                raise ValueError("Сначала подготовьте SEO-контекст и выполните SEO-разметку для SKU")
            if source["context_hash"] != source["allocation_context_hash"]:
                raise ValueError("SEO-контекст изменился; сначала повторите SEO-разметку для SKU")
            context = source.get("context_json") or {}
            allocation = source.get("allocation_json") or {}
            settings = _ai_script_models(conn, "content_generation")
        ai_payload = analyze_one(context, allocation, settings)
        raw = ai_payload.get("draft") if isinstance(ai_payload, dict) else None
        draft = _validate_content_draft(context, allocation, raw)
        draft["rules_version"] = rules_version
        draft["draft_version"] = prompt_version
        refinement = {"attempted": False}
        usage_check = (draft.get("seo_writing_checks") or {}).get("keyword_usage") or {}
        coverage_check = (draft.get("seo_writing_checks") or {}).get("keyword_coverage") or {}
        stuffing_check = (draft.get("seo_writing_checks") or {}).get("keyword_stuffing") or {}
        if not coverage_check.get("passed") and stuffing_check.get("passed"):
            audit_rows = draft.get("keyword_usage_audit") or []
            missing_exact = [row.get("query") for row in audit_rows if row.get("exact_required") and not row.get("passed")]
            uncovered = [row.get("query") for row in audit_rows if not row.get("exact_required") and not row.get("passed")]
            missing_bases = []
            for row in audit_rows:
                for base in row.get("unmatched_bases") or []:
                    if base not in missing_bases:
                        missing_bases.append(base)
            refinement_allocation = dict(allocation)
            refinement_allocation["generation_feedback"] = {
                "reason": "Не достигнуто точное включение основного интента или покрытие морфологических основ",
                "exact_required": int(coverage_check.get("exact_required") or 0),
                "exact_passed": int(coverage_check.get("exact_passed") or 0),
                "base_coverage": coverage_check.get("base_coverage"),
                "minimum_base_coverage": 0.8,
                "density_limit_per_1000_chars": 6.0,
                "missing_exact_keywords": missing_exact[:1],
                "uncovered_support_queries": uncovered[:12],
                "missing_morphological_bases": missing_bases[:24],
                "current_title": draft.get("title") or "",
                "current_description_sections": draft.get("description_sections") or [],
            }
            refinement = {
                "attempted": True,
                "initial_exact_passed": int(coverage_check.get("exact_passed") or 0),
                "initial_base_coverage": float(coverage_check.get("base_coverage") or 0),
            }
            try:
                refined_payload = analyze_one(context, refinement_allocation, settings)
                refined_raw = refined_payload.get("draft") if isinstance(refined_payload, dict) else None
                refined_draft = _validate_content_draft(context, allocation, refined_raw)
                refined_usage = (refined_draft.get("seo_writing_checks") or {}).get("keyword_usage") or {}
                refined_coverage = (refined_draft.get("seo_writing_checks") or {}).get("keyword_coverage") or {}
                refined_stuffing = (refined_draft.get("seo_writing_checks") or {}).get("keyword_stuffing") or {}
                initial_score = int(coverage_check.get("exact_passed") or 0) * 10 + float(coverage_check.get("base_coverage") or 0)
                refined_score = int(refined_coverage.get("exact_passed") or 0) * 10 + float(refined_coverage.get("base_coverage") or 0)
                refinement.update({
                    "refined_used_keywords": int(refined_usage.get("used") or 0),
                    "refined_exact_passed": int(refined_coverage.get("exact_passed") or 0),
                    "refined_base_coverage": float(refined_coverage.get("base_coverage") or 0),
                    "accepted": bool(
                        refined_stuffing.get("passed")
                        and refined_score > initial_score
                    ),
                })
                if refinement["accepted"]:
                    draft = refined_draft
                    ai_payload = refined_payload
            except Exception as exc:
                refinement.update({"accepted": False, "error": _clean_text(exc, 300)})
        model = _clean_text((ai_payload or {}).get("model"), 200)
        allocation_hash = _allocation_hash(allocation)
        status = "partial" if draft["warnings"] else "ok"
        audit = (ai_payload or {}).get("audit") if isinstance(ai_payload, dict) else {}
        audit = dict(audit) if isinstance(audit, dict) else {}
        audit.update({
            "validated_result": draft,
            "validation_status": status,
            "rules_version": rules_version,
            "prompt_version": prompt_version,
            "refinement": refinement,
        })
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO public.seo_generation_content_drafts
                   (project_id, sku, draft_json, audit_json, allocation_hash, rules_version, prompt_version,
                    status, last_error, model)
                   VALUES (%s,%s,%s::jsonb,%s::jsonb,%s,%s,%s,%s,NULL,%s)
                   ON CONFLICT(project_id, sku) DO UPDATE SET draft_json=excluded.draft_json,
                     audit_json=excluded.audit_json,
                     allocation_hash=excluded.allocation_hash, rules_version=excluded.rules_version,
                     prompt_version=excluded.prompt_version, status=excluded.status,
                     last_error=NULL, model=excluded.model, generated_at=now()""",
                (project_id, sku, json.dumps(draft, ensure_ascii=False, default=str),
                 json.dumps(audit, ensure_ascii=False, default=str), allocation_hash,
                 rules_version, prompt_version, status, model),
            )
        conn.commit()
    return {"ok": True, "project_id": project_id, "sku": sku, "status": status,
            "model": model, "draft": draft, "audit": audit}


def project_content_draft(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                """SELECT draft_json, audit_json, allocation_hash, rules_version, prompt_version, status,
                          last_error, model, generated_at
                   FROM public.seo_generation_content_drafts WHERE project_id=%s AND sku=%s""",
                (project_id, sku),
            )
            row = cur.fetchone()
    if not row:
        return {"ok": True, "project_id": project_id, "sku": sku, "available": False}
    return {"ok": True, "project_id": project_id, "sku": sku, "available": True, **_serialize(row)}


def _content_review_source_facts(context):
    """Return exact source-backed facts that an expert review must not contradict."""
    product = context.get("product") or {}
    rows = product.get("seo_characteristics") or []
    facts, seen = [], set()
    for row in rows:
        if not isinstance(row, dict):
            continue
        name = _clean_text(row.get("name"), 120)
        value = _clean_text(row.get("value"), 500)
        key = (_seo_normalize(name), _seo_normalize(value))
        if not name or not value or key in seen:
            continue
        facts.append({"name": name, "value": value})
        seen.add(key)
    return facts


def _content_review_required_fact_groups(source_facts):
    groups = {"gender": [], "age_audience": [], "material": [], "color": []}
    for row in source_facts:
        name = _seo_normalize(row.get("name"))
        if name == "пол" or "гендер" in name:
            groups["gender"].append(row)
        if "возраст" in name or "аудитор" in name:
            groups["age_audience"].append(row)
        if "материал" in name or "состав" in name:
            groups["material"].append(row)
        if "цвет" in name or "оттенок" in name:
            groups["color"].append(row)
    return groups


def _validate_content_review(context, draft, raw_review):
    """Validate the expert verdict and apply deterministic blocking guards."""
    if not isinstance(raw_review, dict):
        raise ValueError("Экспертная модель вернула результат неподдерживаемого типа")
    verdict = str(raw_review.get("verdict") or "").strip().lower()
    if verdict not in {"approved", "needs_revision", "blocked"}:
        raise ValueError("Экспертная модель не вернула допустимый verdict")
    try:
        score = max(0, min(100, int(round(float(raw_review.get("score"))))))
    except (TypeError, ValueError):
        raise ValueError("Экспертная модель не вернула числовой score")
    checks = raw_review.get("checks")
    if not isinstance(checks, dict):
        raise ValueError("Экспертная модель не вернула таблицу проверок")
    required_checks = (
        "russian_language", "phrase_logic", "prohibited_information", "gender", "age_audience",
        "material", "color", "other_characteristics",
    )
    normalized_checks = {}
    for key in required_checks:
        raw = checks.get(key)
        if not isinstance(raw, dict):
            raise ValueError(f"Экспертная модель не вернула проверку {key}")
        status = str(raw.get("status") or "").strip().lower()
        if status not in {"pass", "review", "fail", "not_applicable"}:
            raise ValueError(f"Недопустимый статус проверки {key}")
        normalized_checks[key] = {
            "status": status,
            "summary": _clean_text(raw.get("summary"), 800),
            "evidence": [_clean_text(item, 500) for item in (raw.get("evidence") or [])[:12] if _clean_text(item, 500)],
        }
    source_facts = _content_review_source_facts(context)
    fact_groups = _content_review_required_fact_groups(source_facts)
    hard_failures = []
    for key in ("gender", "age_audience", "material", "color"):
        if fact_groups[key] and normalized_checks[key]["status"] != "pass":
            hard_failures.append(key)
    text = _seo_normalize(f"{draft.get('title') or ''} {draft.get('description') or ''}")
    gender_values = " ".join(row["value"] for row in fact_groups["gender"]).lower()
    if ("муж" in gender_values and re.search(r"\b(?:женск\w*|девоч\w*|для\s+женщин)\b", text)) or (
        ("жен" in gender_values or "девоч" in gender_values) and re.search(r"\b(?:мужск\w*|мальчик\w*|для\s+мужчин)\b", text)
    ):
        if "gender" not in hard_failures:
            hard_failures.append("gender")
        normalized_checks["gender"] = {
            "status": "fail",
            "summary": "Текст содержит явное обозначение пола, конфликтующее с исходной карточкой.",
            "evidence": normalized_checks["gender"].get("evidence") or [],
        }
    issues = []
    for raw in (raw_review.get("issues") or [])[:50]:
        if not isinstance(raw, dict):
            continue
        severity = str(raw.get("severity") or "minor").strip().lower()
        field = str(raw.get("field") or "description").strip().lower()
        issues.append({
            "severity": severity if severity in {"critical", "major", "minor"} else "minor",
            "field": field if field in {"title", "description", "hashtags", "all"} else "all",
            "category": _clean_text(raw.get("category"), 120),
            "fragment": _clean_text(raw.get("fragment"), 500),
            "message": _clean_text(raw.get("message"), 1000),
            "suggestion": _clean_text(raw.get("suggestion"), 1000),
        })
    suggested_edits = []
    for raw in (raw_review.get("edits") or [])[:12]:
        if not isinstance(raw, dict):
            continue
        field = str(raw.get("field") or "").strip().lower()
        find = _clean_text(raw.get("find"), 1000)
        replace = _clean_text(raw.get("replace"), 1000)
        if field not in {"title", "description", "hashtags"} or not find or find == replace:
            continue
        suggested_edits.append({
            "field": field,
            "find": find,
            "replace": replace,
            "reason": _clean_text(raw.get("reason"), 500),
        })
    if hard_failures:
        verdict = "blocked"
        score = min(score, 49)
    elif any(row["status"] == "fail" for row in normalized_checks.values()):
        verdict = "blocked"
        score = min(score, 49)
    elif any(row["status"] == "review" for row in normalized_checks.values()) and verdict == "approved":
        verdict = "needs_revision"
        score = min(score, 79)
    return {
        "review_version": SEO_CONTENT_REVIEW_PROMPT_VERSION,
        "rules_version": SEO_CONTENT_REVIEW_RULES_VERSION,
        "sku": str(context.get("sku") or ""),
        "verdict": verdict,
        "score": score,
        "summary": _clean_text(raw_review.get("summary"), 1500),
        "checks": normalized_checks,
        "issues": issues,
        "suggested_edits": suggested_edits,
        "recommended_changes": [_clean_text(item, 1000) for item in (raw_review.get("recommended_changes") or [])[:20] if _clean_text(item, 1000)],
        "source_facts": source_facts,
        "hard_failures": hard_failures,
        "publication_allowed": verdict == "approved" and not hard_failures,
        "publication_status": "expert_approved_draft" if verdict == "approved" and not hard_failures else "blocked_draft",
    }


def _apply_content_review_edits(context, allocation, draft, edits):
    """Apply only uniquely anchored local replacements and rebuild the canonical draft audit."""
    candidate = json.loads(json.dumps(draft or {}, ensure_ascii=False, default=str))
    applied, skipped = [], []
    changed_budget = {"title": 0, "description": 0}
    field_limits = {"title": 120, "description": max(300, int(len(str(candidate.get("description") or "")) * 0.2))}
    hashtag_edits = []
    for edit in (edits or [])[:12]:
        field = str(edit.get("field") or "").strip().lower()
        find = str(edit.get("find") or "")
        replace = str(edit.get("replace") or "")
        if field not in {"title", "description", "hashtags"} or not find or find == replace:
            skipped.append({**edit, "status": "skipped", "reason": "Некорректная точечная замена"})
            continue
        if field == "hashtags":
            hashtags = [str(value) for value in (candidate.get("hashtags") or [])]
            if hashtags.count(find) != 1:
                skipped.append({**edit, "status": "skipped", "reason": "Хештег найден не ровно один раз"})
                continue
            if replace and (not replace.startswith("#") or any(char.isspace() for char in replace)):
                skipped.append({**edit, "status": "skipped", "reason": "Исправленный хештег имеет неверный формат"})
                continue
            hashtag_edits.append(edit)
            applied.append({**edit, "status": "applied"})
            continue
        current = str(candidate.get(field) or "")
        if current.count(find) != 1:
            skipped.append({**edit, "status": "skipped", "reason": "Исходный фрагмент найден не ровно один раз"})
            continue
        edit_size = max(len(find), len(replace))
        if changed_budget[field] + edit_size > field_limits[field]:
            skipped.append({**edit, "status": "skipped", "reason": "Превышен лимит точечной редакции поля"})
            continue
        updated = current.replace(find, replace, 1)
        title_limit = 60 if str(context.get("marketplace") or "").lower() == "wb" else 200
        if field == "title" and len(updated) > title_limit:
            skipped.append({**edit, "status": "skipped", "reason": "Название превысит допустимую длину черновика"})
            continue
        if field == "description" and candidate.get("description_sections"):
            matching_sections = [
                row for row in candidate["description_sections"]
                if isinstance(row, dict) and str(row.get("text") or "").count(find) == 1
            ]
            replacement_sections = [
                row for row in candidate["description_sections"]
                if replace and isinstance(row, dict) and str(row.get("text") or "").count(replace) == 1
            ]
            if len(matching_sections) == 1:
                matching_sections[0]["text"] = str(matching_sections[0].get("text") or "").replace(find, replace, 1)
            elif len(matching_sections) == 0 and len(replacement_sections) == 1:
                # The flattened description can be stale while the canonical
                # structured block already contains the requested replacement.
                # Validation below rebuilds the flattened copy from sections.
                pass
            else:
                skipped.append({**edit, "status": "skipped", "reason": "Фрагмент не привязан к одному смысловому блоку"})
                continue
        candidate[field] = updated
        changed_budget[field] += edit_size
        applied.append({**edit, "status": "applied"})
    if not applied:
        return draft, applied, skipped
    # Recompute keyword, claim, characteristic and structural audits from the patched copy.
    validated = _validate_content_draft(context, allocation, candidate)
    if hashtag_edits:
        hashtags = list(validated.get("hashtags") or [])
        details = list(validated.get("hashtag_details") or [])
        for edit in hashtag_edits:
            find = str(edit.get("find") or "")
            replace = str(edit.get("replace") or "")
            index = hashtags.index(find)
            if replace:
                hashtags[index] = replace
            else:
                hashtags.pop(index)
            replacement_details = []
            for row in details:
                if str((row or {}).get("hashtag") or "") != find:
                    replacement_details.append(row)
                elif replace:
                    replacement_details.append({
                        **row,
                        "hashtag": replace,
                        "source_type": "expert_revision",
                        "revision_reason": _clean_text(edit.get("reason"), 500),
                    })
            details = replacement_details
        validated["hashtags"] = hashtags
        validated["hashtag_details"] = details
        keyword_candidates = {}
        for placement, key in (("title", "title_keywords"), ("description", "description_keywords")):
            for row in allocation.get(key) or []:
                query = _clean_text((row or {}).get("query"), 300)
                normalized = _seo_normalize(query)
                if normalized:
                    keyword_candidates[normalized] = {**row, "query": query, "placement": placement}
        usage, coverage = _audit_keyword_strategy(
            keyword_candidates,
            validated.get("title") or "",
            validated.get("description") or "",
            hashtags,
            _clean_text(((context or {}).get("product") or {}).get("intent"), 300),
        )
        validated["keyword_usage_audit"] = usage
        validated["seo_writing_checks"]["keyword_coverage"] = coverage
        warnings = [
            value for value in (validated.get("warnings") or [])
            if not str(value).startswith("Сформировано ") or "качественных хештегов" not in str(value)
        ]
        hashtag_target = min(8, len(keyword_candidates))
        if len(hashtags) < hashtag_target:
            warnings.append(
                f"После экспертной правки сохранено {len(hashtags)} хештегов из целевых {hashtag_target}"
            )
        validated["warnings"] = warnings
    return validated, applied, skipped


def _content_review_ai_call(phase, model, call_audit):
    call_audit = call_audit if isinstance(call_audit, dict) else {}
    return {
        "phase": phase,
        "model": _clean_text(model, 200),
        "duration_ms": max(0, int(call_audit.get("duration_ms") or 0)),
        "attempt_count": max(1, int(call_audit.get("attempt_count") or 1)),
        "fallback_used": bool(call_audit.get("fallback_used")),
    }


def _scope_content_revision_edits(expert_review, revision_edits):
    """Keep only replacements anchored to the exact fragments approved by the expert."""
    allowed = {
        (str(row.get("field") or "").strip().lower(), str(row.get("find") or "")): row
        for row in (expert_review.get("suggested_edits") or [])
        if isinstance(row, dict)
    }
    scoped, rejected = [], []
    for raw in (revision_edits or [])[:12]:
        if not isinstance(raw, dict):
            continue
        field = str(raw.get("field") or "").strip().lower()
        find = str(raw.get("find") or "")
        replace = str(raw.get("replace") or "")
        if (field, find) not in allowed:
            rejected.append({**raw, "status": "rejected", "reason": "Фрагмент не входил в задание эксперта"})
            continue
        scoped.append({
            "field": field,
            "find": find,
            "replace": replace,
            "reason": _clean_text(raw.get("reason") or allowed[(field, find)].get("reason"), 500),
        })
    return scoped, rejected


def _compiled_content_review(initial_review, patched_draft, applied_edits, skipped_edits):
    """Close the review after the expert-directed edits were compiled exactly once."""
    expected = {
        (str(row.get("field") or "").strip().lower(), str(row.get("find") or ""))
        for row in (initial_review.get("suggested_edits") or [])
        if isinstance(row, dict)
    }
    applied = {
        (str(row.get("field") or "").strip().lower(), str(row.get("find") or ""))
        for row in (applied_edits or [])
        if isinstance(row, dict)
    }
    # The compiler already proves that every expected anchor was found exactly
    # once and replaced. Do not require the old anchor to be absent as a
    # substring of its replacement: valid edits can extend an anchor, e.g.
    # ``#джинсы_утепленные`` -> ``#джинсы_утепленные_для_девочки``.
    addressed = bool(expected) and expected.issubset(applied) and not skipped_edits
    if not addressed:
        return initial_review, False, "not_all_expert_edits_were_compiled"
    result = json.loads(json.dumps(initial_review, ensure_ascii=False, default=str))
    result["verdict"] = "approved"
    result["score"] = max(80, int(result.get("score") or 0))
    result["summary"] = "Точечные правки эксперта применены быстрой моделью; остальной текст не изменён."
    for check in (result.get("checks") or {}).values():
        if isinstance(check, dict) and check.get("status") in {"review", "fail"}:
            check["status"] = "pass"
            check["summary"] = "Исправлено точечной заменой по заданию эксперта."
    result["issues"] = []
    result["suggested_edits"] = []
    result["recommended_changes"] = []
    result["hard_failures"] = []
    result["publication_allowed"] = True
    result["publication_status"] = "expert_directed_revision_completed"
    return result, True, "all_expert_edits_compiled"


def generate_project_content_review(config, payload, analyze_one, revise_one=None):
    """Review once, delegate exact fixes once, compile them, and never publish."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            marketplace = project.get("marketplace")
            rules_version, prompt_version = _content_stage_versions(marketplace, "review")
            cur.execute(
                """SELECT c.context_json, a.allocation_json, d.draft_json, d.audit_json AS draft_audit_json,
                          d.generated_at
                   FROM public.seo_generation_semantic_contexts c
                   JOIN public.seo_generation_content_allocations a ON a.project_id=c.project_id AND a.sku=c.sku
                   JOIN public.seo_generation_content_drafts d ON d.project_id=c.project_id AND d.sku=c.sku
                   WHERE c.project_id=%s AND c.sku=%s""", (project_id, sku),
            )
            source = cur.fetchone()
            if not source:
                raise ValueError("Сначала подготовьте SEO-контекст и сформируйте SEO-текст для SKU")
            context = source.get("context_json") or {}
            allocation = source.get("allocation_json") or {}
            draft = source.get("draft_json") or {}
            settings = _ai_script_models(conn, "content_expert_review")
            revision_settings = _ai_script_models(conn, "content_generation")
        review_started = time.monotonic()
        ai_payload = analyze_one(context, draft, settings)
        raw = ai_payload.get("review") if isinstance(ai_payload, dict) else None
        initial_review = _validate_content_review(context, draft, raw)
        initial_review["rules_version"] = rules_version
        initial_review["review_version"] = prompt_version
        review = initial_review
        model = _clean_text((ai_payload or {}).get("model"), 200)
        audit = (ai_payload or {}).get("audit") if isinstance(ai_payload, dict) else {}
        audit = dict(audit) if isinstance(audit, dict) else {}
        ai_calls = [_content_review_ai_call("initial", model, audit)]
        auto_fix = payload.get("apply_fixes", True) not in {False, 0, "0", "false", "False"}
        remediation = {
            "requested": auto_fix,
            "policy": "one_expert_then_one_targeted_revision_no_recheck",
            "attempted": False,
            "accepted": False,
            "attempted_edits": [],
            "applied_edits": [],
            "rejected_edits": [],
            "skipped_edits": [],
            "initial_verdict": initial_review["verdict"],
            "initial_score": initial_review["score"],
            "suggested_edits": initial_review.get("suggested_edits") or [],
            "rounds": [],
        }
        draft_audit = source.get("draft_audit_json") or {}
        draft_audit = dict(draft_audit) if isinstance(draft_audit, dict) else {}
        if auto_fix and initial_review["verdict"] != "approved" and initial_review.get("suggested_edits"):
            remediation["attempted"] = True
            revision_payload = (
                revise_one(context, draft, initial_review, revision_settings)
                if revise_one else {"edits": initial_review.get("suggested_edits") or [], "model": "deterministic"}
            )
            revision_audit = (revision_payload or {}).get("audit") if isinstance(revision_payload, dict) else {}
            revision_model = _clean_text((revision_payload or {}).get("model"), 200)
            ai_calls.append(_content_review_ai_call("targeted_revision", revision_model, revision_audit))
            scoped_edits, rejected = _scope_content_revision_edits(
                initial_review, (revision_payload or {}).get("edits") or []
            )
            patched_draft, applied, skipped = _apply_content_review_edits(
                context, allocation, draft, scoped_edits
            )
            skipped.extend(rejected)
            review, accepted, acceptance_reason = _compiled_content_review(
                initial_review, patched_draft, applied, skipped
            )
            remediation["attempted_edits"].extend(scoped_edits)
            remediation["applied_edits"].extend(applied if accepted else [])
            remediation["rejected_edits"].extend(rejected)
            remediation["skipped_edits"].extend(skipped)
            remediation["accepted"] = accepted
            remediation["rounds"].append({
                "round": 1,
                "kind": "targeted_revision",
                "accepted": accepted,
                "acceptance_reason": acceptance_reason,
                "applied_edits": applied,
                "skipped_edits": skipped,
                "revision_model": revision_model,
            })
            remediation["revision_model"] = revision_model
            remediation["revision_audit"] = revision_audit
            if accepted:
                draft = patched_draft
                model = f"{model} + {revision_model}" if revision_model else model
            remediation.update({
                "acceptance_reason": acceptance_reason,
                "final_verdict": review["verdict"],
                "final_score": review["score"],
            })
        if remediation["accepted"]:
            draft_audit["expert_auto_fix"] = {
                "review_version": SEO_CONTENT_REVIEW_PROMPT_VERSION,
                "original_draft_hash": _allocation_hash(source.get("draft_json") or {}),
                "applied_edits": remediation["applied_edits"],
                "rounds": remediation["rounds"],
                "final_verdict": review["verdict"],
                "final_score": review["score"],
            }
            draft_status = "partial" if draft.get("warnings") else "ok"
            with conn.cursor() as cur:
                cur.execute(
                    """UPDATE public.seo_generation_content_drafts
                       SET draft_json=%s::jsonb, audit_json=%s::jsonb, status=%s,
                           last_error=NULL, generated_at=now()
                       WHERE project_id=%s AND sku=%s""",
                    (json.dumps(draft, ensure_ascii=False, default=str),
                     json.dumps(draft_audit, ensure_ascii=False, default=str),
                     draft_status, project_id, sku),
                )
        draft_hash = _allocation_hash(draft)
        status = "ok" if review["publication_allowed"] else "partial"
        audit.update({
            "validated_result": review,
            "validation_status": status,
            "draft_hash": draft_hash,
            "rules_version": rules_version,
            "prompt_version": prompt_version,
            "remediation": remediation,
            "performance": {
                "policy": remediation["policy"],
                "initial_ai_calls": 1,
                "revision_ai_calls": max(0, len(ai_calls) - 1),
                "expert_recheck_ai_calls": 0,
                "total_ai_calls": len(ai_calls),
                "provider_attempts": sum(row["attempt_count"] for row in ai_calls),
                "fallback_calls": sum(1 for row in ai_calls if row["fallback_used"]),
                "ai_duration_ms": sum(row["duration_ms"] for row in ai_calls),
                "wall_duration_ms": round((time.monotonic() - review_started) * 1000),
                "calls": ai_calls,
            },
        })
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO public.seo_generation_content_reviews
                   (project_id, sku, review_json, audit_json, draft_hash, rules_version, prompt_version,
                    status, last_error, model)
                   VALUES (%s,%s,%s::jsonb,%s::jsonb,%s,%s,%s,%s,NULL,%s)
                   ON CONFLICT(project_id, sku) DO UPDATE SET review_json=excluded.review_json,
                     audit_json=excluded.audit_json, draft_hash=excluded.draft_hash,
                     rules_version=excluded.rules_version, prompt_version=excluded.prompt_version,
                     status=excluded.status, last_error=NULL, model=excluded.model, reviewed_at=now()""",
                (project_id, sku, json.dumps(review, ensure_ascii=False, default=str),
                 json.dumps(audit, ensure_ascii=False, default=str), draft_hash,
                 rules_version, prompt_version, status, model),
            )
        conn.commit()
    return {"ok": True, "project_id": project_id, "sku": sku, "status": status,
            "model": model, "review": review, "audit": audit, "remediation": remediation,
            "draft_updated": bool(remediation.get("accepted"))}


def project_content_review(config, payload):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    sku = _clean_text(payload.get("sku"), 200)
    if not sku:
        raise ValueError("SKU не передан")
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                """SELECT r.review_json, r.audit_json, r.draft_hash, r.rules_version, r.prompt_version,
                          r.status, r.last_error, r.model, r.reviewed_at, d.draft_json
                   FROM public.seo_generation_content_reviews r
                   LEFT JOIN public.seo_generation_content_drafts d ON d.project_id=r.project_id AND d.sku=r.sku
                   WHERE r.project_id=%s AND r.sku=%s""", (project_id, sku),
            )
            row = cur.fetchone()
    if not row:
        return {"ok": True, "project_id": project_id, "sku": sku, "available": False}
    current_hash = _allocation_hash(row.get("draft_json") or {})
    payload_row = _serialize(row)
    payload_row.pop("draft_json", None)
    return {"ok": True, "project_id": project_id, "sku": sku, "available": True,
            "stale": current_hash != row.get("draft_hash"), **payload_row}


def refresh_project_mpstats(config, payload, mpstats_fetcher=None):
    """Collect project keywords from the existing MPStats integration in small UI batches."""
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = {_clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)}
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind)
            cur.execute(
                "SELECT sku, product_name FROM public.seo_monitoring_project_skus WHERE project_id=%s ORDER BY sku",
                (project_id,),
            )
            sku_rows = [dict(row) for row in cur.fetchall()]
    if requested_skus:
        sku_rows = [row for row in sku_rows if row["sku"] in requested_skus]
    if not sku_rows:
        raise ValueError("В проекте нет выбранных SKU")
    if len(sku_rows) > MPSTATS_BATCH_LIMIT:
        raise ValueError(f"За один запрос MPStats можно обработать не более {MPSTATS_BATCH_LIMIT} SKU")
    period_from = payload.get("date_from") or project.get("date_from") or date.today()
    period_to = payload.get("date_to") or project.get("date_to") or date.today()
    snapshot_date = payload.get("snapshot_date") or date.today().isoformat()
    fetcher = mpstats_fetcher or fetch_mpstats_project_keywords
    rows = []
    errors = []
    empty_skus = []
    for sku_row in sku_rows:
        try:
            words = fetcher(project["marketplace"], sku_row["sku"], period_from, period_to)
            normalized = _normalize_mpstats_keyword_rows(project["marketplace"], sku_row, words)
            rows.extend(normalized)
            if not normalized:
                empty_skus.append(sku_row["sku"])
        except Exception as exc:
            errors.append({"sku": sku_row["sku"], "error": str(exc)[:500]})
    source = f"mpstats_{project['marketplace']}_item_keywords"
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            _assert_project_kind(cur, project_id, project_kind, for_update=True)
            imported = _insert_snapshot_rows(cur, project_id, snapshot_date, period_from, period_to, source, rows)
            status = "error" if errors and not imported else ("partial" if errors else "ready")
            last_error = "; ".join(f"{item['sku']}: {item['error']}" for item in errors) or None
            cur.execute(
                """UPDATE public.seo_monitoring_projects
                   SET status=%s, last_refreshed_at=now(), updated_at=now(), last_error=%s
                   WHERE project_id=%s""",
                (status, last_error, project_id),
            )
    return {
        "ok": True, "project_id": project_id, "source": source,
        "sku_count": len(sku_rows), "row_count": imported,
        "empty_count": len(empty_skus), "empty_skus": empty_skus, "errors": errors, "status": status,
        "snapshot_date": str(snapshot_date), "date_from": str(period_from), "date_to": str(period_to),
    }


def latest_available_day(today=None):
    """Newest day Ozon has already calculated."""
    return (today or date.today()) - timedelta(days=OZON_ANALYTICS_LAG_DAYS)


def default_backfill_range(today=None):
    """First run collects the last month day by day."""
    last_day = latest_available_day(today)
    return last_day - timedelta(days=OZON_DAILY_HISTORY_DAYS - 1), last_day


def project_snapshot_days(config, project_id):
    project_id = _project_id(project_id)
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                "SELECT DISTINCT snapshot_date FROM public.seo_monitoring_keyword_snapshots WHERE project_id=%s ORDER BY snapshot_date",
                (project_id,),
            )
            return [row["snapshot_date"] for row in cur.fetchall()]


def _as_date(value):
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    return date.fromisoformat(str(value)[:10])


def collect_project_days(config, payload, ozon_fetcher=None, progress=None):
    """Collect one snapshot per day, so the project keeps a daily monitoring history.

    Days already present are skipped unless `force` is set; each day is stored with
    `snapshot_date` equal to that day and the period narrowed to it.
    """
    project_id = _project_id(payload.get("project_id"))
    force = str(payload.get("force") or "").strip().lower() in {"1", "true", "yes"}
    date_from, date_to = default_backfill_range()
    if payload.get("date_from"):
        date_from = _as_date(payload["date_from"])
    if payload.get("date_to"):
        date_to = _as_date(payload["date_to"])
    if date_to > latest_available_day():
        date_to = latest_available_day()
    if date_from > date_to:
        raise ValueError("Нет доступных дней: Ozon считает аналитику с задержкой 1–2 дня")
    existing = set() if force else set(project_snapshot_days(config, project_id))
    days = []
    current = date_from
    while current <= date_to:
        if force or current not in existing:
            days.append(current)
        current += timedelta(days=1)
    results = []
    rows_total = 0
    errors = []
    for index, day in enumerate(days, start=1):
        if progress:
            progress(index, len(days), day, rows_total, len(errors))
        day_payload = {
            **{key: value for key, value in payload.items() if key not in {"date_from", "date_to", "snapshot_date", "force"}},
            "project_id": project_id,
            "date_from": day.isoformat(),
            "date_to": day.isoformat(),
            "snapshot_date": day.isoformat(),
        }
        try:
            result = refresh_project(config, day_payload, ozon_fetcher)
        except Exception as exc:  # one day must not abort the rest of the backfill
            errors.append({"day": day.isoformat(), "error": str(exc)[:500]})
            continue
        rows_total += int(result.get("row_count") or 0)
        errors.extend({**item, "day": day.isoformat()} for item in (result.get("errors") or []))
        results.append({"day": day.isoformat(), "row_count": result.get("row_count"), "sku_count": result.get("sku_count")})
    return {
        "ok": True, "project_id": project_id, "days_planned": len(days),
        "days_done": len(results), "days_skipped": 0 if force else len(existing & {*days}),
        "row_count": rows_total, "errors": errors,
        "date_from": date_from.isoformat(), "date_to": date_to.isoformat(),
        "latest_available_day": latest_available_day().isoformat(),
    }


def refresh_project(config, payload, ozon_fetcher=None):
    project_id = _project_id(payload.get("project_id"))
    project_kind = _project_kind(payload.get("project_kind"))
    requested_skus = {_clean_text(value, 200) for value in (payload.get("skus") or []) if _clean_text(value, 200)}
    with _conn(config) as conn:
        ensure_schema(conn)
        with conn.cursor() as cur:
            project = _assert_project_kind(cur, project_id, project_kind, for_update=True)
            cur.execute("SELECT sku, product_name FROM public.seo_monitoring_project_skus WHERE project_id=%s ORDER BY sku", (project_id,))
            sku_rows = [dict(row) for row in cur.fetchall()]
            if requested_skus:
                sku_rows = [row for row in sku_rows if row["sku"] in requested_skus]
            if not sku_rows:
                raise ValueError("В проекте нет выбранных SKU")
            if len(sku_rows) > 1000:
                raise ValueError("За один запуск можно актуализировать не более 1000 SKU")
            period_from = payload.get("date_from") or project.get("date_from") or date.today()
            period_to = payload.get("date_to") or project.get("date_to") or date.today()
            snapshot_date = payload.get("snapshot_date") or date.today().isoformat()
            try:
                limit_by_sku = int(payload.get("limit_by_sku") or 15)
            except (TypeError, ValueError) as exc:
                raise ValueError("Ключей на SKU должно быть числом") from exc
            if limit_by_sku < 1 or limit_by_sku > 15:
                raise ValueError("Ozon отдаёт от 1 до 15 фраз по SKU за запуск")
            sort_by = _clean_text(payload.get("sort_by"), 40).upper() or "BY_SEARCHES"
            # Documented values; BY_ORDERS does not exist and Ozon silently falls back to BY_SEARCHES.
            if sort_by not in OZON_QUERY_SORTS:
                raise ValueError("Сортировка поддерживает " + ", ".join(sorted(OZON_QUERY_SORTS)))
            requested_sorts = payload.get("sort_by_list")
            if isinstance(requested_sorts, (list, tuple)):
                sort_plan = [value for value in (_clean_text(item, 40).upper() for item in requested_sorts) if value in OZON_QUERY_SORTS]
            elif str(payload.get("all_sorts") or "").strip().lower() in {"1", "true", "yes"}:
                sort_plan = list(OZON_QUERY_SORTS_ORDER)
            else:
                sort_plan = []
            if not sort_plan:
                sort_plan = [sort_by]
            # Every sort returns its own top-`limit_by_sku` slice, so several passes widen the phrase set.
            seen_sorts = []
            for value in sort_plan:
                if value not in seen_sorts:
                    seen_sorts.append(value)
            sort_plan = seen_sorts
            sort_dir = _clean_text(payload.get("sort_dir"), 20).upper() or "DESCENDING"
            if sort_dir not in {"ASCENDING", "DESCENDING"}:
                raise ValueError("Направление сортировки поддерживает ASCENDING или DESCENDING")
            rows = []
            errors = []
            if project["marketplace"] == "wb":
                sku_values = [row["sku"] for row in sku_rows]
                cur.execute(
                    """SELECT wb_sku::text AS sku, max(product_name) AS product_name, search_query,
                              avg(average_position) AS average_position, sum(query_count) AS search_demand,
                              sum(card_visits) AS traffic, sum(cart_adds) AS cart_adds,
                              sum(ordered_units) AS orders, NULL::numeric AS revenue_rub
                       FROM public.wb_search_queries_daily
                       WHERE report_date BETWEEN %s AND %s AND wb_sku::text = ANY(%s)
                       GROUP BY wb_sku, search_query""",
                    (period_from, period_to, sku_values),
                )
                rows = [dict(row, raw_payload={"source": "wb_search_queries_daily"}) for row in cur.fetchall()]
                source = "wb_search_queries_daily"
            else:
                if not ozon_fetcher:
                    raise RuntimeError("Ozon SEO fetcher не настроен")
                source = "ozon_seller_api_product_queries_details"
                names_by_sku = {row["sku"]: row.get("product_name") for row in sku_rows}
                sku_values = [row["sku"] for row in sku_rows]
                batches = [sku_values[index:index + OZON_SKUS_PER_REQUEST] for index in range(0, len(sku_values), OZON_SKUS_PER_REQUEST)]
                for batch in batches:
                    for current_sort in sort_plan:
                        page = 0
                        while page < OZON_MAX_PAGES:
                            result = None
                            last_error = None
                            for attempt in range(OZON_PAGE_RETRIES):
                                try:
                                    result = ozon_fetcher({
                                        "skus": batch, "date_from": str(period_from), "date_to": str(period_to),
                                        "limit_by_sku": limit_by_sku, "page": page, "page_size": OZON_PAGE_SIZE,
                                        "sort_by": current_sort, "sort_dir": sort_dir,
                                    })
                                    break
                                except Exception as exc:  # 429 and transient failures deserve a retry
                                    last_error = exc
                                    message = str(exc)
                                    retryable = "429" in message or "rate limit" in message.lower() or "timed out" in message.lower()
                                    if not retryable or attempt == OZON_PAGE_RETRIES - 1:
                                        break
                                    time.sleep(OZON_RETRY_BACKOFF_SECONDS * (attempt + 1))
                            if result is None:
                                errors.append({"sku": f"{batch[0]}…{batch[-1]}", "sort_by": current_sort, "page": page, "error": str(last_error)[:500]})
                                break
                            items = extract_ozon_items(result.get("response") if isinstance(result, dict) else result)
                            for item in items:
                                row = normalize_ozon_item(item)
                                if not row:
                                    continue
                                if not row.get("product_name"):
                                    row["product_name"] = names_by_sku.get(row["sku"])
                                row["raw_payload"] = {**(row.get("raw_payload") or {}), "sort_by": current_sort}
                                rows.append(row)
                            if len(items) < OZON_PAGE_SIZE:
                                break
                            page += 1
                            if OZON_PAGE_PAUSE_SECONDS:
                                time.sleep(OZON_PAGE_PAUSE_SECONDS)
            imported = _insert_snapshot_rows(cur, project_id, snapshot_date, period_from, period_to, source, rows)
            status = "error" if errors and not imported else ("partial" if errors else "ready")
            last_error = "; ".join(f"{item['sku']}: {item['error']}" for item in errors) or None
            cur.execute(
                """UPDATE public.seo_monitoring_projects
                   SET status=%s, last_refreshed_at=now(), updated_at=now(), last_error=%s
                   WHERE project_id=%s""",
                (status, last_error, project_id),
            )
    return {
        "ok": True, "project_id": project_id, "sku_count": len(sku_rows), "row_count": imported,
        "errors": errors, "status": status, "snapshot_date": str(snapshot_date),
        "date_from": str(period_from), "date_to": str(period_to),
        "limit_by_sku": limit_by_sku, "sort_by": sort_by, "sort_dir": sort_dir,
        "sorts_used": sort_plan if project["marketplace"] != "wb" else [],
    }
