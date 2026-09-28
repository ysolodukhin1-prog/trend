"""Read-only marketplace P&L selection and consolidation of existing reports."""
from datetime import date, timedelta
from decimal import Decimal

LABELS = {'ozon': 'Ozon', 'wb': 'WB', 'yandex': 'Яндекс Маркет', 'all': 'Итого'}
AMOUNTS = ('revenue', 'seller_revenue', 'units', 'marketplace_net', 'ozon_costs',
           'cogs', 'gross_profit', 'seller_unit_costs', 'seller_costs', 'manual_expenses',
           'vat', 'tax', 'taxes', 'profit_before_tax', 'management_result', 'net_profit')


def connected_marketplaces(config, client, configured):
    # Yandex cabinets are registered separately from the legacy Ozon/WB selector.
    from km_trade_finance import connect_km
    supported = list(configured)
    with connect_km(config) as conn, conn.cursor() as cur:
        cur.execute("SELECT to_regclass('public.yandex_dim_store') relation")
        if cur.fetchone()['relation']:
            cur.execute('SELECT 1 FROM yandex_dim_store WHERE client_key=%s LIMIT 1', (client,))
            if cur.fetchone() and 'yandex' not in supported:
                supported.append('yandex')
    return supported


def total(values):
    values = list(values)
    if not values or any(v is None for v in values):
        return None
    numbers = [Decimal(str(v)) for v in values]
    return float(sum(numbers)) if all(v.is_finite() for v in numbers) else None


def pct(amount, revenue):
    return amount / revenue * 100 if amount is not None and revenue is not None and revenue > 0 else None


def statement(t):
    spec = [('revenue', 'Продажи и возвраты', 1), ('ozon_costs', 'Расходы площадок за вычетом корректировок', -1),
            ('marketplace_net', 'После услуг площадок, до себестоимости и налогов', 1),
            ('cogs', 'Себестоимость', -1), ('gross_profit', 'Валовая прибыль', 1),
            ('seller_unit_costs', 'Расходы продавца', -1), ('manual_expenses', 'Дополнительные расходы', -1),
            ('vat', 'НДС модели', -1), ('tax', 'Налог модели', -1),
            ('management_result', 'Результат модели', 1)]
    return [dict(key=k, label=label, amount=t.get(k) * sign if t.get(k) is not None else None,
                 revenue_pct=pct(t.get(k) * sign if t.get(k) is not None else None, t.get('revenue')),
                 kind='total' if k == 'management_result' else 'subtotal' if k in ('marketplace_net', 'gross_profit') else 'expense')
            for k, label, sign in spec]


def unavailable(client, market, start, end, message):
    return dict(client=client, marketplace=market, available=False, message=message,
                date_from=start, date_to=end, totals={}, statement=[], products=[], months=[], monthly=[],
                period=dict(date_from=start, date_to=end), model_read_only=True)


def yandex_payload(config, start, end, client):
    from km_trade_finance import connect_km
    from unit_economics_workspace import _yandex
    with connect_km(config) as conn, conn.cursor() as cur:
        cur.execute("SELECT to_regclass('public.yandex_fact_realization') r, to_regclass('public.yandex_fact_services') s")
        tables = cur.fetchone()
        if not tables['r'] or not tables['s']:
            return unavailable(client, 'yandex', start, end, 'Реализации и услуги Яндекс Маркета ещё не загружены.')
        cur.execute('''SELECT min(dt) first, max(dt) last FROM (
            SELECT event_date dt FROM yandex_fact_realization WHERE client_key=%s
            UNION ALL SELECT coalesce(service_date,act_date) FROM yandex_fact_services WHERE client_key=%s) d''', (client, client))
        bounds = cur.fetchone()
        if not bounds['last']:
            return unavailable(client, 'yandex', start, end, 'Реализации и услуги Яндекс Маркета ещё не загружены.')
        start_date, end_date = date.fromisoformat(start), date.fromisoformat(end)
        rows, _ = _yandex(cur, start_date, end_date, client)
        def summarize(items):
            unknown = any(r.get('missing_amount') for r in items)
            revenue = None if unknown else total(r.get('revenue') or 0 for r in items)
            net = None if unknown else total(r.get('net') or 0 for r in items)
            return dict(revenue=revenue, seller_revenue=revenue,
                        units=total(r.get('units') or 0 for r in items), marketplace_net=net,
                        ozon_costs=revenue-net if revenue is not None and net is not None else None,
                        cogs=None, gross_profit=None, management_result=None, net_profit=None,
                        tax_configured=False, profit_ready=False)
        totals = summarize(rows)
        months = []
        cursor = start_date
        while cursor <= end_date:
            following = (cursor.replace(day=28) + timedelta(days=4)).replace(day=1)
            month_rows, _ = _yandex(cur, cursor, min(end_date, following-timedelta(days=1)), client)
            months.append(dict(month=cursor.strftime('%Y-%m'), **summarize(month_rows)))
            cursor = following
        products = []
        for r in rows:
            t = summarize([r])
            products.append(dict(t, sku=r['sku'], article=r['sku'], name=r['sku'],
                full_name=r.get('cabinet_label', r['cabinet']), cabinet=r['cabinet'], marketplace='yandex',
                cogs_known=False, cogs_source='missing', commission=total([r.get('commission_cost')]),
                delivery=total([r.get('logistics_cost')]), storage=total([r.get('storage_cost')]), advertising=total([r.get('advertising_cost')]),
                other_ozon=total([r.get('other_cost')]), seller_costs=None, income_tax=None, vat=None))
    return dict(client=client, marketplace='yandex', available=bool(rows),
        message='' if rows else 'За выбранный период нет загруженных операций Яндекс Маркета.',
        date_from=start, date_to=end, available_date_from=str(bounds['first']), available_date_to=str(bounds['last']),
        period=dict(date_from=start, date_to=end, available_from=str(bounds['first']), available_to=str(bounds['last'])),
        totals=totals, statement=statement(totals), monthly=months, months=months, products=products,
        model_read_only=True, model_notice='Яндекс Маркет: реализации, услуги, компенсации и сторно из загруженных отчётов. Себестоимость и налоги в P&L не подтверждены; прибыль не рассчитана.',
        methodology=dict(source='Все кабинеты текущего аккаунта; услуги без SKU сохранены в итогах.',
                         coverage='Полнота периода и закрывающих документов требует сверки.'))


def consolidate(reports, client, start, end):
    complete = bool(reports) and all(r.get('available') is not False for r in reports)
    def aggregate(items):
        result = {k: total(r.get(k) for r in items) if complete else None for k in AMOUNTS}
        result['management_margin_pct'] = pct(result['management_result'], result['revenue'])
        result['margin_pct'] = result['management_margin_pct']
        result['tax_configured'] = all(r.get('tax_configured') is True for r in items)
        return result
    totals = aggregate([r.get('totals', {}) for r in reports])
    month_keys = sorted({str(m['month'])[:7] for r in reports for m in r.get('monthly', r.get('months', []))})
    months = []
    for month in month_keys:
        # A missing market/month is unknown; never silently add it as zero.
        items = [next((m for m in r.get('monthly', r.get('months', [])) if str(m['month'])[:7] == month), {}) for r in reports]
        months.append(dict(month=month, **aggregate(items)))
    products = [dict(p, marketplace=r['marketplace']) for r in reports for p in r.get('products', [])]
    coverage = [dict(marketplace=r['marketplace'], label=LABELS[r['marketplace']], available=r.get('available') is not False,
                     message=r.get('message', ''), **r.get('totals', {})) for r in reports]
    return dict(client=client, marketplace='all', available=any(r.get('available') is not False for r in reports),
        partial=not complete, date_from=start, date_to=end, totals=totals, statement=statement(totals),
        monthly=months, months=months, products=products, marketplace_totals=coverage, model_read_only=True,
        period=dict(date_from=start, date_to=end),
        model_notice='Итого по подключённым площадкам. Расчётные показатели Ozon/WB сохраняют модель себестоимости; неизвестные суммы и прибыль не заменяются нулями.',
        methodology=dict(aggregation='Денежные статьи суммируются за один период. Маржа = суммарный результат / суммарная выручка. Расходы продавца учитываются один раз в своей площадке.'),
        available_date_from=min((r.get('available_date_from', start) for r in reports), default=start),
        available_date_to=max((r.get('available_date_to', end) for r in reports), default=end))


def selected_payload(config, raw_from, raw_to, client, marketplace, supported):
    from km_trade_finance import pl_payload
    end = date.fromisoformat(raw_to) if raw_to else date.today()-timedelta(days=1)
    start = date.fromisoformat(raw_from) if raw_from else end-timedelta(days=29)
    if start > end or (end-start).days > 366:
        raise ValueError('Период должен быть от 1 до 367 дней')
    start, end = str(start), str(end)
    if marketplace == 'all':
        reports = []
        for market in ('ozon', 'wb', 'yandex'):
            if market not in supported:
                continue
            try:
                reports.append(selected_payload(config, start, end, client, market, supported))
            except Exception:
                reports.append(unavailable(client, market, start, end, 'Не удалось загрузить P&L площадки. Повторите запрос.'))
        return consolidate(reports, client, start, end)
    if marketplace not in supported or marketplace not in LABELS:
        return unavailable(client, marketplace, start, end, 'Площадка не подключена к этому аккаунту.')
    if marketplace == 'yandex':
        return yandex_payload(config, start, end, client)
    return pl_payload(config, raw_from, raw_to, client, marketplace)


def workbook_bytes(data):
    """Keep the existing export and add unambiguous market/cabinet provenance."""
    from io import BytesIO
    from openpyxl import load_workbook
    from km_trade_finance import pl_workbook_bytes
    workbook = load_workbook(BytesIO(pl_workbook_bytes(data)))
    products = workbook['По товарам']
    products.insert_cols(1, 2)
    products.cell(1, 1, 'Площадка'); products.cell(1, 2, 'Кабинет')
    for index, row in enumerate(data['products'], 2):
        products.cell(index, 1, LABELS.get(row.get('marketplace', data['marketplace']), data['marketplace']))
        products.cell(index, 2, str(row.get('cabinet', '')))
    if data.get('marketplace_totals'):
        markets = workbook.create_sheet('По площадкам')
        keys = ('label', 'revenue', 'ozon_costs', 'marketplace_net', 'management_result', 'message')
        markets.append(['Площадка', 'Выручка', 'Расходы площадки', 'После услуг площадки', 'Результат модели', 'Данные'])
        for row in data['marketplace_totals']:
            markets.append([row.get(k) for k in keys])
    output = BytesIO(); workbook.save(output)
    return output.getvalue()
