"""Source-backed marketplace reviews dashboard for KOKOC BI."""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from typing import Any, Callable
from urllib.parse import parse_qs

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPORT_ID = "reviews"
SOURCE_LABEL = "API кабинетов · WB Feedbacks / Ozon Seller"
PAGE_SIZE_MAX = 100
EXPORT_ROW_LIMIT = 200_000


def _first_value(row: Any) -> Any:
    if isinstance(row, dict):
        return next(iter(row.values()), None)
    return row[0] if row else None


def _json_safe(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, dict):
        return {key: _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    return value


def _scalar(value: list[str] | None, default: str = "") -> str:
    return str((value or [default])[0] or default).strip()


def _parse_date(value: str) -> date | None:
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return None


def _percent(numerator: int | float | None, denominator: int | float | None) -> float | None:
    if not denominator:
        return None
    return round(float(numerator or 0) / float(denominator) * 100, 2)


def _table_available(conn) -> bool:
    with conn.cursor() as cur:
        cur.execute("SELECT to_regclass('public.marketplace_reviews') IS NOT NULL")
        return bool(_first_value(cur.fetchone()))


def _filter_context(parsed) -> dict[str, Any]:
    params = parse_qs(parsed.query)
    marketplace = _scalar(params.get("marketplace"), "total").lower()
    if marketplace not in {"total", "wb", "ozon"}:
        marketplace = "total"
    rating = _scalar(params.get("rating"), "all").lower()
    if rating not in {"all", "positive", "negative", "1", "2", "3", "4", "5"}:
        rating = "all"
    response_status = _scalar(params.get("response_status"), "all").lower()
    if response_status not in {"all", "answered", "unanswered", "unknown"}:
        response_status = "all"
    try:
        page = max(1, int(_scalar(params.get("page"), "1")))
    except ValueError:
        page = 1
    try:
        page_size = min(PAGE_SIZE_MAX, max(10, int(_scalar(params.get("page_size"), "25"))))
    except ValueError:
        page_size = 25
    return {
        "marketplace": marketplace,
        "rating": rating,
        "response_status": response_status,
        "date_from": _parse_date(_scalar(params.get("date_from"))),
        "date_to": _parse_date(_scalar(params.get("date_to"))),
        "product_id": _scalar(params.get("product_id")),
        "q": _scalar(params.get("q"))[:200],
        "page": page,
        "page_size": page_size,
    }


def _where(filters: dict[str, Any], *, include_dates: bool = True) -> tuple[str, list[Any]]:
    clauses = ["TRUE"]
    values: list[Any] = []
    if filters["marketplace"] in {"wb", "ozon"}:
        clauses.append("marketplace = %s")
        values.append(filters["marketplace"])
    if include_dates and filters["date_from"]:
        clauses.append("review_date >= %s")
        values.append(filters["date_from"])
    if include_dates and filters["date_to"]:
        clauses.append("review_date <= %s")
        values.append(filters["date_to"])
    if filters["product_id"]:
        clauses.append("product_id = %s")
        values.append(filters["product_id"])
    if filters["rating"] == "positive":
        clauses.append("rating >= 4")
    elif filters["rating"] == "negative":
        clauses.append("rating <= 3")
    elif filters["rating"] in {"1", "2", "3", "4", "5"}:
        clauses.append("rating = %s")
        values.append(int(filters["rating"]))
    if filters["response_status"] == "answered":
        clauses.append("answer_available AND answered IS TRUE")
    elif filters["response_status"] == "unanswered":
        clauses.append("answer_available AND answered IS FALSE")
    elif filters["response_status"] == "unknown":
        clauses.append("NOT answer_available")
    if filters["q"]:
        clauses.append(
            "(product_name ILIKE %s OR product_id ILIKE %s OR seller_article ILIKE %s OR review_text ILIKE %s)"
        )
        term = f"%{filters['q']}%"
        values.extend([term, term, term, term])
    return " AND ".join(clauses), values


def _empty_payload(filters: dict[str, Any], reason: str) -> dict[str, Any]:
    return {
        "ok": True,
        "report": REPORT_ID,
        "data_status": "empty",
        "message": reason,
        "source": {"label": SOURCE_LABEL, "synced_at": None, "last_run": None},
        "period": {"date_from": None, "date_to": None},
        "filters": {
            "applied": _serialize_filters(filters),
            "marketplaces": [{"id": "total", "label": "Все"}, {"id": "ozon", "label": "Ozon"}, {"id": "wb", "label": "WB"}],
            "products": [],
        },
        "summary": _summary_payload({}),
        "rating_distribution": [],
        "trend": {"grain": "month", "rows": []},
        "top_products": [],
        "feed": {"rows": [], "page": 1, "page_size": filters["page_size"], "total": 0, "total_pages": 1},
        "quality": {"status": "empty", "notes": [reason], "checks": []},
    }


def _serialize_filters(filters: dict[str, Any]) -> dict[str, Any]:
    return {
        **filters,
        "date_from": filters["date_from"].isoformat() if filters.get("date_from") else None,
        "date_to": filters["date_to"].isoformat() if filters.get("date_to") else None,
    }


def _summary_payload(row: dict[str, Any]) -> dict[str, Any]:
    total = int(row.get("total_reviews") or 0)
    negative = int(row.get("negative_reviews") or 0)
    positive = int(row.get("positive_reviews") or 0)
    text_reviews = int(row.get("text_reviews") or 0)
    answer_observed = int(row.get("answer_observed") or 0)
    answered = int(row.get("answered_reviews") or 0)
    unanswered = int(row.get("unanswered_reviews") or 0)
    return {
        "total_reviews": total,
        "average_rating": round(float(row["average_rating"]), 2) if row.get("average_rating") is not None else None,
        "negative_reviews": negative,
        "negative_share_pct": _percent(negative, total),
        "positive_reviews": positive,
        "positive_share_pct": _percent(positive, total),
        "text_reviews": text_reviews,
        "text_share_pct": _percent(text_reviews, total),
        "answer_observed": answer_observed,
        "answered_reviews": answered,
        "unanswered_reviews": unanswered,
        "response_rate_pct": _percent(answered, answer_observed),
        "unanswered_negative_reviews": int(row.get("unanswered_negative_reviews") or 0),
        "products_count": int(row.get("products_count") or 0),
    }


def _trend_grain(date_from: date | None, date_to: date | None) -> str:
    if not date_from or not date_to:
        return "month"
    days = max(0, (date_to - date_from).days)
    if days <= 45:
        return "day"
    if days <= 240:
        return "week"
    return "month"


def _last_run(conn, marketplace='total') -> dict[str, Any] | None:
    with conn.cursor() as cur:
        cur.execute("SELECT to_regclass('public.marketplace_review_sync_runs') IS NOT NULL")
        if not _first_value(cur.fetchone()):
            return None
        cur.execute(
            """
            SELECT run_id, started_at, finished_at, status, marketplaces,
                   products_total, products_processed, reviews_loaded, products_failed,
                   error_summary
            FROM public.marketplace_review_sync_runs
            WHERE (%s = 'total' OR %s = ANY(marketplaces))
            ORDER BY run_id DESC LIMIT 1
            """, (marketplace, marketplace)
        )
        row = cur.fetchone()
        return dict(row) if row else None


def dashboard_payload(parsed, get_conn: Callable, client_key: str) -> dict[str, Any]:
    filters = _filter_context(parsed)
    analytics_only = (parse_qs(parsed.query).get('view') or [''])[0] == 'analytics'
    with get_conn() as conn:
        if not _table_available(conn):
            return _empty_payload(filters, "Витрина отзывов ещё не загружена.")
        where_sql, values = _where(filters)
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT MIN(review_date) AS date_from, MAX(review_date) AS date_to,
                       MAX(source_synced_at) AS synced_at, COUNT(*) AS source_rows
                FROM public.marketplace_reviews
                WHERE (%s = 'total' OR marketplace = %s)
                """, (filters['marketplace'], filters['marketplace'])
            )
            source = dict(cur.fetchone())
            cur.execute(
                f"""
                SELECT COUNT(*) AS total_reviews,
                       AVG(rating) AS average_rating,
                       COUNT(*) FILTER (WHERE rating <= 3) AS negative_reviews,
                       COUNT(*) FILTER (WHERE rating >= 4) AS positive_reviews,
                       COUNT(*) FILTER (WHERE NULLIF(BTRIM(review_text), '') IS NOT NULL) AS text_reviews,
                       COUNT(*) FILTER (WHERE answer_available) AS answer_observed,
                       COUNT(*) FILTER (WHERE answer_available AND answered IS TRUE) AS answered_reviews,
                       COUNT(*) FILTER (WHERE answer_available AND answered IS FALSE) AS unanswered_reviews,
                       COUNT(*) FILTER (WHERE answer_available AND answered IS FALSE AND rating <= 3) AS unanswered_negative_reviews,
                       COUNT(DISTINCT (marketplace, product_id)) AS products_count
                FROM public.marketplace_reviews WHERE {where_sql}
                """,
                values,
            )
            summary_row = dict(cur.fetchone())
            cur.execute(
                f"""
                SELECT rating::integer AS rating, COUNT(*) AS review_count
                FROM public.marketplace_reviews
                WHERE {where_sql} AND rating BETWEEN 1 AND 5
                GROUP BY rating::integer ORDER BY rating::integer DESC
                """,
                values,
            )
            rating_counts = {int(row["rating"]): int(row["review_count"]) for row in cur.fetchall()}
            total_rated = sum(rating_counts.values())
            rating_distribution = [
                {"rating": rating, "review_count": rating_counts.get(rating, 0), "share_pct": _percent(rating_counts.get(rating, 0), total_rated) or 0}
                for rating in range(5, 0, -1)
            ]

            effective_from = filters["date_from"] or source.get("date_from")
            effective_to = filters["date_to"] or source.get("date_to")
            grain = _trend_grain(effective_from, effective_to)
            cur.execute(
                f"""
                SELECT date_trunc('{grain}', review_date)::date AS period,
                       COUNT(*) AS review_count,
                       AVG(rating) AS average_rating,
                       COUNT(*) FILTER (WHERE rating <= 3) AS negative_reviews,
                       COUNT(*) FILTER (WHERE answer_available) AS answer_observed,
                       COUNT(*) FILTER (WHERE answer_available AND answered IS TRUE) AS answered_reviews
                FROM public.marketplace_reviews
                WHERE {where_sql} AND review_date IS NOT NULL
                GROUP BY 1 ORDER BY 1
                """,
                values,
            )
            trend_rows = []
            for row in cur.fetchall():
                count = int(row["review_count"] or 0)
                observed = int(row["answer_observed"] or 0)
                trend_rows.append(
                    {
                        "period": row["period"].isoformat(),
                        "review_count": count,
                        "average_rating": round(float(row["average_rating"]), 2) if row["average_rating"] is not None else None,
                        "negative_share_pct": _percent(row["negative_reviews"], count),
                        "response_rate_pct": _percent(row["answered_reviews"], observed),
                    }
                )
            cur.execute(
                f"""
                SELECT marketplace, product_id, MAX(seller_article) AS seller_article,
                       MAX(product_name) AS product_name, MAX(category_name) AS category_name,
                       COUNT(*) AS review_count, AVG(rating) AS average_rating,
                       COUNT(*) FILTER (WHERE rating <= 3) AS negative_reviews,
                       COUNT(*) FILTER (WHERE answer_available AND answered IS FALSE AND rating <= 3) AS unanswered_negative_reviews
                FROM public.marketplace_reviews
                WHERE {where_sql}
                GROUP BY marketplace, product_id
                ORDER BY negative_reviews DESC, average_rating ASC NULLS LAST, review_count DESC
                LIMIT 100
                """,
                values,
            )
            top_products = []
            for row in cur.fetchall():
                count = int(row["review_count"] or 0)
                top_products.append(
                    {
                        **dict(row),
                        "review_count": count,
                        "average_rating": round(float(row["average_rating"]), 2) if row["average_rating"] is not None else None,
                        "negative_reviews": int(row["negative_reviews"] or 0),
                        "negative_share_pct": _percent(row["negative_reviews"], count),
                        "unanswered_negative_reviews": int(row["unanswered_negative_reviews"] or 0),
                    }
                )
            cur.execute(
                f"""
                SELECT marketplace, product_id, MAX(product_name) AS product_name,
                       MAX(seller_article) AS seller_article, COUNT(*) AS review_count
                FROM public.marketplace_reviews
                WHERE {_where(filters, include_dates=False)[0]}
                GROUP BY marketplace, product_id
                ORDER BY product_name NULLS LAST, marketplace, product_id
                """,
                _where(filters, include_dates=False)[1],
            )
            product_options = [dict(row) for row in cur.fetchall()]
            feed_total = int(summary_row['total_reviews'] or 0)
            offset = (filters["page"] - 1) * filters["page_size"]
            cur.execute(
                f"""
                SELECT review_key, marketplace, product_id, seller_article, product_name,
                       category_name, brand, review_date, rating, review_text,
                       pros, cons, answer_text, answer_available, answered, has_photo, likes,
                       order_status, source_synced_at
                FROM public.marketplace_reviews
                WHERE {where_sql}
                ORDER BY review_date DESC NULLS LAST, source_synced_at DESC, review_key
                LIMIT %s OFFSET %s
                """,
                [*values, 0 if analytics_only else filters["page_size"], offset],
            )
            feed_rows = [dict(row) for row in cur.fetchall()]
            cur.execute(
                """
                SELECT marketplace, COUNT(*) AS rows,
                       COUNT(*) FILTER (WHERE answer_available) AS answer_observed,
                       COUNT(*) FILTER (WHERE review_date IS NULL) AS missing_dates,
                       COUNT(*) FILTER (WHERE rating IS NULL OR rating < 1 OR rating > 5) AS invalid_ratings
                FROM public.marketplace_reviews WHERE (%s = 'total' OR marketplace = %s)
                GROUP BY marketplace ORDER BY marketplace
                """, (filters['marketplace'], filters['marketplace'])
            )
            checks = [dict(row) for row in cur.fetchall()]

        synced_at = source.get("synced_at")
        freshness_hours = None
        if synced_at:
            current = datetime.now(timezone.utc)
            aware_synced_at = synced_at if synced_at.tzinfo else synced_at.replace(tzinfo=timezone.utc)
            freshness_hours = round((current - aware_synced_at).total_seconds() / 3600, 1)
        notes = [
            "WB передаёт текст ответа продавца; для Ozon в списке отзывов доступен статус обработки, но не текст комментария.",
            "Каждый отзыв хранится по исходному ID кабинета продавца; WB и Ozon не смешиваются.",
        ]
        invalid_rows = sum(int(item["missing_dates"] or 0) + int(item["invalid_ratings"] or 0) for item in checks)
        quality_status = "attention" if invalid_rows or (freshness_hours is not None and freshness_hours > 48) else "good"
        last_run = _last_run(conn, filters['marketplace'])
        no_rows = not summary_row.get("total_reviews")
        blocked = no_rows and last_run and last_run.get("status") == "blocked"
        return _json_safe({
            "ok": True,
            "report": REPORT_ID,
            "data_status": "blocked" if blocked else "ready" if not no_rows else "empty",
            "message": "API кабинетов пока не выдали доступ к отзывам." if blocked else "",
            "source": {
                "label": SOURCE_LABEL,
                "synced_at": synced_at,
                "freshness_hours": freshness_hours,
                "source_rows": int(source.get("source_rows") or 0),
                "last_run": last_run,
            },
            "period": {
                "date_from": source["date_from"].isoformat() if source.get("date_from") else None,
                "date_to": source["date_to"].isoformat() if source.get("date_to") else None,
            },
            "filters": {
                "applied": _serialize_filters(filters),
                "marketplaces": [{"id": "total", "label": "Все"}, {"id": "ozon", "label": "Ozon"}, {"id": "wb", "label": "WB"}],
                "products": product_options,
            },
            "summary": _summary_payload(summary_row),
            "rating_distribution": rating_distribution,
            "trend": {"grain": grain, "rows": trend_rows},
            "top_products": top_products,
            "feed": {
                "rows": feed_rows,
                "page": filters["page"],
                "page_size": filters["page_size"],
                "total": feed_total,
                "total_pages": max(1, (feed_total + filters["page_size"] - 1) // filters["page_size"]),
            },
            "quality": {"status": quality_status, "notes": notes, "checks": checks},
        })


def _export_rows(conn, filters: dict[str, Any]) -> list[dict[str, Any]]:
    where_sql, values = _where(filters)
    with conn.cursor() as cur:
        cur.execute(
            f"""
            SELECT marketplace, product_id, seller_article, product_name, category_name,
                   brand, review_date, rating, review_text, pros, cons, answer_text,
                   answer_available, answered, has_photo, likes, order_status, source_synced_at
            FROM public.marketplace_reviews
            WHERE {where_sql}
            ORDER BY review_date DESC NULLS LAST, product_name, product_id
            LIMIT {EXPORT_ROW_LIMIT}
            """,
            values,
        )
        return [dict(row) for row in cur.fetchall()]


def export_workbook(parsed, get_conn: Callable, client_key: str) -> tuple[bytes, str]:
    filters = _filter_context(parsed)
    payload = dashboard_payload(parsed, get_conn, client_key)
    with get_conn() as conn:
        rows = _export_rows(conn, filters) if _table_available(conn) else []
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Сводка"
    summary_rows = [
        ("Показатель", "Значение"),
        ("Клиент", client_key),
        ("Источник", SOURCE_LABEL),
        ("Всего отзывов", payload["summary"]["total_reviews"]),
        ("Средняя оценка", payload["summary"]["average_rating"]),
        ("Доля негативных, %", payload["summary"]["negative_share_pct"]),
        ("Доля с ответом, % (только доступный статус)", payload["summary"]["response_rate_pct"]),
        ("Негативных без ответа", payload["summary"]["unanswered_negative_reviews"]),
        ("Дата от", filters["date_from"]),
        ("Дата до", filters["date_to"]),
        ("Площадка", filters["marketplace"]),
        ("Фильтр оценки", filters["rating"]),
        ("Фильтр ответа", filters["response_status"]),
    ]
    for row in summary_rows:
        summary_sheet.append(row)
    reviews_sheet = workbook.create_sheet("Отзывы")
    headers = [
        "Площадка", "ID товара", "Артикул продавца", "Товар", "Категория", "Бренд",
        "Дата", "Оценка", "Текст отзыва", "Достоинства", "Недостатки", "Ответ продавца",
        "Статус ответа доступен", "Отвечен", "Есть фото", "Лайки", "Статус заказа", "Обновлено",
    ]
    reviews_sheet.append(headers)
    for row in rows:
        reviews_sheet.append(
            [
                row["marketplace"].upper(), row["product_id"], row["seller_article"], row["product_name"],
                row["category_name"], row["brand"], row["review_date"], row["rating"], row["review_text"],
                row["pros"], row["cons"], row["answer_text"], row["answer_available"], row["answered"],
                row["has_photo"], row["likes"], row["order_status"],
                row["source_synced_at"].replace(tzinfo=None) if row["source_synced_at"] else None,
            ]
        )
    products_sheet = workbook.create_sheet("Товары")
    product_headers = ["Площадка", "ID товара", "Артикул", "Товар", "Категория", "Отзывов", "Средняя оценка", "Доля негатива, %", "Негатив без ответа"]
    products_sheet.append(product_headers)
    for row in payload["top_products"]:
        products_sheet.append(
            [row["marketplace"].upper(), row["product_id"], row["seller_article"], row["product_name"], row["category_name"], row["review_count"], row["average_rating"], row["negative_share_pct"], row["unanswered_negative_reviews"]]
        )
    notes_sheet = workbook.create_sheet("Методология")
    notes_sheet.append(["Поле", "Описание"])
    notes_sheet.append(["Гранулярность", "Один отзыв по исходному ID API кабинета продавца"])
    notes_sheet.append(["Негатив", "Отзывы с оценкой 1–3"])
    notes_sheet.append(["Позитив", "Отзывы с оценкой 4–5"])
    notes_sheet.append(["Ответы", "WB: текст и статус ответа; Ozon: статус обработки из списка отзывов, без текста комментария"])
    notes_sheet.append(["Ограничение", f"Одна выгрузка ограничена {EXPORT_ROW_LIMIT:,} строками".replace(",", " ")])

    header_fill = PatternFill("solid", fgColor="D9F4F2")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="173D3B")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            longest = min(60, max(10, max(len(str(cell.value or "")) for cell in column_cells[:200]) + 2))
            sheet.column_dimensions[letter].width = longest
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {4, 5, 9, 10})
    output = BytesIO()
    workbook.save(output)
    date_suffix = datetime.now().strftime("%Y%m%d")
    return output.getvalue(), f"km_trade_reviews_{date_suffix}.xlsx"
