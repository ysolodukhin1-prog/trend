"""Read-only Avito Ads BI payloads for the legacy PULSE interface."""
from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import parse_qs

import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from psycopg2.extras import RealDictCursor


AVITO_DASHBOARDS = {
    "avitoOverview": "overview",
    "avitoCampaigns": "campaigns",
    "avitoGroups": "groups",
    "avitoCreatives": "creatives",
    "avitoDaily": "daily",
}

SECTION_COLUMNS = {
    "campaigns": [
        ("campaign_id", "ID кампании"), ("name", "Кампания"), ("status", "Статус"),
        ("campaign_type", "Тип"), ("payment_model", "Модель оплаты"),
        ("budget_rub", "Бюджет, ₽"), ("views", "Показы"), ("clicks", "Клики"),
        ("ctr_pct", "CTR, %"), ("spend_rub", "Расход, ₽"),
        ("bonus_spend_rub", "Бонусы, ₽"), ("cpc_rub", "CPC, ₽"), ("vtr_pct", "VTR, %"),
    ],
    "groups": [
        ("group_id", "ID группы"), ("campaign_id", "ID кампании"),
        ("campaign_name", "Кампания"), ("name", "Группа"), ("status", "Статус"),
        ("budget_rub", "Бюджет, ₽"), ("price_rub", "Цена, ₽"),
        ("views", "Показы"), ("clicks", "Клики"), ("ctr_pct", "CTR, %"),
        ("spend_rub", "Расход, ₽"), ("cpc_rub", "CPC, ₽"), ("vtr_pct", "VTR, %"),
    ],
    "creatives": [
        ("creative_id", "ID креатива"), ("campaign_id", "ID кампании"),
        ("campaign_name", "Кампания"), ("group_id", "ID группы"),
        ("group_name", "Группа"), ("name", "Креатив"), ("status", "Статус"),
        ("views", "Показы"), ("clicks", "Клики"), ("ctr_pct", "CTR, %"),
        ("spend_rub", "Расход, ₽"), ("cpc_rub", "CPC, ₽"), ("vtr_pct", "VTR, %"),
    ],
    "daily": [
        ("report_date", "Дата"), ("campaign_id", "ID кампании"),
        ("campaign_name", "Кампания"), ("views", "Показы"), ("clicks", "Клики"),
        ("ctr_pct", "CTR, %"), ("spend_rub", "Расход, ₽"),
        ("bonus_spend_rub", "Бонусы, ₽"), ("cpm_rub", "CPM, ₽"),
        ("cpc_rub", "CPC, ₽"), ("video_views_25", "Видео 25%"),
        ("video_views_50", "Видео 50%"), ("video_views_75", "Видео 75%"),
        ("video_views_100", "Видео 100%"), ("q25", "Q25"), ("q50", "Q50"),
        ("q75", "Q75"), ("vtr_pct", "VTR, %"),
    ],
}


def _iso_date(value: str | None, fallback: date) -> date:
    try:
        return date.fromisoformat(str(value or ""))
    except ValueError:
        return fallback


def _plain(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date,)):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: _plain(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_plain(item) for item in value]
    return value


def _missing_ranges(present_dates: set[date], date_from: date, date_to: date) -> list[dict]:
    ranges: list[dict] = []
    cursor = date_from
    start = None
    while cursor <= date_to:
        if cursor not in present_dates and start is None:
            start = cursor
        if cursor in present_dates and start is not None:
            ranges.append({"date_from": start.isoformat(), "date_to": (cursor - timedelta(days=1)).isoformat()})
            start = None
        cursor += timedelta(days=1)
    if start is not None:
        ranges.append({"date_from": start.isoformat(), "date_to": date_to.isoformat()})
    return ranges


def _covered_dates_from_runs(runs: list[dict], date_from: date, date_to: date) -> set[date]:
    covered: set[date] = set()
    for run in runs:
        run_from = run.get("date_from")
        run_to = run.get("date_to")
        if not isinstance(run_from, date) or not isinstance(run_to, date):
            continue
        cursor = max(date_from, run_from)
        end = min(date_to, run_to)
        while cursor <= end:
            covered.add(cursor)
            cursor += timedelta(days=1)
    return covered


def _metric_sql(alias="s") -> str:
    return f"""
        sum({alias}.views)::bigint AS views,
        sum({alias}.clicks)::bigint AS clicks,
        CASE WHEN COALESCE(sum({alias}.views), 0) > 0
             THEN round(sum({alias}.clicks)::numeric * 100 / sum({alias}.views), 2) END AS ctr_pct,
        round(sum({alias}.spend_kopeks)::numeric / 100, 2) AS spend_rub,
        round(sum({alias}.spend_bonus_kopeks)::numeric / 100, 2) AS bonus_spend_rub,
        CASE WHEN COALESCE(sum({alias}.views), 0) > 0
             THEN round(sum({alias}.spend_kopeks)::numeric * 10 / sum({alias}.views), 2) END AS cpm_rub,
        CASE WHEN COALESCE(sum({alias}.clicks), 0) > 0
             THEN round(sum({alias}.spend_kopeks)::numeric / 100 / sum({alias}.clicks), 2) END AS cpc_rub,
        sum({alias}.video_views_25)::bigint AS video_views_25,
        sum({alias}.video_views_50)::bigint AS video_views_50,
        sum({alias}.video_views_75)::bigint AS video_views_75,
        sum({alias}.video_views_100)::bigint AS video_views_100,
        CASE WHEN COALESCE(sum({alias}.views), 0) > 0
             THEN round(sum({alias}.video_views_100)::numeric * 100 / sum({alias}.views), 2) END AS vtr_pct
    """


def _tables_ready(cur) -> bool:
    cur.execute(
        """SELECT to_regclass('public.avito_ads_stats_daily') AS stats,
                  to_regclass('public.avito_ads_campaigns') AS campaigns,
                  to_regclass('public.avito_ads_groups') AS groups,
                  to_regclass('public.avito_ads_creatives') AS creatives"""
    )
    row = cur.fetchone() or {}
    return all(row.get(key) for key in ("stats", "campaigns", "groups", "creatives"))


def _source_status(cur, date_from: date, date_to: date) -> dict:
    cur.execute(
        """SELECT min(report_date) AS data_from, max(report_date) AS data_to,
                  max(synced_at) AS synced_at,
                  array_agg(DISTINCT report_date ORDER BY report_date) FILTER (
                      WHERE entity_level='campaign' AND report_date BETWEEN %s AND %s
                  ) AS present_dates
           FROM public.avito_ads_stats_daily""",
        (date_from, date_to),
    )
    row = cur.fetchone() or {}
    present_dates = set(row.get("present_dates") or [])
    requested_days = (date_to - date_from).days + 1
    cur.execute("SELECT to_regclass('public.avito_ads_api_runs') AS runs")
    latest_run = None
    successful_runs: list[dict] = []
    if (cur.fetchone() or {}).get("runs"):
        cur.execute(
            """SELECT status, date_from, date_to, rows_loaded, finished_at
               FROM public.avito_ads_api_runs ORDER BY started_at DESC LIMIT 1"""
        )
        latest_run = dict(cur.fetchone() or {}) or None
        cur.execute(
            """SELECT date_from, date_to
               FROM public.avito_ads_api_runs
               WHERE status='ok' AND date_to >= %s AND date_from <= %s""",
            (date_from, date_to),
        )
        successful_runs = [dict(item) for item in cur.fetchall()]
    covered_dates = _covered_dates_from_runs(successful_runs, date_from, date_to)
    if not covered_dates and present_dates:
        covered_dates = set(present_dates)
    missing = _missing_ranges(covered_dates, date_from, date_to)
    status = "complete" if covered_dates and not missing else ("partial" if covered_dates else ("empty" if latest_run else "not_started"))
    return {
        "status": status,
        "requested_days": requested_days,
        "covered_days": len(covered_dates),
        "days_with_data": len(present_dates),
        "coverage_pct": round(len(covered_dates) * 100 / requested_days, 1) if requested_days else None,
        "missing_ranges": missing,
        "data_from": row.get("data_from"),
        "data_to": row.get("data_to"),
        "synced_at": row.get("synced_at"),
        "last_run": latest_run,
        "note": "Покрытие считается по успешным окнам загрузки; дни без строк остаются подтверждённо пустыми и не заменяются нулём.",
    }


def _summary(cur, date_from: date, date_to: date) -> dict:
    cur.execute(
        f"""SELECT {_metric_sql()}
             FROM public.avito_ads_stats_daily s
             WHERE s.entity_level='campaign' AND s.report_date BETWEEN %s AND %s""",
        (date_from, date_to),
    )
    result = dict(cur.fetchone() or {})
    cur.execute(
        """SELECT
             (SELECT count(*) FROM public.avito_ads_campaigns) AS campaigns,
             (SELECT count(*) FROM public.avito_ads_groups) AS groups,
             (SELECT count(*) FROM public.avito_ads_creatives) AS creatives,
             (SELECT round(balance_kopeks::numeric / 100, 2)
                FROM public.avito_ads_balances_daily ORDER BY snapshot_date DESC LIMIT 1) AS balance_rub,
             (SELECT round(bonus_balance_kopeks::numeric / 100, 2)
                FROM public.avito_ads_balances_daily ORDER BY snapshot_date DESC LIMIT 1) AS bonus_balance_rub"""
    )
    result.update(dict(cur.fetchone() or {}))
    return result


def _daily(cur, date_from: date, date_to: date) -> list[dict]:
    cur.execute(
        f"""SELECT s.report_date, {_metric_sql()}
             FROM public.avito_ads_stats_daily s
             WHERE s.entity_level='campaign' AND s.report_date BETWEEN %s AND %s
             GROUP BY s.report_date ORDER BY s.report_date""",
        (date_from, date_to),
    )
    return [dict(row) for row in cur.fetchall()]


def _detail_rows(cur, section: str, date_from: date, date_to: date) -> list[dict]:
    metrics = _metric_sql()
    if section == "campaigns":
        query = f"""SELECT c.campaign_id, c.name, c.status, c.campaign_type, c.payment_model,
                            round(c.budget_kopeks::numeric / 100, 2) AS budget_rub, {metrics}
                     FROM public.avito_ads_campaigns c
                     LEFT JOIN public.avito_ads_stats_daily s ON s.account_id=c.account_id
                       AND s.campaign_id=c.campaign_id AND s.entity_level='campaign'
                       AND s.report_date BETWEEN %s AND %s
                     GROUP BY c.account_id, c.campaign_id, c.name, c.status, c.campaign_type,
                              c.payment_model, c.budget_kopeks
                     ORDER BY COALESCE(sum(s.spend_kopeks), 0) DESC, c.campaign_id LIMIT 1000"""
    elif section == "groups":
        query = f"""SELECT g.group_id, g.campaign_id, c.name AS campaign_name, g.name, g.status,
                            round(g.budget_kopeks::numeric / 100, 2) AS budget_rub,
                            round(g.price_kopeks::numeric / 100, 2) AS price_rub, {metrics}
                     FROM public.avito_ads_groups g
                     LEFT JOIN public.avito_ads_campaigns c ON c.account_id=g.account_id AND c.campaign_id=g.campaign_id
                     LEFT JOIN public.avito_ads_stats_daily s ON s.account_id=g.account_id
                       AND s.group_id=g.group_id AND s.entity_level='group'
                       AND s.report_date BETWEEN %s AND %s
                     GROUP BY g.account_id, g.group_id, g.campaign_id, c.name, g.name, g.status,
                              g.budget_kopeks, g.price_kopeks
                     ORDER BY COALESCE(sum(s.spend_kopeks), 0) DESC, g.group_id LIMIT 1000"""
    elif section == "creatives":
        query = f"""SELECT cr.creative_id, cr.campaign_id, c.name AS campaign_name,
                            cr.group_id, g.name AS group_name, cr.name, cr.status, {metrics}
                     FROM public.avito_ads_creatives cr
                     LEFT JOIN public.avito_ads_campaigns c ON c.account_id=cr.account_id AND c.campaign_id=cr.campaign_id
                     LEFT JOIN public.avito_ads_groups g ON g.account_id=cr.account_id AND g.group_id=cr.group_id
                     LEFT JOIN public.avito_ads_stats_daily s ON s.account_id=cr.account_id
                       AND s.creative_id=cr.creative_id AND s.entity_level='creative'
                       AND s.report_date BETWEEN %s AND %s
                     GROUP BY cr.account_id, cr.creative_id, cr.campaign_id, c.name,
                              cr.group_id, g.name, cr.name, cr.status
                     ORDER BY COALESCE(sum(s.spend_kopeks), 0) DESC, cr.creative_id LIMIT 1000"""
    else:
        query = """SELECT s.report_date, s.campaign_id, c.name AS campaign_name,
                          s.views, s.clicks,
                          CASE WHEN s.views > 0 THEN round(s.clicks::numeric * 100 / s.views, 2) END AS ctr_pct,
                          round(s.spend_kopeks::numeric / 100, 2) AS spend_rub,
                          round(s.spend_bonus_kopeks::numeric / 100, 2) AS bonus_spend_rub,
                          CASE WHEN s.views > 0 THEN round(s.spend_kopeks::numeric * 10 / s.views, 2) END AS cpm_rub,
                          CASE WHEN s.clicks > 0 THEN round(s.spend_kopeks::numeric / 100 / s.clicks, 2) END AS cpc_rub,
                          s.video_views_25, s.video_views_50, s.video_views_75, s.video_views_100,
                          s.q25, s.q50, s.q75,
                          CASE WHEN s.views > 0 THEN round(s.video_views_100::numeric * 100 / s.views, 2) END AS vtr_pct
                   FROM public.avito_ads_stats_daily s
                   LEFT JOIN public.avito_ads_campaigns c ON c.account_id=s.account_id AND c.campaign_id=s.campaign_id
                   WHERE s.entity_level='campaign' AND s.report_date BETWEEN %s AND %s
                   ORDER BY s.report_date DESC, s.spend_kopeks DESC NULLS LAST, s.campaign_id LIMIT 2000"""
    cur.execute(query, (date_from, date_to))
    return [dict(row) for row in cur.fetchall()]


def dashboard_payload(parsed, db_config: dict, *, connect=psycopg2.connect) -> dict:
    params = parse_qs(parsed.query)
    dashboard = str((params.get("dashboard") or ["avitoOverview"])[0])
    section = AVITO_DASHBOARDS.get(dashboard, "overview")
    today = date.today()
    date_to = _iso_date((params.get("date_to") or [None])[0], today)
    date_from = _iso_date((params.get("date_from") or [None])[0], date_to - timedelta(days=29))
    if date_from > date_to:
        raise ValueError("Дата начала Avito-отчёта позже даты окончания")
    with connect(**db_config, cursor_factory=RealDictCursor) as conn, conn.cursor() as cur:
        if not _tables_ready(cur):
            return {
                "ok": True, "available": False, "section": section,
                "date_from": date_from, "date_to": date_to,
                "summary": {}, "daily": [], "trend_rows": [], "rows": [], "columns": [],
                "coverage": {"status": "not_started", "note": "Инфраструктура Avito Ads ещё не создана для этого клиента."},
            }
        coverage = _source_status(cur, date_from, date_to)
        summary = _summary(cur, date_from, date_to)
        if not coverage.get("last_run"):
            for key in ("campaigns", "groups", "creatives"):
                if summary.get(key) == 0:
                    summary[key] = None
        daily = _daily(cur, date_from, date_to)
        detail_section = "campaigns" if section == "overview" else section
        rows = _detail_rows(cur, detail_section, date_from, date_to)
        trend_rows = rows if detail_section == "daily" else _detail_rows(cur, "daily", date_from, date_to)
        columns = SECTION_COLUMNS[detail_section]
    return _plain({
        "ok": True, "available": True, "section": section,
        "date_from": date_from, "date_to": date_to,
        "summary": summary, "daily": daily, "trend_rows": trend_rows, "rows": rows,
        "columns": [{"key": key, "label": label} for key, label in columns],
        "coverage": coverage,
        "grain": "campaign" if section in {"overview", "campaigns", "daily"} else section.rstrip("s"),
    })


def _filtered_export_rows(payload: dict, params: dict) -> list[dict]:
    rows = list(payload.get("rows") or [])
    text_query = str((params.get("q") or [""])[0]).strip().casefold()
    filters = {
        "status": str((params.get("status") or [""])[0]).strip(),
        "campaign_id": str((params.get("campaign_id") or [""])[0]).strip(),
        "campaign_name": str((params.get("campaign_name") or [""])[0]).strip(),
        "group_name": str((params.get("group_name") or [""])[0]).strip(),
        "campaign_type": str((params.get("campaign_type") or [""])[0]).strip(),
        "payment_model": str((params.get("payment_model") or [""])[0]).strip(),
    }
    result = []
    for row in rows:
        if text_query and text_query not in " ".join(str(value or "") for value in row.values()).casefold():
            continue
        if any(expected and str(row.get(key) or "") != expected for key, expected in filters.items()):
            continue
        result.append(row)
    return result


def export_workbook(parsed, db_config: dict, *, connect=psycopg2.connect) -> tuple[bytes, str]:
    """Build a read-only XLSX from the exact Avito report filters in the query."""
    payload = dashboard_payload(parsed, db_config, connect=connect)
    params = parse_qs(parsed.query)
    rows = _filtered_export_rows(payload, params)
    columns = payload.get("columns") or []
    workbook = Workbook()
    report = workbook.active
    report.title = "Данные"
    report.append([column["label"] for column in columns])
    for row in rows:
        report.append([row.get(column["key"]) for column in columns])
    report.freeze_panes = "A2"
    report.auto_filter.ref = report.dimensions
    header_fill = PatternFill("solid", fgColor="E9F3F0")
    for cell in report[1]:
        cell.font = Font(name="Montserrat", size=10, bold=True, color="24322E")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for row in report.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Montserrat", size=9, color="24322E")
            cell.alignment = Alignment(vertical="top")
    for index, column in enumerate(columns, start=1):
        values = [str(column["label"])] + [str(row.get(column["key"]) or "") for row in rows[:250]]
        report.column_dimensions[get_column_letter(index)].width = min(max(max(map(len, values)) + 2, 11), 42)

    meta = workbook.create_sheet("Параметры")
    meta_rows = [
        ("Отчёт", payload.get("section") or "—"),
        ("Уровень", payload.get("grain") or "—"),
        ("Период", f"{payload.get('date_from')} — {payload.get('date_to')}"),
        ("Строк после фильтров", len(rows)),
        ("Покрытие", (payload.get("coverage") or {}).get("status") or "—"),
        ("Проверено дней", f"{(payload.get('coverage') or {}).get('covered_days', 0)}/{(payload.get('coverage') or {}).get('requested_days', 0)}"),
        ("Источник", "Avito Ads API · read-only"),
    ]
    for key, values in params.items():
        if key in {"client", "dashboard", "date_from", "date_to"} or not values or not values[0]:
            continue
        meta_rows.append((f"Фильтр: {key}", values[0]))
    for item in meta_rows:
        meta.append(item)
    meta.column_dimensions["A"].width = 26
    meta.column_dimensions["B"].width = 52
    for cell in meta[1]:
        cell.font = Font(name="Montserrat", bold=True)

    output = BytesIO()
    workbook.save(output)
    dashboard = str((params.get("dashboard") or ["avitoOverview"])[0])
    return output.getvalue(), f"{dashboard}_{payload.get('date_from')}_{payload.get('date_to')}.xlsx"
