"""Read-only P&L cost provenance, dated/current scenarios and WB operation drilldown.

The existing /pl contract remains the default. Only explicit workbench requests
use these models; current operational ex-VAT prices never become historical facts.
"""
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from pathlib import PureWindowsPath

D = lambda value: Decimal(str(value or 0))
SIGN = "CASE WHEN lower(seller_oper_name) IN ('возврат','return') THEN -1 WHEN lower(seller_oper_name) IN ('продажа','sale') THEN 1 ELSE 0 END"
EXPRESSIONS = {
    "revenue": f"retail_amount*({SIGN})",
    "commission": f"-(retail_amount-for_pay-acquiring_fee)*({SIGN})",
    "acquiring": f"-acquiring_fee*({SIGN})",
    "delivery": "-delivery_service", "storage": "-paid_storage",
    "acceptance": "-paid_acceptance", "deduction": "-deduction",
    "penalty": "-penalty", "additional_payment": "additional_payment-cashback_amount",
    "cashback_discount": "coalesce(nullif(coalesce(raw_payload->>'cashbackDiscount',raw_payload->>'cashback_discount'),'')::numeric,0)*CASE WHEN lower(coalesce(raw_payload->>'docTypeName',raw_payload->>'doc_type_name')) IN ('возврат','return') THEN -1 ELSE 1 END",
    "rebill_logistics": "-rebill_logistic_cost",
}
MODES = {"dated": "По датам операций", "current": "По текущему прайсу", "estimate": "Оценка: выручка ÷ 3"}
ARTICLE_LABELS = dict(revenue="Продажи и возвраты", commission="Комиссия", acquiring="Эквайринг",
    delivery="Логистика", storage="Хранение", acceptance="Приёмка", deduction="Удержания",
    penalty="Штрафы", additional_payment="Доплаты за вычетом участия в лояльности", cashback_discount="Компенсация скидки по программе лояльности", rebill_logistics="Возмещение логистических издержек")


def separate_unallocated_products(payload):
    """WB nm_id=0 is a source attribution gap, never a sellable product."""
    products = payload.get("products", [])
    payload["unallocated_products"] = [p for p in products if not str(p.get("sku") or "").strip() or str(p["sku"]).strip() == "0"]
    payload["products"] = [p for p in products if p not in payload["unallocated_products"]]
    payload["products_total"] = len(payload["products"])


def safe_source(raw):
    parts = str(raw or "").split("|")
    return " · ".join([PureWindowsPath(parts[0]).name] + [p for p in parts[1:] if p.startswith("TDSheet!") or p == "operational_ex_vat"])


def choose_cost(versions, day, mode, returned=False):
    if mode == "dated" and returned:
        return None, "return_basis_missing"
    eligible = [r for r in versions if mode == "current" or str(r["valid_from"]) <= str(day)]
    if not eligible:
        return None, "before_first_cost" if versions else "no_match"
    latest = max(str(r["valid_from"]) for r in eligible)
    rows = [r for r in eligible if str(r["valid_from"]) == latest]
    if len({D(r["amount"]) for r in rows}) != 1:
        return None, "conflict"
    return rows[0], "sourced"


def summarize(rows):
    sales = [r for r in rows if r["sale_units"] > 0]
    sold = sum((D(r["sale_units"]) for r in sales), Decimal(0))
    covered = sum((D(r["covered_sale_units"]) for r in sales), Decimal(0))
    revenue = sum((D(r["sale_revenue"]) for r in sales), Decimal(0))
    covered_revenue = sum((D(r["covered_sale_revenue"]) for r in sales), Decimal(0))
    skus = {r["sku"] for r in sales}
    covered_skus = {sku for sku in skus if all(r["cost_complete"] for r in sales if r["sku"] == sku)}
    return dict(sold_units=float(sold), covered_units=float(covered),
                units_pct=float(covered/sold*100) if sold else None,
                sale_revenue=float(revenue), covered_revenue=float(covered_revenue),
                revenue_pct=float(covered_revenue/revenue*100) if revenue else None,
                sku_total=len(skus), sku_covered=len(covered_skus),
                sku_pct=len(covered_skus)/len(skus)*100 if skus else None,
                return_units=sum(r["return_units"] for r in rows),
                missing_return_units=sum(r["return_units"] for r in rows if not r["cost_complete"]),
                complete=all(r["cost_complete"] for r in rows))


def cost_rows(events, registry, mode):
    """Keep exact WB size barcode identities; count sales and returns separately."""
    groups = {}
    daily = defaultdict(lambda: {"known_cogs": Decimal(0), "complete": True})
    for r in events:
        sku, barcode, day = str(r["sku"] or ""), str(r["barcode"] or ""), str(r["day"])
        key = (sku, barcode)
        if key not in groups:
            groups[key] = dict(sku=sku, barcode=barcode, article=r.get("article") or sku,
                name=r.get("name") or r.get("article") or sku, units=0., revenue=0.,
                known_cogs=0., sale_units=0., covered_sale_units=0., sale_revenue=0.,
                covered_sale_revenue=0., return_units=0., cost_complete=True,
                sources=set(), dates=set(), statuses=set(), unit_costs=set())
        g = groups[key]
        units, revenue = D(r["units"]), D(r["revenue"])
        returned = int(r["sign"]) < 0
        signed_units = units * int(r["sign"])
        g["units"] += float(signed_units); g["revenue"] += float(revenue)
        if returned: g["return_units"] += float(abs(units))
        else:
            g["sale_units"] += float(abs(units)); g["sale_revenue"] += float(abs(revenue))
        version, status = choose_cost(registry.get(barcode, []), day, mode, returned)
        if mode == "estimate":
            cost, status = revenue/3, "estimated"
        elif units <= 0:
            cost, status = None, "invalid_quantity"
        else:
            cost = D(version["amount"])*signed_units if version else None
        if version and mode != "estimate":
            g["dates"].add(str(version["valid_from"]))
            g["sources"].add(safe_source(version["source_ref"]))
            g["unit_costs"].add(float(version["amount"]))
        g["statuses"].add(status)
        if cost is None:
            g["cost_complete"] = False; daily[day]["complete"] = False
        else:
            g["known_cogs"] += float(cost); daily[day]["known_cogs"] += cost
            if not returned:
                g["covered_sale_units"] += float(abs(units))
                g["covered_sale_revenue"] += float(abs(revenue))
    result = []
    for g in groups.values():
        g["cogs"] = g["known_cogs"] if g["cost_complete"] else None
        for key in ("sources", "dates", "statuses", "unit_costs"): g[key] = sorted(g[key])
        result.append(g)
    return result, daily


def attach_workbench(config, payload, mode="dated"):
    if mode not in MODES:
        raise ValueError("Неизвестный режим себестоимости")
    if payload.get("available") is False or payload.get("marketplace") != "wb":
        return payload
    if config.get("database") != payload.get("client"):
        raise ValueError("Клиент P&L не совпадает с базой данных")
    from km_trade_finance import connect_km
    start, end = date.fromisoformat(payload["date_from"]), date.fromisoformat(payload["date_to"])
    if (end-start).days > 366 or end < start:
        raise ValueError("Выберите период до 367 дней")
    registry = defaultdict(list)
    with connect_km(config) as conn, conn.cursor() as cur:
        cur.execute("SELECT to_regclass('public.unit_cogs_versions') AS relation")
        if cur.fetchone()["relation"]:
            cur.execute("SELECT barcode,amount,valid_from,source_ref FROM public.unit_cogs_versions ORDER BY barcode,valid_from,created_at")
            for r in cur.fetchall(): registry[str(r["barcode"])].append(dict(r))
        cur.execute(f"""SELECT operation_date::text AS day,nm_id::text AS sku,sku AS barcode,
            max(vendor_code) AS article,max(title) AS name,({SIGN}) AS sign,
            sum(quantity) AS units,sum(retail_amount*({SIGN})) AS revenue
            FROM public.wb_finance_lines
            WHERE operation_date BETWEEN %s AND %s AND (retail_amount<>0 OR for_pay<>0)
            GROUP BY operation_date,nm_id,sku,({SIGN})""", (start, end))
        events = [dict(r) for r in cur.fetchall()]
        columns = ",".join(f"sum({expr}) AS {key}" for key, expr in EXPRESSIONS.items())
        cur.execute(f"SELECT operation_date::text AS day,count(*) events,{columns} FROM public.wb_finance_lines WHERE operation_date BETWEEN %s AND %s GROUP BY operation_date ORDER BY operation_date", (start, end))
        days = [dict(r) for r in cur.fetchall()]
        cur.execute(f"SELECT {columns} FROM public.wb_finance_lines WHERE operation_date BETWEEN %s AND %s AND (nm_id IS NULL OR nm_id=0)", (start, end))
        unattributed = cur.fetchone() or {}
        unallocated_articles = [dict(key=k,label=ARTICLE_LABELS[k],amount=float(D(unattributed.get(k)))) for k in EXPRESSIONS if D(unattributed.get(k)) != 0]
    if not days:
        payload.update(available=False, message="Нет операций WB за выбранный период")
        return payload
    costs, cost_days = cost_rows(events, registry, mode)
    coverage = summarize(costs)
    t = payload["totals"]
    original = {k:t.get(k) for k in ("cogs", "management_result", "management_margin_pct")}
    totals_cost = sum((D(r["known_cogs"]) for r in costs), Decimal(0))
    cogs = float(totals_cost) if coverage["complete"] else None
    before = (D(t["marketplace_net"])-D(cogs)-D(t.get("seller_unit_costs"))-D(t.get("manual_expenses"))-D(t.get("vat"))) if cogs is not None else None
    result = float(before-D(t.get("tax"))) if before is not None else None
    # Historical returns lack an original-sale cost reference. Do not credit
    # a cost from the return date as if it were the original cost of goods.
    notice = ("Расчёт по версиям затрат на даты операций. Возвраты без исходной себестоимости не оценены."
              if mode == "dated" else "Сценарий по текущему прайсу; не фактическая прибыль прошлого периода."
              if mode == "current" else "Оценка себестоимости: выручка ÷ 3.")
    t.update(cogs=cogs, known_cogs=float(totals_cost), gross_profit=float(D(t["revenue"])-D(cogs)) if cogs is not None else None,
             profit_before_tax=float(before) if before is not None else None, management_result=result,
             management_margin_pct=result/float(t["revenue"])*100 if result is not None and t["revenue"] else None,
             cogs_coverage_units_pct=coverage["units_pct"], cogs_model_coverage_pct=coverage["units_pct"],
             estimated_sku_count=len({r["sku"] for r in costs}) if mode=="estimate" else 0,
             net_profit=None, margin_pct=None, profit_ready=False)
    payload["model_notice"] = notice
    payload.setdefault("methodology", {}).update(cogs=notice,
        warning="Источник операционный без НДС. Состав затрат и налоговая применимость не подтверждены; модель не является фактической чистой прибылью.",
        coverage="Покрытие единиц и выручки считается по продажам до возвратов; возвраты и их пропуски показаны отдельно. SKU покрыт только при покрытии всех его продаж.")
    by_sku = defaultdict(list)
    for r in costs: by_sku[r["sku"]].append(r)
    for p in payload["products"]:
        matched = by_sku[str(p["sku"])]
        complete = all(r["cost_complete"] for r in matched)
        cost = sum(r["known_cogs"] for r in matched) if complete else None
        value = float(D(p["revenue"])-D(p["marketplace_costs_before_cogs"])-D(cost)-D(p.get("seller_costs"))-D(p.get("vat"))-D(p.get("income_tax"))) if cost is not None else None
        p.update(cogs=cost, known_cogs=sum(r["known_cogs"] for r in matched), cogs_known=complete and mode!="estimate",
                 cogs_source=mode if complete else "missing", gross_profit=float(D(p["revenue"])-D(cost)) if cost is not None else None,
                 management_result=value, profit_before_common_costs=value, net_profit=None,
                 margin_pct=value/float(p["revenue"])*100 if value is not None and p["revenue"] else None)
    # Preserve missing days. A day without source events is unknown, not zero.
    daily = []
    present = {r["day"]:r for r in days}
    expenses = payload.get("all_expenses", [])
    from pulse_financial_model import allocate_expense
    day = start
    while day <= end:
        row = present.get(str(day))
        if row:
            revenue = D(row["revenue"])
            net = sum((D(row[k]) for k in EXPRESSIONS), Decimal(0))
            dc = cost_days.get(str(day), {"known_cogs":Decimal(0), "complete":True})
            manual = sum((allocate_expense(x["source_amount"],date.fromisoformat(x["date_from"]),date.fromisoformat(x["date_to"]),day,day) for x in expenses),Decimal(0))
            rate = D(payload.get("taxes",{}).get("income_tax_pct"))+D(payload.get("taxes",{}).get("vat_pct"))
            value = float(net-dc["known_cogs"]-manual-revenue*rate/100) if dc["complete"] else None
            daily.append(dict(date=str(day),revenue=float(revenue),marketplace_net=float(net),expenses=float(revenue-net),
                              cogs=float(dc["known_cogs"]) if dc["complete"] else None,management_result=value,events=row["events"]))
        else:
            daily.append(dict(date=str(day),revenue=None,marketplace_net=None,expenses=None,cogs=None,management_result=None,events=None))
        day += timedelta(days=1)
    for month in payload.get("monthly",[]):
        ds = [r for r in daily if r["date"].startswith(str(month["month"])[:7]) and r["events"] is not None]
        cm = sum(r["cogs"] for r in ds) if all(r["cogs"] is not None for r in ds) else None
        pre = (D(month["marketplace_net"])-D(cm)-D(month.get("manual_expenses"))-D(month.get("seller_costs"))-D(month.get("vat"))) if cm is not None else None
        mr = float(pre-D(month["revenue"])*D(payload.get("taxes",{}).get("income_tax_pct"))/100) if pre is not None else None
        month.update(cogs=cm,cogs_complete=cm is not None,profit_before_tax=float(pre) if pre is not None else None,
                     management_result=mr,net_profit=None,gross_profit=float(D(month["revenue"])-D(cm)) if cm is not None else None,
                     margin_pct=mr/float(month["revenue"])*100 if mr is not None and month["revenue"] else None)
    values = {"cogs":-cogs if cogs is not None else None,"gross_profit":t["gross_profit"],
              "profit_before_tax":t["profit_before_tax"],"net_profit":None,"management_result":result}
    for r in payload["statement"]:
        if r["key"] in values: r["amount"] = values[r["key"]]
        if r["key"] == "cogs": r.update(label="Себестоимость",is_partial=not coverage["complete"],is_estimated=mode!="dated")
        if r["key"] == "management_result": r["label"]="Результат модели" + (" до налога" if not t.get("tax_configured") else "")
        r["revenue_pct"] = r["amount"]/float(t["revenue"])*100 if r.get("amount") is not None and t["revenue"] else None
    payload["months"] = payload.get("monthly", [])
    separate_unallocated_products(payload)
    product_net = sum((D(p["revenue"])-D(p["marketplace_costs_before_cogs"]) for p in payload["products"]),Decimal(0))
    payload["workbench"] = dict(mode=mode,label=MODES[mode],coverage=coverage,costs=costs,daily=daily,
        registry_barcodes=len(registry),registry_dates=sorted({str(r["valid_from"]) for rs in registry.values() for r in rs}),
        previous_estimate=original,
        unallocated_marketplace_net=float(D(t["marketplace_net"])-product_net),
        unallocated_articles=unallocated_articles,
        source_days=sum(r["events"] is not None for r in daily),selected_days=len(daily),
        completeness_confirmed=False,
        source_notice="Дни с операциями не доказывают полноту загрузки. Сравнение наблюдаемых сумм требует проверки источника.",
        cost_notice="Операционная себестоимость без НДС: состав затрат требует проверки для исключения двойного учёта.",
        returns_notice="Для расчёта по датам возврат требует себестоимости исходной продажи; текущий прайс используется только в сценарии.")
    payload["reconciliation"] = dict(monthly_management_result_difference=(round(result-sum(m["management_result"] for m in payload.get("monthly",[])),6)
        if result is not None and all(m["management_result"] is not None for m in payload.get("monthly",[])) else None))
    return payload


def operations(config, start, end, article="deduction", query="", sku="", page=1, limit=100):
    from km_trade_finance import connect_km
    if article not in EXPRESSIONS: raise ValueError("Для этой статьи нет операций WB")
    start, end = date.fromisoformat(start), date.fromisoformat(end)
    if end < start or (end-start).days>366: raise ValueError("Выберите период до 367 дней")
    page, limit = max(1,int(page)),max(1,min(200,int(limit)))
    expr = EXPRESSIONS[article]
    where = f"operation_date BETWEEN %s AND %s AND ({expr})<>0"
    args = [start,end]
    if sku == "unallocated": where += " AND (nm_id IS NULL OR nm_id=0)"
    elif sku: where += " AND nm_id::text=%s"; args.append(str(sku))
    if query:
        where += " AND concat_ws(' ',vendor_code,title,rrd_id::text,report_id::text,seller_oper_name) ILIKE %s"
        args.append("%"+query[:200]+"%")
    with connect_km(config) as conn, conn.cursor() as cur:
        cur.execute(f"SELECT count(*) count,coalesce(sum({expr}),0) amount FROM wb_finance_lines WHERE {where}",args)
        total = dict(cur.fetchone())
        cur.execute(f"""SELECT rrd_id::text id,report_id::text report_id,operation_date::text date,
            nm_id::text sku,sku barcode,vendor_code article,title name,seller_oper_name operation,
            ({expr}) amount FROM wb_finance_lines WHERE {where}
            ORDER BY operation_date DESC,rrd_id DESC LIMIT %s OFFSET %s""",args+[limit,(page-1)*limit])
        rows = [dict(r, amount=float(r["amount"])) for r in cur.fetchall()]
    return dict(rows=rows,total=int(total["count"]),amount=float(total["amount"]),page=page,limit=limit,article=article,
                source="WB Finance API · строки отчёта реализации")


def workbench_workbook(data):
    """Export the exact selected model and its dated cost provenance."""
    from io import BytesIO
    from openpyxl import load_workbook
    from marketplace_pl import workbook_bytes
    if data.get("financial_model") and not data.get("workbench"):
        from pl_multi_model import model_workbook
        return model_workbook(data)
    raw = workbook_bytes(data)
    if not data.get("workbench"):
        return raw
    wb = load_workbook(BytesIO(raw))
    w = data["workbench"]
    common = wb.create_sheet("Без привязки к SKU")
    common.append(["Операции WB без указанного товара; уже учтены в итогах P&L"])
    common.append(["Статья", "Влияние на результат, ₽"])
    for row in w.get("unallocated_articles", []): common.append([row["label"], row["amount"]])
    common.append(["После услуг WB: вне товаров", w["unallocated_marketplace_net"]])
    common.freeze_panes = "A3"
    products = wb["По товарам"]
    for index, row in enumerate(data["products"], 2):
        products.cell(index, products.max_column, w["label"] if row.get("cogs") is not None else "Не полная себестоимость")
    provenance = wb.create_sheet("Основания затрат")
    provenance.append(["Режим", w["label"]])
    provenance.append(["Ограничение", data["methodology"]["warning"]])
    provenance.append(["Возвраты", w["returns_notice"]])
    provenance.append(["Покрытие проданных единиц, %",w["coverage"]["units_pct"]])
    provenance.append(["Источник", "EAN", "SKU", "Артикул", "Даты действия", "Нетто шт", "Себестоимость ₽", "Известная часть ₽", "Состояние"])
    for row in w["costs"]:
        provenance.append(["; ".join(row["sources"]),row["barcode"],row["sku"],row["article"],"; ".join(row["dates"]),row["units"],row["cogs"],row["known_cogs"],"; ".join(row["statuses"])])
    provenance.freeze_panes = "A6"
    provenance.auto_filter.ref = f"A5:I{provenance.max_row}"
    for row in provenance:
        for cell in row:
            if isinstance(cell.value,str): cell.data_type = "s"
    wb["P&L"].append([data["methodology"]["warning"]])
    if data.get("financial_model"):
        model=data["financial_model"]
        # Supersede the old incomplete summary with the same model shown in the UI.
        summary=wb["P&L"];summary.delete_rows(1,summary.max_row)
        summary.append(["Статья","Итого, ₽ / шт / %"])
        for row in model["rows"]: summary.append([row["label"],model["total"]["values"].get(row["key"])])
        summary.append(["Текущий справочник; расчёт по параметрам юнитки. Операционные расходы — внесённые записи."])
        if "По месяцам" in wb: del wb["По месяцам"]
        provenance.title="Диагностика EAN"
        trace=wb.create_sheet("Себестоимость расчёта")
        trace.append(["Дата","SKU","Артикул","EAN","Нетто, шт","Цена, ₽","Себестоимость, ₽","Источник"])
        for r in model.get("cost_events",[]): trace.append([r.get(k) for k in ("date","sku","article","barcode","units","unit_cost","cost","source")])
        trace.freeze_panes="A2"
        for row in trace:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type="s"
        for kind,title in (("day","Финмодель по дням"),("week","Финмодель по неделям"),("month","Финмодель по месяцам")):
            sheet=wb.create_sheet(title,0 if kind=="week" else len(wb.sheetnames))
            periods=model["periods"][kind]
            sheet.append(["Статья","Итого"]+[p["from_date"]+" — "+p["to_date"] for p in periods])
            for row in model["rows"]:
                sheet.append([row["label"],model["total"]["values"].get(row["key"])]+[p["values"].get(row["key"]) for p in periods])
            sheet.append(["Основа: текущий справочник; налоги — параметры юнитки. Дни без источника не считаются подтверждёнными нулями."])
            sheet.append(["Без цены, шт",model["total"]["missing_cost_units"]])
            sheet.append(["По единой цене SKU, шт",model["total"]["sku_cost_units"]])
            sheet.freeze_panes="C2";sheet.column_dimensions["A"].width=65
            from openpyxl.utils import get_column_letter
            for c in range(2,sheet.max_column+1):sheet.column_dimensions[get_column_letter(c)].width=24
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value,str):cell.data_type="s"
                    elif isinstance(cell.value,(int,float)):cell.number_format='#,##0.00;[Red]-#,##0.00'
        wb.active=0
    out=BytesIO(); wb.save(out)
    return out.getvalue()
