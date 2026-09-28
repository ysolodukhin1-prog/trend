"""Runtime XLSX export using the same openpyxl stack as existing PULSE exports."""
from io import BytesIO
from decimal import Decimal,ROUND_HALF_UP
import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter


def export_plan(record):
    data=record['payload'];snapshot=data['snapshot'];result=data['result']
    workbook=Workbook();summary=workbook.active;summary.title='План поставок'
    def write(ws,row):
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'  # SKU =,+,-,@ never becomes an Excel formula
    write(summary,['PULSE · внутренний план поставок',record['plan_id']])
    for label,value in [('Версия',record['revision']),('Клиент',snapshot['client']),('Автор',record['actor']),
        ('Дата расчёта',result['as_of']),('Дата отгрузки',result['ship_date']),('Горизонт, дней',result['horizon_days']),
        ('Основа',result['basis']),('Тип','Сценарная оценка до налогов'),('Модель',snapshot['model_version']),
        ('Контрольная сумма',data['fingerprint']),('Количество, шт.',result['quantity']),
        ('Бюджет партии, ₽',float(result['cash_required'])),('Результат до налогов, ₽',float(result['result'])),
        ('Δ к центру, ₽',None if result['delta_center'] is None else float(result['delta_center'])),
        ('Поиск','Лучший из проверенных кандидатов; глобальный оптимум не доказан')]:write(summary,[label,value])
    lines=workbook.create_sheet('Состав и экономика')
    headers=['Площадка','Схема','SKU','Размер','Физический SKU','Кластер','Маршрут','Количество, шт.',
        'Выкупы, шт.','Остаток в конце, шт.','Выручка, ₽','Все расходы без фрахта, ₽',
        'Доля фрахта партии, ₽','Результат строки, ₽','Источник спроса','Источник запаса','Источник финансов']
    write(lines,headers)
    # Allocate freight once by quantity, with an exact cent residual on the last
    # member. Allocation explains the group total; it is not an extra charge.
    shares={};selections=result['selection']
    for rid,load in result['route_loads'].items():
        members=[i for i,o in enumerate(selections) if o['route_id']==rid and o['quantity']]
        remainder=Decimal(load['freight']).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
        for position,index in enumerate(members):
            value=remainder if position==len(members)-1 else (Decimal(load['freight'])*selections[index]['quantity']/load['units']).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
            shares[index]=value;remainder-=value
    ledger=workbook.create_sheet('Статьи расходов');write(ledger,['Строка','SKU','Статья','Сумма, ₽'])
    for index,(item,option) in enumerate(zip(snapshot['input']['items'],selections)):
        costs=sum((Decimal(v) for v in option['costs'].values()),Decimal(0))
        freight=shares.get(index,Decimal(0));profit=Decimal(option['result'])-freight
        write(lines,[item['marketplace'],item['scheme'],item['sku'],item.get('size',''),item['inventory_id'],item['cluster'],
            option['route_id'],option['quantity'],float(option['simulation']['retained_units']),float(option['simulation']['end_stock']),
            float(option['revenue']),float(costs),float(freight),float(profit),
            *[item['evidence'][k]['ref'] for k in ('demand','stock','finance')]])
        for key,value in option['costs'].items():write(ledger,[index+1,item['sku'],key,float(value)])
        write(ledger,[index+1,item['sku'],'allocated_batch_freight',float(freight)])
    routes=workbook.create_sheet('Условия маршрутов')
    write(routes,['Маршрут','Версия','Площадка','Схема','Откуда','Куда','Кластер','Действует с','Действует до','Источник','Владелец','Условия JSON'])
    for route in snapshot['routes']:
        write(routes,[route['route_id'],route['revision'],route['marketplace'],route['scheme'],route['origin'],route['destination'],
            route['cluster'],route['valid_from'],route['valid_to'],route['source'],route['owner'],json.dumps(route,ensure_ascii=False)])
    daily=workbook.create_sheet('Дневной баланс');write(daily,['Строка','SKU','День','Заказы','Исполнено','Выкупы','Потеряно','Приход','Возврат в продажу','Остаток'])
    for index,(item,option) in enumerate(zip(snapshot['input']['items'],selections)):
        for row in option['simulation']['daily']:
            write(daily,[index+1,item['sku'],row['day'],*[float(row[k]) for k in ('orders','fulfilled','retained','lost','arrivals','restocked','stock')]])
    inputs=workbook.create_sheet('Вводные сценария');write(inputs,['Параметр / строка','Значение'])
    for key,value in snapshot['input'].items():
        if key=='items':
            for i,item in enumerate(value):write(inputs,[f'Строка {i+1}',json.dumps(item,ensure_ascii=False)])
        else:write(inputs,[key,json.dumps(value,ensure_ascii=False)])
    for ws in workbook:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(name='Arial',size=11,bold=True,color='18202B');cell.fill=PatternFill('solid',fgColor='F2F4F7')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font=Font(name='Arial',size=10,color='18202B');cell.alignment=Alignment(vertical='top',wrap_text=False)
                if isinstance(cell.value,(int,float)):cell.number_format='#,##0.00;[Red]-#,##0.00;0.00'
        for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=24 if col<6 else 20
        ws.sheet_view.showGridLines=False
    summary.column_dimensions['A'].width=34;summary.column_dimensions['B'].width=75
    stream=BytesIO();workbook.save(stream)
    return stream.getvalue(),f"pulse-supply-{record['plan_id']}.xlsx"
